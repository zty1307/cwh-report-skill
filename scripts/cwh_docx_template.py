"""Packaged Word skeleton: scripts fill named slots, never ask models to format Word."""
from __future__ import annotations

from copy import deepcopy
import hashlib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'templates/cwh_fill_in_template.docx'
SLOTS = ('title', 'opening', 'propagation', 'topics', 'domestic', 'comments',
         'hotwords', 'overseas_media', 'overseas_comments', 'overseas_table', 'wechat_table')
SECTION_SLOTS = {
    '（一）总事件传播情况': 'propagation', '（二）子议题传播情况': 'topics',
    '（一）境内媒体自媒体情况': 'domestic', '（二）网民评论情况': 'comments',
    '（三）热词分布情况': 'hotwords', '（一）境外媒体情况': 'overseas_media',
    '（二）境外网民评论': 'overseas_comments', '（一）外媒报道列表（部分）': 'overseas_table',
    '（二）传播量较大的公众号文章': 'wechat_table',
}
CHAPTERS = ('一、舆情传播情况', '二、境内舆论情况', '三、境外舆论情况', '附录')


def element_text(element):
    return ''.join(node.text or '' for node in element.iter(qn('w:t')))


def slot_name(element):
    names = [node.get(qn('w:name'), '')[4:] for node in element.iter(qn('w:bookmarkStart'))
             if node.get(qn('w:name'), '').startswith('CWH_')]
    if len(names) > 1:
        raise ValueError('Multiple template slots in one block')
    return names[0] if names else None


def validate_template(document):
    found = [slot_name(node) for node in document._element.body if slot_name(node)]
    if tuple(found) != SLOTS:
        raise ValueError('Missing, duplicated or reordered Word template slots: ' + repr(found))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith('Heading')]
    expected = [CHAPTERS[0], *list(SECTION_SLOTS)[:2], CHAPTERS[1], *list(SECTION_SLOTS)[2:5],
                CHAPTERS[2], *list(SECTION_SLOTS)[5:7], CHAPTERS[3], *list(SECTION_SLOTS)[7:]]
    if headings != expected:
        raise ValueError('Fixed Word template section order changed')
    for node in document._element.body:
        text = element_text(node).strip()
        if text and not slot_name(node) and text not in CHAPTERS and text not in SECTION_SLOTS:
            raise ValueError('Unbound template text could leak into a report: ' + text[:100])


def import_reviewed_template(source, output=TEMPLATE):
    """Import the approved Word without touching it; consume its editing notes."""
    from cwh_writing_rules import editorial_profile, writing_rules
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve():
        raise ValueError('The human-reviewed source must never be overwritten')
    expected = editorial_profile()['source_template_sha256']
    if hashlib.sha256(source.read_bytes()).hexdigest() != expected:
        raise ValueError('Review source changed: read and approve the new version before importing')
    doc = Document(source)
    notes = {'3.这里有几个子议题就有几个点', '三是（依然是有几个子议题就几个点）'}
    for p in list(doc.paragraphs):
        if p.text.strip() in notes:
            p._p.getparent().remove(p._p)
            continue
        for run in p.runs:
            for old in ('{{具体判断及其必要理由或条件}}', '{{同一判断的互补依据或具体建议}}', '{{判断、依据及适用范围}}'):
                run.text = run.text.replace(old, '{{具体观点}}')
        if slot_name(p._p) == 'comments':
            for run in list(p.runs):
                run._r.getparent().remove(run._r)
            p.add_run(writing_rules()['comments']['lead_template'].format(stance_summaries='{{实际评论观点}}'))
            p.add_run('\n一是{{本议题评论观点}}。网民称，“{{真实原话一}}”“{{真实原话二}}”。'
                      '\n二是{{下一议题评论观点}}。网民称，“{{真实原话}}”。')
    validate_template(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    if hashlib.sha256(source.read_bytes()).hexdigest() != expected:
        raise ValueError('Source changed during import')
    return output


def apply_template(document):
    """Move generated OOXML into slots in the SAME package; keep link/image relations."""
    skeleton = Document(TEMPLATE)
    validate_template(skeleton)
    slots = {key: [] for key in SLOTS}
    active = 'title'
    found = []
    for node in list(document._element.body):
        if node.tag == qn('w:sectPr'):
            continue
        text = element_text(node)
        if node.tag == qn('w:p') and text in SECTION_SLOTS:
            active = SECTION_SLOTS[text]
            found.append(active)
        elif node.tag == qn('w:p') and text in CHAPTERS:
            continue
        else:
            slots[active].append(node)
            if active == 'title':
                active = 'opening'
    if found != list(SECTION_SLOTS.values()):
        raise ValueError('Generated report does not match the fixed section contract')
    body = document._element.body
    for node in list(body):
        if node.tag != qn('w:sectPr'):
            body.remove(node)
    for node in skeleton._element.body:
        if node.tag == qn('w:sectPr'):
            continue
        key = slot_name(node)
        if key:
            for content in slots[key]:
                body.insert(len(body) - 1, content)
        else:
            body.insert(len(body) - 1, deepcopy(node))
    return {'path': 'templates/' + TEMPLATE.name, 'sha256': hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
            'filled_slots': list(SLOTS), 'mode': 'fixed_skeleton', 'unfilled_slots': []}


def build_template(output=TEMPLATE):
    """Development-only reproducible builder. Never called by a report model worker."""
    from cwh_writing_rules import writing_rules
    frames = writing_rules()
    legacy = Document(ROOT / 'templates/formal_report_template_complete_20260714.docx')
    doc = Document()
    # Retain the measured baseline page system and established style definitions.
    section = deepcopy(legacy.sections[0]._sectPr)
    for tag in ('w:headerReference', 'w:footerReference'):
        for node in list(section.iter(qn(tag))):
            node.getparent().remove(node)
    doc._element.body.replace(doc.sections[0]._sectPr, section)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        original = doc.styles[name]._element
        original.getparent().replace(original, deepcopy(legacy.styles[name]._element))
    roles = {'Normal': ('仿宋_GB2312', 16, False), 'Title': ('华文中宋', 22, True),
             'Heading 1': ('黑体', 16, False), 'Heading 2': ('楷体_GB2312', 16, True),
             'Heading 3': ('仿宋_GB2312', 16, True)}
    for name, (font, size, bold) in roles.items():
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.bold = font, Pt(size), bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + key), font)
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + key), None)
        pf = style.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.first_line_indent = Pt(0 if name == 'Title' else 32)
        pf.line_spacing = 1
        pf.keep_together = False
        pf.page_break_before = False
        pf.keep_with_next = name.startswith('Heading')
    doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Page numbers are native fields, with no report-specific header/footer text.
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    def slot(key, text, style='Normal'):
        p = doc.add_paragraph(text, style)
        if style == 'Normal':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(SLOTS.index(key) + 1)); start.set(qn('w:name'), 'CWH_' + key)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(SLOTS.index(key) + 1))
        p._p.insert(1 if p._p.pPr is not None else 0, start); p._p.append(end)
        return p

    slot('title', '{{会议日期}}国务院常务会议舆情综述', 'Title')
    slot('opening', frames['document']['opening_with_chair'].format(
        chair_name=frames['document']['approved_chair_default'], date='{{会议日期}}', agenda='{{本期完整议程}}'))
    doc.add_heading(CHAPTERS[0], 1); doc.add_heading(list(SECTION_SLOTS)[0], 2)
    propagation = frames['propagation']
    slot('propagation', '\n'.join([
        propagation['total_template'].format(total_text='{{总量}}', focus_sentence=propagation['focus_template'].format(focus='{{已核实话题}}'), peak_sentence=propagation['peak_template'].format(peak_date='{{峰值日期}}')),
        propagation['domestic_template'].format(domestic_label='境内主流媒体', count='{{数量}}'),
        propagation['new_media_template'].format(wechat='{{微信量}}', weibo='{{微博量}}', video='{{视频号量}}', other='{{其他量}}'),
        propagation['overseas_template'].format(examples=propagation['overseas_examples_template'].format(sources='{{来源}}'), count='{{境外量}}'),
        '{{传播走势图}}',
    ]))
    doc.add_heading(list(SECTION_SLOTS)[1], 2)
    slot('topics', '{{脚本生成全部子议题统计表及传播量图，议题数量不限}}')
    doc.add_heading(CHAPTERS[1], 1); doc.add_heading(list(SECTION_SLOTS)[2], 2)
    domestic = slot('domestic', '1.{{本议题有依据的总判断}}\n一是{{共同判断甲}}。{{机构、职务、姓名}}认为，{{具体观点}}。{{媒体或自媒体账号全名}}称，{{具体观点}}。\n二是{{另一个不同判断}}。{{完整发言主体}}指出，{{具体观点}}。\n2.{{下一议题的总判断}}。{{只有一个成熟观点组时，直接接完整主体及观点，不单列“一是”}}')
    # Render real bold lead-ins in the preview, not merely a prose instruction about bold.
    import re
    for run in list(domestic.runs):
        run._r.getparent().remove(run._r)
    for index, line in enumerate([
        '1.{{本议题有依据的总判断}}',
        '一是{{共同判断甲}}。{{机构、职务、姓名}}认为，{{具体观点}}。{{媒体或自媒体账号全名}}称，{{具体观点}}。',
        '二是{{另一个不同判断}}。{{完整发言主体}}指出，{{具体观点}}。',
        '2.{{下一议题的总判断}}。{{只有一个成熟观点组时，直接接完整主体及观点，不单列“一是”}}',
    ]):
        if index:
            domestic.add_run().add_break()
        lead = re.match(r'^(?:\d+\.|[一二]是)\{\{.*?\}\}[。]?', line)
        if lead:
            domestic.add_run(lead[0]).bold = True
            domestic.add_run(line[lead.end():])
        else:
            domestic.add_run(line)
    doc.add_heading(list(SECTION_SLOTS)[3], 2)
    slot('comments', frames['comments']['lead_template'].format(stance_summaries='{{实际评论观点}}') + '\n一是{{原话支持的判断或诉求}}。网民称，“{{真实原话一}}”“{{真实原话二}}”。\n二是{{另一组评论的实际观点}}。网民称，“{{真实原话}}”。')
    doc.add_heading(list(SECTION_SLOTS)[4], 2)
    slot('hotwords', '从热词分布来看，“{{热词甲}}”“{{热词乙}}”等词主要涉及{{本期议题}}，相关观点{{已审核的具体判断}}。此外，“{{热词丙}}”“{{热词丁}}”等词主要涉及{{另一议题}}，相关观点{{已审核的建议或判断}}。\n{{腾讯字体词云，限高且自然分页}}')
    doc.add_heading(CHAPTERS[2], 1); doc.add_heading(list(SECTION_SLOTS)[5], 2)
    slot('overseas_media', '数据周期内，境外媒体{{依据实际类别分布填写报道概况}}。如{{来源甲}}文章《{{简体标题甲}}》、{{来源乙}}文章《{{简体标题乙}}》等，{{有解读时接“相关解读如下”；无解读时接“暂未发现可引用的评论性文章”}}。\n{{有解读才填此段：中心判断。来源＋文章标题＋原文判断及必要依据、条件。}}')
    doc.add_heading(list(SECTION_SLOTS)[6], 2)
    slot('overseas_comments', '{{有可引用评论时逐组填入观点与译文；无可引用评论时采用下列固定句}}\n' + frames['overseas']['no_comment_evidence_sentence'])
    doc.add_heading(CHAPTERS[3], 1); doc.add_heading(list(SECTION_SLOTS)[7], 2)
    slot('overseas_table', '{{脚本填表：序号｜来源｜报道日期｜超链接标题}}')
    doc.add_heading(list(SECTION_SLOTS)[8], 2)
    slot('wechat_table', '{{脚本填表：微信公号文章阅读量TOP；序号｜账号｜标题｜阅读量｜在看量}}')
    # Table prototypes use the exact production builders; no second style implementation.
    from formalize_cwh_report import add_topic_table, add_overseas_table, add_wechat_top_table
    for key, factory in (
        ('topics', lambda: add_topic_table(doc, [], True)),
        ('overseas_table', lambda: add_overseas_table(doc, [])),
        ('wechat_table', lambda: add_wechat_top_table(doc, [])),
    ):
        placeholder = next(node for node in doc._element.body if slot_name(node) == key)
        table = factory()
        paragraph = table.cell(0, 0).paragraphs[0]._p
        for mark in list(placeholder.iter(qn('w:bookmarkStart'))) + list(placeholder.iter(qn('w:bookmarkEnd'))):
            paragraph.append(deepcopy(mark))
        doc._element.body.replace(placeholder, table._tbl)
    doc.core_properties.title = '国务院常务会议舆情报告通用填空模板'
    doc.core_properties.author = ''
    validate_template(doc)
    output = Path(output); output.parent.mkdir(parents=True, exist_ok=True); doc.save(output)
    return output


if __name__ == '__main__':
    print(build_template())
