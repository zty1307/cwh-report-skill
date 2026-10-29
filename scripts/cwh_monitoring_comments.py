"""Source-row verified comment exports. Never simulate a platform API response."""
import hashlib,json,os,re
from pathlib import Path
import openpyxl

def read(p):return json.loads(Path(p).read_text('utf-8-sig'))
def write(p,v):Path(p).write_text(json.dumps(v,ensure_ascii=False,indent=2),'utf-8')
def clean(v):return '' if v is None or str(v).strip()=='-' else str(v).strip()
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def validate(capture):
    from run_cwh_sentiment_stage import stable_sample_id
    from cwh_weibo_capture import timestamp,window
    left,right=window(capture['monitoring_period']['start'],capture['monitoring_period']['end'])
    grouped={};seen=set();parents={p['id']:p for p in capture['posts']}
    for row in capture['rows']:grouped.setdefault(row['raw_file'],[]).append(row)
    for path,rows in grouped.items():
        if sha(path)!=capture['source_workbooks'][path]:raise ValueError('Monitoring comment workbook changed')
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        bysheet={}
        for row in rows:bysheet.setdefault(row['raw_sheet'],{})[int(row['raw_row'])]=row
        for sheet,wanted in bysheet.items():
            ws=wb[sheet];ws.reset_dimensions();it=ws.iter_rows(values_only=True);headers=list(next(it))
            for n,values in enumerate(it,2):
                if n not in wanted:continue
                r=dict(zip(headers,values));row=wanted[n];post=parents[row['parent_post_id']]
                parent_url=clean(r.get('原贴url')) or clean(r.get('发文链接'))
                parent_text=clean(r.get('原帖正文')) or clean(r.get('标题'))
                if clean(r.get('帖子类型')) not in {'评论','回复'}:raise ValueError('Article row cannot become comment')
                expected=(clean(r.get('正文内容')),clean(r.get('发表时间')),parent_url,parent_text)
                got=(row['text'],row['published_at'],row['url'],post['text'])
                if got!=expected:raise ValueError('Comment/parent differs from original workbook row')
                native=clean(r.get('发文ID'))
                expected_id=native or 'monitoring-row-'+hashlib.sha256(f'{sha(path)}|{sheet}|{n}'.encode()).hexdigest()[:24]
                if row['comment_id']!=expected_id:raise ValueError('Changed export record identity')
                if not left<=timestamp(row['published_at']).astimezone(left.tzinfo)<=right:raise ValueError('Comment outside period')
                if row['sample_id']!=stable_sample_id(row) or row['sample_id'] in seen:raise ValueError('Invalid sample identity')
                seen.add(row['sample_id'])
        wb.close()
    if len(seen)!=len(capture['rows']):raise ValueError('Comment source row missing')

def prepare_task(task_path):
    from run_cwh_sentiment_stage import stable_sample_id
    from cwh_comment_filters import is_procedural_only
    task=read(task_path);out=Path(task['stage_workspace']);out.mkdir(parents=True,exist_ok=True)
    plan=read(task['inputs']['research_plan']);registry=read(task['inputs']['source_registry'])
    topics=[r['topic'] for r in plan['topics']];period=plan['monitoring_period']
    raw=Path(os.environ['CWH_RAW_SAMPLE_DIR']);rows=[];posts={};files={};deferred=[];seen=set()
    for path in sorted(raw.glob('总-*.xlsx')):
        if any(x in path.name for x in ['热度分析','境外新闻','推特']):continue
        digest=sha(path);files[str(path)]=digest
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        for ws in wb:
            ws.reset_dimensions();it=ws.iter_rows(values_only=True);headers=list(next(it,()))
            if '帖子类型' not in headers:continue
            for n,values in enumerate(it,2):
                r=dict(zip(headers,values))
                if clean(r.get('帖子类型')) not in {'评论','回复'}:continue
                text=clean(r.get('正文内容'));url=clean(r.get('原贴url')) or clean(r.get('发文链接'))
                parent=clean(r.get('原帖正文')) or clean(r.get('标题'));date=clean(r.get('发表时间'))
                if not text or not url or not parent or not period['start'][:10]<=date[:10]<=period['end'][:10]:
                    deferred.append({'file':str(path),'sheet':ws.title,'row':n,'reason':'missing_text_parent_url_or_in_window_time'});continue
                platform=clean(r.get('站点名称'));key=(platform,url,text,date)
                if key in seen:continue
                seen.add(key)
                if is_procedural_only(text) or len(re.findall('[\u4e00-\u9fff]',text))<12:
                    deferred.append({'file':str(path),'sheet':ws.title,'row':n,'reason':'procedural_or_short_reading_reserve'});continue
                pid='export-parent-'+hashlib.sha256((url+'\n'+parent).encode()).hexdigest()[:20]
                posts[pid]={'id':pid,'url':url,'text':parent,'source':clean(r.get('原帖用户昵称')),
                    'published_at':clean(r.get('原贴发布时间')),'context_kind':'monitoring_export_parent_field'}
                native=clean(r.get('发文ID'));cid=native or 'monitoring-row-'+hashlib.sha256(f'{digest}|{ws.title}|{n}'.encode()).hexdigest()[:24]
                row={'text':text,'topic':'','platform':platform,'source':clean(r.get('昵称')) or clean(r.get('用户名')),
                    'published_at':date,'url':url,'comment_id':cid,'parent_post_id':pid,
                    'comment_identifier_kind':'platform_native_id' if native else 'monitoring_export_row_id_not_platform_id',
                    'like_count':clean(r.get('（拇指）点赞量')) or 0,'raw_file':str(path),'raw_sheet':ws.title,'raw_row':n,
                    'parent_post_title':parent,'parent_post_source':posts[pid]['source'],
                    'parent_post_published_at':posts[pid]['published_at']}
                row['sample_id']=stable_sample_id(row);rows.append(row)
        wb.close()
    write(out/'all_export_comment_candidates.json',rows)
    # Diversify real parent posts; metadata priority is not semantic approval.
    ranked=sorted(rows,key=lambda r:(-min(len(r['text']),180),r['sample_id']))
    groups={};chosen=[];body_chars=0
    for row in ranked:groups.setdefault(row['parent_post_id'],[]).append(row)
    for k in range(8):
        for pid,group in groups.items():
            if k>=len(group) or len(chosen)>=96:continue
            row=group[k];additional=0 if any(r['parent_post_id']==pid for r in chosen) else len(posts[pid]['text'])
            if body_chars+additional>110000:continue
            chosen.append(row);body_chars+=additional
    used={r['parent_post_id'] for r in chosen}
    capture={'schema_version':'1.0','capture_format':'monitoring_export_xlsx_v1','monitoring_period':period,
        'rows':chosen,'posts':[p for k,p in posts.items() if k in used],'raw_captures':[],
        'source_workbooks':files,'excluded':[],'sampling':{'all_candidates':len(rows),'selected':len(chosen),
            'unread':len(rows)-len(chosen),'rule':'parent_diversity_and_complete_parent_body_budget',
            'scope':'Sample review; unread reserve is not semantically excluded; no population sentiment estimate'}}
    validate(capture)
    cp=out/'monitoring_comment_capture.json';write(cp,capture);write(out/'script_comment_reserve.json',deferred)
    source_platforms={'weibo_comments':'微博','toutiao_public_comments':'今日头条','douyin_comments':'抖音app','bilibili_comments':'哔哩哔哩'}
    checks=[]
    for s in registry['sources']:
        if not(s.get('must_check') and s.get('tier')=='comment_platform'):continue
        count=sum(r['platform']==source_platforms.get(s['id']) for r in chosen)
        checks.append({'source_id':s['id'],'status':'hit' if count else 'access_failed',
            'execution_mode':'provided_monitoring_export_scan','queries_or_seed_urls':[],
            'eligible_comment_ids':[],'captured_rows':count,
            'blocker':'' if count else '本次导出扫描未取得该平台可审评论；本适配器未声称执行在线平台采集'})
    topic_checks=[{**c,'status':'access_failed','eligible_comment_ids':[],
        'blocker':'仅完成全局导出扫描，尚待模型逐议题判断；未声称在线采集覆盖'} for c in checks]
    audit={'registry_version':registry['version'],'checks':checks,'multi_topic_parents':True,
        'coverage_by_topic':[{'topic':t,'checks':topic_checks} for t in topics],
        'scope':'Provided export scan only. Topic relevance is assigned by model against full exported parent context.',
        'export_scan_candidates':len(rows),'selected_for_native_review':len(chosen),
        'identifier_note':'monitoring-row identifiers locate immutable exported rows; they are not invented native platform comment IDs'}
    ap=out/'monitoring_comment_collection_audit.json';write(ap,audit)
    task['inputs'].update(comment_capture=str(cp),comment_collection_audit=str(ap));write(task_path,task)
