from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

from cwh_pipeline_runtime import ArtifactSpec, PipelineRunner, StageOutcome, StageSpec, atomic_write_json
from domestic_evidence_mapping import (
    analysis_bundle_sha256,
    apply_semantic_review_packet,
    mapping_problem_messages,
    validate_analysis_mapping,
    validate_release_mapping,
)
from cwh_model_contract import build_task_payload, execution_profile, stage_budget_seconds
from normalize_cwh_analysis import normalize_analysis
from complete_cwh_evidence_structure import complete_analysis_structure
from cwh_viewpoint_gate import cluster_density_result
from report_rules import domestic_viewpoint_quality_issues


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_ROOT = SCRIPT_DIR.parent
SOURCE_REGISTRY_PATH = SKILL_ROOT / "config" / "source_registry.v1.json"
VALID_COVERAGE_STATES = {"hit", "no_relevant_result", "access_failed", "waiting_login", "not_applicable"}
VALID_DOMESTIC_CANDIDATE_DECISIONS = {"eligible", "duplicate", "excluded"}
VALID_DOMESTIC_DISCOVERY_ORIGINS = {
    "raw_monitoring",
    "fixed_registry_web",
    "open_web",
    "public_platform_web",
}
DOMESTIC_SATURATION_STOPS = {
    "two_consecutive_rounds_no_material_new_independent_viewpoint": 2,
    "coverage_minimum_then_one_zero_new_round_or_budget_exhausted": 1,
}
# Backward-compatible name for imported validators and legacy audit fixtures.
DOMESTIC_SATURATION_STOP = "two_consecutive_rounds_no_material_new_independent_viewpoint"


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def file_contract(path: str) -> dict[str, Any]:
    if not path:
        return {}
    resolved = Path(path).resolve()
    return {
        "path": str(resolved),
        "exists": resolved.exists(),
        "size": resolved.stat().st_size if resolved.exists() and resolved.is_file() else 0,
        "mtime_ns": resolved.stat().st_mtime_ns if resolved.exists() else 0,
    }


def directory_contract(path: str) -> dict[str, Any]:
    if not path:
        return {}
    resolved = Path(path).resolve()
    files = []
    if resolved.is_dir():
        files = [
            {
                "path": str(item.relative_to(resolved)),
                "size": item.stat().st_size,
                "mtime_ns": item.stat().st_mtime_ns,
            }
            for item in sorted(resolved.rglob("*"))
            if item.is_file()
        ]
    return {"path": str(resolved), "exists": resolved.is_dir(), "files": files}


def topic_titles(workbook: Path) -> list[str]:
    sys.path.insert(0, str(SCRIPT_DIR))
    from ingest_monitoring_workbook import ingest_workbook

    data = ingest_workbook(str(workbook))
    return [str(row.get("title") or "").strip() for row in data.get("topics") or [] if str(row.get("title") or "").strip()]


def _cell_int(value: Any) -> int:
    if value in (None, ""):
        return 0
    try:
        return int(float(str(value).replace(",", "")))
    except (TypeError, ValueError):
        return 0


def validate_data_workbook(path: Path, report_data: dict[str, Any]) -> list[str]:
    """Cross-check the delivered analyst workbook against report_data authority."""
    if not path.exists():
        return ["缺少cwh_data_workbook.xlsx"]
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        return [f"cwh_data_workbook.xlsx无法解析：{exc}"]

    problems: list[str] = []
    topics = report_data.get("topic_stats") or []
    system_data = report_data.get("system_data") or {}
    expected_sheets = ["总事件", *[f"子事件{idx}" for idx in range(1, len(topics) + 1)], "子事件数据汇总"]
    for sheet_name in expected_sheets:
        if sheet_name not in workbook.sheetnames:
            problems.append(f"数据工作簿缺少工作表：{sheet_name}")

    if "子事件数据汇总" in workbook.sheetnames:
        summary = workbook["子事件数据汇总"]
        for idx, topic in enumerate(topics, 1):
            row = idx + 2
            expected = [
                _cell_int(topic.get("domestic_media")),
                _cell_int(topic.get("self_media")),
                _cell_int(topic.get("overseas_media")),
                _cell_int(topic.get("spread_count", topic.get("total_samples"))),
            ]
            actual = [_cell_int(summary.cell(row, col).value) for col in range(3, 7)]
            if actual != expected:
                problems.append(f"子事件{idx}汇总C:F与报告权威数据不一致：{actual}≠{expected}")
            if sum(actual[:3]) != actual[3]:
                problems.append(f"子事件{idx}汇总总量不等于三类传播量之和")
            if not topic.get("sentiment_formal_ready"):
                sentiment_values = [summary.cell(row, col).value for col in range(7, 11)]
                if any(value not in (None, "") for value in sentiment_values):
                    problems.append(f"子事件{idx}情感样本未达门槛但工作簿G:J仍有值")

    def validate_trend(sheet_name: str, event: dict[str, Any], is_total: bool) -> None:
        if sheet_name not in workbook.sheetnames:
            return
        sheet = workbook[sheet_name]
        daily = event.get("daily") or []
        if not daily:
            problems.append(f"{sheet_name}缺少权威逐日数据")
            return
        keys = (
            ["domestic_mainstream", "overseas_media", "wechat_public", "weibo", "video_account", "new_media", "total_spread"]
            if is_total
            else ["domestic_mainstream", "overseas_media", "new_media", "total_spread"]
        )
        for offset, item in enumerate(daily, 4):
            expected = [_cell_int(item.get(key)) for key in keys]
            actual = [_cell_int(sheet.cell(offset, col).value) for col in range(2, 2 + len(keys))]
            if actual != expected:
                problems.append(f"{sheet_name}第{offset}行逐日数据不一致：{actual}≠{expected}")
        total_row = max(12, len(daily) + 5)
        expected_totals = [_cell_int((event.get("totals") or {}).get(key)) for key in keys]
        actual_totals = [_cell_int(sheet.cell(total_row, col).value) for col in range(2, 2 + len(keys))]
        if actual_totals != expected_totals:
            problems.append(f"{sheet_name}合计行与权威逐日合计不一致：{actual_totals}≠{expected_totals}")

    validate_trend("总事件", system_data.get("overall") or {}, True)
    subevents = system_data.get("subevents") or []
    for idx in range(1, len(topics) + 1):
        event = next((item for item in subevents if _cell_int(item.get("index")) == idx), {})
        validate_trend(f"子事件{idx}", event, False)
    workbook.close()
    return list(dict.fromkeys(problems))


def ai_task(
    runner: PipelineRunner,
    stage_id: str,
    *,
    task_type: str,
    expected_output: Path,
    inputs: dict[str, Any],
    rules: list[str],
    output_schema: dict[str, Any] | None = None,
) -> Path:
    path = runner.root / "tasks" / f"{stage_id}.json"
    payload = build_task_payload(
        stage_id=stage_id,
        task_type=task_type,
        expected_output=expected_output,
        inputs=inputs,
        rules=rules,
        profile_name=str(runner.input_contract.get("execution_profile") or ""),
        output_schema=output_schema,
        stage_workspace=runner.root / "worker" / stage_id,
    )
    Path(payload["stage_workspace"]).mkdir(parents=True, exist_ok=True)
    atomic_write_json(path, payload)
    return path


def maybe_run_ai_worker(
    runner: PipelineRunner,
    spec: StageSpec,
    task: Path,
    output: Path,
    *,
    allow_missing_output: bool = False,
) -> StageOutcome | None:
    template = runner.input_contract.get("ai_worker_command") or []
    if not isinstance(template, list) or not template:
        return None
    remaining = runner.remaining_budget_seconds(spec.stage_id)
    if remaining is not None:
        task_data = read_json(task)
        task_data["remaining_budget_seconds"] = max(1, int(remaining))
        atomic_write_json(task, task_data)
    command = [
        str(item).format(
            stage=spec.stage_id,
            task=str(task),
            output=str(output),
            job_dir=str(runner.root),
            skill_root=str(SKILL_ROOT),
        )
        for item in template
    ]
    timeout = stage_budget_seconds(spec.stage_id, str(runner.input_contract.get("execution_profile") or ""))
    code, log_path = runner.run_command(spec.stage_id, command, cwd=runner.root, timeout_seconds=timeout)
    if code != 0:
        return StageOutcome.failed(
            f"AI工作器退出码为{code}，详见{log_path}",
            retryable=(code in spec.transient_exit_codes and code != 124) or code == 65,
            error_code=f"ai_worker_exit_{code}",
            details={"task": str(task), "log_path": str(log_path)},
        )
    if not output.exists() and not allow_missing_output:
        return StageOutcome.failed(
            f"AI工作器结束但未生成约定产物：{output}",
            error_code="ai_output_missing",
            details={"task": str(task), "log_path": str(log_path)},
        )
    return None


def validate_analysis_bundle(path: Path, topics: list[str], *, require_semantic_review: bool = True, allow_deferred_corpus: bool = False) -> list[str]:
    problems: list[str] = []
    try:
        data = read_json(path)
    except Exception as exc:
        return [f"analysis_bundle.json无法解析：{type(exc).__name__}: {exc}"]
    by_topic = ((data.get("viewpoints") or {}).get("by_topic") or [])
    rows = {str(item.get("topic") or "").strip(): item for item in by_topic}
    registry = read_json(SOURCE_REGISTRY_PATH)
    registry_required_sources = {
        str(row.get("id"))
        for row in registry.get("sources") or []
        if row.get("must_check") and row.get("region") == "domestic" and row.get("tier") != "comment_platform"
    }
    research = ((data.get("research_audit") or {}).get("domestic_media_research") or {})
    required_sources = {
        str(value).strip()
        for value in research.get("required_source_ids") or registry_required_sources
        if str(value).strip()
    }
    public_corpus_path = path.parent / "public_article_evidence.json"
    public_corpus_ids: set[str] = set()
    public_corpus_ids_by_topic: dict[str, set[str]] = {}
    public_corpus_retained_by_topic: dict[str, set[str]] = {}
    public_corpus_by_id: dict[str, dict[str, Any]] = {}
    if public_corpus_path.exists():
        try:
            public_corpus = read_json(public_corpus_path)
        except Exception as exc:
            problems.append(f"public_article_evidence.json无法解析：{type(exc).__name__}: {exc}")
            public_corpus = {}
        corpus_rows = [row for row in public_corpus.get("candidates") or [] if isinstance(row, dict)]
        public_corpus_by_id = {str(row.get("record_id")): row for row in corpus_rows}
        public_corpus_ids = {
            str(row.get("record_id") or "").strip() for row in corpus_rows if str(row.get("record_id") or "").strip()
        }
        for topic_index, topic in enumerate(topics, 1):
            public_corpus_ids_by_topic[topic] = {
                str(row.get("record_id") or "").strip()
                for row in corpus_rows
                if topic_index in (row.get("topic_hits") or []) and str(row.get("record_id") or "").strip()
            }
        corpus_review = research.get("public_article_corpus_review") or {}
        indexed_ids = {}
        index_path = path.parent / "public_corpus_index.json"
        if allow_deferred_corpus and index_path.exists():
            index = read_json(index_path)
            if index.get("source_sha256") == hashlib.sha256(public_corpus_path.read_bytes()).hexdigest():
                indexed_ids = {row["topic"]: set(row["record_ids"]) for row in index.get("topics", [])}
        topic_reviews = {
            str(row.get("topic") or "").strip(): row
            for row in corpus_review.get("topic_reviews") or []
            if isinstance(row, dict)
        }
        for topic in topics:
            expected_ids = public_corpus_ids_by_topic.get(topic, set())
            review_row = topic_reviews.get(topic) or {}
            reviewed_ids = {str(value or "").strip() for value in review_row.get("reviewed_record_ids") or []}
            retained_ids = {str(value or "").strip() for value in review_row.get("retained_record_ids") or []}
            public_corpus_retained_by_topic[topic] = retained_ids
            excluded_rows = [row for row in review_row.get("excluded") or [] if isinstance(row, dict)]
            excluded_ids = {str(row.get("record_id") or "").strip() for row in excluded_rows}
            missing_review = sorted(expected_ids - reviewed_ids)
            deferred = set(review_row.get("deferred_record_ids") or [])
            valid_deferral = (allow_deferred_corpus and indexed_ids.get(topic) == expected_ids
                              and deferred == set(missing_review) and len(reviewed_ids) >= min(12, len(expected_ids))
                              and bool(str(review_row.get("deferral_reason") or "").strip()))
            if deferred and not valid_deferral:
                problems.append(f"{topic}公众文章待审清单缺少完整索引、最低实审量或与未审记录不一致")
            if missing_review and not valid_deferral:
                problems.append(f"{topic}原始公众文章证据池有{len(missing_review)}条未逐条审核")
            if retained_ids | excluded_ids != reviewed_ids:
                problems.append(f"{topic}公众文章证据池审核记录未将全部已审记录明确分为保留或排除")
            if any(not str(row.get("reason") or "").strip() for row in excluded_rows):
                problems.append(f"{topic}公众文章证据池排除记录缺少理由")
            invalid_ids = (reviewed_ids | retained_ids | excluded_ids) - expected_ids
            if invalid_ids:
                problems.append(f"{topic}公众文章证据池审核引用了不属于该议题的记录ID")
    if research.get("registry_version") != registry.get("version"):
        problems.append("境内媒体检索审计未使用当前稳定来源库版本")
    if research.get("open_search_completed") is not True:
        problems.append("稳定来源检索后尚未完成开放检索补充")
    source_by_id = {str(row.get("id") or ""): row for row in registry.get("sources") or []}
    platform_specific_tiers = set(
        (registry.get("execution") or {}).get("platform_specific_search_required_for_tiers") or []
    )
    coverage = {
        str(item.get("topic") or "").strip(): {
            str(check.get("source_id") or ""): check
            for check in item.get("checks") or []
            if isinstance(check, dict)
        }
        for item in research.get("coverage_by_topic") or []
    }
    candidate_pools = {
        str(item.get("topic") or "").strip(): item
        for item in research.get("candidate_pool_by_topic") or []
        if isinstance(item, dict)
    }
    research_pool_summary = research.get("pool_summary") or {}
    if public_corpus_path.exists() and int(research_pool_summary.get("raw_monitoring_candidate_count") or 0) <= 0:
        problems.append("境内媒体检索审计未单独统计原始监测全文池，不能用全网候选池替代")
    for summary_field in (
        "fixed_registry_web_candidate_count",
        "open_web_candidate_count",
        "public_platform_web_candidate_count",
    ):
        try:
            int(research_pool_summary.get(summary_field))
        except (TypeError, ValueError):
            problems.append(f"境内媒体检索审计缺少有效池计数：{summary_field}")
    for topic in topics:
        checks = coverage.get(topic, {})
        missing_sources = sorted(required_sources - set(checks))
        invalid_sources = sorted(
            source_id
            for source_id, check in checks.items()
            if source_id in required_sources and str(check.get("status") or "") not in VALID_COVERAGE_STATES
        )
        if missing_sources:
            problems.append(f"{topic}稳定来源未逐项留痕：" + "、".join(missing_sources))
        if invalid_sources:
            problems.append(f"{topic}来源状态无效：" + "、".join(invalid_sources))
        pool = candidate_pools.get(topic)
        candidates = [row for row in (pool or {}).get("candidates") or [] if isinstance(row, dict)]
        candidates_by_id = {
            str(row.get("candidate_id") or "").strip(): row
            for row in candidates
            if str(row.get("candidate_id") or "").strip()
        }
        for source_id in sorted(required_sources & set(checks)):
            check = checks[source_id]
            source = source_by_id.get(source_id) or {}
            status = str(check.get("status") or "")
            mode = str(check.get("execution_mode") or "")
            executed_queries = check.get("queries") or check.get("queries_or_urls") or []
            if status in {"hit", "no_relevant_result"} and not executed_queries:
                problems.append(f"{topic}{source.get('name') or source_id}没有实际查询留痕，不能记为{status}")
            if str(source.get("tier") or "") in platform_specific_tiers and status in {"hit", "no_relevant_result"} and mode not in {
                "platform_specific",
                "platform_search",
                "site_restricted_search",
                "platform_adapter",
                "browser_platform_search",
                "weread_platform_search",
            }:
                problems.append(f"{topic}{source.get('name') or source_id}没有平台定向检索及查询留痕，不能记为{status}")
            if status == "hit":
                hit_candidate_ids = check.get("candidate_ids")
                if not isinstance(hit_candidate_ids, list) or not hit_candidate_ids:
                    problems.append(f"{topic}{source.get('name') or source_id}标记hit但没有web候选ID")
                else:
                    for candidate_id in hit_candidate_ids:
                        candidate = candidates_by_id.get(str(candidate_id))
                        if not candidate:
                            problems.append(f"{topic}{source.get('name') or source_id}的hit候选不存在：{candidate_id}")
                            continue
                        if str(candidate.get("discovery_origin") or "") == "raw_monitoring":
                            problems.append(f"{topic}{source.get('name') or source_id}错误地用原始监测池候选证明全网命中")
                        discovered_source_id = str(candidate.get("discovered_source_id") or "").strip()
                        if discovered_source_id and discovered_source_id != source_id:
                            problems.append(f"{topic}{source.get('name') or source_id}的hit候选来源ID不一致：{candidate_id}")
            if status == "access_failed" and not str(check.get("blocker") or check.get("error") or "").strip():
                problems.append(f"{topic}{source.get('name') or source_id}访问失败但未记录阻断原因")
            if status == "waiting_login":
                if check.get("terminal") is not False:
                    problems.append(f"{topic}{source.get('name') or source_id}登录等待未明确terminal=false")
                if not str(check.get("blocker") or check.get("error") or "").strip():
                    problems.append(f"{topic}{source.get('name') or source_id}登录等待未记录验证码或登录阻断")
                if not (check.get("checkpoint") or check.get("resume_action")):
                    problems.append(f"{topic}{source.get('name') or source_id}登录等待缺少断点或恢复动作")
        eligible_candidates: dict[str, dict[str, Any]] = {}
        eligible_cluster_keys: set[str] = set()
        bounded_profile = False
        if not pool:
            problems.append(f"{topic}缺少全网候选池及饱和检索审计")
        else:
            if not candidates:
                problems.append(f"{topic}全网候选池为空")
            seen_candidate_ids: set[str] = set()
            for candidate_index, candidate in enumerate(candidates, 1):
                candidate_id = str(candidate.get("candidate_id") or "").strip()
                decision = str(candidate.get("decision") or "").strip()
                if not candidate_id:
                    problems.append(f"{topic}第{candidate_index}个候选缺少candidate_id")
                    continue
                if candidate_id in seen_candidate_ids:
                    problems.append(f"{topic}候选ID重复：{candidate_id}")
                seen_candidate_ids.add(candidate_id)
                if decision not in VALID_DOMESTIC_CANDIDATE_DECISIONS:
                    problems.append(f"{topic}候选{candidate_id}缺少有效decision")
                discovery_origin = str(candidate.get("discovery_origin") or "").strip()
                if discovery_origin not in VALID_DOMESTIC_DISCOVERY_ORIGINS:
                    problems.append(f"{topic}候选{candidate_id}缺少有效discovery_origin")
                if not str(candidate.get("decision_reason") or "").strip():
                    problems.append(f"{topic}候选{candidate_id}缺少筛选或排除理由")
                if not str(candidate.get("discovery_query_id") or "").strip():
                    problems.append(f"{topic}候选{candidate_id or candidate_index}缺少discovery_query_id，无法回溯实际查询")
                if not str(candidate.get("first_seen_round") or "").strip():
                    problems.append(f"{topic}候选{candidate_id or candidate_index}缺少first_seen_round")
                raw_record_id = str(candidate.get("raw_evidence_record_id") or "").strip()
                raw_source = public_corpus_by_id.get(raw_record_id)
                if raw_source and (candidate.get("url") != raw_source.get("url") or
                                   (candidate.get("source_snapshot") or {}).get("source_text") != raw_source.get("content")):
                    problems.append(f"{topic}监测来源候选的URL或全文与原始记录不一致：{raw_record_id}")
                if raw_record_id and raw_record_id not in public_corpus_ids:
                    problems.append(f"{topic}候选{candidate_id or candidate_index}引用了不存在的原始公众文章记录")
                if decision == "eligible":
                    missing = [
                        field
                        for field in (
                            "source",
                            "title",
                            "url",
                            "published_at",
                            "published_at_source_text",
                            "discovery_route",
                            "content_summary",
                            "source_type",
                            "source_tier",
                        )
                        if not str(candidate.get(field) or "").strip()
                    ]
                    if missing:
                        problems.append(f"{topic}合格候选{candidate_id}缺少字段：" + "、".join(missing))
                    if not str(candidate.get("url") or "").startswith(("http://", "https://")):
                        problems.append(f"{topic}合格候选{candidate_id}没有原始链接")
                    if candidate.get("published_at_verified_from_source") is not True:
                        problems.append(f"{topic}合格候选{candidate_id}的发布时间未从原始页面或原始监测表复核")
                    published_date_match = re.match(r"^(\d{4}-\d{2}-\d{2})", str(candidate.get("published_at") or ""))
                    monitoring_start = str((data.get("metadata") or {}).get("monitoring_start") or "")[:10]
                    monitoring_end = str((data.get("metadata") or {}).get("monitoring_end") or "")[:10]
                    if not published_date_match:
                        problems.append(f"{topic}合格候选{candidate_id}发布时间格式无效")
                    elif monitoring_start and monitoring_end:
                        published_date = published_date_match.group(1)
                        if not monitoring_start <= published_date <= monitoring_end:
                            problems.append(
                                f"{topic}合格候选{candidate_id}发布时间{published_date}超出监测期"
                                f"{monitoring_start}至{monitoring_end}"
                            )
                    cluster_key = str(candidate.get("viewpoint_cluster_key") or "").strip()
                    if not cluster_key:
                        problems.append(f"{topic}合格候选{candidate_id}缺少viewpoint_cluster_key")
                    else:
                        eligible_cluster_keys.add(cluster_key)
                    eligible_candidates[candidate_id] = candidate
            linked_raw_ids = {
                str(candidate.get("raw_evidence_record_id") or "").strip()
                for candidate in candidates
                if str(candidate.get("raw_evidence_record_id") or "").strip()
            }
            missing_linked = sorted(public_corpus_retained_by_topic.get(topic, set()) - linked_raw_ids)
            if missing_linked:
                problems.append(f"{topic}公众文章证据池有{len(missing_linked)}条保留记录未进入候选池")
            saturation = pool.get("saturation") or {}
            rounds = [row for row in saturation.get("rounds") or [] if isinstance(row, dict)]
            query_executions: dict[str, dict[str, Any]] = {}
            if saturation.get("completed") is not True:
                problems.append(f"{topic}尚未完成候选池饱和检索")
            stop_reason = str(saturation.get("stop_reason") or "")
            bounded_profile = stop_reason == "coverage_minimum_then_one_zero_new_round_or_budget_exhausted"
            required_zero_rounds = DOMESTIC_SATURATION_STOPS.get(stop_reason)
            if required_zero_rounds is None:
                problems.append(f"{topic}候选池没有按受支持的检索停止规则结束")
                required_zero_rounds = 2
            if len(rounds) < required_zero_rounds or any(
                int(row.get("new_independent_viewpoints") or 0) != 0
                for row in rounds[-required_zero_rounds:]
            ):
                problems.append(f"{topic}缺少{required_zero_rounds}轮无新增高相关独立观点的检索记录")
            for round_index, round_row in enumerate(rounds, 1):
                executions = [row for row in round_row.get("executions") or [] if isinstance(row, dict)]
                if not executions:
                    problems.append(f"{topic}第{round_index}轮没有保存具体查询执行和结果URL快照")
                    continue
                for execution_index, execution in enumerate(executions, 1):
                    query_id = str(execution.get("query_id") or "").strip()
                    query = str(execution.get("query") or "").strip()
                    backend = str(execution.get("backend") or "").strip()
                    route = str(execution.get("route") or "").strip()
                    status = str(execution.get("status") or "").strip()
                    executed_at = str(execution.get("executed_at") or "").strip()
                    result_urls = execution.get("result_urls")
                    retained_ids = execution.get("retained_candidate_ids")
                    if not query_id:
                        problems.append(f"{topic}第{round_index}轮第{execution_index}个查询缺少query_id")
                    elif query_id in query_executions:
                        problems.append(f"{topic}查询ID重复：{query_id}")
                    else:
                        query_executions[query_id] = execution
                    if len(query) < 8 or query in {"stable_registry", "general_open_search", "toutiao_site_search", "wechat_site_search"}:
                        problems.append(f"{topic}第{round_index}轮第{execution_index}个查询不是可复核的具体查询串")
                    if not backend or route not in {"stable_registry", "open_web", "public_platform"} or not executed_at:
                        problems.append(f"{topic}查询{query_id or execution_index}缺少backend、有效route或executed_at")
                    if status not in {"completed", "access_failed", "waiting_login"}:
                        problems.append(f"{topic}查询{query_id or execution_index}状态无效：{status or 'missing'}")
                    if status in {"access_failed", "waiting_login"}:
                        if not str(execution.get("blocker") or execution.get("error") or "").strip():
                            problems.append(f"{topic}查询{query_id or execution_index}{status}但未记录阻断原因")
                        if status == "waiting_login" and execution.get("terminal") is not False:
                            problems.append(f"{topic}查询{query_id or execution_index}错误地把waiting_login当成终点")
                    else:
                        if not isinstance(result_urls, list) or not isinstance(retained_ids, list):
                            problems.append(f"{topic}查询{query_id or execution_index}未保存result_urls和retained_candidate_ids数组")
                        try:
                            result_count = int(execution.get("result_count"))
                        except (TypeError, ValueError):
                            problems.append(f"{topic}查询{query_id or execution_index}缺少有效result_count")
                        else:
                            if isinstance(result_urls, list) and result_count < len(result_urls):
                                problems.append(f"{topic}查询{query_id or execution_index}的result_count小于URL快照数量")
            for candidate in candidates:
                candidate_id = str(candidate.get("candidate_id") or "").strip()
                query_id = str(candidate.get("discovery_query_id") or "").strip()
                execution = query_executions.get(query_id)
                if query_id and not execution:
                    problems.append(f"{topic}候选{candidate_id}引用了不存在的查询执行：{query_id}")
                    continue
                candidate_url = str(candidate.get("url") or "").strip()
                if candidate_url and execution and execution.get("status") == "completed":
                    result_urls = [str(value) for value in execution.get("result_urls") or []]
                    if candidate_url not in result_urls:
                        problems.append(f"{topic}候选{candidate_id}的原始链接不在查询{query_id}的结果快照中")
            mapped_result_urls = {
                (str(candidate.get("discovery_query_id") or "").strip(), str(candidate.get("url") or "").strip())
                for candidate in candidates
                if str(candidate.get("url") or "").strip()
            }
            for query_id, execution in query_executions.items():
                if execution.get("status") != "completed":
                    continue
                for result_url in [str(value).strip() for value in execution.get("result_urls") or [] if str(value).strip()]:
                    if (query_id, result_url) not in mapped_result_urls:
                        problems.append(f"{topic}查询{query_id}的结果URL未进入候选判定：{result_url}")
            route_coverage = {
                str(row.get("route") or ""): str(row.get("status") or "")
                for row in saturation.get("route_coverage") or []
                if isinstance(row, dict)
            }
            for required_route in ("open_web", "public_platform"):
                if route_coverage.get(required_route) not in {"completed", "access_failed", "waiting_login"}:
                    problems.append(f"{topic}候选池停止前未完成{required_route}检索或记录真实阻断")
                route_executions = [
                    row
                    for row in query_executions.values()
                    if str(row.get("route") or "") == required_route
                ]
                if route_coverage.get(required_route) == "completed" and not any(
                    str(row.get("status") or "") == "completed" for row in route_executions
                ):
                    problems.append(f"{topic}{required_route}被标记completed但没有成功执行记录")
        item = rows.get(topic)
        if not item:
            problems.append(f"缺少子议题观点：{topic}")
            continue
        clusters = item.get("clusters") or []
        if not clusters:
            problems.append(f"子议题没有观点簇：{topic}")
            continue
        if len(clusters) < 2:
            exception = item.get("single_cluster_exception") or {}
            if len(eligible_cluster_keys) >= 2:
                problems.append(f"{topic}合格候选池已有{len(eligible_cluster_keys)}个观点家族，但正式分析仅形成1个观点簇")
            elif not all(str(exception.get(field) or "").strip() for field in ("reason", "search_evidence", "reviewed_by")):
                problems.append(f"{topic}仅形成1个观点簇，且缺少可复核的single_cluster_exception")
        formal_candidate_ids: set[str] = set()
        for index, cluster in enumerate(clusters, 1):
            if not str(cluster.get("summary") or "").strip():
                problems.append(f"{topic}第{index}个观点簇缺少态度型结论")
            evidence = cluster.get("evidence") or []
            if not evidence:
                problems.append(f"{topic}第{index}个观点簇没有证据")
                continue
            if not any(str(row.get("url") or "").startswith(("http://", "https://")) for row in evidence):
                problems.append(f"{topic}第{index}个观点簇没有原始链接")
            if not str(cluster.get("details") or "").strip():
                problems.append(f"{topic}第{index}个观点簇没有成文说明")
            details_text = str(cluster.get("details") or cluster.get("analysis") or "")
            density = cluster_density_result(cluster)
            if not density["passed"]:
                problems.append(f"{topic}第{index}个观点簇{density['message']}")
            for evidence_index, evidence_row in enumerate(evidence, 1):
                candidate_id = str(evidence_row.get("candidate_id") or "").strip()
                if candidate_id:
                    formal_candidate_ids.add(candidate_id)
                if not candidate_id:
                    problems.append(f"{topic}第{index}个观点簇第{evidence_index}条证据缺少candidate_id")
                elif candidate_id not in eligible_candidates:
                    problems.append(f"{topic}第{index}个观点簇引用了未入选候选池的证据：{candidate_id}")
                else:
                    candidate_url = str(eligible_candidates[candidate_id].get("url") or "").strip()
                    evidence_url = str(evidence_row.get("url") or "").strip()
                    if candidate_url and evidence_url and candidate_url != evidence_url:
                        problems.append(f"{topic}候选{candidate_id}与成文证据的原始链接不一致")
                missing_evidence_fields = [
                    field
                    for field in (
                        "attribution",
                        "attribution_status",
                        "source_excerpt",
                        "formal_claim",
                        "wording_fidelity",
                        "selection_reason",
                    )
                    if not str(evidence_row.get(field) or "").strip()
                ]
                if missing_evidence_fields:
                    problems.append(
                        f"{topic}第{index}个观点簇第{evidence_index}条证据缺少字段："
                        + "、".join(missing_evidence_fields)
                    )
                attribution = str(evidence_row.get("attribution") or "").strip()
                source = str(evidence_row.get("source") or "").strip()
                attribution_status = str(evidence_row.get("attribution_status") or "").lower()
                mention_keys = [attribution, source]
                if attribution_status == "named_person":
                    person_match = re.search(r"([\u4e00-\u9fff]{2,4})$", attribution)
                    if person_match:
                        mention_keys.insert(0, person_match.group(1))
                if details_text and not any(key and key in details_text for key in mention_keys):
                    problems.append(
                        f"{topic}第{index}个观点簇的成文说明未写入候选{candidate_id or evidence_index}的具体人物、媒体或账号主体"
                    )
        reserve_candidates = {
            candidate_id
            for candidate_id, candidate in eligible_candidates.items()
            if str(candidate.get("formal_use") or "").strip().lower() == "reserve"
        }
        if reserve_candidates and not bounded_profile:
            problems.append(f"{topic}仅bounded限时档允许使用formal_use=reserve")
        for candidate_id in sorted(reserve_candidates):
            if not str(eligible_candidates[candidate_id].get("reserve_reason") or "").strip():
                problems.append(f"{topic}候选{candidate_id}标记为reserve但缺少reserve_reason")
        missing_formal_candidates = sorted(set(eligible_candidates) - formal_candidate_ids - reserve_candidates)
        if missing_formal_candidates:
            problems.append(
                f"{topic}仍有{len(missing_formal_candidates)}条合格独立样本未归入正式观点和正文："
                + "、".join(missing_formal_candidates[:12])
            )
        if bounded_profile:
            if len(formal_candidate_ids) > 12:
                problems.append(f"{topic}bounded限时档正式声音{len(formal_candidate_ids)}条，超过12条硬上限")
            if len(formal_candidate_ids) < 4:
                shortfall = item.get("evidence_shortfall") or {}
                if not all(str(shortfall.get(field) or "").strip() for field in ("reason", "search_evidence", "reviewed_by")):
                    problems.append(
                        f"{topic}bounded限时档正式声音仅{len(formal_candidate_ids)}条，"
                        "低于4条且缺少evidence_shortfall审计"
                    )
    for issue in domestic_viewpoint_quality_issues(data):
        if issue.get("severity") == "error":
            problems.append(str(issue.get("message") or issue.get("code") or "境内观点质量错误"))
    problems.extend(mapping_problem_messages(validate_analysis_mapping(data, require_semantic_review=require_semantic_review)))
    return problems


def validate_hotword_audit(path: Path, topics: list[str]) -> list[str]:
    try:
        data = read_json(path)
    except Exception as exc:
        return [f"热词审核文件无法解析：{type(exc).__name__}: {exc}"]
    method = str(data.get("review_method") or data.get("method") or "")
    if "ai" not in method.lower() and "human_reviewed" not in method.lower():
        return ["热词审核未标明AI语义审核或人工复核"]
    if "human_reviewed" not in method.lower() and not data.get("second_pass_completed"):
        return ["热词AI二次复核尚未完成"]
    selected = data.get("selected") or []
    minimum = int(((data.get("settings") or {}).get("minimum_term_count") or 36))
    if len(selected) < minimum:
        return [f"热词仅{len(selected)}个，低于本期配置下限{minimum}个"]
    words = [str(row.get("term") or row.get("word") or "").strip() for row in selected]
    if any(not value for value in words) or len(set(words)) != len(words):
        return ["热词存在空值或重复值"]
    return []


def validate_sentiment(summary_path: Path, handoff_path: Path, topics: list[str]) -> list[str]:
    problems: list[str] = []
    try:
        summary = read_json(summary_path)
    except Exception as exc:
        return [f"情感汇总无法解析：{type(exc).__name__}: {exc}"]
    by_topic = {str(row.get("title") or "").strip(): row for row in summary.get("topics") or []}
    registry = read_json(SOURCE_REGISTRY_PATH)
    required_comment_sources = {
        str(row.get("id"))
        for row in registry.get("sources") or []
        if row.get("must_check") and row.get("tier") == "comment_platform"
    }
    collection = summary.get("collection_audit") or {}
    if collection.get("waiting_login_terminal") is True or str(collection.get("terminal_status") or "") == "waiting_login":
        problems.append("评论采集错误地把waiting_login当成终点；必须继续已批准的免登录或系统兜底并记录结果")
    if collection.get("registry_version") != registry.get("version"):
        problems.append("评论采集审计未使用当前稳定来源库版本")
    checks = {str(row.get("source_id") or ""): str(row.get("status") or "") for row in collection.get("checks") or []}
    missing = sorted(required_comment_sources - set(checks))
    invalid = sorted(source_id for source_id, status in checks.items() if source_id in required_comment_sources and status not in VALID_COVERAGE_STATES)
    if missing:
        problems.append("评论平台未逐项留痕：" + "、".join(missing))
    if invalid:
        problems.append("评论平台状态无效：" + "、".join(invalid))
    source_by_id = {str(row.get("id") or ""): row for row in registry.get("sources") or []}
    coverage_by_topic = {
        str(row.get("topic") or "").strip(): {
            str(check.get("source_id") or ""): check
            for check in row.get("checks") or []
            if isinstance(check, dict)
        }
        for row in collection.get("coverage_by_topic") or []
        if isinstance(row, dict)
    }
    for topic in topics:
        topic_checks = coverage_by_topic.get(topic, {})
        missing_topic_sources = sorted(required_comment_sources - set(topic_checks))
        if missing_topic_sources:
            problems.append(f"{topic}评论采集缺少逐平台留痕：" + "、".join(missing_topic_sources))
        for source_id in sorted(required_comment_sources & set(topic_checks)):
            check = topic_checks[source_id]
            source = source_by_id.get(source_id) or {}
            status = str(check.get("status") or "")
            if status not in VALID_COVERAGE_STATES:
                problems.append(f"{topic}{source.get('name') or source_id}评论采集状态无效：{status or 'missing'}")
                continue
            if status in {"hit", "no_relevant_result"}:
                mode = str(check.get("execution_mode") or "").strip()
                query_evidence = check.get("queries_or_seed_urls") or check.get("queries") or []
                if not mode or not isinstance(query_evidence, list) or not query_evidence:
                    problems.append(f"{topic}{source.get('name') or source_id}没有查询或原帖种子留痕")
                if not isinstance(check.get("eligible_comment_ids"), list):
                    problems.append(f"{topic}{source.get('name') or source_id}没有eligible_comment_ids数组")
                try:
                    int(check.get("result_count"))
                except (TypeError, ValueError):
                    problems.append(f"{topic}{source.get('name') or source_id}缺少有效result_count")
            if status == "access_failed" and not str(check.get("blocker") or check.get("error") or "").strip():
                problems.append(f"{topic}{source.get('name') or source_id}访问失败但未记录阻断原因")
        row = by_topic.get(topic)
        if not row:
            problems.append(f"缺少子议题情感结果：{topic}")
            continue
        status = str(row.get("status") or "").strip()
        denominator = int(row.get("denominator") or 0)
        if status == "ready" and denominator <= 0:
            problems.append(f"子议题情感结果标记ready但分母无效：{topic}")
        elif status == "insufficient_sample" and denominator <= 0:
            problems.append(f"子议题情感结果标记样本不足但没有观察样本：{topic}")
        elif status not in {"ready", "insufficient_sample", "pending", "no_public_evidence"}:
            problems.append(f"子议题情感状态无效：{topic}={status or 'missing'}")
    try:
        with handoff_path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    except Exception as exc:
        problems.append(f"评论交接表无法读取：{type(exc).__name__}: {exc}")
        return problems
    if not rows:
        problems.append("评论交接表为空")
    required = {"sample_id", "topic", "content", "url", "ai_formal_include", "topic_comment_heading"}
    fields = set(rows[0]) if rows else set()
    missing = required - fields
    if missing:
        problems.append("评论交接表缺少字段：" + "、".join(sorted(missing)))
    formal_by_topic: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        if str(row.get("ai_formal_include") or "").strip().lower() in {"true", "1", "yes", "y"}:
            formal_by_topic.setdefault(str(row.get("topic") or "").strip(), []).append(row)
    for topic, topic_rows in formal_by_topic.items():
        headings = {str(row.get("topic_comment_heading") or "").strip() for row in topic_rows}
        headings.discard("")
        if len(topic_rows) > 1 and (len(headings) != 1 or any(not str(row.get("topic_comment_heading") or "").strip() for row in topic_rows)):
            problems.append(f"{topic}的多条正式评论未填写唯一一致的topic_comment_heading")
    return problems


def validate_foreign(audit_path: Path, supplements_path: Path) -> list[str]:
    problems: list[str] = []
    try:
        audit = read_json(audit_path)
    except Exception as exc:
        return [f"境外采集审计无法解析：{type(exc).__name__}: {exc}"]
    attempted = bool(
        audit.get("attempted")
        or audit.get("run_attempted")
        or audit.get("completed")
        or audit.get("collection_completed")
    )
    exit_code = audit.get("command_exit_code", audit.get("exit_code", 0))
    verified_completed = bool(
        audit.get("collection_completed") is True
        and int(audit.get("collector_exit_code", 0) or 0) == 0
    )
    if not attempted:
        problems.append("境外媒体与境外网民评论尚未实际执行采集")
    if exit_code not in (None, 0, "0") and not verified_completed:
        problems.append(f"境外采集退出码为{exit_code}")
    try:
        supplements = read_json(supplements_path)
    except Exception as exc:
        problems.append(f"境外补充证据无法解析：{type(exc).__name__}: {exc}")
        return problems
    if not isinstance(supplements.get("media", []), list) or not isinstance(supplements.get("comments", []), list):
        problems.append("境外补充证据必须分别包含media和comments数组")
        return problems
    for index, row in enumerate(supplements.get("media") or [], 1):
        if row.get("formal_include") is not True:
            continue
        missing = [
            field
            for field in ("title_cn_simplified", "source_cn_simplified", "summary_cn_simplified")
            if not str(row.get(field) or "").strip()
        ]
        if missing:
            problems.append(f"第{index}条拟入正式集合的境外报道缺少简体审核字段：{'、'.join(missing)}")
        if row.get("simplified_chinese_reviewed") is not True:
            problems.append(f"第{index}条拟入正式集合的境外报道尚未完成简体中文AI复核")
    return problems


def build_specs() -> list[StageSpec]:
    return [
        StageSpec("preflight", "运行环境与Skill完整性预检", artifacts=(ArtifactSpec("preflight", "artifacts/preflight.json"),), max_attempts=1),
        StageSpec("intake", "输入识别与契约固化", dependencies=("preflight",), artifacts=(ArtifactSpec("intake", "artifacts/intake.json"),), max_attempts=1),
        StageSpec(
            "workbook",
            "标准总表生成与审核",
            dependencies=("intake",),
            artifacts=(ArtifactSpec("workbook", "artifacts/CWH舆情情况_标准总表.xlsx", minimum_bytes=1024),
                       ArtifactSpec("public_article_evidence", "artifacts/public_article_evidence.json", required=False)),
            max_attempts=2,
        ),
        StageSpec(
            "research_plan",
            "六类检索任务与证据口径",
            dependencies=("workbook",),
            artifacts=(ArtifactSpec("research_plan", "artifacts/research_plan.json"),),
            max_attempts=2,
        ),
        StageSpec(
            "domestic_viewpoints",
            "境内媒体自媒体观点检索与AI审核",
            dependencies=("research_plan",),
            artifacts=(ArtifactSpec("analysis_bundle", "artifacts/analysis_bundle.json"),),
            max_attempts=3,
            kind="ai",
        ),
        StageSpec(
            "domestic_evidence_verification",
            "境内观点原文映射与独立语义复核",
            dependencies=("domestic_viewpoints",),
            artifacts=(
                ArtifactSpec("domestic_evidence_semantic_review", "artifacts/domestic_evidence_semantic_review.json"),
                ArtifactSpec("analysis_bundle_verified", "artifacts/analysis_bundle_verified.json"),
                ArtifactSpec("domestic_evidence_mapping_audit", "artifacts/domestic_evidence_mapping_audit.json"),
            ),
            max_attempts=3,
            kind="ai",
        ),
        StageSpec(
            "domestic_comments_sentiment",
            "真实网民评论与自有情感分析",
            dependencies=("research_plan",),
            artifacts=(
                ArtifactSpec("comment_handoff", "artifacts/report_comment_handoff.csv"),
                ArtifactSpec("sentiment_results", "artifacts/sentiment_results.csv"),
                ArtifactSpec("sentiment_summary", "artifacts/sentiment_workbook_summary.json"),
            ),
            max_attempts=3,
            kind="collector_ai",
        ),
        StageSpec(
            "overseas_evidence",
            "境外媒体与境外网民评论采集审核",
            dependencies=("research_plan",),
            artifacts=(
                ArtifactSpec("overseas_supplements", "artifacts/public_overseas_supplements.json"),
                ArtifactSpec("foreign_collection_audit", "artifacts/foreign_collection_audit.json"),
            ),
            max_attempts=3,
            kind="collector_ai",
        ),
        StageSpec(
            "hotwords",
            "热词证据审核与词云准备",
            dependencies=("domestic_evidence_verification", "domestic_comments_sentiment"),
            artifacts=(ArtifactSpec("hotword_audit", "artifacts/hotword_audit.json"),),
            max_attempts=3,
            kind="ai",
        ),
        StageSpec(
            "render",
            "正式Word、Excel与工作台生成",
            dependencies=("workbook", "domestic_evidence_verification", "domestic_comments_sentiment", "overseas_evidence", "hotwords"),
            artifacts=(
                ArtifactSpec("report_data", "report/report_data.json"),
                ArtifactSpec("word", "report/cwh_formal_report.docx", minimum_bytes=1024),
                ArtifactSpec("dashboard", "report/cwh_dashboard.html", minimum_bytes=1024),
                ArtifactSpec("data_workbook", "report/cwh_data_workbook.xlsx", minimum_bytes=1024),
                ArtifactSpec("cwh_audit", "report/cwh_audit.json"),
            ),
            max_attempts=2,
        ),
        StageSpec(
            "delivery_gate",
            "成品一致性与正式交付门禁",
            dependencies=("render",),
            artifacts=(
                ArtifactSpec("delivery_audit", "artifacts/delivery_gate.json"),
                ArtifactSpec("release_mapping_audit", "report/domestic_evidence_mapping_audit.json"),
                ArtifactSpec("cwh_audit", "report/cwh_audit.json"),
            ),
            max_attempts=2,
        ),
    ]


class CwhPipeline:
    def __init__(self, job_dir: Path, contract: dict[str, Any]):
        self.job_dir = job_dir.resolve()
        self.contract = contract
        self.artifacts = self.job_dir / "artifacts"
        self.tasks = self.job_dir / "tasks"
        self.report = self.job_dir / "report"
        for directory in (self.artifacts, self.tasks, self.report):
            directory.mkdir(parents=True, exist_ok=True)
        executors = {
            "preflight": self.preflight,
            "intake": self.intake,
            "workbook": self.workbook,
            "research_plan": self.research_plan,
            "domestic_viewpoints": self.domestic_viewpoints,
            "domestic_evidence_verification": self.domestic_evidence_verification,
            "domestic_comments_sentiment": self.domestic_comments_sentiment,
            "overseas_evidence": self.overseas_evidence,
            "hotwords": self.hotwords,
            "render": self.render,
            "delivery_gate": self.delivery_gate,
        }
        self.runner = PipelineRunner(
            self.job_dir,
            build_specs(),
            executors,
            pipeline_name="cwh-formal-report",
            input_contract=contract,
        )

    def preflight(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        target = self.artifacts / "preflight.json"
        command = [
            sys.executable,
            str(SCRIPT_DIR / "cwh_preflight.py"),
            "--skill-root",
            str(SKILL_ROOT),
            "--output",
            str(target),
        ]
        code, log_path = runner.run_command(
            spec.stage_id,
            command,
            cwd=SKILL_ROOT,
            timeout_seconds=stage_budget_seconds("preflight", str(self.contract.get("execution_profile") or "")),
        )
        if code != 0:
            payload = read_json(target) if target.exists() else {}
            repair = str(payload.get("repair_command") or "")
            return StageOutcome.failed(
                "运行环境预检未通过" + (f"；建议执行：{repair}" if repair else ""),
                error_code="preflight_failed",
                details={"log_path": str(log_path), "preflight": str(target), "problems": payload.get("problems") or []},
            )
        return StageOutcome.succeeded("运行依赖与Skill关键文件预检通过。", details={"log_path": str(log_path)})

    def intake(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        agenda = str(self.contract.get("agenda") or "").strip()
        workbook = str(self.contract.get("system_workbook") or "").strip()
        raw_dir = str(self.contract.get("raw_input_dir") or "").strip()
        problems = []
        if not agenda:
            problems.append("缺少会议日期或完整议程")
        if bool(workbook) == bool(raw_dir):
            problems.append("必须且只能提供system_workbook或raw_input_dir之一")
        if workbook and not Path(workbook).is_file():
            problems.append(f"系统总表不存在：{workbook}")
        if raw_dir and not Path(raw_dir).is_dir():
            problems.append(f"原始数据目录不存在：{raw_dir}")
        if problems:
            return StageOutcome.failed("；".join(problems), error_code="invalid_input")
        payload = {
            "agenda": agenda,
            "input_mode": "standard_workbook" if workbook else "raw_bundle",
            "system_workbook": workbook,
            "raw_input_dir": raw_dir,
            "metadata": self.contract.get("metadata") or "",
            "contract_fingerprint": runner.state["input_fingerprint"],
        }
        target = self.artifacts / "intake.json"
        atomic_write_json(target, payload)
        return StageOutcome.succeeded("输入已固化，后续节点只能消费本任务产物。")

    def workbook(self, runner: PipelineRunner, spec: StageSpec, *, review_applied: bool = False) -> StageOutcome:
        target = self.artifacts / "CWH舆情情况_标准总表.xlsx"
        source = str(self.contract.get("system_workbook") or "").strip()
        supplied_public_evidence = str(self.contract.get("public_article_evidence") or "").strip()
        public_evidence_target = self.artifacts / "public_article_evidence.json"
        if supplied_public_evidence and not public_evidence_target.exists():
            evidence_source = Path(supplied_public_evidence)
            if not evidence_source.is_file():
                return StageOutcome.failed(
                    f"public_article_evidence不存在：{evidence_source}",
                    error_code="missing_public_article_evidence",
                )
            shutil.copy2(evidence_source, public_evidence_target)
        if source:
            if not target.exists() or target.stat().st_size != Path(source).stat().st_size:
                shutil.copy2(source, target)
            try:
                topics = topic_titles(target)
            except Exception as exc:
                return StageOutcome.failed(f"标准总表结构校验失败：{exc}", error_code="invalid_workbook")
            if not topics:
                return StageOutcome.failed("标准总表未识别到子议题", error_code="missing_topics")
            return StageOutcome.succeeded("已识别标准总表，未重复计算系统统计。", details={"topics": topics})

        metadata = Path(str(self.contract.get("metadata") or ""))
        if not metadata.is_file():
            return StageOutcome.waiting(
                "waiting_review",
                "原始表模式需要本期议题元数据，以便可靠匹配子事件表。",
                details={"required_input": "metadata", "stage": spec.stage_id},
            )
        raw_dir = Path(str(self.contract["raw_input_dir"]))
        raw_run = self.job_dir / "raw_workbook"
        raw_run.mkdir(parents=True, exist_ok=True)
        audit = raw_run / "comparison_audit.md"
        command = [
            sys.executable,
            str(SCRIPT_DIR / "raw_system_workbook_pipeline.py"),
            "--input-dir",
            str(raw_dir),
            "--metadata",
            str(metadata),
            "--output",
            str(target),
            "--audit",
            str(audit),
        ]
        hotword_review = self.artifacts / "hotword_ai_review.json"
        overseas_review = self.artifacts / "overseas_ai_review.json"
        public_review = self.artifacts / "public_top_ai_review.json"
        if hotword_review.exists():
            command.extend(["--hotword-review", str(hotword_review)])
        if overseas_review.exists():
            command.extend(["--overseas-review", str(overseas_review)])
        if public_review.exists():
            command.extend(["--public-review", str(public_review)])
        code, log_path = runner.run_command(spec.stage_id, command, cwd=SKILL_ROOT.parent)
        # The raw pipeline locates its run directory beside --output, not --audit.
        packet_dir = target.parent / "run"
        if code == 0 and target.exists():
            raw_public_evidence = packet_dir / "public_article_evidence.json"
            if raw_public_evidence.exists():
                shutil.copy2(raw_public_evidence, public_evidence_target)
            return StageOutcome.succeeded(
                "原始监测表已经AI审核、汇总并生成标准总表。",
                details={"log_path": str(log_path), "audit": str(audit)},
            )
        packets = {
            "hotword": packet_dir / "hotword_review_packet.json",
            "overseas": packet_dir / "overseas_review_packet.json",
            "public_top": packet_dir / "public_top_review_packet.json",
        }
        if any(path.exists() for path in packets.values()):
            if review_applied:
                return StageOutcome.failed(
                    f"审核结果未被原始表处理器接受，详见{log_path}",
                    retryable=True, error_code="raw_review_rejected",
                    details={"log_path": str(log_path)},
                )
            task = ai_task(
                runner,
                spec.stage_id,
                task_type="raw_workbook_semantic_reviews",
                expected_output=hotword_review,
                inputs={
                    "validation_log": str(log_path),
                    **{key: str(value) for key, value in packets.items() if value.exists()},
                    "expected_outputs": {
                        "hotword_ai_review": str(hotword_review),
                        "overseas_ai_review": str(overseas_review),
                        "public_top_ai_review": str(public_review),
                    },
                },
                rules=[
                    "分别完成热词语义审核、境外报道逐条相关性审核，以及公众文章TOP临界候选全文审核。",
                    "境外审核必须覆盖全部候选；热词审核必须保留证据并完成二次自审。",
                    "公众文章TOP审核只决定附录榜单，不得删除或缩减独立的观点分析全文证据池。",
                    f"热词输出写入{hotword_review}，境外输出写入{overseas_review}，公众TOP输出写入{public_review}。",
                ],
            )
            # Packets can be emitted incrementally. Do not require the hotword
            # output before the corresponding packet is available.
            worker_failure = maybe_run_ai_worker(runner, spec, task, hotword_review, allow_missing_output=True)
            if worker_failure:
                return worker_failure
            if hotword_review.exists() and overseas_review.exists() and public_review.exists():
                return self.workbook(runner, spec, review_applied=True)
            return StageOutcome.waiting(
                "waiting_ai",
                "标准总表已生成审核包，等待AI完成热词、境外和公众TOP逐条审核后从本节点续跑。",
                details={
                    "task": str(task),
                    "expected_outputs": [str(hotword_review), str(overseas_review), str(public_review)],
                },
            )
        return StageOutcome.failed(
            f"原始表处理失败，且未生成可继续的审核包；详见{log_path}",
            retryable=code in spec.transient_exit_codes,
            error_code=f"workbook_exit_{code}",
        )

    def research_plan(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        workbook = self.artifacts / "CWH舆情情况_标准总表.xlsx"
        target = self.artifacts / "research_plan.json"
        command = [
            sys.executable,
            str(SCRIPT_DIR / "build_research_plan.py"),
            str(workbook),
            "--agenda",
            str(self.contract.get("agenda") or ""),
            "--output",
            str(target),
            "--execution-profile",
            str(self.contract.get("execution_profile") or execution_profile()[0]),
        ]
        code, log_path = runner.run_command(spec.stage_id, command, cwd=SKILL_ROOT.parent)
        if code != 0:
            return StageOutcome.failed(
                f"检索计划生成失败，详见{log_path}",
                retryable=code in spec.transient_exit_codes,
                error_code=f"research_plan_exit_{code}",
            )
        data = read_json(target)
        if not data.get("topics"):
            return StageOutcome.failed("检索计划没有子议题", error_code="empty_research_plan")
        return StageOutcome.succeeded("已按系统子议题生成独立检索任务和证据规则。")

    def domestic_viewpoints(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        target = self.artifacts / "analysis_bundle.json"
        supplied = str(self.contract.get("analysis_bundle") or "").strip()
        if supplied and not target.exists():
            shutil.copy2(supplied, target)
        topics = topic_titles(self.artifacts / "CWH舆情情况_标准总表.xlsx")

        corpus_path = self.artifacts / "public_article_evidence.json"
        corpus = read_json(corpus_path) if corpus_path.exists() else {}
        index_path = None
        if corpus:
            from prepare_cwh_corpus_index import prepare_corpus_index
            index_path = prepare_corpus_index(corpus_path, corpus, topics)
        allow_deferred = self.contract.get("execution_profile") in {"bounded_40m", "bounded_60m"}

        def evaluate_draft() -> list[str]:
            if not target.exists():
                return ["尚未生成analysis_bundle.json"]
            try:
                draft = read_json(target)
                if not isinstance(draft, dict):
                    return ["analysis_bundle.json顶层必须是JSON对象"]
                if allow_deferred and corpus:
                    from prepare_cwh_corpus_index import complete_corpus_deferrals
                    draft = complete_corpus_deferrals(draft, corpus, topics)
                normalized = normalize_analysis(complete_analysis_structure(draft, corpus))
            except (ValueError, TypeError, AttributeError) as exc:
                return [f"analysis_bundle.json无法解析或结构无效：{type(exc).__name__}: {exc}"]
            atomic_write_json(target, normalized)
            try:
                return validate_analysis_bundle(target, topics, require_semantic_review=False, allow_deferred_corpus=allow_deferred)
            except (ValueError, TypeError, AttributeError) as exc:
                return [f"analysis_bundle.json字段类型无效：{type(exc).__name__}: {exc}"]

        problems = evaluate_draft()
        if not problems:
            return StageOutcome.succeeded("境内媒体自媒体观点已覆盖全部子议题并通过证据门禁。")
        task = ai_task(
            runner,
            spec.stage_id,
            task_type="domestic_viewpoint_research_and_review",
            expected_output=target,
            inputs={
                "research_plan": str(self.artifacts / "research_plan.json"),
                "existing_output": str(target) if target.exists() else "",
                "public_corpus_index": str(index_path) if index_path else "",
                "public_corpus_reading_indexes": read_json(index_path)["reading_indexes"] if index_path else [],
                "validation_problems": problems,
                "source_registry": str(SOURCE_REGISTRY_PATH),
                "execution_policy": str(SKILL_ROOT / "config" / "execution_policy.v1.json"),
                "formal_writing_rules": str(SKILL_ROOT / "config" / "formal_writing_rules.v1.json"),
                "analysis_schema": str(SKILL_ROOT / "references" / "analysis_bundle_schema.md"),
                "public_article_evidence": (
                    str(self.artifacts / "public_article_evidence.json")
                    if (self.artifacts / "public_article_evidence.json").exists()
                    else ""
                ),
            },
            rules=[
                "严格读取research_plan、execution_policy和analysis_schema；按其中议题、通道、查询/抓取上限、停止规则执行，不得自行扩大范围。",
                "若有public_corpus_index，按各议题shortlist优先读取完整文章，需要时从全量索引补充；排序只是阅读顺序，不是语义审核。公众TOP10不是观点证据上限。",
                "bounded档每议题至少实际审核min(12,该议题全文池数量)篇，已审明确保留或排除；控制器自动计算并记录未审清单deferred_record_ids，不必抄写上千ID，绝不能冒称已审或没有观点。exhaustive仍须全文池逐条审核。",
                "监测来源候选写raw_evidence_record_id；缺失的原文快照和原始标题、URL、发布时间由控制器从原始全文池回填，不要重复抄写全文或计算哈希。",
                "每个实际检索结果都进入候选池并标记eligible、duplicate或excluded及理由；保存查询ID、结果URL快照和真实阻断。",
                "eligible候选须核验监测期内发布时间并保存完整原文快照及SHA-256；搜索摘要不能冒充原文。",
                "一篇文章中的不同发言主体分成独立证据；同一主体同一观点的转载只保留一个正式候选。",
                "每条入选证据保留candidate_id、evidence_id、单一speaker_name、连续source_excerpt及字符位置、45至120汉字的formal_claim；不得增强原文结论。",
                "bounded限时档每议题选择6至12个、最多12个代表性独立声音，其余有效候选设formal_use=reserve并写reserve_reason；不足4个时须写含reason、search_evidence、reviewed_by的evidence_shortfall。exhaustive才全部成文。",
                "本节点不填写semantic_review，也不写最终正文；脚本将从formal_claim机械生成cluster.details，独立下一节点再逐命题复核。",
                "只有证据不足或平台真实受阻时才返回结构化blocker；不得编造链接、引文、人物、ID、快照或状态。",
            ],
        )
        worker_failure = maybe_run_ai_worker(runner, spec, task, target)
        if worker_failure:
            return worker_failure
        if target.exists():
            problems = evaluate_draft()
            if not problems:
                return StageOutcome.succeeded("AI工作器输出已通过境内观点证据门禁。")
            return StageOutcome.failed(
                "AI输出未通过观点门禁：" + "；".join(problems[:12]),
                retryable=True,
                error_code="analysis_bundle_invalid",
                details={"task": str(task), "problems": problems},
            )
        return StageOutcome.waiting(
            "waiting_ai",
            "等待AI按检索计划生成境内观点证据包；完成后从本节点继续。",
            details={"task": str(task), "expected_output": str(target), "problems": problems},
        )

    def domestic_evidence_verification(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        source = self.artifacts / "analysis_bundle.json"
        review_path = self.artifacts / "domestic_evidence_semantic_review.json"
        verified_path = self.artifacts / "analysis_bundle_verified.json"
        audit_path = self.artifacts / "domestic_evidence_mapping_audit.json"
        topics = topic_titles(self.artifacts / "CWH舆情情况_标准总表.xlsx")
        source_hash = analysis_bundle_sha256(source)

        def evaluate_review() -> list[str]:
            if not review_path.exists():
                return ["尚未生成独立境内观点语义复核文件"]
            try:
                packet = read_json(review_path)
            except Exception as exc:
                return [f"独立境内观点语义复核文件无法解析：{exc}"]
            source_data = read_json(source)
            verified_data, review_issues = apply_semantic_review_packet(
                source_data,
                packet,
                source_bundle_sha256=source_hash,
            )
            if review_issues:
                return mapping_problem_messages({"issues": review_issues})
            atomic_write_json(verified_path, verified_data)
            mapping_audit = validate_analysis_mapping(verified_data)
            atomic_write_json(audit_path, mapping_audit)
            return validate_analysis_bundle(verified_path, topics, allow_deferred_corpus=self.contract.get("execution_profile") in {"bounded_40m", "bounded_60m"})

        problems = evaluate_review()
        if not problems:
            return StageOutcome.succeeded("境内观点已完成独立第二遍语义复核并通过原文映射门禁。")
        task = ai_task(
            runner,
            spec.stage_id,
            task_type="independent_domestic_evidence_verification",
            expected_output=review_path,
            inputs={
                "analysis_bundle": str(source),
                "source_bundle_sha256": source_hash,
                "validation_problems": problems,
                "expected_verified_bundle": str(verified_path),
                "expected_mapping_audit": str(audit_path),
            },
            rules=[
                "这是与观点撰写分离的独立第二遍核验，不得改写观点或补造原文；使用新的reviewer_run_id，且不得等于analysis_bundle.metadata.authoring_run_id。",
                "输出review_version=1.0、review_pass=independent_second_pass、reviewer_run_id、source_bundle_sha256和reviews数组；每个evidence_id必须且只能审核一次。",
                "逐条对照候选source_snapshot.source_text、speaker_name、article_title、source_excerpt及位置，确认映射的是同一篇文章、同一发言主体和连续原文。",
                "将formal_claim完整拆分为命题，不得遗漏任何因果、程度、预测、数字、专名、限定或政策效果。每个命题都要引用同一source_excerpt内的连续source_quote和精确字符位置。",
                "只有全文含义完全支持时才写fully_supported；partially_supported、unsupported或uncertain必须如实返回，本节点随后阻止正式报告生成。",
                "每条review保存reviewed_by、reviewed_at、rationale和propositions；每条proposition保存text、verdict、source_quote、source_quote_start、source_quote_end和rationale。",
            ],
        )
        worker_failure = maybe_run_ai_worker(runner, spec, task, review_path)
        if worker_failure:
            return worker_failure
        if review_path.exists():
            problems = evaluate_review()
            if not problems:
                return StageOutcome.succeeded("AI工作器的独立语义复核已通过原文映射门禁。")
            return StageOutcome.failed(
                "独立境内观点语义复核未通过：" + "；".join(problems[:12]),
                retryable=True,
                error_code="domestic_evidence_verification_invalid",
                details={"task": str(task), "problems": problems},
            )
        return StageOutcome.waiting(
            "waiting_ai",
            "等待独立第二遍境内观点语义复核；完成后从本节点继续。",
            details={"task": str(task), "expected_output": str(review_path), "problems": problems},
        )

    def domestic_comments_sentiment(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        targets = {
            "comment_handoff": self.artifacts / "report_comment_handoff.csv",
            "sentiment_results": self.artifacts / "sentiment_results.csv",
            "sentiment_summary": self.artifacts / "sentiment_workbook_summary.json",
        }
        for key, target in targets.items():
            supplied = str(self.contract.get(key) or "").strip()
            if supplied and not target.exists():
                shutil.copy2(supplied, target)
        topics = topic_titles(self.artifacts / "CWH舆情情况_标准总表.xlsx")
        if all(path.exists() for path in targets.values()):
            problems = validate_sentiment(targets["sentiment_summary"], targets["comment_handoff"], topics)
            if not problems:
                return StageOutcome.succeeded("真实评论、逐条AI引用审核及自有情感结果已通过门禁。")
        else:
            problems = [f"缺少{key}" for key, path in targets.items() if not path.exists()]
        task = ai_task(
            runner,
            spec.stage_id,
            task_type="domestic_comment_collection_and_sentiment",
            expected_output=targets["sentiment_summary"],
            inputs={
                "research_plan": str(self.artifacts / "research_plan.json"),
                "workbook": str(self.artifacts / "CWH舆情情况_标准总表.xlsx"),
                "validation_problems": problems,
                "expected_outputs": {key: str(value) for key, value in targets.items()},
                "source_registry": str(SOURCE_REGISTRY_PATH),
            },
            rules=[
                "只接收监测期内、带真实原话、平台、URL、评论标识的评论区数据。",
                "搜索摘要、文章正文、帖子正文和评论数量都不能冒充网民评论。",
                "每条评论必须由AI判断是否有实质观点、是否适合正式引用，并保留审核理由。",
                "情感比例只能来自自有评论级审核结果，系统原情感标签不进入分母。",
                "评论采集必须单独输出collection_audit，按每个子议题×每个必查平台记录状态、执行方式、查询或原帖种子、结果数和合格评论ID；全局平台状态不能代替逐议题覆盖。",
                "今日头条公开评论是无需登录的必查适配器；微博、抖音、哔哩哔哩如无获准登录态，应按每个子议题记录access_failed及真实阻断，不得误作完成，也不得把waiting_login当成流程终点。",
                "文章检索命中不能代替评论采集成功；无评论结论必须来自实际原帖评论接口或平台定向检索的零结果证据。",
            ],
        )
        worker_failure = maybe_run_ai_worker(
            runner,
            spec,
            task,
            targets["sentiment_summary"],
            allow_missing_output=True,
        )
        if worker_failure:
            return worker_failure
        if all(path.exists() for path in targets.values()):
            problems = validate_sentiment(targets["sentiment_summary"], targets["comment_handoff"], topics)
            if not problems:
                return StageOutcome.succeeded("AI工作器已完成评论采集、引用审核和自有情感分析。")
            return StageOutcome.failed(
                "评论与情感产物未通过门禁：" + "；".join(problems[:12]),
                retryable=True,
                error_code="comment_sentiment_invalid",
                details={"task": str(task), "problems": problems},
            )
        if not any(str(self.contract.get(key) or "").strip() for key in targets):
            return StageOutcome.waiting(
                "waiting_ai",
                "尚无可核验评论明细。应先按任务包尝试系统明细、已授权平台以及免登录公开备选源；只有具体采集器确认登录受阻时才进入等待登录。",
                details={
                    "task": str(task),
                    "expected_outputs": [str(path) for path in targets.values()],
                    "fallback_order": ["系统评论明细", "已授权平台", "免登录头条等公开评论备选源"],
                },
            )
        return StageOutcome.waiting(
            "waiting_ai",
            "评论或情感审核产物尚未通过门禁，需按任务包补齐后继续。",
            details={"task": str(task), "problems": problems},
        )

    def overseas_evidence(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        supplements = self.artifacts / "public_overseas_supplements.json"
        audit = self.artifacts / "foreign_collection_audit.json"
        for key, target in (("overseas_supplements", supplements), ("foreign_collection_audit", audit)):
            supplied = str(self.contract.get(key) or "").strip()
            if supplied and not target.exists():
                shutil.copy2(supplied, target)
        problems = validate_foreign(audit, supplements) if audit.exists() and supplements.exists() else ["尚未生成境外采集与审核产物"]
        if not problems:
            return StageOutcome.succeeded("境外媒体与境外网民评论已分开采集、审核并保留来源。")
        task = ai_task(
            runner,
            spec.stage_id,
            task_type="overseas_media_and_comment_collection_review",
            expected_output=supplements,
            inputs={
                "research_plan": str(self.artifacts / "research_plan.json"),
                "validation_problems": problems,
                "expected_audit": str(audit),
            },
            rules=[
                "境外媒体和境外网民评论必须分开采集、分开审核。",
                "只有实质报道或解读本次会议的报道可纳入；宽泛议题重合不能纳入。",
                "每条拟入正式集合的境外报道都必须生成title_cn_simplified、source_cn_simplified、summary_cn_simplified并设置simplified_chinese_reviewed=true；繁体中文先用完整OpenCC转简体，外文再翻译。",
                "AI复核转换或翻译后的专名、机构名、政策名、数字和语义；原文只保留在结构化证据和工作台，Word/Markdown一律使用简体中文。",
                "外媒附录从完整审核、简体化和报道级去重后的合格集合中选代表性TOP10；只有合格集合确实少于10条时才允许不足10条。",
                "评论应归纳观点并排除明显反党反政府辱骂，但不能过滤普通政策批评。",
                "采集完成但0条与采集失败是两种不同状态，审计必须明确记录。",
            ],
        )
        worker_failure = maybe_run_ai_worker(runner, spec, task, supplements)
        if worker_failure:
            return worker_failure
        if supplements.exists() and audit.exists():
            problems = validate_foreign(audit, supplements)
            if not problems:
                return StageOutcome.succeeded("AI工作器已完成境外媒体与网民评论的独立采集审核。")
            return StageOutcome.failed(
                "境外采集审核产物未通过门禁：" + "；".join(problems[:12]),
                retryable=True,
                error_code="foreign_evidence_invalid",
                details={"task": str(task), "problems": problems},
            )
        return StageOutcome.waiting(
            "waiting_ai",
            "等待境外检索、逐条语义审核和分类完成；完成后从本节点继续。",
            details={"task": str(task), "expected_outputs": [str(supplements), str(audit)], "problems": problems},
        )

    def hotwords(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        target = self.artifacts / "hotword_audit.json"
        supplied = str(self.contract.get("hotword_audit") or "").strip()
        if supplied and not target.exists():
            shutil.copy2(supplied, target)
        topics = topic_titles(self.artifacts / "CWH舆情情况_标准总表.xlsx")
        problems = validate_hotword_audit(target, topics) if target.exists() else ["尚未生成热词审核文件"]
        if not problems:
            return StageOutcome.succeeded("热词已由证据提取并完成AI或人工二次复核。")
        task = ai_task(
            runner,
            spec.stage_id,
            task_type="hotword_evidence_review",
            expected_output=target,
            inputs={
                "analysis_bundle": str(self.artifacts / "analysis_bundle_verified.json"),
                "comment_handoff": str(self.artifacts / "report_comment_handoff.csv"),
                "validation_problems": problems,
            },
            rules=[
                "从境内媒体、自媒体和合格网民评论证据中按子议题提取候选短语。",
                "合并同义词，排除套话、碎片、纯地名、人名、机构名和无证据词。",
                "默认形成36至42个有证据热词并保证各子议题覆盖，随后执行第二轮AI自审。",
                "展示权重只控制词云字号，不得表述为全网精确热度。",
            ],
        )
        worker_failure = maybe_run_ai_worker(runner, spec, task, target)
        if worker_failure:
            return worker_failure
        if target.exists():
            problems = validate_hotword_audit(target, topics)
            if not problems:
                return StageOutcome.succeeded("AI工作器已完成热词证据审核与二次自审。")
            return StageOutcome.failed(
                "热词审核产物未通过门禁：" + "；".join(problems[:12]),
                retryable=True,
                error_code="hotword_audit_invalid",
                details={"task": str(task), "problems": problems},
            )
        return StageOutcome.waiting(
            "waiting_ai",
            "等待热词证据审核与二次自审完成；完成后从本节点继续。",
            details={"task": str(task), "expected_output": str(target), "problems": problems},
        )

    def render(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        input_path = self.artifacts / "agenda.txt"
        input_path.write_text(str(self.contract.get("agenda") or "国务院常务会议"), encoding="utf-8")
        command = [
            sys.executable,
            str(SCRIPT_DIR / "cwh_orchestrator.py"),
            "--input-file",
            str(input_path),
            "--system-workbook",
            str(self.artifacts / "CWH舆情情况_标准总表.xlsx"),
            "--analysis-bundle",
            str(self.artifacts / "analysis_bundle_verified.json"),
            "--comment-handoff",
            str(self.artifacts / "report_comment_handoff.csv"),
            "--sentiment-results",
            str(self.artifacts / "sentiment_results.csv"),
            "--sentiment-summary",
            str(self.artifacts / "sentiment_workbook_summary.json"),
            "--overseas-supplements",
            str(self.artifacts / "public_overseas_supplements.json"),
            "--foreign-collection-audit",
            str(self.artifacts / "foreign_collection_audit.json"),
            "--hotword-audit",
            str(self.artifacts / "hotword_audit.json"),
            "--out-dir",
            str(self.report),
            "--data-mode",
            "system",
            "--no-tasks",
            "--no-agent-reach-plan",
        ]
        code, log_path = runner.run_command(spec.stage_id, command, cwd=SKILL_ROOT.parent)
        if code not in (0, 1):
            return StageOutcome.failed(
                f"报告渲染异常退出，详见{log_path}",
                retryable=code in spec.transient_exit_codes,
                error_code=f"render_exit_{code}",
            )
        required = [
            self.report / "report_data.json",
            self.report / "cwh_formal_report.docx",
            self.report / "cwh_dashboard.html",
            self.report / "cwh_data_workbook.xlsx",
        ]
        if not all(path.exists() for path in required):
            return StageOutcome.failed("报告渲染结束但成品不完整", error_code="render_artifacts_missing")
        return StageOutcome.succeeded("正式报告成品已生成，进入最终一致性门禁。", details={"log_path": str(log_path), "exit_code": code})

    def delivery_gate(self, runner: PipelineRunner, spec: StageSpec) -> StageOutcome:
        audit_path = self.report / "cwh_audit.json"
        report_data_path = self.report / "report_data.json"
        problems = []
        if not audit_path.exists():
            problems.append("缺少cwh_audit.json")
            audit = {}
        else:
            audit = read_json(audit_path)
        acceptance = audit.get("acceptance") or {}
        if not acceptance.get("ready_for_formal_delivery"):
            problems.extend(str(item) for item in acceptance.get("blockers") or ["正式交付门禁未通过"])
        try:
            report_data = read_json(report_data_path)
        except Exception as exc:
            report_data = {}
            problems.append(f"report_data.json无法解析：{exc}")
        data_workbook_path = self.report / "cwh_data_workbook.xlsx"
        if report_data:
            problems.extend(validate_data_workbook(data_workbook_path, report_data))
            try:
                upstream_analysis = read_json(self.artifacts / "analysis_bundle_verified.json")
                mapping_audit = validate_release_mapping(self.report, upstream_analysis=upstream_analysis)
            except Exception as exc:
                mapping_audit = {
                    "schema_version": "1.0",
                    "status": "blocked",
                    "issues": [{"code": "mapping_audit_failed", "severity": "error", "message": f"境内证据映射审计无法完成：{exc}"}],
                }
            mapping_audit_path = self.report / "domestic_evidence_mapping_audit.json"
            atomic_write_json(mapping_audit_path, mapping_audit)
            problems.extend(mapping_problem_messages(mapping_audit))
        artifacts = report_data.get("artifacts") or {}
        for name in ("formal_docx", "dashboard", "data_workbook"):
            value = str(artifacts.get(name) or "")
            if value and not Path(value).exists():
                problems.append(f"report_data中的{name}路径不存在")
        payload = {
            "status": "passed" if not problems else "blocked",
            "problems": list(dict.fromkeys(problems)),
            "checked_artifacts": [
                str(self.report / "cwh_formal_report.docx"),
                str(self.report / "cwh_dashboard.html"),
                str(data_workbook_path),
                str(report_data_path),
                str(audit_path),
                str(self.report / "domestic_evidence_mapping_audit.json"),
            ],
        }
        target = self.artifacts / "delivery_gate.json"
        atomic_write_json(target, payload)
        if problems:
            return StageOutcome.failed(
                "正式交付门禁未通过：" + "；".join(payload["problems"][:12]),
                error_code="delivery_gate_blocked",
                artifacts={"delivery_audit": str(target)},
                details={"problems": payload["problems"]},
            )
        return StageOutcome.succeeded(
            "全部成品与质量门禁通过，可以归档交付。",
            artifacts={
                "delivery_audit": str(target),
                "release_mapping_audit": str(self.report / "domestic_evidence_mapping_audit.json"),
                "cwh_audit": str(audit_path),
            },
        )


def contract_from_args(args: argparse.Namespace) -> dict[str, Any]:
    profile_name, profile = execution_profile(str(args.execution_profile or ""))
    return {
        "execution_profile": profile_name,
        "wall_clock_budget_seconds": int(profile.get("wall_clock_budget_seconds") or 0),
        "research_deadline_seconds": int(profile.get("research_deadline_seconds") or 0),
        "stage_timeouts_seconds": dict(profile.get("stage_budgets_seconds") or {}),
        "agenda": str(args.agenda or "").strip(),
        "system_workbook": str(Path(args.system_workbook).resolve()) if args.system_workbook else "",
        "raw_input_dir": str(Path(args.raw_input_dir).resolve()) if args.raw_input_dir else "",
        "metadata": str(Path(args.metadata).resolve()) if args.metadata else "",
        "analysis_bundle": str(Path(args.analysis_bundle).resolve()) if args.analysis_bundle else "",
        "public_article_evidence": str(Path(args.public_article_evidence).resolve()) if args.public_article_evidence else "",
        "comment_handoff": str(Path(args.comment_handoff).resolve()) if args.comment_handoff else "",
        "sentiment_results": str(Path(args.sentiment_results).resolve()) if args.sentiment_results else "",
        "sentiment_summary": str(Path(args.sentiment_summary).resolve()) if args.sentiment_summary else "",
        "overseas_supplements": str(Path(args.overseas_supplements).resolve()) if args.overseas_supplements else "",
        "foreign_collection_audit": str(Path(args.foreign_collection_audit).resolve()) if args.foreign_collection_audit else "",
        "hotword_audit": str(Path(args.hotword_audit).resolve()) if args.hotword_audit else "",
        "ai_worker_command": json.loads(args.ai_worker_command_json) if args.ai_worker_command_json else [],
        "source_contracts": {
            "system_workbook": file_contract(args.system_workbook),
            "raw_input_dir": directory_contract(args.raw_input_dir),
            "metadata": file_contract(args.metadata),
            "analysis_bundle": file_contract(args.analysis_bundle),
            "public_article_evidence": file_contract(args.public_article_evidence),
            "comment_handoff": file_contract(args.comment_handoff),
            "sentiment_results": file_contract(args.sentiment_results),
            "sentiment_summary": file_contract(args.sentiment_summary),
            "overseas_supplements": file_contract(args.overseas_supplements),
            "foreign_collection_audit": file_contract(args.foreign_collection_audit),
            "hotword_audit": file_contract(args.hotword_audit),
        },
    }


def launch_local_workbench(job_dir: Path) -> dict[str, Any]:
    dashboard = job_dir / "report" / "cwh_dashboard.html"
    if not dashboard.exists():
        dashboards = sorted(job_dir.rglob("cwh_dashboard.html"), key=lambda path: path.stat().st_mtime_ns, reverse=True)
        dashboard = dashboards[0] if dashboards else dashboard
    if not dashboard.exists():
        return {"launched": False, "reason": "dashboard_missing"}
    port = 8788
    while port < 8818:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as probe:
            try:
                probe.bind(("127.0.0.1", port))
            except OSError:
                port += 1
                continue
        break
    log_path = job_dir / "workbench_server.log"
    command = [
        sys.executable,
        str(SCRIPT_DIR / "serve_dashboard.py"),
        str(dashboard),
        "--library-root",
        str(job_dir.parent),
        "--workspace",
        str(SKILL_ROOT.parent),
        "--report-runner",
        "pipeline",
        "--port",
        str(port),
    ]
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0) | getattr(subprocess, "DETACHED_PROCESS", 0)
    with log_path.open("a", encoding="utf-8", errors="replace") as log:
        process = subprocess.Popen(
            command,
            cwd=str(SKILL_ROOT.parent),
            stdout=log,
            stderr=subprocess.STDOUT,
            creationflags=flags,
        )
    url = f"http://127.0.0.1:{port}/cwh_dashboard.html"
    required_endpoints = {
        "dashboard": url,
        "archive": f"http://127.0.0.1:{port}/api/archive",
    }
    endpoint_status: dict[str, int | str] = {}
    endpoint_elapsed_seconds: dict[str, float] = {}
    error = "http_health_timeout"
    dashboard_started = time.monotonic()
    for _ in range(60):
        exit_code = process.poll()
        if exit_code is not None:
            error = f"server_exited_before_http_ready:{exit_code}"
            break
        try:
            with urllib.request.urlopen(required_endpoints["dashboard"], timeout=0.5) as response:
                endpoint_status["dashboard"] = int(response.status)
                if response.status == 200:
                    endpoint_elapsed_seconds["dashboard"] = round(time.monotonic() - dashboard_started, 3)
                    break
        except urllib.error.HTTPError as exc:
            endpoint_status["dashboard"] = int(exc.code)
        except (urllib.error.URLError, TimeoutError, ConnectionError) as exc:
            endpoint_status["dashboard"] = type(exc).__name__
        time.sleep(0.25)
    else:
        error = "http_health_timeout:dashboard"

    if endpoint_status.get("dashboard") == 200 and process.poll() is None:
        archive_started = time.monotonic()
        try:
            # The first archive request performs one cold, idempotent library
            # scan. Wait for that single request instead of creating dozens of
            # short-lived requests that queue behind the same sync lock.
            with urllib.request.urlopen(required_endpoints["archive"], timeout=120.0) as response:
                endpoint_status["archive"] = int(response.status)
                endpoint_elapsed_seconds["archive"] = round(time.monotonic() - archive_started, 3)
        except urllib.error.HTTPError as exc:
            endpoint_status["archive"] = int(exc.code)
            endpoint_elapsed_seconds["archive"] = round(time.monotonic() - archive_started, 3)
            error = "required_endpoint_failed:archive"
        except (urllib.error.URLError, TimeoutError, ConnectionError) as exc:
            endpoint_status["archive"] = type(exc).__name__
            endpoint_elapsed_seconds["archive"] = round(time.monotonic() - archive_started, 3)
            error = "required_endpoint_timeout:archive"

    if (
        endpoint_status.get("dashboard") == 200
        and endpoint_status.get("archive") == 200
        and process.poll() is None
    ):
        result = {
            "launched": True,
            "url": url,
            "pid": process.pid,
            "log_path": str(log_path),
            "http_status": endpoint_status["dashboard"],
            "required_endpoints": endpoint_status,
            "endpoint_elapsed_seconds": endpoint_elapsed_seconds,
        }
        atomic_write_json(job_dir / "workbench_launch.json", result)
        return result
    if process.poll() is None:
        process.terminate()
    result = {
        "launched": False,
        "url": url,
        "pid": process.pid,
        "log_path": str(log_path),
        "reason": error,
        "partially_available": endpoint_status.get("dashboard") == 200,
        "required_endpoints": endpoint_status,
        "endpoint_elapsed_seconds": endpoint_elapsed_seconds,
    }
    atomic_write_json(job_dir / "workbench_launch.json", result)
    return result


def should_launch_local_workbench(
    state: dict[str, Any], *, no_open: bool, environment: dict[str, str] | None = None
) -> bool:
    env = environment if environment is not None else os.environ
    return bool(
        state.get("status") == "succeeded"
        and not no_open
        and not env.get("WITH_PROJECT_ID")
        and env.get("CWH_RUNTIME", "").lower() != "with"
    )


def cli_state_summary(state: dict[str, Any], job_dir: Path) -> dict[str, Any]:
    """Keep routine handoffs small without changing the durable full state."""
    summary = {key: state.get(key) for key in (
        "pipeline_id", "status", "current_stage", "next_action",
        "wall_clock_elapsed_seconds", "completed_elapsed_seconds", "review_delivery",
    ) if key in state}
    summary["state_path"] = str((job_dir / "pipeline_state.json").resolve())
    summary["stage_statuses"] = {row["stage_id"]: row["status"] for row in state.get("stages") or []}
    if state.get("workbench"):
        summary["workbench"] = state["workbench"]
    current = next((row for row in state.get("stages") or [] if row.get("stage_id") == state.get("current_stage")), {})
    if current.get("last_error"):
        summary["current_error"] = current["last_error"]
    return summary


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description="Run the CWH formal-report workflow as a durable resumable pipeline.")
    result.add_argument("action", choices=["run", "status", "invalidate"])
    result.add_argument("--job-dir", required=True)
    result.add_argument("--agenda", default="")
    result.add_argument("--system-workbook", default="")
    result.add_argument("--raw-input-dir", default="")
    result.add_argument("--metadata", default="")
    result.add_argument("--analysis-bundle", default="")
    result.add_argument("--public-article-evidence", default="")
    result.add_argument("--comment-handoff", default="")
    result.add_argument("--sentiment-results", default="")
    result.add_argument("--sentiment-summary", default="")
    result.add_argument("--overseas-supplements", default="")
    result.add_argument("--foreign-collection-audit", default="")
    result.add_argument("--hotword-audit", default="")
    result.add_argument("--ai-worker-command-json", default=os.environ.get("CWH_PIPELINE_AI_COMMAND_JSON", ""))
    result.add_argument(
        "--execution-profile",
        choices=["bounded_40m", "bounded_60m", "exhaustive"],
        default=execution_profile()[0],
        help="Default bounded_60m targets one hour and reserves final delivery time; bounded_40m is a tighter option.",
    )
    result.add_argument("--until-stage", default="")
    result.add_argument("--invalidate-from", default="")
    result.add_argument("--no-open", action="store_true", help="Do not launch the local workbench after success.")
    result.add_argument("--no-review-delivery", action="store_true", help="Skip the labelled review snapshot for diagnostic or worker-only invocations.")
    result.add_argument("--output-format", choices=["compact", "full"], default="compact", help="Compact stdout saves model context; full durable state is always saved in pipeline_state.json.")
    return result


def main() -> None:
    args = parser().parse_args()
    job_dir = Path(args.job_dir)
    state_path = job_dir / "pipeline_state.json"
    if state_path.exists():
        existing = read_json(state_path)
        contract = existing.get("input_contract") or {}
    else:
        contract = contract_from_args(args)
    pipeline = CwhPipeline(job_dir, contract)
    if args.action == "status":
        state = pipeline.runner.state
    elif args.action == "invalidate":
        if not args.invalidate_from:
            raise ValueError("--invalidate-from is required for invalidate")
        pipeline.runner.invalidate_from(args.invalidate_from, reason="operator_requested")
        state = pipeline.runner.state
    else:
        state = pipeline.runner.run(until_stage=args.until_stage)
        if state.get("status") != "succeeded" and not args.until_stage and not args.no_review_delivery:
            # Separate review artifacts never change a failed gate or overwrite
            # report/. They remain useful even if a later node is unavailable.
            try:
                from build_cwh_review_delivery import build_review_delivery
                state["review_delivery"] = build_review_delivery(job_dir)
            except Exception as exc:
                state["review_delivery"] = {"status": "failed", "error": f"{type(exc).__name__}: {exc}"}
        if should_launch_local_workbench(state, no_open=args.no_open):
            state["workbench"] = launch_local_workbench(job_dir.resolve())
    output = state if args.output_format == "full" else cli_state_summary(state, job_dir)
    print(json.dumps(output, ensure_ascii=False, indent=2))
    if state.get("status") == "succeeded":
        raise SystemExit(0)
    if state.get("status") in {"waiting_ai", "waiting_login", "waiting_review", "blocked", "paused", "interrupted"}:
        raise SystemExit(20)
    raise SystemExit(1)


if __name__ == "__main__":
    main()
