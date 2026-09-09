"""Build and validate the independent CWH comment-sentiment stage.

This module deliberately sits after the raw-workbook normalization stage.  It
only accepts traceable row-level comments into the sentiment denominator; an
article body plus a numeric comment-count column is never treated as comment
text.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import subprocess
import sys
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any, Iterable


SUPPORTED_SUFFIXES = {".csv", ".json", ".xlsx", ".xlsm"}
OUTPUT_FIELDS = [
    "sample_id",
    "text",
    "topic",
    "platform",
    "source",
    "published_at",
    "url",
    "comment_id",
    "reply_id",
    "parent_comment_id",
    "like_count",
    "parent_post_title",
    "parent_post_published_at",
    "parent_post_source",
    "raw_file",
    "raw_sheet",
    "raw_row",
]

REPORT_HANDOFF_FIELDS = [
    "id",
    "sample_id",
    "topic",
    "platform",
    "source_type",
    "region",
    "source",
    "title",
    "content",
    "url",
    "published_at",
    "is_comment",
    "quote_verified",
    "evidence_mode",
    "comment_id",
    "reply_id",
    "parent_comment_id",
    "like_count",
    "raw_file",
    "raw_sheet",
    "raw_row",
    "sentiment",
    "sentiment_source",
    "sentiment_status",
    "in_sentiment_denominator",
    "sentiment_exclusion_reason",
    "needs_review",
    "ai_formal_include",
    "ai_semantic_quality",
    "ai_formal_reason",
    "comment_heading",
    "topic_comment_heading",
]


def normalized_header(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).lower()
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", text)


ALIASES = {
    "comment_text": {
        "评论内容", "评论原文", "评论正文", "回复内容", "回复正文", "留言内容",
        "commenttext", "commentcontent", "replytext", "replycontent",
    },
    "generic_text": {"内容", "正文", "正文内容", "文本", "content", "text"},
    "comment_id": {"评论id", "评论编号", "commentid"},
    "reply_id": {"回复id", "回复编号", "replyid"},
    "parent_comment_id": {"父评论id", "上级评论id", "原评论id", "parentcommentid"},
    "topic": {"子事件", "子议题", "议题", "事件名称", "事件", "topic"},
    "platform": {"平台", "渠道", "来源平台", "platform"},
    "source": {"来源", "媒体名称", "媒体", "账号", "作者", "昵称", "source"},
    "published_at": {"发布时间", "发布日期", "时间", "日期", "publishedat", "date"},
    "url": {"链接", "原文链接", "发布地址", "url", "link", "articleurl"},
    "like_count": {"点赞量", "点赞数", "赞数", "likecount", "likes"},
    "parent_post_title": {"父原帖标题", "原帖标题", "parentposttitle"},
    "parent_post_published_at": {"父原帖发布时间", "原帖发布时间", "parentpostpublishedat"},
    "parent_post_source": {"父原帖来源", "原帖来源", "parentpostsource"},
    "row_id": {"id", "数据id", "信息id", "记录id", "rowid"},
    "marker": {"iscomment", "是否评论", "信息类型", "来源类型", "内容类型", "数据类型", "sourcetype"},
}

COMMENT_WORD_RE = re.compile(r"评论|回复|留言|comment|reply", re.I)
COMMENT_FILE_RE = re.compile(r"评论(?:明细|详情|数据)|回复(?:明细|详情|数据)|comment.?detail|reply.?detail", re.I)
AGGREGATE_COMMENT_RE = re.compile(r"评论(?:量|数|总数)|精选评论量|commentcount|replycount", re.I)
TRUE_VALUES = {"1", "true", "yes", "y", "是", "评论", "回复", "留言"}
FALSE_VALUES = {"0", "false", "no", "n", "否"}
VALID_LABELS = {"positive", "neutral", "negative"}
VALID_LABEL_SOURCES = {"ai_reviewed", "human_reviewed", "classifier", "active_learning_classifier"}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def first_value(row: dict[str, Any], aliases: set[str]) -> Any:
    for key, value in row.items():
        if normalized_header(key) in aliases and value not in (None, ""):
            return value
    return ""


def canonical_map(headers: Iterable[Any]) -> dict[str, str]:
    mapped: dict[str, str] = {}
    for header in headers:
        normalized = normalized_header(header)
        for field, aliases in ALIASES.items():
            if normalized in aliases and field not in mapped:
                mapped[field] = str(header)
    return mapped


def row_is_explicit_comment(row: dict[str, Any], mapping: dict[str, str], file_marker: bool) -> bool:
    if file_marker:
        return True
    marker_header = mapping.get("marker")
    if marker_header:
        value = clean_text(row.get(marker_header)).lower()
        if value in TRUE_VALUES or COMMENT_WORD_RE.search(value):
            return True
    return bool(mapping.get("comment_id") or mapping.get("reply_id"))


def stable_sample_id(record: dict[str, Any]) -> str:
    explicit = clean_text(record.get("comment_id") or record.get("reply_id"))
    if explicit:
        namespace = clean_text(record.get("platform") or record.get("source") or "comment")
        return "comment-" + hashlib.sha256(f"{namespace}|{explicit}".encode("utf-8")).hexdigest()[:20]
    payload = "|".join(
        clean_text(record.get(field))
        for field in ("raw_file", "raw_sheet", "raw_row", "published_at", "source", "url", "text")
    )
    return "comment-" + hashlib.sha256(payload.encode("utf-8")).hexdigest()[:20]


def detect_aggregate_columns(headers: Iterable[Any]) -> list[str]:
    result = []
    for header in headers:
        text = str(header or "")
        if AGGREGATE_COMMENT_RE.search(normalized_header(text)):
            result.append(text)
    return result


def parse_rows(
    rows: list[dict[str, Any]],
    source_path: Path,
    sheet_name: str,
    starting_row: int,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    headers = list(rows[0].keys()) if rows else []
    mapping = canonical_map(headers)
    file_marker = bool(COMMENT_FILE_RE.search(source_path.stem))
    direct_header = mapping.get("comment_text")
    generic_header = mapping.get("generic_text")
    aggregate_columns = detect_aggregate_columns(headers)
    aggregate_sums: dict[str, float] = {name: 0.0 for name in aggregate_columns}
    aggregate_numeric_rows: dict[str, int] = {name: 0 for name in aggregate_columns}
    accepted: list[dict[str, Any]] = []
    rejected_non_comment_text = 0
    blank_comment_text = 0

    for offset, row in enumerate(rows, start=starting_row):
        for name in aggregate_columns:
            value = row.get(name)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                aggregate_sums[name] += float(value)
                aggregate_numeric_rows[name] += 1

        explicit = row_is_explicit_comment(row, mapping, file_marker)
        text_value = row.get(direct_header) if direct_header else None
        if not direct_header and generic_header and explicit:
            text_value = row.get(generic_header)
        elif not direct_header and generic_header and clean_text(row.get(generic_header)):
            rejected_non_comment_text += 1
        text = clean_text(text_value)
        if not text:
            if direct_header or explicit:
                blank_comment_text += 1
            continue

        record = {
            "text": text,
            "topic": clean_text(row.get(mapping.get("topic", ""))),
            "platform": clean_text(row.get(mapping.get("platform", ""))),
            "source": clean_text(row.get(mapping.get("source", ""))),
            "published_at": clean_text(row.get(mapping.get("published_at", ""))),
            "url": clean_text(row.get(mapping.get("url", ""))),
            "comment_id": clean_text(row.get(mapping.get("comment_id", ""))),
            "reply_id": clean_text(row.get(mapping.get("reply_id", ""))),
            "parent_comment_id": clean_text(row.get(mapping.get("parent_comment_id", ""))),
            "like_count": clean_text(row.get(mapping.get("like_count", ""))),
            "parent_post_title": clean_text(row.get(mapping.get("parent_post_title", ""))),
            "parent_post_published_at": clean_text(row.get(mapping.get("parent_post_published_at", ""))),
            "parent_post_source": clean_text(row.get(mapping.get("parent_post_source", ""))),
            "raw_file": source_path.name,
            "raw_sheet": sheet_name,
            "raw_row": offset,
        }
        record["sample_id"] = stable_sample_id(record)
        accepted.append(record)

    audit = {
        "sheet": sheet_name,
        "rows_seen": len(rows),
        "headers": [str(item) for item in headers],
        "direct_comment_text_field": direct_header or "",
        "generic_text_field": generic_header or "",
        "explicit_comment_file": file_marker,
        "accepted_comment_rows": len(accepted),
        "rejected_generic_text_rows": rejected_non_comment_text,
        "blank_comment_text_rows": blank_comment_text,
        "aggregate_comment_fields": [
            {"field": name, "numeric_rows": aggregate_numeric_rows[name], "sum": aggregate_sums[name]}
            for name in aggregate_columns
        ],
    }
    return accepted, audit


def choose_header_row(values: list[list[Any]]) -> int:
    best_index = 0
    best_score = -1
    known = set().union(*ALIASES.values())
    for index, row in enumerate(values[:20]):
        nonempty = sum(value not in (None, "") for value in row)
        recognized = sum(normalized_header(value) in known for value in row if value not in (None, ""))
        commentish = sum(bool(COMMENT_WORD_RE.search(str(value))) for value in row if value not in (None, ""))
        score = nonempty + recognized * 8 + commentish * 2
        if score > best_score:
            best_index, best_score = index, score
    return best_index


def unique_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    used: Counter[str] = Counter()
    for index, value in enumerate(values, start=1):
        base = clean_text(value) or f"column_{index}"
        used[base] += 1
        headers.append(base if used[base] == 1 else f"{base}_{used[base]}")
    return headers


def table_from_matrix(matrix: list[list[Any]]) -> tuple[list[dict[str, Any]], int]:
    if not matrix:
        return [], 1
    header_index = choose_header_row(matrix)
    headers = unique_headers(matrix[header_index])
    rows: list[dict[str, Any]] = []
    for values in matrix[header_index + 1 :]:
        if not any(value not in (None, "") for value in values):
            continue
        padded = values + [None] * max(0, len(headers) - len(values))
        rows.append(dict(zip(headers, padded[: len(headers)])))
    return rows, header_index + 2


def read_csv_tables(path: Path) -> list[tuple[str, list[dict[str, Any]], int]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                matrix = [list(row) for row in csv.reader(handle)]
            rows, start = table_from_matrix(matrix)
            return [("CSV", rows, start)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Unable to decode CSV {path}: {last_error}")


def read_json_tables(path: Path) -> list[tuple[str, list[dict[str, Any]], int]]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(payload, dict):
        for key in ("rows", "data", "items", "comments", "results"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
        else:
            payload = [payload]
    if not isinstance(payload, list):
        raise ValueError("JSON input must be a row list or contain rows/data/items/comments/results")
    return [("JSON", [item for item in payload if isinstance(item, dict)], 1)]


def read_excel_tables(path: Path) -> list[tuple[str, list[dict[str, Any]], int]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[tuple[str, list[dict[str, Any]], int]] = []
    try:
        for sheet in workbook.worksheets:
            # Monitoring exports sometimes declare A1:A1 despite containing a table.
            sheet.reset_dimensions()
            sheet.calculate_dimension(force=True)
            matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
            rows, start = table_from_matrix(matrix)
            tables.append((sheet.title, rows, start))
    finally:
        workbook.close()
    return tables


def read_tables(path: Path) -> list[tuple[str, list[dict[str, Any]], int]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv_tables(path)
    if suffix == ".json":
        return read_json_tables(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_excel_tables(path)
    raise ValueError(f"Unsupported input type: {path.suffix}")


def discover_inputs(explicit: list[str], input_dir: str | None, system_workbook: str | None) -> list[Path]:
    candidates = [Path(item).expanduser().resolve() for item in explicit]
    if input_dir:
        directory = Path(input_dir).expanduser().resolve()
        candidates.extend(path for path in directory.iterdir() if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES)
    system = Path(system_workbook).expanduser().resolve() if system_workbook else None
    unique: list[Path] = []
    seen: set[Path] = set()
    for path in sorted(candidates, key=lambda item: str(item).lower()):
        if system and path == system:
            continue
        if path not in seen:
            unique.append(path)
            seen.add(path)
    return unique


def workbook_metadata(path: str | None) -> dict[str, Any]:
    if not path:
        return {"source_file": "", "topics": [], "monitoring_period": {}}
    scripts_dir = Path(__file__).resolve().parent
    if str(scripts_dir) not in sys.path:
        sys.path.insert(0, str(scripts_dir))
    try:
        from ingest_monitoring_workbook import ingest_workbook

        data = ingest_workbook(path)
        topic_names = []
        for topic in data.get("topics") or []:
            if isinstance(topic, dict):
                name = topic.get("name") or topic.get("topic") or topic.get("title")
            else:
                name = topic
            if name:
                topic_names.append(str(name))
        return {
            "source_file": str(Path(path).resolve()),
            "master_event": data.get("master_event") or {},
            "topics": topic_names,
            "monitoring_period": data.get("monitoring_period") or {},
        }
    except Exception as exc:  # Metadata helps the audit but must not hide comment extraction.
        return {"source_file": str(Path(path).resolve()), "topics": [], "monitoring_period": {}, "metadata_error": str(exc)}


def parse_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    normalized = clean_text(value).lower()
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    return None


def read_result_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if path.suffix.lower() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(payload, dict):
            payload = payload.get("rows") or payload.get("results") or payload.get("data") or []
        if isinstance(payload, list):
            return [row for row in payload if isinstance(row, dict)]
    raise ValueError("Sentiment results must be a CSV or JSON row list")


def match_topic_indexes(topic_value: Any, topics: list[str]) -> list[int]:
    text = normalized_header(topic_value)
    if not text:
        return []
    matches = [index for index, topic in enumerate(topics, start=1) if normalized_header(topic) in text or text in normalized_header(topic)]
    if matches:
        return matches
    for pattern in (r"子(?:事件|议题)?(\d+)", r"议题(\d+)"):
        found = re.search(pattern, clean_text(topic_value), re.I)
        if found and 1 <= int(found.group(1)) <= len(topics):
            return [int(found.group(1))]
    return []


def build_workbook_summary(
    comments: list[dict[str, Any]],
    metadata: dict[str, Any],
    results_path: str | None,
    min_topic_denominator: int = 1,
) -> tuple[dict[str, Any], list[str]]:
    topics = [str(item) for item in metadata.get("topics") or []]
    summary_topics = [
        {
            "index": index,
            "title": title,
            "status": "pending",
            "denominator": 0,
            "positive": None,
            "neutral": None,
            "negative": None,
        }
        for index, title in enumerate(topics, start=1)
    ]
    blockers: list[str] = []
    if not comments:
        return {"schema_version": "1.0", "status": "missing_comment_detail", "topics": summary_topics}, blockers
    if not results_path:
        return {"schema_version": "1.0", "status": "awaiting_sentiment_results", "topics": summary_topics}, blockers

    result_rows = read_result_rows(Path(results_path).expanduser().resolve())
    by_id: dict[str, dict[str, Any]] = {}
    duplicate_ids: set[str] = set()
    for row in result_rows:
        sample_id = clean_text(row.get("sample_id"))
        if not sample_id:
            blockers.append("情感结果中存在空 sample_id。")
            continue
        if sample_id in by_id:
            duplicate_ids.add(sample_id)
        by_id[sample_id] = row
    if duplicate_ids:
        blockers.append(f"情感结果存在 {len(duplicate_ids)} 个重复 sample_id。")
    comment_ids = {row["sample_id"] for row in comments}
    missing = sorted(comment_ids - set(by_id))
    unknown = sorted(set(by_id) - comment_ids)
    if missing:
        blockers.append(f"情感结果缺少 {len(missing)} 个评论 sample_id。")
    if unknown:
        blockers.append(f"情感结果包含 {len(unknown)} 个输入中不存在的 sample_id。")

    counts: dict[int, Counter[str]] = {index: Counter() for index in range(1, len(topics) + 1)}
    for comment in comments:
        result = by_id.get(comment["sample_id"])
        if not result:
            continue
        needs_review = parse_bool(result.get("needs_review"))
        if needs_review is not False:
            blockers.append(f"{comment['sample_id']} 的 needs_review 不是 false。")
            continue
        in_denominator = parse_bool(result.get("in_sentiment_denominator"))
        if in_denominator is False:
            if not clean_text(result.get("exclusion_reason")):
                blockers.append(f"{comment['sample_id']} 被排除但缺少 exclusion_reason。")
            continue
        if in_denominator is not True:
            blockers.append(f"{comment['sample_id']} 缺少有效的 in_sentiment_denominator。")
            continue
        label = clean_text(result.get("label") or result.get("predicted_label")).lower()
        label_source = clean_text(result.get("label_source")).lower()
        if label not in VALID_LABELS:
            blockers.append(f"{comment['sample_id']} 的情感标签无效。")
            continue
        if label_source not in VALID_LABEL_SOURCES:
            blockers.append(f"{comment['sample_id']} 的 label_source 无效。")
            continue
        indexes = match_topic_indexes(comment.get("topic") or result.get("topic"), topics)
        if not indexes:
            blockers.append(f"{comment['sample_id']} 无法映射到子议题。")
            continue
        for index in indexes:
            counts[index][label] += 1

    for item in summary_topics:
        topic_counts = counts[item["index"]]
        denominator = sum(topic_counts.values())
        if denominator >= min_topic_denominator:
            item.update(
                {
                    "status": "ready",
                    "denominator": denominator,
                    "positive": topic_counts["positive"] / denominator,
                    "neutral": topic_counts["neutral"] / denominator,
                    "negative": topic_counts["negative"] / denominator,
                    "counts": dict(topic_counts),
                }
            )
        elif denominator:
            # Preserve auditable observed counts, but do not turn absence in a
            # small convenience sample into a formal 0% workbook finding.
            item.update(
                {
                    "status": "insufficient_sample",
                    "denominator": denominator,
                    "counts": dict(topic_counts),
                    "observed_rates": {
                        "positive": topic_counts["positive"] / denominator,
                        "neutral": topic_counts["neutral"] / denominator,
                        "negative": topic_counts["negative"] / denominator,
                    },
                    "interpretation": "样本不足；仅描述本批观察结果，不回填正式百分比。",
                }
            )
    ready_topics = sum(item["status"] == "ready" for item in summary_topics)
    if blockers:
        status = "blocked_invalid_results"
    elif ready_topics == len(summary_topics) and ready_topics:
        status = "ready_for_workbook_backfill"
    elif ready_topics:
        status = "partial_workbook_backfill"
    else:
        status = "no_topic_denominator"
    return {
        "schema_version": "1.0",
        "status": status,
        "result_file": str(Path(results_path).resolve()),
        "minimum_topic_denominator_for_backfill": min_topic_denominator,
        "topics": summary_topics,
    }, blockers


def build_report_comment_handoff(
    comments: list[dict[str, Any]],
    results_path: str | None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Create the strict comment-evidence handoff consumed by report rendering.

    Only fully reviewed denominator rows are exposed to the dashboard and formal
    report.  Excluded/unresolved rows remain in the sentiment result and audit
    files, but cannot accidentally become quoted netizen comments.
    """
    if not results_path:
        return [], {
            "status": "awaiting_sentiment_results",
            "eligible_rows": 0,
            "excluded_rows": 0,
            "unresolved_rows": len(comments),
        }

    result_rows = read_result_rows(Path(results_path).expanduser().resolve())
    by_id = {
        clean_text(row.get("sample_id")): row
        for row in result_rows
        if clean_text(row.get("sample_id"))
    }
    eligible: list[dict[str, Any]] = []
    excluded = 0
    unresolved = 0
    missing_traceability = 0
    missing_formal_review = 0
    for comment in comments:
        result = by_id.get(comment["sample_id"])
        if not result or parse_bool(result.get("needs_review")) is not False:
            unresolved += 1
            continue
        if parse_bool(result.get("in_sentiment_denominator")) is not True:
            excluded += 1
            continue
        label = clean_text(result.get("label") or result.get("predicted_label")).lower()
        label_source = clean_text(result.get("label_source")).lower()
        if label not in VALID_LABELS or label_source not in VALID_LABEL_SOURCES:
            unresolved += 1
            continue
        has_identity = bool(comment.get("comment_id") or comment.get("reply_id"))
        if not clean_text(comment.get("url")) or not has_identity:
            missing_traceability += 1
            continue
        ai_formal_include = parse_bool(result.get("ai_formal_include"))
        ai_semantic_quality = clean_text(result.get("ai_semantic_quality"))
        ai_formal_reason = clean_text(result.get("ai_formal_reason"))
        comment_heading = clean_text(result.get("comment_heading"))
        topic_comment_heading = clean_text(result.get("topic_comment_heading"))
        if ai_formal_include is None or not ai_semantic_quality or not ai_formal_reason:
            missing_formal_review += 1
            continue
        if ai_formal_include and not comment_heading:
            missing_formal_review += 1
            continue
        eligible.append(
            {
                "id": comment["sample_id"],
                "sample_id": comment["sample_id"],
                "topic": comment.get("topic") or result.get("topic", ""),
                "platform": comment.get("platform", ""),
                "source_type": "netizen_comment",
                "region": "domestic",
                "source": comment.get("source", ""),
                "title": comment.get("text", ""),
                "content": comment.get("text", ""),
                "url": comment.get("url", ""),
                "published_at": comment.get("published_at", ""),
                "is_comment": "true",
                "quote_verified": "true",
                "evidence_mode": "verbatim_public_comment",
                "comment_id": comment.get("comment_id", ""),
                "reply_id": comment.get("reply_id", ""),
                "parent_comment_id": comment.get("parent_comment_id", ""),
                "like_count": comment.get("like_count", ""),
                "raw_file": comment.get("raw_file", ""),
                "raw_sheet": comment.get("raw_sheet", ""),
                "raw_row": comment.get("raw_row", ""),
                "sentiment": label,
                "sentiment_source": label_source,
                "sentiment_status": "classified",
                "in_sentiment_denominator": "true",
                "sentiment_exclusion_reason": "",
                "needs_review": "false",
                "ai_formal_include": "true" if ai_formal_include else "false",
                "ai_semantic_quality": ai_semantic_quality,
                "ai_formal_reason": ai_formal_reason,
                "comment_heading": comment_heading,
                "topic_comment_heading": topic_comment_heading,
            }
        )
    formal_by_topic: dict[str, list[dict[str, Any]]] = {}
    for row in eligible:
        if parse_bool(row.get("ai_formal_include")) is True:
            formal_by_topic.setdefault(clean_text(row.get("topic")), []).append(row)
    for topic, rows in formal_by_topic.items():
        if len(rows) == 1 and not clean_text(rows[0].get("topic_comment_heading")):
            rows[0]["topic_comment_heading"] = clean_text(rows[0].get("comment_heading"))
            continue
        topic_headings = {clean_text(row.get("topic_comment_heading")) for row in rows if clean_text(row.get("topic_comment_heading"))}
        if len(topic_headings) != 1 or any(not clean_text(row.get("topic_comment_heading")) for row in rows):
            missing_formal_review += len(rows)
    return eligible, {
        "status": "ready" if eligible and not unresolved and not missing_traceability and not missing_formal_review else "partial",
        "eligible_rows": len(eligible),
        "excluded_rows": excluded,
        "unresolved_rows": unresolved,
        "missing_traceability_rows": missing_traceability,
        "missing_formal_review_rows": missing_formal_review,
        "result_file": str(Path(results_path).expanduser().resolve()),
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def label_schema_text() -> str:
    return """# CWH 网民情感标注口径

仅标注监测时段内、可追溯的逐条网民评论。新闻稿、公众号文章正文、政策文本和仅有评论量的汇总行不进入分母。

## 标签

- `positive`：明确支持、认可、称赞、期待政策或会议议题。
- `neutral`：事实陈述、询问、信息补充，或态度不明显。
- `negative`：明确反对、不满、质疑、担忧、讽刺或批评。

## 结果字段

每个 `sample_id` 必须有一条结果。纳入分母的行填写 `label`、`label_source`、`in_sentiment_denominator=true`、`needs_review=false`；排除行填写 `in_sentiment_denominator=false` 和明确的 `exclusion_reason`。文章正文不得作为网民评论标注。

每条可追溯评论还必须单独完成正式引用审核：填写 `ai_formal_include=true/false`、`ai_semantic_quality` 和 `ai_formal_reason`。仅当 `ai_formal_include=true` 时填写表达具体立场、诉求或判断的 `comment_heading`；禁止使用“关注+议题名称”代替观点判断。

同一子议题有多条 `ai_formal_include=true` 的评论时，还必须为这些行填写完全一致的 `topic_comment_heading`，作为正文的唯一并列小标题。只有一条正式评论时可与 `comment_heading` 相同。
"""


def locate_skill_dir(explicit: str | None) -> Path:
    candidates = [
        explicit,
        os.environ.get("CWH_SENTIMENT_SKILL_DIR"),
        str(Path.home() / ".codex" / "skills" / "large-scale-sentiment-analysis"),
    ]
    for item in candidates:
        if item and (Path(item).expanduser() / "scripts" / "prepare_text_data.py").exists():
            return Path(item).expanduser().resolve()
    raise FileNotFoundError("large-scale-sentiment-analysis skill scripts were not found")


def run_preparation(input_csv: Path, out_dir: Path, skill_dir: Path, row_count: int) -> dict[str, Any]:
    scripts = skill_dir / "scripts"
    prepared_dir = out_dir / "prepared"
    quality_dir = out_dir / "quality_profile"
    batch_dir = out_dir / "seed_batches"
    commands = [
        [sys.executable, str(scripts / "prepare_text_data.py"), "--input", str(input_csv), "--output-dir", str(prepared_dir), "--text-col", "text", "--date-col", "published_at", "--source-col", "sample_id", "--dedupe-mode", "none"],
        [sys.executable, str(scripts / "profile_text_quality.py"), "--input", str(input_csv), "--output-dir", str(quality_dir), "--text-col", "text", "--sample-size", str(min(1000, row_count))],
    ]
    logs = []
    for command in commands:
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
        logs.append({"command": command, "returncode": result.returncode, "stdout": result.stdout.strip(), "stderr": result.stderr.strip()})
        if result.returncode:
            raise RuntimeError(f"Sentiment preparation failed: {result.stderr or result.stdout}")

    prepared_path = prepared_dir / "prepared_texts.csv"
    with input_csv.open("r", encoding="utf-8-sig", newline="") as handle:
        originals = {row["sample_id"]: row for row in csv.DictReader(handle)}
    with prepared_path.open("r", encoding="utf-8-sig", newline="") as handle:
        prepared = list(csv.DictReader(handle))
    enriched = []
    for row in prepared:
        sample_id = row.get("source", "")
        original = originals.get(sample_id, {})
        row.update({"sample_id": sample_id, "topic": original.get("topic", ""), "platform": original.get("platform", ""), "comment_source": original.get("source", "")})
        enriched.append(row)
    fields = list(enriched[0].keys()) if enriched else ["sample_id", "clean_text", "topic", "platform"]
    write_csv(prepared_path, enriched, fields)

    batch_command = [
        sys.executable, str(scripts / "make_llm_batches.py"), "--input", str(prepared_path), "--output-dir", str(batch_dir),
        "--sample-n", str(min(1000, len(enriched))), "--batch-size", "100", "--strata-cols", "topic,platform,text_len_bucket",
    ]
    result = subprocess.run(batch_command, capture_output=True, text=True, encoding="utf-8", errors="replace")
    logs.append({"command": batch_command, "returncode": result.returncode, "stdout": result.stdout.strip(), "stderr": result.stderr.strip()})
    if result.returncode:
        raise RuntimeError(f"LLM batch creation failed: {result.stderr or result.stdout}")
    return {"skill_dir": str(skill_dir), "prepared_rows": len(enriched), "logs": logs}


def markdown_audit(audit: dict[str, Any]) -> str:
    lines = [
        "# CWH 情感分析阶段审计",
        "",
        f"- 状态：`{audit['status']}`",
        f"- 扫描文件：{audit['files_scanned']} 个",
        f"- 识别逐条评论：{audit['comment_rows']} 条",
        f"- 唯一 sample_id：{audit['unique_sample_ids']} 个",
        f"- 仅有评论量的汇总字段：{audit['aggregate_comment_field_count']} 个",
        "",
        "## 结论",
        "",
    ]
    if audit["status"] == "missing_comment_detail":
        lines.extend([
            "当前输入没有可追溯的逐条评论正文，不能计算正式网民情感比例。文章正文和评论量汇总均未进入情感分母。",
            "",
            "需要从监测系统补导评论/回复明细，至少包含评论内容，并尽量包含评论 ID、平台、发布时间、原文链接和议题字段。",
        ])
    elif audit["status"] == "ready_for_preparation":
        lines.extend([
            "评论明细已通过入口门禁，可继续进行清洗、分层抽样、AI 复核、分类器训练和低置信度复核。正式比例仍须等逐条结果完整回收后计算。",
        ])
    elif audit["status"] == "ai_review_complete_sample_only":
        lines.extend([
            "逐条AI审核已完成，合格评论可作为可追溯样本观察和报告引文；各议题有效样本不足总体比例门槛，不得回填或表述为总体情绪比例。",
        ])
    else:
        lines.extend(["逐条审核与结果门禁已完成；是否允许回填总体比例以工作簿摘要状态为准。"])
    lines.extend(["", "## 文件检查", ""])
    for item in audit["file_audits"]:
        lines.append(f"- `{item['file']}`：识别 {item['accepted_comment_rows']} 条逐条评论；{item.get('error') or '读取正常'}")
    lines.extend(["", "## 门禁", ""])
    for blocker in audit["blockers"]:
        lines.append(f"- {blocker}")
    if not audit["blockers"]:
        lines.append("- 评论明细入口门禁通过；情感结果门禁尚待后续标注流程完成。")
    return "\n".join(lines) + "\n"


def run(args: argparse.Namespace) -> dict[str, Any]:
    out_dir = Path(args.output_dir).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = discover_inputs(args.comment_detail or [], args.input_dir, args.system_workbook)
    all_comments: list[dict[str, Any]] = []
    file_audits: list[dict[str, Any]] = []
    for path in paths:
        file_audit = {"file": str(path), "accepted_comment_rows": 0, "sheets": []}
        try:
            for sheet_name, rows, starting_row in read_tables(path):
                comments, sheet_audit = parse_rows(rows, path, sheet_name, starting_row)
                all_comments.extend(comments)
                file_audit["accepted_comment_rows"] += len(comments)
                file_audit["sheets"].append(sheet_audit)
        except Exception as exc:
            file_audit["error"] = str(exc)
        file_audits.append(file_audit)

    duplicates = Counter(item["sample_id"] for item in all_comments)
    duplicate_ids = {key for key, count in duplicates.items() if count > 1}
    # Preserve the first provenance row for identical explicit comment IDs.
    unique_comments = list({item["sample_id"]: item for item in reversed(all_comments)}.values())
    unique_comments.sort(key=lambda row: (row["raw_file"], row["raw_sheet"], int(row["raw_row"])))
    input_csv = out_dir / "sentiment_input.csv"
    write_csv(input_csv, unique_comments, OUTPUT_FIELDS)
    (out_dir / "label_schema.md").write_text(label_schema_text(), encoding="utf-8")

    aggregate_field_count = sum(len(sheet["aggregate_comment_fields"]) for file in file_audits for sheet in file.get("sheets", []))
    blockers = []
    if not unique_comments:
        blockers.append("缺少逐条评论正文；正式情感比例保持待分析。")
    if duplicate_ids:
        blockers.append(f"发现 {len(duplicate_ids)} 个重复 sample_id，已按 ID 保留首条并在 JSON 审计中记录。")
    metadata = workbook_metadata(args.system_workbook)
    workbook_summary, result_blockers = build_workbook_summary(
        unique_comments,
        metadata,
        args.sentiment_results,
        args.min_topic_denominator,
    )
    blockers.extend(result_blockers)
    summary_path = out_dir / "sentiment_workbook_summary.json"
    summary_path.write_text(json.dumps(workbook_summary, ensure_ascii=False, indent=2), encoding="utf-8")
    report_handoff_rows, report_handoff_audit = build_report_comment_handoff(
        unique_comments, args.sentiment_results
    )
    report_handoff_path = out_dir / "report_comment_handoff.csv"
    write_csv(report_handoff_path, report_handoff_rows, REPORT_HANDOFF_FIELDS)
    if not unique_comments:
        stage_status = "missing_comment_detail"
        percentage_status = "pending"
    elif not args.sentiment_results:
        stage_status = "ready_for_preparation"
        percentage_status = "pending"
    elif report_handoff_audit.get("status") == "ready":
        if workbook_summary.get("status") == "no_topic_denominator":
            stage_status = "ai_review_complete_sample_only"
            percentage_status = "sample_only"
        else:
            stage_status = "ai_review_complete"
            percentage_status = "available"
    else:
        stage_status = "ai_review_incomplete"
        percentage_status = "pending"
    audit: dict[str, Any] = {
        "schema_version": "1.0",
        "status": stage_status,
        "files_scanned": len(paths),
        "comment_rows_before_id_dedupe": len(all_comments),
        "comment_rows": len(unique_comments),
        "unique_sample_ids": len({item["sample_id"] for item in unique_comments}),
        "duplicate_sample_ids": sorted(duplicate_ids),
        "aggregate_comment_field_count": aggregate_field_count,
        "formal_sentiment_percentages": percentage_status,
        "metadata": metadata,
        "file_audits": file_audits,
        "blockers": blockers,
        "workbook_backfill_status": workbook_summary["status"],
        "report_comment_handoff": report_handoff_audit,
        "outputs": {
            "sentiment_input": str(input_csv),
            "label_schema": str(out_dir / "label_schema.md"),
            "workbook_summary": str(summary_path),
            "report_comment_handoff": str(report_handoff_path),
        },
    }
    if args.run_preparation and unique_comments:
        skill_dir = locate_skill_dir(args.sentiment_skill_dir)
        audit["preparation"] = run_preparation(input_csv, out_dir, skill_dir, len(unique_comments))
        audit["status"] = "prepared_for_seed_review"
    elif args.run_preparation:
        audit["preparation"] = {"status": "not_run", "reason": "no row-level comment text"}

    if args.backfill_workbook:
        if workbook_summary["status"] in {"ready_for_workbook_backfill", "partial_workbook_backfill"} and not result_blockers:
            if not args.backfill_output:
                raise ValueError("--backfill-output is required with --backfill-workbook")
            verification = out_dir / "sentiment_backfill_verification.json"
            backfill_script = Path(__file__).resolve().with_name("backfill_cwh_sentiment_workbook.ps1")
            command = [
                "powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(backfill_script),
                "-BaseWorkbook", str(Path(args.backfill_workbook).resolve()),
                "-SentimentSummary", str(summary_path),
                "-Output", str(Path(args.backfill_output).resolve()),
                "-Verification", str(verification),
            ]
            result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
            if result.returncode:
                raise RuntimeError(f"Workbook sentiment backfill failed: {result.stderr or result.stdout}")
            audit["workbook_backfill"] = {"status": "completed", "output": str(Path(args.backfill_output).resolve()), "verification": str(verification)}
        else:
            audit["workbook_backfill"] = {"status": "not_run", "reason": "no audited topic sentiment denominator"}

    json_path = out_dir / "sentiment_stage_audit.json"
    md_path = out_dir / "sentiment_stage_audit.md"
    audit["outputs"].update({"audit_json": str(json_path), "audit_markdown": str(md_path)})
    json_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    md_path.write_text(markdown_audit(audit), encoding="utf-8")
    return audit


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Extract and gate row-level comments for the independent CWH sentiment stage.")
    parser.add_argument("--system-workbook", help="Generated/standard CWH workbook, used for topic and monitoring metadata.")
    parser.add_argument("--comment-detail", action="append", default=[], help="Explicit CSV/JSON/XLSX comment-detail file; may be repeated.")
    parser.add_argument("--input-dir", help="Directory to audit and auto-discover supported detail files.")
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--run-preparation", action="store_true", help="Run cleaning, quality profiling and stratified seed-batch scripts when comments exist.")
    parser.add_argument("--sentiment-skill-dir", help="Override the installed large-scale-sentiment-analysis skill directory.")
    parser.add_argument("--sentiment-results", help="Reviewed row-level CSV/JSON results to validate and aggregate for workbook backfill.")
    parser.add_argument(
        "--min-topic-denominator",
        type=int,
        default=20,
        help="Minimum reviewed comments per topic before formal percentages may be backfilled (default: 20).",
    )
    parser.add_argument("--backfill-workbook", help="Accepted standard workbook whose sentiment cells may be filled after the result gate passes.")
    parser.add_argument("--backfill-output", help="Output workbook path for the chart-preserving sentiment copy.")
    return parser


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    audit = run(build_parser().parse_args())
    print(json.dumps({"status": audit["status"], "comment_rows": audit["comment_rows"], "output_dir": str(Path(audit["outputs"]["sentiment_input"]).parent)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
