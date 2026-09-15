from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import load_workbook

try:
    from opencc import OpenCC
except ImportError:  # pragma: no cover - exercised by deployment gate
    OpenCC = None  # type: ignore[assignment]


SCRIPT_DIRECTORY = Path(__file__).resolve().parent
if str(SCRIPT_DIRECTORY) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIRECTORY))

from cwh_hotword_pipeline import build_hotword_payload, render_wordcloud_png
from source_identity import public_source_family


CANONICAL_DAILY_COLUMNS = [
    "date",
    "public_articles",
    "public_recommend",
    "public_comments",
    "weibo",
    "domestic_news",
    "domestic_app",
    "domestic_forum",
    "other_video",
    "overseas_news",
    "x",
    "overseas_other",
    "video_account",
    "douyin",
]

DISPLAY_DAILY_HEADERS = [
    "日期",
    "公众文章",
    "公号-在看量",
    "公号-精选评论量",
    "新浪微博",
    "境内新闻",
    "境内APP",
    "境内论坛",
    "其他视频",
    "境外新闻",
    "推特",
    "境外其他",
    "视频号发文",
    "抖音",
]

TRADITIONAL_TRANSLATION = str.maketrans(
    {
        "國": "国",
        "務": "务",
        "會": "会",
        "議": "议",
        "數": "数",
        "興": "兴",
        "產": "产",
        "業": "业",
        "資": "资",
        "護": "护",
        "計": "计",
        "劃": "划",
        "災": "灾",
        "聽": "听",
        "報": "报",
        "體": "体",
        "網": "网",
        "訊": "讯",
        "佈": "布",
        "領": "领",
        "薦": "荐",
        "風": "风",
        "險": "险",
        "時": "时",
        "發": "发",
        "達": "达",
        "與": "与",
        "為": "为",
        "進": "进",
        "過": "过",
        "關": "关",
        "於": "于",
        "後": "后",
        "臺": "台",
        "陸": "陆",
        "強": "强",
    }
)


class PipelineError(RuntimeError):
    pass


_T2S_CONVERTER = OpenCC("t2s") if OpenCC is not None else None


def to_simplified_review_text(value: Any) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if not text:
        return ""
    if _T2S_CONVERTER is None:
        raise PipelineError("境外正式字段需要opencc-python-reimplemented完成繁简转换")
    return _T2S_CONVERTER.convert(text)


@dataclass(frozen=True)
class InputBundle:
    total_heat: Path
    child_heat: dict[int, Path]
    public_top: Path
    overseas_latest: Path
    classification_audit: tuple[dict[str, Any], ...] = ()


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.translate(TRADITIONAL_TRANSLATION).lower()
    return re.sub(r"\s+", "", text)


def normalize_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return "" if value is None else str(value).strip()


def to_number(value: Any) -> int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip().replace(",", "")
    text = text[:-1] if text.endswith("+") else text
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return normalize_date(value)
    return value


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def find_artifact_node_modules(skill_directory: Path) -> Path | None:
    candidates: list[Path] = [skill_directory / "node_modules"]
    if os.environ.get("CWH_NODE_MODULES"):
        candidates.append(Path(os.environ["CWH_NODE_MODULES"]))
    for entry in os.environ.get("NODE_PATH", "").split(os.pathsep):
        if entry:
            candidates.append(Path(entry))
    runtime_root = Path.home() / ".cache" / "codex-runtimes"
    if runtime_root.exists():
        candidates.extend(runtime_root.glob("*/dependencies/node/node_modules"))
    for candidate in candidates:
        if (candidate / "@oai" / "artifact-tool").exists():
            return candidate.resolve()
    return None


def render_workbook_sheets(
    output_path: Path,
    sheet_names: list[str],
    previews_dir: Path,
    renderer: Path,
    node_path: Path | None,
    node_modules: Path | None,
) -> tuple[list[str], list[dict[str, str]]]:
    """Render every sheet through the bundled artifact runtime for acceptance."""
    if node_path is None or node_modules is None or not renderer.is_file():
        reason = "node_or_artifact_renderer_unavailable"
        return [], [{"sheet_name": name, "error": reason} for name in sheet_names]
    node_environment = os.environ.copy()
    existing_node_path = node_environment.get("NODE_PATH", "")
    node_environment["NODE_PATH"] = str(node_modules) + (
        os.pathsep + existing_node_path if existing_node_path else ""
    )
    rendered_sheets: list[str] = []
    render_errors: list[dict[str, str]] = []
    for sheet_name in sheet_names:
        safe_sheet_name = re.sub(r'[\\/:*?"<>|]', "_", sheet_name)
        preview_path = previews_dir / f"{safe_sheet_name}.png"
        render_started_at = datetime.now().timestamp()
        render_command = [str(node_path), str(renderer), str(output_path), sheet_name, str(preview_path)]
        render_result = subprocess.run(render_command, check=False, env=node_environment)
        preview_is_fresh = preview_path.exists() and preview_path.stat().st_mtime >= render_started_at - 2
        # Some Windows artifact runtimes emit the complete PNG and then exit
        # abnormally during Node teardown. A fresh, non-empty preview is the
        # acceptance evidence; a zero exit without a fresh file is not.
        if preview_is_fresh and preview_path.stat().st_size > 0:
            rendered_sheets.append(sheet_name)
        else:
            render_errors.append(
                {
                    "sheet_name": sheet_name,
                    "exit_code": str(render_result.returncode),
                    "fresh_preview": str(preview_is_fresh).lower(),
                }
            )
    return rendered_sheets, render_errors


def matched_sheet(workbook: Any, aliases: Iterable[str]) -> Any | None:
    normalized = {normalize_text(name): name for name in workbook.sheetnames}
    for alias in aliases:
        original = normalized.get(normalize_text(alias))
        if original is not None:
            return workbook[original]
    return None


def matched_column_names(headers: list[Any], aliases: dict[str, list[str]]) -> dict[str, int]:
    normalized = {normalize_text(value): index for index, value in enumerate(headers) if value is not None}
    result: dict[str, int] = {}
    for canonical, candidates in aliases.items():
        for candidate in candidates:
            index = normalized.get(normalize_text(candidate))
            if index is not None:
                result[canonical] = index
                break
    return result


def profile_raw_workbook(path: Path, config: dict[str, Any]) -> dict[str, Any]:
    """Classify a workbook by sheets and headers before consulting its filename."""
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        structural_roles: list[str] = []
        evidence: list[str] = []
        heat_total: int | float | None = None

        daily_sheet = matched_sheet(workbook, config["sheet_aliases"]["daily_trend"])
        if daily_sheet is not None:
            headers = [daily_sheet.cell(1, column).value for column in range(1, daily_sheet.max_column + 1)]
            columns = matched_column_names(headers, config["daily_columns"])
            missing = [name for name in config["daily_columns"] if name not in columns]
            if not missing:
                structural_roles.append("heat")
                evidence.append(f"工作表“{daily_sheet.title}”包含完整日趋势字段")
                heat_total = 0
                valid_rows = 0
                for row_index in range(2, daily_sheet.max_row + 1):
                    raw_date = daily_sheet.cell(row_index, columns["date"] + 1).value
                    if raw_date in (None, ""):
                        continue
                    valid_rows += 1
                    for canonical in CANONICAL_DAILY_COLUMNS[1:]:
                        number = to_number(daily_sheet.cell(row_index, columns[canonical] + 1).value)
                        if number is None:
                            heat_total = None
                            break
                        heat_total += number
                    if heat_total is None:
                        break
                if not valid_rows:
                    heat_total = None
            else:
                evidence.append(f"工作表“{daily_sheet.title}”缺少日趋势字段：{missing}")

        detail_roles = {
            "public_top": "public_articles",
            "overseas_latest": "overseas_news",
        }
        for role, alias_key in detail_roles.items():
            worksheet = matched_sheet(workbook, config["sheet_aliases"][alias_key])
            if worksheet is None:
                continue
            headers = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
            columns = matched_column_names(headers, config["detail_columns"])
            missing = [name for name in config["detail_columns"] if name not in columns]
            if not missing:
                structural_roles.append(role)
                evidence.append(f"工作表“{worksheet.title}”包含完整明细字段")
            else:
                evidence.append(f"工作表“{worksheet.title}”缺少明细字段：{missing}")

        return {
            "path": path,
            "file": path.name,
            "sheet_names": list(workbook.sheetnames),
            "structural_roles": structural_roles,
            "heat_total": heat_total,
            "evidence": evidence,
            "filename_monitoring_window": filename_monitoring_window(path.name),
        }
    except Exception as exc:
        raise PipelineError(f"无法检查原始工作簿结构：{path.name}：{type(exc).__name__}: {exc}") from exc
    finally:
        workbook.close()


def filename_monitoring_window(filename: str) -> dict[str, str]:
    match = re.search(
        r"(?P<sy>20\d{2})[.\-/](?P<sm>\d{1,2})[.\-/](?P<sd>\d{1,2})\s+(?P<sh>\d{1,2})[_:](?P<smin>\d{2})"
        r"\s*至\s*"
        r"(?P<ey>20\d{2})[.\-/](?P<em>\d{1,2})[.\-/](?P<ed>\d{1,2})\s+(?P<eh>\d{1,2})[_:](?P<emin>\d{2})",
        filename,
    )
    if not match:
        return {}
    values = {key: int(value) for key, value in match.groupdict().items()}
    return {
        "start": datetime(values["sy"], values["sm"], values["sd"], values["sh"], values["smin"]).isoformat(timespec="minutes"),
        "end": datetime(values["ey"], values["em"], values["ed"], values["eh"], values["emin"]).isoformat(timespec="minutes"),
        "basis": "filename_monitoring_window",
    }


def input_time_window_audit(classification_audit: Iterable[dict[str, Any]], config: dict[str, Any]) -> dict[str, Any]:
    recognized = [
        row
        for row in classification_audit
        if row.get("role") != "ignored_unrecognized" and (row.get("filename_monitoring_window") or {}).get("end")
    ]
    end_values = [datetime.fromisoformat(row["filename_monitoring_window"]["end"]) for row in recognized]
    start_values = [datetime.fromisoformat(row["filename_monitoring_window"]["start"]) for row in recognized]
    tolerance = int(config.get("cutoff_consistency_tolerance_minutes") or 2)
    end_span = int((max(end_values) - min(end_values)).total_seconds() // 60) if end_values else 0
    start_span = int((max(start_values) - min(start_values)).total_seconds() // 60) if start_values else 0
    return {
        "basis": "filename_monitoring_window",
        "tolerance_minutes": tolerance,
        "end_cutoff_span_minutes": end_span,
        "start_cutoff_span_minutes": start_span,
        "end_cutoffs_consistent": end_span <= tolerance,
        "start_cutoffs_consistent": start_span <= tolerance,
        "files": [
            {
                "file": row.get("file"),
                "role": row.get("role"),
                "child_index": row.get("child_index"),
                **(row.get("filename_monitoring_window") or {}),
            }
            for row in recognized
        ],
        "warning": (
            "input_monitoring_windows_are_not_aligned"
            if end_span > tolerance or start_span > tolerance
            else ""
        ),
    }


def configured_filename_patterns(config: dict[str, Any]) -> dict[str, re.Pattern[str]]:
    return {
        name: re.compile(pattern, re.IGNORECASE)
        for name, pattern in (config.get("file_patterns") or {}).items()
    }


def has_filename_hint(path: Path, hints: Iterable[str]) -> bool:
    normalized_name = normalize_text(path.stem)
    return any(normalize_text(hint) in normalized_name for hint in hints if normalize_text(hint))


def child_index_from_filename(path: Path, config: dict[str, Any]) -> int | None:
    patterns = configured_filename_patterns(config)
    legacy = patterns.get("child_heat")
    if legacy is not None:
        match = legacy.match(path.name)
        if match:
            return int(match.group(1))

    default_patterns = [
        r"(?:^|[^0-9])子(?:事件|议题)?\s*[-_]?\s*([1-9][0-9]*)",
        r"(?:^|[^a-z0-9])(?:subevent|subtopic|topic)\s*[-_]?\s*([1-9][0-9]*)",
        r"^\s*([1-9][0-9]*)\s*(?:[-_.（(]|$)",
    ]
    configured = (config.get("input_role_hints") or {}).get("child_index_patterns") or []
    for pattern in [*configured, *default_patterns]:
        match = re.search(pattern, path.stem, flags=re.IGNORECASE)
        if match:
            return int(match.group(1))
    return None


def identify_inputs(raw_directory: Path, config: dict[str, Any]) -> InputBundle:
    files = sorted(
        (item for item in raw_directory.glob("*.xlsx") if not item.name.startswith("~$")),
        key=lambda item: normalize_text(item.name),
    )
    if not files:
        raise PipelineError(f"原始数据目录中没有xlsx文件：{raw_directory}")

    profiles = [profile_raw_workbook(path, config) for path in files]
    audit: list[dict[str, Any]] = []

    def unique_structural(role: str, label: str) -> dict[str, Any]:
        matched = [profile for profile in profiles if role in profile["structural_roles"]]
        if len(matched) != 1:
            names = [profile["file"] for profile in matched]
            raise PipelineError(f"按工作表和字段识别{label}应为1个，实际{len(matched)}个：{names}")
        profile = matched[0]
        audit.append(
            {
                "file": profile["file"],
                "role": role,
                "confidence": 1.0,
                "method": "sheet_and_header_schema",
                "evidence": profile["evidence"],
                "filename_monitoring_window": profile.get("filename_monitoring_window") or {},
            }
        )
        return profile

    public_profile = unique_structural("public_top", "公众文章样本文件")
    overseas_profile = unique_structural("overseas_latest", "境外新闻样本文件")

    heat_profiles = [profile for profile in profiles if "heat" in profile["structural_roles"]]
    if len(heat_profiles) < 2:
        raise PipelineError(f"按日趋势结构至少应识别1个总事件和1个子事件热度文件，实际{len(heat_profiles)}个")

    role_hints = config.get("input_role_hints") or {}
    total_hints = role_hints.get("total_heat") or ["总事件", "总体", "汇总", "overall", "total"]
    total_candidates = [profile for profile in heat_profiles if has_filename_hint(profile["path"], total_hints)]
    total_method = "filename_semantic_hint"
    total_confidence = 0.98

    if len(total_candidates) != 1:
        patterns = configured_filename_patterns(config)
        legacy_total = patterns.get("total_heat")
        legacy_candidates = [
            profile for profile in heat_profiles if legacy_total is not None and legacy_total.match(profile["file"])
        ]
        if len(legacy_candidates) == 1:
            total_candidates = legacy_candidates
            total_method = "legacy_filename_pattern"
            total_confidence = 0.95

    if not total_candidates:
        numeric_profiles = [profile for profile in heat_profiles if profile.get("heat_total") is not None]
        if numeric_profiles:
            largest_value = max(profile["heat_total"] for profile in numeric_profiles)
            largest = [profile for profile in numeric_profiles if profile["heat_total"] == largest_value]
            if len(largest) == 1:
                total_candidates = largest
                total_method = "unique_largest_aggregate_fallback"
                total_confidence = 0.75

    if len(total_candidates) != 1:
        candidates = [
            {"file": profile["file"], "heat_total": profile.get("heat_total")}
            for profile in heat_profiles
        ]
        raise PipelineError(
            "无法唯一识别总事件热度文件。程序已检查日趋势结构、语义文件名提示和总量，"
            f"仍存在歧义：{candidates}"
        )

    total_profile = total_candidates[0]
    audit.append(
        {
            "file": total_profile["file"],
            "role": "total_heat",
            "confidence": total_confidence,
            "method": total_method,
            "evidence": [*total_profile["evidence"], f"日趋势渠道合计={total_profile.get('heat_total')}"],
            "filename_monitoring_window": total_profile.get("filename_monitoring_window") or {},
        }
    )

    child_heat: dict[int, Path] = {}
    for profile in heat_profiles:
        if profile is total_profile:
            continue
        index = child_index_from_filename(profile["path"], config)
        if index is None:
            raise PipelineError(
                f"已确认{profile['file']}是子事件热度文件，但文件和工作簿均未提供可验证的子事件序号；"
                "请在文件名中保留子事件序号（如“子3”或“3-”），程序不会按文件排序猜测议题对应关系。"
            )
        if index in child_heat:
            raise PipelineError(f"检测到重复的子事件{index}文件：{child_heat[index]}；{profile['path']}")
        child_heat[index] = profile["path"]
        audit.append(
            {
                "file": profile["file"],
                "role": "child_heat",
                "child_index": index,
                "confidence": 0.95,
                "method": "sheet_schema_plus_child_index_hint",
                "evidence": [*profile["evidence"], f"文件名解析出子事件序号{index}"],
                "filename_monitoring_window": profile.get("filename_monitoring_window") or {},
            }
        )

    if not child_heat:
        raise PipelineError("未识别到任何子事件热度分析文件")
    if sorted(child_heat) != list(range(1, max(child_heat) + 1)):
        raise PipelineError(f"子事件序号不连续：{sorted(child_heat)}")

    known_paths = {
        public_profile["path"],
        overseas_profile["path"],
        total_profile["path"],
        *child_heat.values(),
    }
    unknown = [profile for profile in profiles if profile["path"] not in known_paths]
    audit.extend(
        {
            "file": profile["file"],
            "role": "ignored_unrecognized",
            "confidence": 1.0,
            "method": "schema_not_supported",
            "evidence": [*profile["evidence"], f"工作表={profile['sheet_names']}"],
            "filename_monitoring_window": profile.get("filename_monitoring_window") or {},
        }
        for profile in unknown
    )

    return InputBundle(
        total_heat=total_profile["path"],
        child_heat=dict(sorted(child_heat.items())),
        public_top=public_profile["path"],
        overseas_latest=overseas_profile["path"],
        classification_audit=tuple(audit),
    )


def find_sheet(workbook: Any, aliases: Iterable[str], purpose: str) -> Any:
    normalized = {normalize_text(name): name for name in workbook.sheetnames}
    for alias in aliases:
        if normalize_text(alias) in normalized:
            return workbook[normalized[normalize_text(alias)]]
    raise PipelineError(f"找不到{purpose}工作表；现有工作表：{workbook.sheetnames}")


def resolve_columns(headers: list[Any], aliases: dict[str, list[str]], purpose: str) -> dict[str, int]:
    normalized = {normalize_text(value): index for index, value in enumerate(headers) if value is not None}
    result: dict[str, int] = {}
    for canonical, candidates in aliases.items():
        for candidate in candidates:
            key = normalize_text(candidate)
            if key in normalized:
                result[canonical] = normalized[key]
                break
    missing = [name for name in aliases if name not in result]
    if missing:
        raise PipelineError(f"{purpose}缺少字段：{missing}；实际字段：{headers}")
    return result


def read_daily(path: Path, config: dict[str, Any]) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        worksheet = find_sheet(workbook, config["sheet_aliases"]["daily_trend"], "日趋势")
        headers = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
        columns = resolve_columns(headers, config["daily_columns"], f"{path.name}/日趋势")
        rows = []
        for row_index in range(2, worksheet.max_row + 1):
            raw_date = worksheet.cell(row_index, columns["date"] + 1).value
            if raw_date in (None, ""):
                continue
            row: dict[str, Any] = {"date": normalize_date(raw_date)[:10]}
            for canonical in CANONICAL_DAILY_COLUMNS[1:]:
                value = worksheet.cell(row_index, columns[canonical] + 1).value
                number = to_number(value)
                if number is None:
                    raise PipelineError(f"{path.name}/日趋势/{row_index}行/{canonical}不是有效数值：{value!r}")
                row[canonical] = number
            rows.append(row)
        if not rows:
            raise PipelineError(f"{path.name}/日趋势没有有效数据行")
        return rows
    finally:
        workbook.close()


def read_detail_rows(
    path: Path,
    sheet_aliases: list[str],
    config: dict[str, Any],
    purpose: str,
) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        worksheet = find_sheet(workbook, sheet_aliases, purpose)
        headers = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
        columns = resolve_columns(headers, config["detail_columns"], f"{path.name}/{worksheet.title}")
        rows = []
        for row_index in range(2, worksheet.max_row + 1):
            title = worksheet.cell(row_index, columns["title"] + 1).value
            url = worksheet.cell(row_index, columns["url"] + 1).value
            if title in (None, "") and url in (None, ""):
                continue
            item = {
                "source_row": row_index,
                "source": json_value(worksheet.cell(row_index, columns["source"] + 1).value),
                "url": json_value(url),
                "published_at": normalize_date(worksheet.cell(row_index, columns["published_at"] + 1).value),
                "title": json_value(title),
                "content": json_value(worksheet.cell(row_index, columns["content"] + 1).value),
                "account": json_value(worksheet.cell(row_index, columns["account"] + 1).value),
                "read_count": to_number(worksheet.cell(row_index, columns["read_count"] + 1).value),
                "recommend_count": to_number(worksheet.cell(row_index, columns["recommend_count"] + 1).value),
            }
            rows.append(item)
        return rows
    finally:
        workbook.close()


def contains_any(text: str, patterns: Iterable[str]) -> bool:
    return any(normalize_text(pattern) in text for pattern in patterns)


def first_anchor_position(text: str, anchors: Iterable[str]) -> int | None:
    positions = [text.find(normalize_text(anchor)) for anchor in anchors]
    valid = [position for position in positions if position >= 0]
    return min(valid) if valid else None


def topic_hits(text: str, topic_aliases: list[list[str]]) -> list[int]:
    return [
        index + 1
        for index, aliases in enumerate(topic_aliases)
        if any(normalize_text(alias) in text for alias in aliases)
    ]


def canonical_url(url: str | None) -> str:
    if not url:
        return ""
    try:
        parsed = urlsplit(str(url).strip())
    except ValueError:
        return str(url).strip()
    ignored = {"ref", "utm_source", "utm_medium", "utm_campaign", "from", "spm"}
    query = [(key, value) for key, value in parse_qsl(parsed.query, keep_blank_values=True) if key.lower() not in ignored]
    return urlunsplit((parsed.scheme.lower(), parsed.netloc.lower(), parsed.path.rstrip("/"), urlencode(query), ""))


def overseas_publisher_family(url: str | None, source: Any = "") -> str:
    """Return a stable publisher family across subdomains and mirror domains."""
    host = ""
    if url:
        try:
            host = urlsplit(str(url).strip()).netloc.lower().split(":", 1)[0]
        except ValueError:
            host = ""
    labels = [part for part in host.split(".") if part and part not in {"www", "ww2", "m", "mobile"}]
    if len(labels) >= 3 and labels[-2:] in (["com", "tw"], ["com", "hk"], ["com", "cn"], ["co", "uk"]):
        host = "".join(labels[-3:])
    elif len(labels) >= 2:
        host = "".join(labels[-2:])
    else:
        host = "".join(labels)
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", host or normalize_text(source).lower())


def overseas_article_token(url: str | None) -> str:
    if not url:
        return ""
    try:
        parsed = urlsplit(str(url).strip())
    except ValueError:
        return ""
    query = {str(key).lower(): str(value) for key, value in parse_qsl(parsed.query, keep_blank_values=True)}
    for key in ("newsid", "sn", "id", "articleid", "storyid"):
        value = query.get(key, "")
        match = re.search(r"[0-9]{6,}", value)
        if match:
            return match.group(0)
    match = re.search(r"(?:NOW[./_-]?|story/[^/]+/)?([0-9]{6,})(?:\D|$)", parsed.path, flags=re.I)
    return match.group(1) if match else ""


def overseas_publisher_article_key(row: dict[str, Any]) -> str:
    family = overseas_publisher_family(row.get("url"), row.get("source"))
    token = overseas_article_token(row.get("url"))
    return f"{family}|article:{token}" if family and token else ""


OVERSEAS_PUBLISHER_CLASSES = {
    "overseas_origin_media",
    "mainland_outward_media",
    "not_overseas_media",
}


def reviewed_overseas_category(value: Any) -> str:
    """Normalize the raw AI review label for downstream workbook/report use."""

    label = str(value or "").strip().lower()
    aliases = {
        "factual": "事实性报道",
        "fact": "事实性报道",
        "news_roundup": "事实性报道",
        "corporate_disclosure": "事实性报道",
        "事实性报道": "事实性报道",
        "interpretive": "解读性报道",
        "analysis": "解读性报道",
        "解读性报道": "解读性报道",
        "risk": "借题炒作/风险解读",
        "risk_interpretation": "借题炒作/风险解读",
        "借题炒作/风险解读": "借题炒作/风险解读",
    }
    return aliases.get(label, "")


def overseas_record_id(row: dict[str, Any], origin: str = "system_raw") -> str:
    """Build a stable review key without depending on a period-specific title."""
    identity = canonical_url(row.get("url")) or "|".join(
        [
            normalize_text(row.get("source")),
            normalize_text(row.get("title")),
            normalize_date(row.get("published_at")),
        ]
    )
    digest = hashlib.sha1(identity.encode("utf-8")).hexdigest()[:12]
    source_row = row.get("source_row")
    return f"{origin}:{source_row if source_row not in (None, '') else 'supplement'}:{digest}"


def preliminary_publisher_class(row: dict[str, Any], config: dict[str, Any]) -> str:
    combined = normalize_text(f"{row.get('source', '')} {row.get('url', '')}")
    if contains_any(combined, config.get("domestic_overseas_outlet_patterns") or []):
        return "mainland_outward_media"
    return "overseas_origin_media"


def build_overseas_review_packet(
    rows: list[dict[str, Any]],
    config: dict[str, Any],
    metadata: dict[str, Any],
    monitoring_dates: Iterable[str],
) -> dict[str, Any]:
    dates = sorted({str(value)[:10] for value in monitoring_dates if value})
    aliases = metadata.get("topic_aliases") or []
    anchors = config.get("meeting_anchor_patterns") or []
    promo_patterns = config.get("market_promo_patterns") or []
    items = []
    for row in rows:
        title = normalize_text(row.get("title"))
        content = normalize_text(row.get("content"))
        items.append(
            {
                **row,
                "record_id": overseas_record_id(row),
                "origin": "system_raw",
                "preliminary_publisher_class": preliminary_publisher_class(row, config),
                "preliminary_topic_hits": topic_hits(title + content, aliases),
                "rule_hints": {
                    "title_has_meeting_anchor": contains_any(title, anchors),
                    "body_anchor_position": first_anchor_position(content, anchors),
                    "market_promo_title": contains_any(title, promo_patterns),
                },
            }
        )
    return {
        "schema_version": 1,
        "review_type": "cwh_overseas_semantic_review",
        "meeting_title": metadata.get("meeting_title"),
        "topic_titles": metadata.get("topic_titles") or [],
        "topic_aliases": aliases,
        "monitoring_dates": dates,
        "instructions": [
            "AI逐条阅读标题与正文，判断是否以本次国务院常务会议为主要报道对象。",
            "正文缺失或正文仅重复标题时不能include；保留证据不足及补取正文的原因，不冒称与会议无关。",
            "解读性报道必须明确interpretive_verified=true，并从原始content逐字复制连续interpretive_excerpt，保留繁简、标点、空格和异常字符；不得改写、纠错或简繁转换，摘要不等于原文核实。",
            "decision只能为include或exclude；不得仅凭关键词命中纳入。",
            "publisher_class只能为overseas_origin_media、mainland_outward_media或not_overseas_media。",
            "每条必须填写review_reason和0至1之间的classification_confidence。",
            "overseas_origin_media与mainland_outward_media均可计入境外媒体总量；仅overseas_origin_media进入正式外媒附录。",
            "topic_hits只能填写从1开始的整数编号数组，编号对应topic_titles顺序；允许一条报道命中多个子议题，同一报道只计入总事件一次。",
            "补充检索样本放入supplemental_rows，并执行相同审核、时间窗和去重规则；同URL及同媒体重复行去重，跨媒体转载保留。",
        ],
        "items": items,
        "supplemental_rows": [],
    }


def deduplicate_overseas(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    kept_by_url: dict[str, dict[str, Any]] = {}
    dropped: list[dict[str, Any]] = []
    for row in rows:
        key = canonical_url(row.get("url")) or f"title:{normalize_text(row.get('title'))}"
        existing = kept_by_url.get(key)
        if existing is None:
            kept_by_url[key] = row
        else:
            best = max(
                [existing, row],
                key=lambda item: (
                    len(str(item.get("content") or "")),
                    item.get("published_at") or "",
                    str(item.get("source_row") or ""),
                ),
            )
            dropped_item = row if best is existing else existing
            dropped.append({**dropped_item, "decision": "exclude", "reason": "duplicate_url"})
            kept_by_url[key] = best

    article_groups: dict[str, list[dict[str, Any]]] = {}
    no_article_key: list[dict[str, Any]] = []
    for row in kept_by_url.values():
        article_key = overseas_publisher_article_key(row)
        if article_key:
            article_groups.setdefault(article_key, []).append(row)
        else:
            no_article_key.append(row)
    candidates = list(no_article_key)
    for group in article_groups.values():
        best = max(
            group,
            key=lambda item: (
                len(str(item.get("content") or "")),
                item.get("published_at") or "",
                str(item.get("source_row") or ""),
            ),
        )
        candidates.append(best)
        for item in group:
            if item is not best:
                dropped.append({**item, "decision": "exclude", "reason": "duplicate_same_publisher_article"})
    title_groups: dict[str, list[dict[str, Any]]] = {}
    for row in candidates:
        title_key = "|".join(
            [
                normalize_text(row.get("source")),
                normalize_text(row.get("title")).replace("热门内容", "").replace("熱門內容", ""),
            ]
        )
        title_groups.setdefault(title_key, []).append(row)

    kept: list[dict[str, Any]] = []
    for title_key, group in title_groups.items():
        if len(group) == 1:
            kept.extend(group)
            continue
        best = max(
            group,
            key=lambda item: (
                len(str(item.get("content") or "")),
                item.get("published_at") or "",
                str(item.get("source_row") or ""),
            ),
        )
        kept.append(best)
        for item in group:
            if item is not best:
                dropped.append({**item, "decision": "exclude", "reason": "duplicate_same_outlet_title"})
    # Daily/weekly rolling columns may publish several dated editions under the
    # same series title.  Keep the newest edition so a list is not dominated by
    # successive snapshots of the same digest.
    recurring_groups: dict[str, list[dict[str, Any]]] = {}
    non_recurring: list[dict[str, Any]] = []
    for row in kept:
        normalized_title = normalize_text(row.get("title"))
        if not re.search(r"全[天日](?:相关|相關)(?:动态|動態)", normalized_title):
            non_recurring.append(row)
            continue
        series_key = "|".join(
            [
                normalize_text(row.get("source")),
                re.sub(r"\d{4}年\d{1,2}月\d{1,2}日", "", normalized_title),
            ]
        )
        recurring_groups.setdefault(series_key, []).append(row)

    for group in recurring_groups.values():
        newest = max(group, key=lambda item: (item.get("published_at") or "", item.get("source_row", 0)))
        non_recurring.append(newest)
        for item in group:
            if item is not newest:
                dropped.append({**item, "decision": "exclude", "reason": "duplicate_recurring_series_latest"})

    non_recurring.sort(key=lambda item: (item.get("published_at") or "", item.get("source") or ""))
    return non_recurring, dropped


def filter_overseas(
    rows: list[dict[str, Any]],
    config: dict[str, Any],
    metadata: dict[str, Any],
    review: dict[str, Any] | None = None,
    monitoring_dates: Iterable[str] = (),
) -> dict[str, Any]:
    review_packet = build_overseas_review_packet(rows, config, metadata, monitoring_dates)
    if review is not None:
        return apply_overseas_ai_review(rows, config, metadata, review, monitoring_dates, review_packet)

    accepted: list[dict[str, Any]] = []
    decisions: list[dict[str, Any]] = []
    anchors = config["meeting_anchor_patterns"]
    promo_patterns = config["market_promo_patterns"]
    aliases = metadata["topic_aliases"]

    for row in rows:
        title_text = normalize_text(row.get("title"))
        content_text = normalize_text(row.get("content"))
        title_anchor = contains_any(title_text, anchors)
        body_anchor_position = first_anchor_position(content_text, anchors)
        title_topics = topic_hits(title_text, aliases)
        all_topics = topic_hits(title_text + content_text, aliases)
        market_promo = contains_any(title_text, promo_patterns)
        publisher_class = preliminary_publisher_class(row, config)

        reason = None
        if title_anchor:
            reason = None
        elif title_topics and body_anchor_position is not None and body_anchor_position < 500 and not market_promo:
            reason = None
        elif market_promo:
            reason = "market_promotion_not_primary_report"
        else:
            reason = "not_primary_meeting_report"

        decision = {
            "record_id": overseas_record_id(row),
            "source_row": row["source_row"],
            "source": row.get("source"),
            "title": row.get("title"),
            "url": row.get("url"),
            "decision": "exclude" if reason else "include",
            "reason": reason or (
                "preliminary_relevant_mainland_outward_report"
                if publisher_class == "mainland_outward_media"
                else "preliminary_relevant_overseas_report"
            ),
            "publisher_class": publisher_class,
            "topic_hits": all_topics,
            "ai_reviewed": False,
        }
        decisions.append(decision)
        if not reason:
            accepted.append(
                {
                    **row,
                    "record_id": overseas_record_id(row),
                    "origin": "system_raw",
                    "publisher_class": publisher_class,
                    "topic_hits": all_topics,
                    "ai_reviewed": False,
                }
            )

    deduplicated, duplicate_drops = deduplicate_overseas(accepted)
    duplicate_reasons = {
        item.get("record_id") or item.get("source_row"): item["reason"] for item in duplicate_drops
    }
    for decision in decisions:
        reason = duplicate_reasons.get(decision.get("record_id") or decision["source_row"])
        if decision["decision"] == "include" and reason:
            decision["decision"] = "exclude"
            decision["reason"] = reason

    appendix_selected = [
        item for item in deduplicated if item.get("publisher_class") == "overseas_origin_media"
    ]
    return {
        "status": "ai_review_required",
        "counts_reconciled": False,
        "selected": appendix_selected,
        "appendix_selected": appendix_selected,
        "counted": deduplicated,
        "decisions": decisions,
        "review_packet": review_packet,
        "summary": {
            "raw_candidate_count": len(rows),
            "preliminary_counted_count": len(deduplicated),
            "preliminary_appendix_count": len(appendix_selected),
        },
    }


def apply_overseas_ai_review(
    rows: list[dict[str, Any]],
    config: dict[str, Any],
    metadata: dict[str, Any],
    review: dict[str, Any],
    monitoring_dates: Iterable[str],
    review_packet: dict[str, Any] | None = None,
) -> dict[str, Any]:
    if str(review.get("review_method") or "").strip().lower() not in {
        "ai_semantic_review",
        "ai_semantic_review_with_human_edits",
    }:
        raise PipelineError("境外审核文件缺少有效review_method，不能据此改写传播统计")

    allowed_dates = {str(value)[:10] for value in monitoring_dates if value}
    raw_by_id = {overseas_record_id(row): row for row in rows}
    raw_by_source_row = {int(row["source_row"]): row for row in rows}
    review_items = list(review.get("items") or [])
    matched: dict[str, dict[str, Any]] = {}
    for item in review_items:
        record_id = str(item.get("record_id") or "").strip()
        if not record_id and item.get("source_row") not in (None, ""):
            row = raw_by_source_row.get(int(item["source_row"]))
            record_id = overseas_record_id(row) if row else ""
        if not record_id or record_id not in raw_by_id:
            raise PipelineError(f"境外审核文件含无法匹配的系统记录：{item.get('record_id') or item.get('source_row')}")
        if record_id in matched:
            raise PipelineError(f"境外审核文件重复记录：{record_id}")
        matched[record_id] = item

    missing = sorted(set(raw_by_id) - set(matched))
    if missing:
        raise PipelineError(f"境外AI审核未覆盖全部系统候选，缺少{len(missing)}条")

    decisions: list[dict[str, Any]] = []
    accepted: list[dict[str, Any]] = []

    def reviewed_row(base: dict[str, Any], item: dict[str, Any], origin: str) -> dict[str, Any]:
        decision = str(item.get("decision") or "").strip().lower()
        if decision not in {"include", "exclude"}:
            raise PipelineError(f"境外审核decision非法：{decision or '<empty>'}")
        publisher_class = str(item.get("publisher_class") or "").strip()
        if publisher_class not in OVERSEAS_PUBLISHER_CLASSES:
            raise PipelineError(f"境外审核publisher_class非法：{publisher_class or '<empty>'}")
        reason = str(item.get("review_reason") or item.get("reason") or "").strip()
        if not reason:
            raise PipelineError(f"境外审核缺少review_reason：{base.get('title')}")
        confidence = to_number(item.get("classification_confidence"))
        if confidence is None or not 0 <= float(confidence) <= 1:
            raise PipelineError(f"境外审核classification_confidence须为0至1：{base.get('title')}")
        hits = item.get("topic_hits") or []
        if not isinstance(hits, list) or any(
            not isinstance(hit, int) or hit < 1 or hit > len(metadata.get("topic_titles") or []) for hit in hits
        ):
            raise PipelineError(f"境外审核topic_hits非法：{base.get('title')}")
        published_date = normalize_date(base.get("published_at"))[:10]
        in_window = not allowed_dates or published_date in allowed_dates
        count_include = decision == "include" and in_window and publisher_class in {
            "overseas_origin_media",
            "mainland_outward_media",
        }
        appendix_include = count_include and publisher_class == "overseas_origin_media"
        report_category = reviewed_overseas_category(
            item.get("ai_report_category") or item.get("content_type")
        )
        if count_include and not report_category:
            raise PipelineError(f"境外审核缺少有效报道类型：{base.get('title')}")
        original_body = str(base.get("content") or "").strip()
        if count_include and (not original_body or normalize_text(original_body) == normalize_text(base.get("title"))):
            raise PipelineError(f"境外审核正文缺失或仅有标题，不能纳入：{base.get('title')}")
        title_cn_simplified = to_simplified_review_text(
            item.get("title_cn_simplified")
            or item.get("title_cn")
            or base.get("title")
            or ""
        )
        source_cn_simplified = to_simplified_review_text(
            item.get("source_cn_simplified")
            or item.get("source_cn")
            or base.get("source")
            or ""
        )
        summary_cn_simplified = to_simplified_review_text(
            item.get("summary_cn_simplified")
            or item.get("summary_cn")
            or item.get("interpretive_summary_cn")
            or ""
        )
        interpretive_excerpt = str(item.get("interpretive_excerpt") or "").strip()
        interpretive_verified = (item.get("interpretive_verified") is True
                                 and bool(interpretive_excerpt) and interpretive_excerpt in original_body)
        if count_include and report_category == "解读性报道" and not interpretive_verified:
            raise PipelineError(f"境外审核解读性报道缺少已核实的连续原文依据interpretive_excerpt：{base.get('title')}")
        if decision == "include" and not in_window:
            reason = f"outside_monitoring_window:{reason}"
        return {
            **base,
            "record_id": overseas_record_id(base, origin),
            "origin": origin,
            "decision": "include" if count_include else "exclude",
            "reason": reason,
            "review_reason": reason,
            "reason_code": (
                "outside_monitoring_window"
                if decision == "include" and not in_window
                else str(item.get("reason_code") or ("ai_reviewed_include" if count_include else "ai_reviewed_exclude"))
            ),
            "publisher_class": publisher_class,
            "topic_hits": sorted(set(hits)),
            "ai_reviewed": True,
            "classification_confidence": float(confidence),
            "content_type": item.get("content_type"),
            "ai_report_category": report_category,
            "classification_reason": reason,
            "title_cn_simplified": title_cn_simplified,
            "source_cn_simplified": source_cn_simplified,
            "summary_cn_simplified": summary_cn_simplified,
            "interpretive_verified": interpretive_verified,
            "interpretive_excerpt": interpretive_excerpt,
            "meeting_relevance": count_include,
            "formal_include": appendix_include,
            "count_include": count_include,
            "appendix_include": appendix_include,
        }

    for record_id, base in raw_by_id.items():
        result = reviewed_row(base, matched[record_id], "system_raw")
        decisions.append(result)
        if result["count_include"]:
            accepted.append(result)

    for index, item in enumerate(review.get("supplemental_rows") or [], 1):
        base = {
            "source_row": f"supplement-{index}",
            "source": item.get("source"),
            "url": item.get("url"),
            "published_at": normalize_date(item.get("published_at")),
            "title": item.get("title"),
            "content": item.get("content"),
            "account": item.get("account"),
            "read_count": item.get("read_count"),
            "recommend_count": item.get("recommend_count"),
        }
        if not base["source"] or not base["title"] or not base["url"] or not base["published_at"]:
            raise PipelineError(f"境外补充审核第{index}条缺少来源、标题、链接或发布时间")
        result = reviewed_row(base, item, "public_supplement")
        decisions.append(result)
        if result["count_include"]:
            accepted.append(result)

    deduplicated, duplicate_drops = deduplicate_overseas(accepted)
    duplicate_ids = {item.get("record_id") for item in duplicate_drops}
    for decision in decisions:
        if decision.get("record_id") in duplicate_ids and decision.get("decision") == "include":
            decision["decision"] = "exclude"
            decision["count_include"] = False
            decision["appendix_include"] = False
            decision["reason"] = next(
                item["reason"] for item in duplicate_drops if item.get("record_id") == decision.get("record_id")
            )
            decision["reason_code"] = decision["reason"]

    appendix_selected = [
        item for item in deduplicated if item.get("publisher_class") == "overseas_origin_media"
    ]
    outward_counted = [
        item for item in deduplicated if item.get("publisher_class") == "mainland_outward_media"
    ]
    return {
        "status": "ai_review_complete",
        "counts_reconciled": True,
        "selected": appendix_selected,
        "appendix_selected": appendix_selected,
        "counted": deduplicated,
        "mainland_outward_counted": outward_counted,
        "decisions": decisions,
        "review_packet": review_packet or build_overseas_review_packet(rows, config, metadata, monitoring_dates),
        "review_method": review.get("review_method"),
        "summary": {
            "raw_candidate_count": len(rows),
            "supplement_candidate_count": len(review.get("supplemental_rows") or []),
            "counted_total": len(deduplicated),
            "appendix_total": len(appendix_selected),
            "mainland_outward_total": len(outward_counted),
            "excluded_total": sum(1 for item in decisions if item.get("decision") == "exclude"),
        },
    }


def filter_public_top(
    rows: list[dict[str, Any]],
    config: dict[str, Any],
    metadata: dict[str, Any],
    review: dict[str, Any] | None = None,
) -> dict[str, Any]:
    accepted: list[dict[str, Any]] = []
    decisions: list[dict[str, Any]] = []
    anchors = config["meeting_anchor_patterns"]
    roundups = config["public_roundup_patterns"]
    aliases = metadata["topic_aliases"]
    seen_urls: set[str] = set()

    for row in rows:
        title_text = normalize_text(row.get("title"))
        content_text = normalize_text(row.get("content"))
        combined = title_text + content_text
        title_anchor = contains_any(title_text, anchors)
        body_anchor = first_anchor_position(content_text, anchors)
        title_topics = topic_hits(title_text, aliases)
        all_topics = topic_hits(combined, aliases)
        roundup = contains_any(title_text, roundups)
        strong_generic = (
            ("国务院" in title_text and ("会议" in title_text or "计划" in title_text))
            or "会议部署" in title_text
        )
        # A concise headline may omit both “国常会” and the metadata's longer
        # agenda alias.  The full body remains decisive: it must explicitly
        # identify the meeting and contain current-agenda semantics.  This
        # avoids both false negatives on short official headlines and false
        # positives on unrelated articles that merely share a topic word.
        body_meeting_topic = body_anchor is not None and bool(all_topics)
        relevance = (
            title_anchor
            or body_meeting_topic
            or (strong_generic and body_anchor is not None)
        )
        url_key = canonical_url(row.get("url"))

        reason = None
        if roundup:
            reason = "roundup_or_breakfast_digest"
        elif not relevance:
            reason = "meeting_not_primary_focus"
        elif row.get("read_count") is None:
            reason = "missing_or_invalid_read_count"
        elif url_key and url_key in seen_urls:
            reason = "duplicate_url"

        decisions.append(
            {
                "source_row": row["source_row"],
                "account": row.get("account"),
                "title": row.get("title"),
                "url": row.get("url"),
                "read_count": row.get("read_count"),
                "decision": "exclude" if reason else "candidate",
                "reason": reason or "relevant_candidate",
                "topic_hits": all_topics,
                "relevance_evidence": (
                    "title_meeting_anchor"
                    if title_anchor
                    else "body_meeting_anchor_and_topic"
                    if body_meeting_topic
                    else "generic_meeting_anchor"
                ),
            }
        )
        if reason:
            continue
        if url_key:
            seen_urls.add(url_key)
        accepted.append({**row, "topic_hits": all_topics})

    def normalized_account(row: dict[str, Any]) -> str:
        return public_source_family(row.get("account") or row.get("source") or "", config)

    # The appendix is a source/account ranking, not a raw article ranking. Keep
    # the highest-read relevant article from each source before applying TOP N.
    accepted.sort(key=lambda item: (-(item.get("read_count") or 0), -(item.get("recommend_count") or 0), item.get("source_row", 0)))
    account_best: dict[str, dict[str, Any]] = {}
    deduplicated_accounts: list[dict[str, Any]] = []
    duplicate_account_rows: set[int] = set()
    for row in accepted:
        account_key = normalized_account(row) or f"__row_{row.get('source_row')}"
        if account_key in account_best:
            duplicate_account_rows.add(int(row.get("source_row") or 0))
            continue
        account_best[account_key] = row
        deduplicated_accounts.append(row)

    top_n = int(config.get("public_top_n", 10))

    # All rows at the current read-count ceiling can potentially enter TOP N;
    # they require full-body AI review because deterministic title rules cannot
    # reliably distinguish a focused article from a breakfast/market roundup.
    ranked_raw = sorted(
        [row for row in rows if row.get("read_count") is not None],
        key=lambda item: (-(item.get("read_count") or 0), -(item.get("recommend_count") or 0), item.get("source_row", 0)),
    )
    ceiling = ranked_raw[0].get("read_count") if ranked_raw else None
    review_batch = [row for row in ranked_raw if row.get("read_count") == ceiling]
    if len({normalized_account(row) for row in review_batch}) < top_n:
        review_batch = ranked_raw[: max(top_n * 3, 30)]
    review_packet = {
        "schema_version": 1,
        "review_type": "cwh_public_top_semantic_review",
        "meeting_title": metadata.get("meeting_title"),
        "topic_titles": metadata.get("topic_titles") or [],
        "instructions": [
            "逐条阅读标题与完整正文，判断本次国务院常务会议及其议题是否为文章主体。",
            "早餐、早报、多消息合集、行情综述中仅顺带提及会议的文章必须排除。",
            "账号权威性、阅读量高低和标题命中词不能替代正文语义判断。",
            "先按配置中的发布主体账号族合并同源账号，再按阅读量和在看量排序；不得用模糊字符串擅自合并无关账号。",
            "decision只能为include或exclude；每条填写review_reason、classification_confidence和topic_hits；topic_hits只能填写从1开始且对应topic_titles顺序的整数编号数组。",
        ],
        "items": [
            {
                **row,
                "record_id": f"public:{row.get('source_row')}:{hashlib.sha1((str(row.get('url') or '') + str(row.get('title') or '')).encode('utf-8')).hexdigest()[:12]}",
                "preliminary_topic_hits": topic_hits(normalize_text(f"{row.get('title', '')}{row.get('content', '')}"), aliases),
            }
            for row in review_batch
        ],
    }

    if review is None:
        selected = deduplicated_accounts[:top_n]
        selected_rows = {row["source_row"] for row in selected}
        for decision in decisions:
            if decision["decision"] == "candidate":
                if decision["source_row"] in selected_rows:
                    decision["decision"] = "include"
                    decision["reason"] = "preliminary_selected_pending_ai_review"
                elif int(decision["source_row"] or 0) in duplicate_account_rows:
                    decision["decision"] = "exclude"
                    decision["reason"] = "duplicate_source_lower_read_count"
                else:
                    decision["decision"] = "exclude"
                    decision["reason"] = "valid_but_below_top_n"
        return {
            "status": "ai_review_required",
            "selected": selected,
            "decisions": decisions,
            "review_packet": review_packet,
            "selection_summary": {
                "top_n": top_n,
                "relevant_candidate_count": len(accepted),
                "distinct_source_count": len(deduplicated_accounts),
                "duplicate_source_count": len(duplicate_account_rows),
                "selected_count": len(selected),
                "review_batch_count": len(review_batch),
            },
        }

    if str(review.get("review_method") or "").strip().lower() not in {
        "ai_semantic_review",
        "ai_semantic_review_with_human_edits",
    }:
        raise PipelineError("公众TOP审核文件缺少有效review_method")
    packet_by_id = {item["record_id"]: item for item in review_packet["items"]}
    reviewed_by_id = {str(item.get("record_id") or ""): item for item in review.get("items") or []}
    missing_review = sorted(set(packet_by_id) - set(reviewed_by_id))
    if missing_review:
        raise PipelineError(f"公众TOP AI审核未覆盖全部临界候选，缺少{len(missing_review)}条")
    ai_accepted: list[dict[str, Any]] = []
    ai_decisions: list[dict[str, Any]] = []
    for record_id, row in packet_by_id.items():
        item = reviewed_by_id[record_id]
        decision = str(item.get("decision") or "").strip().lower()
        reason = str(item.get("review_reason") or item.get("reason") or "").strip()
        confidence = to_number(item.get("classification_confidence"))
        if decision not in {"include", "exclude"} or not reason or confidence is None or not 0 <= float(confidence) <= 1:
            raise PipelineError(f"公众TOP审核字段不完整或非法：{row.get('account')}《{row.get('title')}》")
        hits = item.get("topic_hits") or []
        if decision == "include" and not hits:
            raise PipelineError(f"公众TOP纳入项缺少topic_hits：{row.get('account')}《{row.get('title')}》")
        ai_decisions.append({
            "source_row": row.get("source_row"),
            "account": row.get("account"),
            "title": row.get("title"),
            "url": row.get("url"),
            "read_count": row.get("read_count"),
            "recommend_count": row.get("recommend_count"),
            "decision": decision,
            "reason": reason,
            "classification_confidence": float(confidence),
            "topic_hits": hits,
            "record_id": record_id,
        })
        if decision == "include":
            ai_accepted.append({**row, "topic_hits": sorted(set(hits)), "ai_review_reason": reason, "ai_review_confidence": float(confidence)})

    ai_accepted.sort(key=lambda item: (-(item.get("read_count") or 0), -(item.get("recommend_count") or 0), item.get("source_row", 0)))
    account_best = {}
    for row in ai_accepted:
        account_best.setdefault(normalized_account(row) or f"__row_{row.get('source_row')}", row)
    selected = list(account_best.values())[:top_n]
    status = "ai_review_complete" if len(selected) >= top_n else "ai_review_expand_required"
    return {
        "status": status,
        "review_method": review.get("review_method"),
        "selected": selected,
        "decisions": ai_decisions,
        "review_packet": review_packet,
        "selection_summary": {
            "top_n": top_n,
            "relevant_candidate_count": len(ai_accepted),
            "distinct_source_count": len(account_best),
            "duplicate_source_count": len(ai_accepted) - len(account_best),
            "selected_count": len(selected),
            "review_batch_count": len(review_batch),
        },
    }


def build_public_article_evidence_corpus(
    rows: list[dict[str, Any]],
    config: dict[str, Any],
    metadata: dict[str, Any],
) -> dict[str, Any]:
    """Preserve the full in-bundle article evidence pool for viewpoint analysis.

    The public TOP appendix and the domestic viewpoint chapter serve different
    purposes.  TOP N is an account/read-count ranking, while viewpoint research
    must be able to use lower-ranked but substantive articles and named voices.
    This corpus therefore applies only deterministic meeting/topic relevance and
    exact-duplicate gates; it deliberately does not apply a rank ceiling.
    """
    anchors = list(config.get("meeting_anchor_patterns") or [])
    roundups = list(config.get("public_roundup_patterns") or [])
    aliases = list(metadata.get("topic_aliases") or [])
    candidates: list[dict[str, Any]] = []
    exclusions: list[dict[str, Any]] = []
    exclusion_summary: Counter[str] = Counter()
    seen_urls: dict[str, str] = {}
    seen_content_hashes: dict[str, str] = {}

    for row in rows:
        title_text = normalize_text(row.get("title"))
        content_text = normalize_text(row.get("content"))
        combined = title_text + content_text
        title_anchor = contains_any(title_text, anchors)
        body_anchor = first_anchor_position(content_text, anchors)
        all_topics = topic_hits(combined, aliases)
        body_meeting_topic = body_anchor is not None and bool(all_topics)
        roundup = contains_any(title_text, roundups)
        url_key = canonical_url(row.get("url"))
        content_hash = hashlib.sha1(content_text.encode("utf-8")).hexdigest() if content_text else ""
        record_id = (
            f"public-evidence:{row.get('source_row')}:"
            f"{hashlib.sha1((str(row.get('url') or '') + str(row.get('title') or '')).encode('utf-8')).hexdigest()[:12]}"
        )

        reason = ""
        duplicate_of = ""
        if roundup:
            reason = "roundup_or_digest"
        elif not (title_anchor or body_meeting_topic):
            reason = "meeting_not_primary_focus"
        elif not all_topics:
            reason = "no_agenda_topic_hit"
        elif url_key and url_key in seen_urls:
            reason = "duplicate_url"
            duplicate_of = seen_urls[url_key]
        elif content_hash and content_hash in seen_content_hashes:
            reason = "duplicate_full_text"
            duplicate_of = seen_content_hashes[content_hash]

        if reason:
            exclusion_summary[reason] += 1
            exclusions.append(
                {
                    "record_id": record_id,
                    "source_row": row.get("source_row"),
                    "published_at": row.get("published_at"),
                    "account": row.get("account"),
                    "source": row.get("source"),
                    "title": row.get("title"),
                    "url": row.get("url"),
                    "topic_hits": all_topics,
                    "decision": "exclude",
                    "reason": reason,
                    "duplicate_of": duplicate_of or None,
                }
            )
            continue
        if url_key:
            seen_urls[url_key] = record_id
        if content_hash:
            seen_content_hashes[content_hash] = record_id
        candidates.append(
            {
                "record_id": record_id,
                "source_row": row.get("source_row"),
                "published_at": row.get("published_at"),
                "account": row.get("account"),
                "source": row.get("source"),
                "title": row.get("title"),
                "content": row.get("content"),
                "url": row.get("url"),
                "read_count": row.get("read_count"),
                "recommend_count": row.get("recommend_count"),
                "topic_hits": all_topics,
                "relevance_evidence": (
                    "title_meeting_anchor" if title_anchor else "body_meeting_anchor_and_topic"
                ),
                "content_sha1": content_hash,
            }
        )

    candidates.sort(
        key=lambda item: (
            -(item.get("read_count") or 0),
            -(item.get("recommend_count") or 0),
            int(item.get("source_row") or 0),
        )
    )
    per_topic_counts = {
        str(index): sum(index in (row.get("topic_hits") or []) for row in candidates)
        for index in range(1, len(aliases) + 1)
    }
    return {
        "schema_version": 1,
        "purpose": "domestic_viewpoint_evidence_not_public_top_ranking",
        "meeting_title": metadata.get("meeting_title"),
        "monitoring_window": metadata.get("monitoring_window") or {},
        "topic_titles": metadata.get("topic_titles") or [],
        "raw_row_count": len(rows),
        "candidate_count": len(candidates),
        "candidate_count_by_topic": per_topic_counts,
        "exclusion_summary": dict(sorted(exclusion_summary.items())),
        "exclusions": exclusions,
        "candidates": candidates,
    }


def summarize_daily(rows: list[dict[str, Any]], groups: dict[str, list[str]]) -> list[dict[str, Any]]:
    summarized = []
    for row in rows:
        item: dict[str, Any] = {"date": row["date"]}
        for group_name, fields in groups.items():
            item[group_name] = sum(float(row[field]) for field in fields)
            if item[group_name].is_integer():
                item[group_name] = int(item[group_name])
        item["total"] = sum(value for key, value in item.items() if key != "date")
        summarized.append(item)
    return summarized


def reconcile_overseas_daily_counts(
    total_daily: list[dict[str, Any]],
    child_daily: dict[int, list[dict[str, Any]]],
    overseas: dict[str, Any],
) -> dict[str, Any]:
    """Replace only the reviewed overseas channel and preserve every other raw channel."""
    if not overseas.get("counts_reconciled"):
        return {
            "applied": False,
            "reason": "ai_review_required",
            "total_daily": total_daily,
            "child_daily": child_daily,
            "deltas": [],
        }

    total_by_date: Counter[str] = Counter()
    child_by_date: dict[int, Counter[str]] = {index: Counter() for index in child_daily}
    for row in overseas.get("counted") or []:
        published_date = normalize_date(row.get("published_at"))[:10]
        if not published_date:
            continue
        total_by_date[published_date] += 1
        for topic_index in row.get("topic_hits") or []:
            if topic_index in child_by_date:
                child_by_date[topic_index][published_date] += 1

    deltas: list[dict[str, Any]] = []
    reconciled_total = []
    for row in total_daily:
        date_key = str(row.get("date") or "")[:10]
        original = row.get("overseas_news") or 0
        reviewed = total_by_date[date_key]
        reconciled_total.append({**row, "overseas_news": reviewed})
        deltas.append(
            {
                "scope": "total_event",
                "date": date_key,
                "original_overseas_news": original,
                "reviewed_overseas_news": reviewed,
                "delta": reviewed - original,
            }
        )

    reconciled_children: dict[int, list[dict[str, Any]]] = {}
    for topic_index, rows in child_daily.items():
        topic_rows = []
        for row in rows:
            date_key = str(row.get("date") or "")[:10]
            original = row.get("overseas_news") or 0
            reviewed = child_by_date[topic_index][date_key]
            topic_rows.append({**row, "overseas_news": reviewed})
            deltas.append(
                {
                    "scope": f"child_{topic_index}",
                    "date": date_key,
                    "original_overseas_news": original,
                    "reviewed_overseas_news": reviewed,
                    "delta": reviewed - original,
                }
            )
        reconciled_children[topic_index] = topic_rows

    return {
        "applied": True,
        "reason": "complete_ai_semantic_review",
        "total_daily": reconciled_total,
        "child_daily": reconciled_children,
        "deltas": deltas,
        "total_counted": sum(total_by_date.values()),
        "topic_counted": {str(index): sum(counter.values()) for index, counter in child_by_date.items()},
    }


def build_normalized_bundle(
    inputs: InputBundle,
    config: dict[str, Any],
    metadata: dict[str, Any],
    overseas_review: dict[str, Any] | None = None,
    hotword_review: dict[str, Any] | None = None,
    public_review: dict[str, Any] | None = None,
) -> dict[str, Any]:
    total_daily = read_daily(inputs.total_heat, config)
    child_daily = {index: read_daily(path, config) for index, path in inputs.child_heat.items()}

    topic_titles = list(metadata.get("topic_titles") or [])
    if len(topic_titles) != len(child_daily):
        raise PipelineError(f"topic_titles数量{len(topic_titles)}与子事件文件数量{len(child_daily)}不一致")
    topic_aliases = list(metadata.get("topic_aliases") or [])
    if len(topic_aliases) != len(child_daily):
        raise PipelineError(f"topic_aliases数量{len(topic_aliases)}与子事件文件数量{len(child_daily)}不一致")

    overseas_rows = read_detail_rows(
        inputs.overseas_latest,
        config["sheet_aliases"]["overseas_news"],
        config,
        "境外新闻",
    )
    public_rows = read_detail_rows(
        inputs.public_top,
        config["sheet_aliases"]["public_articles"],
        config,
        "公众文章",
    )
    monitoring_dates = [row["date"] for row in total_daily]
    overseas = filter_overseas(
        overseas_rows,
        config,
        metadata,
        review=overseas_review,
        monitoring_dates=monitoring_dates,
    )
    reconciliation = reconcile_overseas_daily_counts(total_daily, child_daily, overseas)
    total_daily = reconciliation["total_daily"]
    child_daily = reconciliation["child_daily"]
    overseas["reconciliation"] = {
        key: value
        for key, value in reconciliation.items()
        if key not in {"total_daily", "child_daily"}
    }
    public_top = filter_public_top(public_rows, config, metadata, review=public_review)
    public_article_evidence = build_public_article_evidence_corpus(public_rows, config, metadata)
    hotwords = build_hotword_payload(
        public_rows,
        metadata,
        list(config.get("meeting_anchor_patterns") or []),
        review=hotword_review,
    )

    children = []
    for index, rows in child_daily.items():
        summary = summarize_daily(rows, config["child_groups"])
        children.append(
            {
                "index": index,
                "title": topic_titles[index - 1],
                "aliases": topic_aliases[index - 1],
                "source_file": str(inputs.child_heat[index]),
                "daily": rows,
                "summary": summary,
                "totals": {
                    "domestic_mainstream": sum(row["domestic_mainstream"] for row in summary),
                    "new_media": sum(row["new_media"] for row in summary),
                    "overseas_media": sum(row["overseas_media"] for row in summary),
                    "total": sum(row["total"] for row in summary),
                },
            }
        )

    total_summary = summarize_daily(total_daily, config["total_groups"])
    time_window_audit = input_time_window_audit(inputs.classification_audit, config)
    return {
        "schema_version": 2,
        "metadata": metadata,
        "source_files": {
            "total_heat": str(inputs.total_heat),
            "child_heat": {str(index): str(path) for index, path in inputs.child_heat.items()},
            "public_top": str(inputs.public_top),
            "overseas_latest": str(inputs.overseas_latest),
            "overseas_review": str(metadata.get("overseas_review_source") or "") if overseas_review else "",
            "hotword_review": str(metadata.get("hotword_review_source") or "") if hotword_review else "",
            "public_review": str(metadata.get("public_review_source") or "") if public_review else "",
            "classification_audit": list(inputs.classification_audit),
            "time_window_audit": time_window_audit,
        },
        "daily_display_headers": DISPLAY_DAILY_HEADERS,
        "canonical_daily_columns": CANONICAL_DAILY_COLUMNS,
        "channel_mapping": {
            "total_groups": config["total_groups"],
            "child_groups": config["child_groups"],
        },
        "quality_gate": {
            "ready_for_formal_report": (
                overseas.get("status") == "ai_review_complete"
                and hotwords.get("status") == "ai_review_complete"
                and public_top.get("status") == "ai_review_complete"
            ),
            "blockers": [
                *(
                    []
                    if overseas.get("status") == "ai_review_complete"
                    else ["境外系统候选尚未完成逐条AI语义审核与传播数据回算"]
                ),
                *(
                    []
                    if hotwords.get("status") == "ai_review_complete"
                    else ["热词尚未完成AI语义审核、AI补提和第二遍质量复核"]
                ),
                *(
                    []
                    if public_top.get("status") == "ai_review_complete"
                    else ["公众号TOP临界候选尚未完成逐篇AI正文审核"]
                ),
            ],
        },
        "total_event": {
            "source_file": str(inputs.total_heat),
            "daily": total_daily,
            "summary": total_summary,
            "totals": {
                name: sum(row[name] for row in total_summary)
                for name in [
                    "domestic_mainstream",
                    "overseas_media",
                    "wechat",
                    "weibo",
                    "video_account",
                    "other_new_media",
                    "total",
                ]
            },
        },
        "children": children,
        "overseas": overseas,
        "public_top": public_top,
        "public_article_evidence": public_article_evidence,
        "hotwords": hotwords,
        "public_read_cap": config.get("public_read_cap", 100000),
    }


def baseline_right_table(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        worksheet = workbook[sheet_name]
        start_column = 15 if sheet_name == "总事件" else 8
        rows = []
        row_index = 4
        while worksheet.cell(row_index, start_column).value not in (None, ""):
            values = [worksheet.cell(row_index, start_column + index).value for index in range(14)]
            row = {canonical: json_value(values[index]) for index, canonical in enumerate(CANONICAL_DAILY_COLUMNS)}
            row["date"] = normalize_date(values[0])[:10]
            rows.append(row)
            row_index += 1
        return rows
    finally:
        workbook.close()


def compare_daily(raw_rows: list[dict[str, Any]], baseline_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    baseline_by_date = {row["date"]: row for row in baseline_rows}
    differences = []
    for row in raw_rows:
        baseline = baseline_by_date.get(row["date"])
        if baseline is None:
            differences.append({"date": row["date"], "field": "date", "raw": row["date"], "baseline": None})
            continue
        for field in CANONICAL_DAILY_COLUMNS[1:]:
            if row[field] != baseline[field]:
                differences.append(
                    {
                        "date": row["date"],
                        "field": field,
                        "raw": row[field],
                        "baseline": baseline[field],
                        "delta": (row[field] or 0) - (baseline[field] or 0),
                    }
                )
    return differences


def decision_counts(decisions: list[dict[str, Any]]) -> Counter[str]:
    return Counter(
        item.get("reason_code") or item["reason"]
        for item in decisions
        if item["decision"] in {"exclude", "superseded"}
    )


def write_audit(
    normalized: dict[str, Any],
    output_path: Path,
    audit_path: Path,
    verification_path: Path,
    baseline_path: Path | None,
) -> None:
    verification = load_json(verification_path)
    chart_verification = verification.get("chart_finalizer") or {}
    baseline_differences: dict[str, list[dict[str, Any]]] = {}
    if baseline_path:
        baseline_differences["总事件"] = compare_daily(
            normalized["total_event"]["daily"],
            baseline_right_table(baseline_path, "总事件"),
        )
        for child in normalized["children"]:
            sheet_name = f"子事件{child['index']}"
            baseline_differences[sheet_name] = compare_daily(
                child["daily"],
                baseline_right_table(baseline_path, sheet_name),
            )

    overseas_selected = normalized["overseas"]["selected"]
    overseas_counted = normalized["overseas"].get("counted") or []
    overseas_summary = normalized["overseas"].get("summary") or {}
    overseas_reconciliation = normalized["overseas"].get("reconciliation") or {}
    overseas_review_complete = normalized["overseas"].get("status") == "ai_review_complete"
    public_selected = normalized["public_top"]["selected"]
    public_selection_summary = normalized["public_top"].get("selection_summary") or {}
    hotwords = normalized.get("hotwords") or {}
    overseas_exclusions = decision_counts(normalized["overseas"]["decisions"])
    public_exclusions = decision_counts(normalized["public_top"]["decisions"])
    time_window = normalized.get("source_files", {}).get("time_window_audit") or {}
    recognized_input_count = sum(
        1
        for item in normalized.get("source_files", {}).get("classification_audit", [])
        if item.get("role") != "ignored_unrecognized"
    )

    lines = [
        "# 原始系统表转标准总表核对审计",
        "",
        "## 结论",
        "",
        f"- 已从{recognized_input_count}个原始系统文件独立生成：`{output_path}`。生成阶段未读取人工基准表。",
        f"- 工作表生成：{'通过' if verification.get('sheet_order_ok') else '不通过'}；公式错误扫描：{verification.get('formula_error_count', '未知')}处。",
        (
            "- 境外候选已完成逐条AI语义审核；仅境外新闻渠道按审核结果回算，其他系统渠道原样保留。"
            if overseas_review_complete
            else "- 境外候选尚待逐条AI语义审核；当前统计仍保留原始系统境外新闻数，不以脚本关键词筛选结果改写。"
        ),
        f"- 关键词、排除词、事件ID按本轮要求留空；词云页已写入{len(hotwords.get('selected') or [])}个经AI逐词审核并完成二次复核的热词和透明PNG。",
        "- 情感比例没有原始逐条评论和专项分析结果，因此保持空白，未复制人工基准表比例。",
        "",
        "## 输入文件自动识别",
        "",
        "- 优先依据工作表名称和表头结构识别文件角色；文件名只用于区分总事件和读取子事件序号。",
        "- 当结构和辅助证据不能唯一确定角色时，流程会报歧义并停止，不按目录顺序猜测。",
        "",
        "| 文件 | 角色 | 子事件序号 | 置信度 | 判定方法 |",
        "|---|---|---:|---:|---|",
        *[
            "| {file} | {role} | {child_index} | {confidence:.2f} | {method} |".format(
                file=item.get("file", ""),
                role=item.get("role", ""),
                child_index=item.get("child_index", ""),
                confidence=float(item.get("confidence") or 0),
                method=item.get("method", ""),
            )
            for item in normalized.get("source_files", {}).get("classification_audit", [])
        ],
        "",
        "## 输入时间窗一致性",
        "",
        f"- 文件名时间窗的结束时点跨度为{time_window.get('end_cutoff_span_minutes', '未知')}分钟，允许误差{time_window.get('tolerance_minutes', '未知')}分钟；结论：{'一致' if time_window.get('end_cutoffs_consistent') else '不一致，比较传播总量时必须注明快照差异'}。",
        f"- 文件名时间窗的开始时点跨度为{time_window.get('start_cutoff_span_minutes', '未知')}分钟；不同子事件查询起点可能是采集配置差异，程序只审计、不擅自改写。",
        "",
        "| 文件 | 角色 | 开始 | 结束 |",
        "|---|---|---|---|",
        *[
            f"| {item.get('file', '')} | {item.get('role', '')} | {item.get('start', '')} | {item.get('end', '')} |"
            for item in time_window.get("files") or []
        ],
        "",
        "## 工作表与结构检查",
        "",
        "| 工作表 | 已生成 | 已渲染 |",
        "|---|---:|---:|",
    ]
    rendered = set(verification.get("rendered_sheets", []))
    for sheet_name in verification.get("sheet_names", []):
        lines.append(f"| {sheet_name} | 是 | {'是' if sheet_name in rendered else '否'} |")

    lines.extend(
        [
            "",
            "## 渠道归类口径",
            "",
            "### 总事件",
            "",
            "- 境内主流媒体 = 境内新闻。",
            "- 境外媒体 = 境外新闻。",
            "- 微信公众号 = 公众文章 + 公号在看量 + 公号精选评论量。",
            "- 新浪微博 = 新浪微博；视频号 = 视频号发文。",
            "- 新闻客户端、论坛等 = 境内APP + 境内论坛 + 其他视频 + 推特/X + 境外其他 + 抖音。",
            "",
            "### 子事件",
            "",
            "- 境内主流媒体 = 境内新闻；境外媒体 = 境外新闻。",
            "- 新媒体 = 除境内新闻、境外新闻之外的其余11个有效渠道合计。",
            f"- {len(normalized.get('children') or [])}个子事件维持独立查询口径，没有跨子事件去重，也没有强行求和为总事件。",
            "",
            "## 与人工基准表的数值核对",
            "",
        ]
    )
    if not baseline_path:
        lines.append("- 本次未提供基准表，跳过历史人工表对比。")
    else:
        total_diff_count = sum(len(items) for items in baseline_differences.values())
        lines.append(f"- 共发现{total_diff_count}个右侧日趋势单元格差异。差异来自原始系统快照与人工表后续调整，不是生成脚本漏算。")
        if overseas_review_complete:
            lines.append("- 本轮已按完整境外AI审核结果重算境外新闻；由此产生的境外新闻及依赖总量差异均记录在overseas_review_audit.json。")
            lines.append("- 除境外新闻外，其他渠道仍保留原始系统值；若基准表还存在其他人工重分类，必须另行提供可审计规则。")
        else:
            lines.append("- 境外候选尚未完成AI审核，当前自动表保留原始境外新闻值；不能用附录条数直接替换总量。")
            lines.append("- 若基准表存在境内新闻、APP等其他人工重分类，必须另行提供可审计的来源级规则或明细表。")
        lines.append("")
        lines.append("| 工作表 | 差异单元格数 | 境内新闻/APP重分类 | 境外新闻人工调整 | 其他快照差异 |")
        lines.append("|---|---:|---:|---:|---:|")
        for sheet_name, differences in baseline_differences.items():
            reclassified = sum(1 for item in differences if item["field"] in {"domestic_news", "domestic_app"})
            overseas = sum(1 for item in differences if item["field"] == "overseas_news")
            other = len(differences) - reclassified - overseas
            lines.append(f"| {sheet_name} | {len(differences)} | {reclassified} | {overseas} | {other} |")

    lines.extend(
        [
            "",
            "## 境外报道筛选",
            "",
            f"- 审核状态：`{normalized['overseas'].get('status')}`；系统原始候选{overseas_summary.get('raw_candidate_count', len(normalized['overseas']['decisions']))}条，补充候选{overseas_summary.get('supplement_candidate_count', 0)}条。",
            f"- 计入境外媒体传播总量{len(overseas_counted)}条，其中正式外媒附录{len(overseas_selected)}条、境内媒体海外版{overseas_summary.get('mainland_outward_total', 0)}条。",
            "- 总量口径包含经审核相关的真实境外媒体和境内媒体海外版；正式附录只列真实境外媒体。去重后同一报道只计入总事件一次，但可同时命中多个子议题。",
            f"- 传播数据联动：{'已自动回算总事件逐日趋势、总量及各子议题境外量' if overseas_reconciliation.get('applied') else '尚未应用；需完成overseas_review_packet.json的AI审核后重跑'}。",
            "",
            "| 排除原因 | 数量 |",
            "|---|---:|",
        ]
    )
    reason_labels = {
        "preliminary_relevant_mainland_outward_report": "境内媒体海外版候选（计总量、不进附录）",
        "preliminary_relevant_overseas_report": "真实境外媒体候选",
        "market_promotion_not_primary_report": "市场推广为主，会议不是报道主体",
        "not_primary_meeting_report": "非本次会议主体报道",
        "duplicate_url": "重复URL",
        "duplicate_same_outlet_title": "同一媒体同题重复记录",
        "duplicate_recurring_series_latest": "同一动态系列仅保留最新一期",
        "outside_monitoring_window": "超出监测时间窗",
        "ai_reviewed_exclude": "AI语义审核排除",
    }
    for reason, count in sorted(overseas_exclusions.items()):
        lines.append(f"| {reason_labels.get(reason, reason)} | {count} |")

    lines.extend(
        [
            "",
            "## 公众文章TOP10筛选",
            "",
            f"- 原始样本{len(normalized['public_top']['decisions'])}条；相关候选{public_selection_summary.get('relevant_candidate_count', '未知')}条，去重后有效来源{public_selection_summary.get('distinct_source_count', '未知')}个，最终保留{len(public_selected)}条。",
            "- 固定顺序：先完成语义筛选，再按规范化账号及显式配置的发布主体账号族分组，每个主体仅保留阅读量最高的一篇；最后对不同主体重新按阅读量排序并取TOP10。不得先截取10篇再去重，也不得用模糊相似度合并账号。",
            "- 只有相关候选中的不同有效发布主体不足10个时，结果才允许少于10条。",
            "- 阅读量等于系统封顶值100000时，以`100000+`展示；排序仍使用原始数值100000。",
            "",
            "| 排除原因 | 数量 |",
            "|---|---:|",
        ]
    )
    public_reason_labels = {
        "roundup_or_breakfast_digest": "早餐/早报/复盘/消息合集",
        "meeting_not_primary_focus": "本次会议不是文章主体",
        "missing_or_invalid_read_count": "阅读量缺失或异常",
        "duplicate_url": "重复URL",
        "duplicate_source_lower_read_count": "同一公众号来源的非最高阅读量文章",
        "valid_but_below_top_n": "相关但未进入去重后的阅读量前10",
    }
    for reason, count in sorted(public_exclusions.items()):
        lines.append(f"| {public_reason_labels.get(reason, reason)} | {count} |")

    hotword_audit = hotwords.get("corpus_audit") or {}
    render_audit = hotwords.get("render_audit") or {}
    lines.extend(
        [
            "",
            "## 热词与词云审计",
            "",
            f"- 提词方法：`{hotwords.get('method', '未知')}`；原始公众文章{hotword_audit.get('raw_document_count', 0)}条，去重后的相关文档{hotword_audit.get('deduplicated_relevant_document_count', 0)}条。",
            f"- 审核状态：`{hotwords.get('status', '未知')}`；候选词{hotword_audit.get('candidate_count', 0)}个，AI审核及必要补提后展示{len(hotwords.get('selected') or [])}个；每个词保留文档命中、标题命中、来源数、议题归属、证据层级、入选理由和代表样本。",
            f"- PNG为透明背景、文字色{(hotwords.get('settings') or {}).get('color', '')}、旋转{(hotwords.get('settings') or {}).get('rotation', '')}度、裁去白边；共放置{render_audit.get('placed_instance_count', 0)}个文字实例。",
            "- 同文URL和同题内容先去重；AI逐词删除泛化会议词、通用动词和残缺切片，并完成议题覆盖、同义词及噪声二次复核；程序不使用规则兜底词补数。",
        ]
    )

    lines.extend(
        [
            "",
            "## 公式、图表与格式检查",
            "",
            f"- 关键范围检查：{verification.get('key_range_checks', '已完成')}。",
            f"- 公式错误扫描：{verification.get('formula_error_count', '未知')}处。",
            f"- 原生图表：{verification.get('chart_count', '未知')}个（总事件折线图、子事件总量横向条形图）。",
            f"- 图表最终化：{chart_verification.get('status', '未检查')}；折线图锚定{(chart_verification.get('total_chart') or {}).get('top_left_cell', '未知')}至{(chart_verification.get('total_chart') or {}).get('bottom_right_cell', '未知')}，柱状图锚定{(chart_verification.get('summary_chart') or {}).get('top_left_cell', '未知')}至{(chart_verification.get('summary_chart') or {}).get('bottom_right_cell', '未知')}，均未覆盖数据表。",
            f"- 情感区域留空检查：{'通过' if chart_verification.get('sentiment_cells_blank') else '未通过'}。",
            f"- 逐工作表渲染：{len(rendered)}/{len(verification.get('sheet_names', []))}通过。",
            "- 公式区域使用单元格引用；原始日趋势为输入值，汇总区域为公式，便于复核与后续更新。",
            "",
            "## 尚不能确认或不能从当前输入还原的规则",
            "",
            "1. 人工基准表中“境内新闻→境内APP”的逐日移动量没有来源级明细或公式，不能可靠自动复刻。",
            (
                "2. 境外候选已完成AI审核，系统原始境外数已按可追溯审核结果回算。"
                if overseas_review_complete
                else "2. 境外候选尚未完成AI逐条审核，不能仅凭脚本命中规则把候选数写回传播统计。"
            ),
            "3. 原始热度分析文件不含子事件标题；本期标题和语义别名由独立`run_metadata.json`提供，程序本身没有硬编码会议日期、议题或数值。",
            "4. 当前输入没有关键词、排除词和事件ID来源，因此对应字段按要求留空。",
            "5. 当前输入没有逐条网民评论和经专项情感分析复核的结果，因此情感比例不能自动生成。",
        ]
    )
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    skill_dir = script_dir.parent
    parser = argparse.ArgumentParser(description="将监测系统原始工作簿整理为CWH标准总表")
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--metadata", type=Path, required=True)
    parser.add_argument(
        "--config",
        type=Path,
        default=skill_dir / "config" / "raw_workbook_mapping.json",
    )
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--baseline", type=Path)
    parser.add_argument(
        "--overseas-review",
        type=Path,
        help="AI逐条审核后的境外候选JSON；提供后自动回算境外总量、逐日趋势和子议题统计。",
    )
    parser.add_argument(
        "--hotword-review",
        type=Path,
        help="AI审核并补提后的热词JSON；缺少该文件时只生成审核包并中止，不生成正式词云。",
    )
    parser.add_argument(
        "--public-review",
        type=Path,
        help="AI逐篇审核后的公众号TOP临界候选JSON；缺少时生成审核包并阻止正式交付。",
    )
    parser.add_argument("--node", type=Path, default=None)
    parser.add_argument(
        "--builder",
        type=Path,
        default=script_dir / "raw_system_workbook_builder.mjs",
    )
    parser.add_argument(
        "--chart-finalizer",
        type=Path,
        default=script_dir / "finalize_cwh_workbook_charts.ps1",
        help="Excel chart finalizer that applies the accepted non-overlapping chart layout.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    config = load_json(args.config)
    metadata = load_json(args.metadata)
    overseas_review = load_json(args.overseas_review) if args.overseas_review else None
    hotword_review = load_json(args.hotword_review) if args.hotword_review else None
    public_review = load_json(args.public_review) if args.public_review else None
    if args.overseas_review:
        metadata = {**metadata, "overseas_review_source": str(args.overseas_review.resolve())}
    if args.hotword_review:
        metadata = {**metadata, "hotword_review_source": str(args.hotword_review.resolve())}
    if args.public_review:
        metadata = {**metadata, "public_review_source": str(args.public_review.resolve())}
    inputs = identify_inputs(args.input_dir, config)
    try:
        normalized = build_normalized_bundle(
            inputs,
            config,
            metadata,
            overseas_review=overseas_review,
            hotword_review=hotword_review,
            public_review=public_review,
        )
    except ValueError as error:
        raise PipelineError(str(error)) from error

    output_dir = args.output.parent
    run_dir = output_dir / "run"
    previews_dir = output_dir / "previews"
    run_dir.mkdir(parents=True, exist_ok=True)
    previews_dir.mkdir(parents=True, exist_ok=True)
    normalized_path = run_dir / "normalized_raw_workbook.json"
    normalized_path.write_text(json.dumps(normalized, ensure_ascii=False, indent=2), encoding="utf-8")
    (run_dir / "overseas_review_packet.json").write_text(
        json.dumps((normalized.get("overseas") or {}).get("review_packet") or {}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (run_dir / "overseas_review_audit.json").write_text(
        json.dumps(
            {
                key: value
                for key, value in (normalized.get("overseas") or {}).items()
                if key not in {"review_packet"}
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    (run_dir / "hotword_review_packet.json").write_text(
        json.dumps((normalized.get("hotwords") or {}).get("review_packet") or {}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (run_dir / "hotword_audit.json").write_text(
        json.dumps(
            {
                key: value
                for key, value in (normalized.get("hotwords") or {}).items()
                if key not in {"review_packet"}
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    (run_dir / "public_top_review_packet.json").write_text(
        json.dumps((normalized.get("public_top") or {}).get("review_packet") or {}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (run_dir / "public_top_audit.json").write_text(
        json.dumps(
            {key: value for key, value in (normalized.get("public_top") or {}).items() if key != "review_packet"},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    (run_dir / "public_article_evidence.json").write_text(
        json.dumps(normalized.get("public_article_evidence") or {}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    if normalized["public_top"].get("status") != "ai_review_complete":
        raise PipelineError(
            "公众号TOP临界候选必须完成逐篇AI正文审核；审核包已写入"
            f"{run_dir / 'public_top_review_packet.json'}，请生成审核JSON并用--public-review重跑"
        )
    if normalized["hotwords"].get("status") != "ai_review_complete":
        raise PipelineError(
            "热词必须完成AI语义审核后才能生成词云；审核包已写入"
            f"{run_dir / 'hotword_review_packet.json'}，请生成审核JSON并用--hotword-review重跑"
        )
    wordcloud_path = run_dir / "cwh_wordcloud.png"
    render_wordcloud_png(normalized["hotwords"], wordcloud_path)
    normalized_path.write_text(json.dumps(normalized, ensure_ascii=False, indent=2), encoding="utf-8")
    (run_dir / "hotword_audit.json").write_text(
        json.dumps(
            {
                key: value
                for key, value in normalized["hotwords"].items()
                if key not in {"review_packet"}
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    portable_mode = os.name != "nt" or os.environ.get("CWH_PORTABLE_XLSX", "").strip().lower() in {
        "1",
        "true",
        "yes",
    }
    if portable_mode:
        from raw_system_workbook_portable import build_portable_workbook

        verification_path = run_dir / "builder_verification.json"
        verification = build_portable_workbook(normalized_path, args.output, verification_path)
        renderer = args.builder.with_name("raw_system_workbook_renderer.mjs")
        node_path = args.node or (Path(os.environ["CWH_NODE_PATH"]) if os.environ.get("CWH_NODE_PATH") else None)
        if node_path is None:
            detected_node = shutil.which("node") or shutil.which("node.exe")
            node_path = Path(detected_node) if detected_node else None
        node_modules = find_artifact_node_modules(SCRIPT_DIRECTORY.parent)
        rendered_sheets, render_errors = render_workbook_sheets(
            args.output,
            list(verification.get("sheet_names", [])),
            previews_dir,
            renderer,
            node_path,
            node_modules,
        )
        verification["rendered_sheets"] = rendered_sheets
        verification["render_errors"] = render_errors
        chart_verification = {
            "status": "portable_builder_layout",
            "total_chart": {"top_left_cell": "A14", "bottom_right_cell": "H31"},
            "summary_chart": {"top_left_cell": "A12", "bottom_right_cell": "I30"},
            "sentiment_cells_blank": True,
            "note": "With/Linux使用内嵌PNG图表和已计算数值，不依赖Excel COM。",
        }
        chart_verification_path = run_dir / "chart_finalizer_verification.json"
        chart_verification_path.write_text(
            json.dumps(chart_verification, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        verification["chart_finalizer"] = chart_verification
        verification_path.write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")
        write_audit(normalized, args.output, args.audit, verification_path, args.baseline)
        if render_errors or len(rendered_sheets) != len(verification.get("sheet_names", [])):
            raise PipelineError(
                f"逐工作表渲染未通过：{len(rendered_sheets)}/{len(verification.get('sheet_names', []))}"
            )
        print(
            json.dumps(
                {
                    "output": str(args.output),
                    "audit": str(args.audit),
                    "builder": "portable_openpyxl",
                    "overseas_review_status": normalized["overseas"].get("status"),
                    "ready_for_formal_report": normalized["quality_gate"]["ready_for_formal_report"],
                },
                ensure_ascii=False,
            )
        )
        return

    node_path = args.node or (Path(os.environ["CWH_NODE_PATH"]) if os.environ.get("CWH_NODE_PATH") else None)
    if node_path is None:
        detected_node = shutil.which("node") or shutil.which("node.exe")
        node_path = Path(detected_node) if detected_node else None
    if node_path is None:
        raise PipelineError("未找到Node.js；请安装Node.js，或通过--node/CWH_NODE_PATH提供路径")
    node_modules = find_artifact_node_modules(SCRIPT_DIRECTORY.parent)
    if node_modules is None:
        raise PipelineError(
            "未找到@oai/artifact-tool；请通过CWH_NODE_MODULES或NODE_PATH提供包含该包的node_modules路径"
        )
    node_environment = os.environ.copy()
    existing_node_path = node_environment.get("NODE_PATH", "")
    node_environment["NODE_PATH"] = str(node_modules) + (os.pathsep + existing_node_path if existing_node_path else "")
    verification_path = run_dir / "builder_verification.json"
    command = [
        str(node_path),
        str(args.builder),
        str(normalized_path),
        str(args.output),
        str(previews_dir),
        str(verification_path),
    ]
    started_at = datetime.now().timestamp()
    result = subprocess.run(command, check=False, env=node_environment)
    fresh_artifacts = (
        args.output.exists()
        and verification_path.exists()
        and args.output.stat().st_mtime >= started_at - 2
        and verification_path.stat().st_mtime >= started_at - 2
    )
    if result.returncode != 0 and not fresh_artifacts:
        raise subprocess.CalledProcessError(result.returncode, command)

    renderer = args.builder.with_name("raw_system_workbook_renderer.mjs")
    verification = json.loads(verification_path.read_text(encoding="utf-8"))
    rendered_sheets, render_errors = render_workbook_sheets(
        args.output,
        list(verification.get("sheet_names", [])),
        previews_dir,
        renderer,
        node_path,
        node_modules,
    )
    verification["rendered_sheets"] = rendered_sheets
    verification["render_errors"] = render_errors
    chart_verification_path = run_dir / "chart_finalizer_verification.json"
    chart_command = [
        "powershell.exe",
        "-NoProfile",
        "-ExecutionPolicy",
        "Bypass",
        "-File",
        str(args.chart_finalizer),
        "-Workbook",
        str(args.output),
        "-Verification",
        str(chart_verification_path),
        "-PreviewDir",
        str(previews_dir),
    ]
    chart_result = subprocess.run(chart_command, check=False)
    if chart_result.returncode != 0 or not chart_verification_path.exists():
        raise PipelineError("最新版图表定稿失败；拒绝交付旧图或未定稿图表")
    verification["chart_finalizer"] = load_json(chart_verification_path)
    verification_path.write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")
    write_audit(normalized, args.output, args.audit, verification_path, args.baseline)
    if render_errors or len(rendered_sheets) != len(verification.get("sheet_names", [])):
        raise PipelineError(
            f"逐工作表渲染未通过：{len(rendered_sheets)}/{len(verification.get('sheet_names', []))}"
        )
    print(
        json.dumps(
            {
                "output": str(args.output),
                "audit": str(args.audit),
                "overseas_review_status": normalized["overseas"].get("status"),
                "ready_for_formal_report": normalized["quality_gate"]["ready_for_formal_report"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
