from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, asdict
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlparse


PLATFORM_PREFIXES = {
    "微信公众号": "wechat_public",
    "微信公号": "wechat_public",
    "头条号": "toutiao_articles",
    "今日头条": "toutiao_articles",
    "百家号": "baijiahao",
    "搜狐号": "sohu_media",
    "网易号": "netease_media",
    "微博账号": "weibo",
    "抖音账号": "douyin",
    "视频号": "video_account",
}

PLATFORM_LABELS = {
    "wechat_public": "微信公众号",
    "toutiao_articles": "今日头条/头条号",
    "baijiahao": "百家号",
    "sohu_media": "搜狐号",
    "netease_media": "网易号",
    "publisher_site": "媒体网站/客户端",
    "named_expert_or_institution": "具名专家/机构",
    "other_named_source": "其他具名来源",
}

DOMAIN_PLATFORM = {
    "mp.weixin.qq.com": "wechat_public",
    "toutiao.com": "toutiao_articles",
    "baijiahao.baidu.com": "baijiahao",
    "mbd.baidu.com": "baijiahao",
    "sohu.com": "sohu_media",
    "163.com": "netease_media",
}

ATTRIBUTION_VERBS = r"称|表示|认为|指出|建议|分析称|分析认为|预计|呼吁|评价称"
EXPLICIT_RE = re.compile(
    r"(?P<prefix>微信公众号|微信公号|头条号|今日头条|百家号|搜狐号|网易号|微博账号|抖音账号|视频号)"
    r"[‘'“\"](?P<source>[^’'”\"]{1,80})[’'”\"]"
)
GENERIC_ATTRIBUTION_RE = re.compile(
    rf"(?P<source>[A-Za-z0-9\u4e00-\u9fff·（）()《》、—\-\s]{{2,90}}?)(?:{ATTRIBUTION_VERBS})"
)


@dataclass
class BaselineSample:
    baseline: str
    source: str
    platform: str
    evidence_text: str
    anchor: str = ""
    extraction_mode: str = "text_attribution"


def normalize(value: Any) -> str:
    text = str(value or "").lower()
    text = re.sub(r"^(?:微信公众号|微信公号|头条号|今日头条|百家号|搜狐号|网易号)[‘'“\"]?", "", text)
    text = re.sub(r"[’'”\"\s\u3000，。；：、（）()《》【】\[\]·—\-]", "", text)
    return text


def parse_assignment(value: str) -> tuple[str, Path]:
    if "=" not in value:
        raise argparse.ArgumentTypeError("参数必须为 LABEL=PATH")
    label, raw_path = value.split("=", 1)
    path = Path(raw_path).expanduser().resolve()
    if not path.exists():
        raise argparse.ArgumentTypeError(f"文件不存在: {path}")
    return label.strip(), path


def read_text(path: Path) -> str:
    if path.suffix.lower() in {".txt", ".md"}:
        return path.read_text(encoding="utf-8-sig", errors="replace")
    if path.suffix.lower() == ".docx":
        from docx import Document

        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(parts)
    raise ValueError(f"暂不直接读取 {path.suffix}；请先转换为 txt 或 docx: {path}")


def domestic_media_section(text: str) -> str:
    starts = [
        text.find("（一）境内媒体自媒体情况"),
        text.find("（一）境内媒体、自媒体情况"),
    ]
    start = next((value for value in starts if value >= 0), -1)
    if start < 0:
        return text
    end_candidates = [
        value
        for marker in ("（二）网民评论情况", "（二）境内网民评论", "（三）热词分布情况")
        if (value := text.find(marker, start + 1)) >= 0
    ]
    end = min(end_candidates) if end_candidates else len(text)
    return text[start:end]


def explicit_platform(source_text: str, url: str = "") -> str | None:
    for prefix, platform in PLATFORM_PREFIXES.items():
        if source_text.startswith(prefix):
            return platform
    if url:
        host = (urlparse(url).hostname or "").lower()
        for domain, platform in DOMAIN_PLATFORM.items():
            if host == domain or host.endswith("." + domain):
                return platform
    return None


def classify_generic_source(source: str, publisher_names: set[str]) -> str:
    compact = normalize(source)
    if any(normalize(name) == compact for name in publisher_names):
        return "publisher_site"
    if re.search(r"网|报|新闻|财经|电视台|电台|杂志|日报|时报|周刊", source):
        return "publisher_site"
    if re.search(r"大学|学院|研究院|研究所|研究中心|协会|学会|委员会|研究员|教授|主任|院长|分析师|经济学家|首席", source):
        return "named_expert_or_institution"
    return "other_named_source"


def clean_generic_source(value: str) -> str:
    value = re.split(r"[。！？；\n\r]", value)[-1]
    value = re.sub(r"^(?:一是|二是|三是|四是|五是|六是|七是|此外|同时|其中|另有|以及)", "", value)
    value = re.sub(r"^.*?[，：]", "", value) if len(value) > 55 and re.search(r"[，：]", value) else value
    return value.strip(" ，：；。\t")


def load_registry_publishers(registry_path: Path | None) -> set[str]:
    if not registry_path or not registry_path.exists():
        return set()
    payload = json.loads(registry_path.read_text(encoding="utf-8-sig"))
    return {str(item.get("name") or "").strip() for item in payload.get("sources") or [] if item.get("name")}


def extract_text_samples(label: str, path: Path, publisher_names: set[str]) -> list[BaselineSample]:
    section = domestic_media_section(read_text(path))
    samples: list[BaselineSample] = []
    occupied: list[tuple[int, int]] = []
    for match in EXPLICIT_RE.finditer(section):
        prefix = match.group("prefix")
        source = match.group("source").strip()
        samples.append(
            BaselineSample(
                baseline=label,
                source=source,
                platform=PLATFORM_PREFIXES[prefix],
                evidence_text=match.group(0),
                extraction_mode="explicit_platform_attribution",
            )
        )
        occupied.append(match.span())

    def overlaps(span: tuple[int, int]) -> bool:
        return any(not (span[1] <= left or span[0] >= right) for left, right in occupied)

    for match in GENERIC_ATTRIBUTION_RE.finditer(section):
        if overlaps(match.span("source")):
            continue
        source = clean_generic_source(match.group("source"))
        if len(normalize(source)) < 2 or len(source) > 70:
            continue
        if source in {"网民", "舆论", "会议", "政策", "规划", "方案", "部署"}:
            continue
        samples.append(
            BaselineSample(
                baseline=label,
                source=source,
                platform=classify_generic_source(source, publisher_names),
                evidence_text=match.group(0).strip(),
            )
        )
    return samples


def extract_inventory_samples(label: str, path: Path, publisher_names: set[str]) -> list[BaselineSample]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    rows = payload.get("media_samples") if isinstance(payload, dict) else payload
    samples: list[BaselineSample] = []
    for row in rows or []:
        raw_source = str(row.get("source") or "").strip()
        platform = explicit_platform(raw_source, str(row.get("url") or ""))
        source = raw_source
        match = EXPLICIT_RE.search(raw_source)
        if match:
            source = match.group("source").strip()
        if not platform:
            platform = classify_generic_source(source, publisher_names)
        samples.append(
            BaselineSample(
                baseline=label,
                source=source,
                platform=platform,
                evidence_text=raw_source,
                anchor=str(row.get("anchor") or row.get("title") or "").strip(),
                extraction_mode="structured_inventory",
            )
        )
    return samples


def iter_analysis_candidates(payload: dict[str, Any]) -> Iterable[dict[str, Any]]:
    research = ((payload.get("research_audit") or {}).get("domestic_media_research") or {})
    for topic in research.get("candidate_pool_by_topic") or []:
        for candidate in topic.get("candidates") or []:
            yield candidate
    for topic in ((payload.get("viewpoints") or {}).get("by_topic") or []):
        for cluster in topic.get("clusters") or []:
            for evidence in cluster.get("evidence") or []:
                yield evidence


def candidate_platform(candidate: dict[str, Any]) -> str:
    platform = str(candidate.get("platform") or "").lower()
    mapping = {
        "wechat": "wechat_public",
        "weixin": "wechat_public",
        "toutiao": "toutiao_articles",
        "baijiahao": "baijiahao",
    }
    if platform in mapping:
        return mapping[platform]
    return explicit_platform(str(candidate.get("source") or ""), str(candidate.get("url") or "")) or ""


def source_similarity(sample: BaselineSample, candidate: dict[str, Any]) -> float:
    source = normalize(sample.source)
    candidate_source = normalize(candidate.get("source"))
    content = normalize(
        " ".join(
            str(candidate.get(key) or "")
            for key in ("title", "content_summary", "content", "source_excerpt", "formal_claim")
        )
    )
    if not source:
        return 0.0
    source_score = 1.0 if source in candidate_source or source in content else SequenceMatcher(None, source, candidate_source).ratio()
    if not sample.anchor:
        return source_score * 0.7
    anchor = normalize(sample.anchor)
    anchor_score = 1.0 if anchor and anchor in content else SequenceMatcher(None, anchor, content).ratio()
    return 0.55 * source_score + 0.45 * anchor_score


def audit_recall(samples: list[BaselineSample], analysis_path: Path) -> dict[str, Any]:
    payload = json.loads(analysis_path.read_text(encoding="utf-8-sig"))
    candidates = list(iter_analysis_candidates(payload))
    results = []
    for sample in samples:
        best_candidate: dict[str, Any] | None = None
        best_score = 0.0
        for candidate in candidates:
            score = source_similarity(sample, candidate)
            if score > best_score:
                best_candidate = candidate
                best_score = score
        threshold = 0.66 if sample.anchor else 0.62
        recalled = best_score >= threshold
        results.append(
            {
                **asdict(sample),
                "recalled": recalled,
                "match_score": round(best_score, 3),
                "matched_candidate_id": (best_candidate or {}).get("candidate_id") if recalled else None,
                "matched_url": (best_candidate or {}).get("url") if recalled else None,
                "match_confidence": "source_and_anchor" if sample.anchor else "source_only_lower_confidence",
            }
        )
    per_platform: dict[str, dict[str, Any]] = {}
    for platform in sorted({sample.platform for sample in samples}):
        rows = [row for row in results if row["platform"] == platform]
        recalled = sum(bool(row["recalled"]) for row in rows)
        per_platform[platform] = {
            "label": PLATFORM_LABELS.get(platform, platform),
            "baseline_samples": len(rows),
            "recalled": recalled,
            "recall_rate": round(recalled / len(rows), 4) if rows else None,
            "missed_sources": [row["source"] for row in rows if not row["recalled"]],
        }
    return {"analysis_path": str(analysis_path), "per_platform": per_platform, "sample_matches": results}


def query_evidence_valid(check: dict[str, Any]) -> bool:
    queries = [str(value).strip() for value in check.get("queries") or [] if str(value).strip()]
    mode = str(check.get("execution_mode") or "")
    if mode not in {"platform_specific", "site_restricted_search", "browser_platform_search", "weread_platform_search"}:
        return False
    if not queries:
        return False
    if str(check.get("status") or "") == "hit" and not (check.get("urls") or check.get("candidate_ids")):
        return False
    return True


def audit_channels(analysis_path: Path) -> dict[str, Any]:
    payload = json.loads(analysis_path.read_text(encoding="utf-8-sig"))
    research = ((payload.get("research_audit") or {}).get("domestic_media_research") or {})
    coverage = research.get("coverage_by_topic") or []
    wanted = ("toutiao_articles", "wechat_public", "baijiahao")
    by_platform: dict[str, Any] = {}
    for source_id in wanted:
        rows = []
        for topic in coverage:
            matches = [check for check in topic.get("checks") or [] if check.get("source_id") == source_id]
            if not matches:
                rows.append({"topic": topic.get("topic"), "status": "missing", "valid_execution_evidence": False, "failure_reason": "missing_platform_check"})
                continue
            check = matches[0]
            status = str(check.get("status") or "missing")
            valid = query_evidence_valid(check)
            reason = ""
            if status in {"access_failed", "waiting_login", "failed"}:
                reason = str(check.get("blocker") or check.get("failure_reason") or status)
            elif not valid:
                reason = "missing_or_invalid_query_snapshot"
            rows.append(
                {
                    "topic": topic.get("topic"),
                    "status": status,
                    "execution_mode": check.get("execution_mode"),
                    "queries": check.get("queries") or [],
                    "urls": check.get("urls") or [],
                    "valid_execution_evidence": valid,
                    "failure_reason": reason,
                }
            )
        valid_rows = [row for row in rows if row["valid_execution_evidence"] and row["status"] in {"hit", "no_relevant_result"}]
        failures = Counter(row["failure_reason"] for row in rows if row["failure_reason"])
        by_platform[source_id] = {
            "label": PLATFORM_LABELS.get(source_id, source_id),
            "topics_expected": len(coverage),
            "topics_with_valid_completed_route": len(valid_rows),
            "route_coverage_rate": round(len(valid_rows) / len(coverage), 4) if coverage else None,
            "status_counts": dict(Counter(row["status"] for row in rows)),
            "failure_reasons": dict(failures),
            "topic_checks": rows,
        }
    return {"analysis_path": str(analysis_path), "topics": len(coverage), "by_platform": by_platform}


def summarize_baselines(samples: list[BaselineSample]) -> dict[str, Any]:
    per_baseline: dict[str, Any] = {}
    for baseline in sorted({sample.baseline for sample in samples}):
        rows = [sample for sample in samples if sample.baseline == baseline]
        counts = Counter(sample.platform for sample in rows)
        per_baseline[baseline] = {
            "total_attributed_samples": len(rows),
            "platform_counts": {
                platform: {"label": PLATFORM_LABELS.get(platform, platform), "count": count}
                for platform, count in sorted(counts.items())
            },
            "sources_by_platform": {
                platform: [sample.source for sample in rows if sample.platform == platform]
                for platform in sorted(counts)
            },
        }
    totals = Counter(sample.platform for sample in samples)
    return {
        "per_baseline": per_baseline,
        "all_baselines": {
            "total_attributed_samples": len(samples),
            "platform_counts": {
                platform: {"label": PLATFORM_LABELS.get(platform, platform), "count": count}
                for platform, count in sorted(totals.items())
            },
        },
    }


def markdown_report(payload: dict[str, Any]) -> str:
    lines = ["# CWH 基准样本平台覆盖审计", "", "## 一、基准样本平台数量", ""]
    for baseline, summary in payload["baseline_summary"]["per_baseline"].items():
        lines.append(f"### {baseline}")
        lines.append("")
        lines.append("| 平台/来源类型 | 样本数 |")
        lines.append("|---|---:|")
        for item in summary["platform_counts"].values():
            lines.append(f"| {item['label']} | {item['count']} |")
        lines.append("")
    lines.extend(["## 二、当前系统平台通道", ""])
    for label, audit in payload.get("channel_audits", {}).items():
        lines.append(f"### {label}")
        lines.append("")
        lines.append("| 平台 | 完成议题/应查议题 | 通道覆盖率 | 状态 | 失败原因 |")
        lines.append("|---|---:|---:|---|---|")
        for row in audit["by_platform"].values():
            failures = "；".join(f"{key}×{value}" for key, value in row["failure_reasons"].items()) or "无"
            statuses = "，".join(f"{key}:{value}" for key, value in row["status_counts"].items())
            rate = "—" if row["route_coverage_rate"] is None else f"{row['route_coverage_rate']:.1%}"
            lines.append(
                f"| {row['label']} | {row['topics_with_valid_completed_route']}/{row['topics_expected']} | {rate} | {statuses} | {failures} |"
            )
        lines.append("")
    lines.extend(["## 三、基准样本召回", ""])
    for label, audit in payload.get("recall_audits", {}).items():
        lines.append(f"### {label}")
        lines.append("")
        lines.append("| 平台/来源类型 | 基准样本 | 已召回 | 召回率 | 未召回来源 |")
        lines.append("|---|---:|---:|---:|---|")
        for row in audit["per_platform"].values():
            rate = "—" if row["recall_rate"] is None else f"{row['recall_rate']:.1%}"
            missed = "、".join(row["missed_sources"]) or "无"
            lines.append(f"| {row['label']} | {row['baseline_samples']} | {row['recalled']} | {rate} | {missed} |")
        lines.append("")
    lines.extend(
        [
            "## 口径说明",
            "",
            "- 通道覆盖率只承认真实平台定向/site 限定或浏览器平台检索，并要求保存查询；`hit` 还必须有 URL 或候选 ID。",
            "- `access_failed`、验证码、登录等待、超时和缺失查询快照均不算已完成通道。",
            "- 结构化清单含原句锚点时按来源+内容匹配；仅有报告文本时为来源级低置信度召回，不能替代逐句人工复核。",
        ]
    )
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="统计历史基准平台样本，并审计当前 CWH 平台通道与样本召回。")
    parser.add_argument("--baseline", action="append", default=[], metavar="LABEL=PATH")
    parser.add_argument("--inventory", action="append", default=[], metavar="LABEL=PATH")
    parser.add_argument("--analysis", action="append", default=[], metavar="LABEL=PATH")
    parser.add_argument("--registry")
    parser.add_argument("--output", required=True)
    parser.add_argument("--markdown")
    args = parser.parse_args()

    registry = Path(args.registry).resolve() if args.registry else Path(__file__).resolve().parent.parent / "config" / "source_registry.v1.json"
    publisher_names = load_registry_publishers(registry)
    samples: list[BaselineSample] = []
    samples_by_label: dict[str, list[BaselineSample]] = defaultdict(list)
    for raw in args.baseline:
        label, path = parse_assignment(raw)
        extracted = extract_text_samples(label, path, publisher_names)
        samples.extend(extracted)
        samples_by_label[label].extend(extracted)
    for raw in args.inventory:
        label, path = parse_assignment(raw)
        extracted = extract_inventory_samples(label, path, publisher_names)
        samples.extend(extracted)
        samples_by_label[label].extend(extracted)

    analyses = [parse_assignment(raw) for raw in args.analysis]
    recall_audits: dict[str, Any] = {}
    channel_audits: dict[str, Any] = {}
    for label, path in analyses:
        channel_audits[label] = audit_channels(path)
        matching_samples = samples_by_label.get(label) or samples
        recall_audits[label] = audit_recall(matching_samples, path)

    payload = {
        "schema_version": "1.0",
        "baseline_summary": summarize_baselines(samples),
        "channel_audits": channel_audits,
        "recall_audits": recall_audits,
        "definitions": {
            "route_coverage_rate": "valid completed topic-platform routes / expected topic-platform routes",
            "candidate_recall_rate": "baseline samples matched in the complete candidate/evidence pool / baseline samples",
            "waiting_login_terminal": False,
        },
    }
    output = Path(args.output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    markdown = Path(args.markdown).resolve() if args.markdown else output.with_suffix(".md")
    markdown.write_text(markdown_report(payload), encoding="utf-8")
    print(json.dumps({"json": str(output), "markdown": str(markdown), "baselines": len(samples_by_label)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
