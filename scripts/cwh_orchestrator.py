from __future__ import annotations

import argparse
import csv
import html
import json
import os
import re
import subprocess
import sys
import unicodedata
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from ingest_monitoring_workbook import ingest_workbook
from report_rules import domestic_viewpoint_quality_issues
from cwh_viewpoint_gate import cluster_density_result
from cwh_writing_rules import opening_paragraph, domestic_media_label

SCHEMA_VERSION = "0.3"
DEFAULT_DATE = "2026-06-29"

TEXT_KEYS = ["content", "desc", "title", "text", "summary", "snippet", "note_desc", "comment_content", "comment_text"]
TITLE_KEYS = ["title", "note_title", "desc", "content", "text"]
AUTHOR_KEYS = ["source", "author", "nickname", "user_nickname", "user_name", "screen_name"]
URL_KEYS = ["url", "link", "source_url", "note_url", "video_url"]
SPREAD_KEYS = ["spread_count", "liked_count", "like_count", "likes", "play_count", "view_count", "read_count", "score"]

PLATFORM_LABELS = {
    "news": "新闻网站",
    "web": "网页",
    "wechat": "微信",
    "weixin": "微信",
    "weibo": "微博",
    "wb": "微博",
    "douyin": "抖音",
    "dy": "抖音",
    "xhs": "小红书",
    "xiaohongshu": "小红书",
    "bili": "B站",
    "bilibili": "B站",
    "zhihu": "知乎",
    "tieba": "贴吧",
    "bing_news": "Bing新闻",
}

SOURCE_TYPE_LABELS = {
    "central_media": "中央媒体",
    "mainstream_media": "主流媒体",
    "commercial_media": "商业媒体",
    "official_media": "政务/官方媒体",
    "self_media": "自媒体",
    "self_media_post": "自媒体",
    "self_media_article": "自媒体",
    "social_note": "社交平台",
    "netizen": "网民评论",
    "netizen_comment": "网民评论",
    "overseas_media": "境外媒体",
    "media": "媒体",
    "news": "媒体",
    "unknown": "未识别",
}

SENTIMENT_ALIASES = {
    "positive": "positive",
    "正面": "positive",
    "积极": "positive",
    "neutral": "neutral",
    "中性": "neutral",
    "客观中立": "neutral",
    "negative": "negative",
    "负面": "negative",
    "消极": "negative",
    "unknown": "unknown",
    "未判定": "unknown",
    "": "unknown",
}
VERIFIED_SENTIMENT_SOURCES = {
    "ai_reviewed",
    "human_reviewed",
    "classifier",
    "active_learning_classifier",
    "large_scale_sentiment_analysis",
}

STOPWORDS = {
    "国务院",
    "常务会议",
    "国务院常务会议",
    "情况",
    "汇报",
    "有关",
    "工作",
    "方案",
    "规划",
    "发展",
    "当前",
    "研究",
    "听取",
    "审议",
    "通过",
    "部署",
    "推进",
    "加强",
    "做好",
    "会议",
    "总理",
    "主持",
    "召开",
}
GENERIC_TERMS = {"十五五", "十四五", "规划", "方案", "行动方案", "有关工作", "情况汇报"}

POSITIVE_WORDS = ["利好", "机遇", "提升", "促进", "推动", "完善", "加快", "支持", "优化", "增长", "稳定", "高质量"]
NEGATIVE_WORDS = ["担忧", "风险", "压力", "质疑", "问题", "挑战", "成本", "不确定", "炒作", "批评", "警惕"]
HYPE_WORDS = ["炒作", "抹黑", "质疑", "威胁", "风险", "批评", "争议", "打压"]

CENTRAL_MEDIA_DOMAINS = ["xinhuanet.com", "people.com.cn", "cctv.com", "cntv.cn", "chinanews.com.cn", "gmw.cn", "ce.cn"]
OFFICIAL_DOMAINS = ["gov.cn", "www.gov.cn"]
OVERSEAS_DOMAINS = [
    "scmp.com", "reuters.com", "bloomberg.com", "rfa.org", "voachinese.com", "bbc.com", "cnn.com", "ft.com",
    "zaobao.com", "zaobao.com.sg", "rthk.hk", "aastocks.com", "hk01.com", "now.com", "nownews.com",
    "881903.com", "metroradio.com.hk", "chinesetoday.com", "macaodaily.com", "waou.com.mo",
    "chinareviewnews.com", "crntt.com", "cna.com.tw", "udn.com", "ltn.com.tw", "ettoday.net",
    "hket.com", "takungpao.com", "wenweipo.com",
]
OVERSEAS_SOURCE_HINTS = [
    "联合早报", "南华早报", "香港电台", "香港商报", "香港01", "新城电台", "商业电台", "now新闻",
    "阿斯达克", "新华澳报", "澳门日报", "中评社", "中评网", "中央社", "联合新闻网", "自由时报",
    "路透", "彭博", "英国金融时报", "bbc", "cnn", "voa", "美国之音",
]
HK_TW_HINTS = ["hk", "hongkong", "taiwan", "tw", "macau", "mo"]


def read_text(path: str | None, inline: str) -> str:
    if path:
        return Path(path).read_text(encoding="utf-8")
    return inline.strip()


def normalize_date(text: str) -> str:
    match = re.search(r"(\d{4})[年/-](\d{1,2})[月/-](\d{1,2})日?", text)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    match = re.search(r"(\d{1,2})月(\d{1,2})日", text)
    if match:
        return f"2026-{int(match.group(1)):02d}-{int(match.group(2)):02d}"
    return DEFAULT_DATE


def clean_topic_text(text: str) -> str:
    value = unicodedata.normalize("NFKC", text or "")
    value = re.sub(r"\s+", "", value)
    value = re.sub(r"^(并)?(进一步部署|听取|研究|审议通过|审议|部署|决定|通过)", "", value)
    value = re.sub(r"(有关工作|情况汇报)$", "", value)
    return value.strip("，。；、和及")


def canonical_topic_key(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).lower()
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", text)


def resolve_topic_label(value: Any, topics: list[str]) -> str:
    text = str(value or "").strip()
    if not text or not topics or text in topics:
        return text
    key = canonical_topic_key(text)
    exact = [topic for topic in topics if canonical_topic_key(topic) == key]
    return exact[0] if len(exact) == 1 else text


def topic_display(topic: str) -> str:
    return clean_topic_text(topic).strip("《》\"'") or topic


def split_compound_approval(topic: str) -> list[str]:
    if "审议通过" not in topic:
        return [topic]
    titles = re.findall(r"《([^》]+)》", topic)
    if len(titles) <= 1:
        return [topic]
    return [f"审议通过《{title}》" for title in titles]


def infer_topics(text: str) -> list[str]:
    compact = re.sub(r"\s+", "", text or "")
    compact = re.sub(r"^.*?国务院常务会议[，,]", "", compact)
    compact = re.sub(r"(请|帮我|完成|生成|撰写|输出).*$", "", compact)
    verb_pattern = r"(进一步部署|部署|听取|研究|审议通过|审议|决定)"
    topics: list[str] = []
    matches = list(re.finditer(verb_pattern, compact))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(compact)
        raw = compact[match.start() : end].strip("，。；")
        for item in split_compound_approval(raw):
            item = item.strip("，。；")
            if item and item not in topics:
                topics.append(item)
    if topics:
        return topics
    chunks = [x.strip("，。；、 ") for x in re.split(r"[，。；;]", compact) if len(x.strip()) >= 6]
    return chunks or ["本次国务院常务会议"]


def extract_terms(text: str) -> list[str]:
    display = topic_display(text)
    terms: list[str] = []
    terms.extend(re.findall(r"《([^》]{2,30})》", text))
    terms.extend(re.findall(r"“([^”]{2,12})”", text))
    cleaned = re.sub(r"[《》“”\"'（）()]", " ", display)
    cleaned = re.sub(r"(听取|研究|审议通过|审议|通过|有关工作|情况汇报)", " ", cleaned)
    pieces = re.split(r"[、，。；和及与的\s]+", cleaned)
    for piece in pieces:
        piece = piece.strip()
        if 2 <= len(piece) <= 12 and piece not in STOPWORDS:
            terms.append(piece)
    for match in re.findall(r"[\u4e00-\u9fffA-Za-z0-9]{2,12}", cleaned):
        if 2 <= len(match) <= 12 and match not in STOPWORDS:
            terms.append(match)
    output: list[str] = []
    for term in terms:
        term = term.strip("：:、，。； ")
        if not term or term in STOPWORDS:
            continue
        if term not in output:
            output.append(term)
    return output[:8] or [display[:12]]


def keywords_for(topic: str) -> list[str]:
    display = topic_display(topic)
    terms = extract_terms(topic)
    queries = [display]
    queries.extend(terms)
    queries.extend([f"国务院常务会议 {term}" for term in terms[:4]])
    return list(dict.fromkeys(q for q in queries if q))[:10]


def safe_name(text: str) -> str:
    value = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "_", text).strip("_")
    return value[:36] or "topic"


def truthy(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "y", "是"}


def intish(value: Any) -> int:
    try:
        return int(float(str(value).replace(",", "").strip()))
    except Exception:
        return 0


def first(row: dict[str, Any], keys: list[str]) -> str:
    for key in keys:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def topic_score(title: str, content: str, topic: str) -> int:
    title_text = title or ""
    content_text = content or ""
    haystack = f"{title_text} {content_text}"
    terms = [topic_display(topic)] + extract_terms(topic)
    primary_terms = [term for term in extract_terms(topic) if term not in GENERIC_TERMS and len(term) >= 3]
    score = 0
    for term in terms:
        if not term:
            continue
        if term in title_text:
            score += 4
        elif term in content_text:
            score += 1
    if primary_terms and any(term in title_text for term in primary_terms):
        score += 3
    if "国务院常务会议" in haystack:
        score += 1
    return score


def match_topic(text: str, topics: list[str]) -> str:
    haystack = text or ""
    best_topic = topics[0] if topics else "未归类"
    best_score = -1
    for topic in topics:
        score = topic_score(haystack[:120], haystack, topic)
        if score > best_score:
            best_score = score
            best_topic = topic
    return best_topic


def force_topic_by_title(title: str, topics: list[str]) -> str:
    title = title or ""
    rules = [
        (["人工智能", "AI", "智能+"], ["人工智能"]),
        (["外贸", "贸易强国", "稳外贸", "进出口"], ["外贸", "贸易强国"]),
        (["碳达峰", "双碳", "绿色低碳"], ["碳达峰"]),
        (["国民健康", "健康规划", "医疗", "医保"], ["国民健康", "健康"]),
    ]
    for title_terms, topic_terms in rules:
        if any(term in title for term in title_terms):
            for topic in topics:
                display = topic_display(topic)
                if any(term in display for term in topic_terms):
                    return topic
    return ""


def domain_from_url(url: str) -> str:
    try:
        return urllib.parse.urlparse(url).netloc.lower()
    except Exception:
        return ""


def canonical_source_url(url: str) -> str:
    value = str(url or "").strip()
    if not value:
        return ""
    try:
        parsed = urllib.parse.urlparse(value)
        if "bing.com" in parsed.netloc.lower():
            query = urllib.parse.parse_qs(parsed.query)
            original = query.get("url", [""])[0]
            if original:
                return urllib.parse.unquote(original)
    except Exception:
        return value
    return value


def overseas_signal(row: dict[str, Any]) -> bool:
    url = canonical_source_url(first(row, URL_KEYS))
    domain = domain_from_url(url)
    source_text = " ".join(str(row.get(key) or "") for key in ["source", "author", "publisher", "site", "media_name"])
    if "中央社会主义学院" in source_text:
        return False
    if any(x in domain for x in OVERSEAS_DOMAINS):
        return True
    if any(hint in domain.split(".") for hint in HK_TW_HINTS):
        return True
    source_lower = source_text.lower()
    for hint in OVERSEAS_SOURCE_HINTS:
        hint_lower = hint.lower()
        if hint == "中央社":
            if re.search(r"(^|[（(\\s])(?:台湾)?中央社($|[）)\\s])", source_text):
                return True
            continue
        if hint_lower in source_lower:
            return True
    return False


def infer_region(row: dict[str, Any]) -> str:
    region = str(row.get("region") or "").strip().lower()
    if region in {"overseas", "foreign", "境外"}:
        return "overseas"
    if overseas_signal(row):
        return "overseas"
    source_type = str(row.get("source_type") or "").strip().lower()
    if source_type == "overseas_media":
        return "overseas"
    return "domestic"


def infer_source_type(row: dict[str, Any]) -> str:
    source_type = str(row.get("source_type") or "").strip()
    if source_type in {"overseas_public_discussion", "overseas_netizen"}:
        return "overseas_public_discussion"
    if overseas_signal(row):
        return "overseas_media"
    if source_type == "overseas_media":
        return "media"
    if source_type:
        return source_type
    if truthy(row.get("is_comment")):
        return "netizen"
    url = canonical_source_url(first(row, URL_KEYS))
    domain = domain_from_url(url)
    if any(x in domain for x in OFFICIAL_DOMAINS):
        return "official_media"
    if any(x in domain for x in CENTRAL_MEDIA_DOMAINS):
        return "central_media"
    return "media"


def infer_platform(row: dict[str, Any]) -> str:
    platform = str(row.get("platform") or "").strip()
    if platform:
        return platform
    url = canonical_source_url(first(row, URL_KEYS))
    domain = domain_from_url(url)
    if "weixin.qq.com" in domain:
        return "wechat"
    if "weibo.com" in domain:
        return "weibo"
    if "bilibili.com" in domain:
        return "bili"
    if "zhihu.com" in domain:
        return "zhihu"
    return "news" if domain else "unknown"


def normalize_sentiment_label(value: Any) -> str:
    return SENTIMENT_ALIASES.get(str(value or "").strip().lower(), "unknown")


def sentiment_metadata(row: dict[str, Any]) -> dict[str, Any]:
    source = str(row.get("sentiment_source") or row.get("label_source") or "").strip().lower()
    label = normalize_sentiment_label(row.get("sentiment") or row.get("label") or row.get("predicted_label"))
    if source in VERIFIED_SENTIMENT_SOURCES and label in {"positive", "neutral", "negative"}:
        in_denominator = truthy(row.get("in_sentiment_denominator", True))
        return {
            "sentiment": label if in_denominator else "unknown",
            "sentiment_source": source,
            "sentiment_status": "classified" if in_denominator else "excluded",
            "in_sentiment_denominator": in_denominator,
            "sentiment_confidence": row.get("sentiment_confidence") or row.get("confidence") or row.get("prediction_confidence") or "",
            "sentiment_exclusion_reason": str(row.get("sentiment_exclusion_reason") or row.get("exclusion_reason") or ""),
        }
    if source in VERIFIED_SENTIMENT_SOURCES and not truthy(row.get("in_sentiment_denominator", True)):
        return {
            "sentiment": "unknown",
            "sentiment_source": source,
            "sentiment_status": "excluded",
            "in_sentiment_denominator": False,
            "sentiment_confidence": row.get("sentiment_confidence") or row.get("confidence") or "",
            "sentiment_exclusion_reason": str(row.get("sentiment_exclusion_reason") or row.get("exclusion_reason") or "低信息、无关或不具备可判定语境"),
        }
    return {
        "sentiment": "unknown",
        "sentiment_source": "unprocessed",
        "sentiment_status": "unprocessed",
        "in_sentiment_denominator": False,
        "sentiment_confidence": "",
        "sentiment_exclusion_reason": "",
    }


def canonicalize_system_row(row: dict[str, Any]) -> dict[str, Any]:
    output = dict(row)
    aliases = {
        "id": ["ID", "数据ID", "信息ID", "文章ID", "评论ID"],
        "topic": ["子事件", "子议题", "议题", "事件名称"],
        "platform": ["平台", "渠道", "来源平台"],
        "source_type": ["来源类型", "媒体类型", "账号类型", "信息类型"],
        "region": ["地区", "境内外", "区域"],
        "source": ["来源", "媒体名称", "账号", "作者", "昵称", "用户名"],
        "title": ["标题", "文章标题", "信息标题"],
        "content": ["内容", "正文", "摘要", "评论内容", "评论原文", "文本"],
        "url": ["链接", "原文链接", "发布地址", "URL", "url"],
        "published_at": ["发布时间", "日期", "时间", "发布时刻"],
        "spread_count": ["传播量", "阅读量", "浏览量", "播放量", "热度", "在看量"],
        "comment_count": ["评论量", "评论数", "精选评论量"],
        "is_comment": ["是否评论", "评论标记"],
        "sentiment": ["情感", "情感倾向", "情感属性"],
    }
    for target, candidates in aliases.items():
        if output.get(target) not in (None, ""):
            continue
        for candidate in candidates:
            if output.get(candidate) not in (None, ""):
                output[target] = output[candidate]
                break
    platform_map = {
        "微信": "wechat",
        "微信公众号": "wechat",
        "微博": "weibo",
        "新浪微博": "weibo",
        "抖音": "douyin",
        "小红书": "xhs",
        "哔哩哔哩": "bili",
        "B站": "bili",
        "知乎": "zhihu",
        "新闻": "news",
    }
    output["platform"] = platform_map.get(str(output.get("platform") or "").strip(), output.get("platform"))
    source_type_text = str(output.get("source_type") or "").strip()
    if "评论" in source_type_text:
        output["source_type"] = "netizen_comment"
        output["is_comment"] = True
    elif "自媒体" in source_type_text:
        output["source_type"] = "self_media"
    elif "境外" in source_type_text:
        output["source_type"] = "overseas_media"
        output["region"] = "overseas"
    elif source_type_text in {"媒体", "新闻", "主流媒体", "境内媒体"}:
        output["source_type"] = "mainstream_media"
    if output.get("sentiment") not in (None, "") and not output.get("sentiment_source"):
        output["sentiment_source"] = "monitoring_system"
    return output


def normalize_sample(row: dict[str, Any], topics: list[str], idx: int) -> dict[str, Any] | None:
    row = canonicalize_system_row(row)
    content = first(row, TEXT_KEYS)
    title = first(row, TITLE_KEYS) or content[:80]
    if not title and not content:
        return None
    full_text = " ".join([str(row.get("topic") or ""), title, content])
    topic = str(row.get("topic") or "").strip() or match_topic(full_text, topics)
    topic = resolve_topic_label(topic, topics)
    forced_topic = force_topic_by_title(title, topics)
    if forced_topic:
        topic = forced_topic
    if topics and not forced_topic:
        scored_topics = [(candidate, topic_score(title, content, candidate)) for candidate in topics]
        best_topic, best_score = max(scored_topics, key=lambda item: item[1])
        current_score = topic_score(title, content, topic) if topic in topics else -1
        if best_score > current_score + 1:
            topic = best_topic
    spread = max([intish(row.get(key)) for key in SPREAD_KEYS] or [0])
    source_type = infer_source_type(row)
    platform = infer_platform(row)
    is_comment_value = truthy(row.get("is_comment")) or source_type in {"netizen", "netizen_comment", "comment"}
    sample = {
        "id": str(row.get("id") or row.get("sample_id") or row.get("raw_id") or f"s{idx:04d}"),
        "topic": topic,
        "platform": platform,
        "source_type": source_type,
        "region": infer_region(row),
        "source": first(row, AUTHOR_KEYS) or domain_from_url(first(row, URL_KEYS)) or platform,
        "title": title,
        "content": content or title,
        "url": canonical_source_url(first(row, URL_KEYS)),
        "published_at": str(row.get("published_at") or row.get("create_date_time") or row.get("time") or row.get("created_at") or "").strip(),
        "spread_count": spread,
        "comment_count": intish(row.get("comment_count")),
        "is_comment": is_comment_value,
        "quote_verified": truthy(row.get("quote_verified")),
        "evidence_mode": str(row.get("evidence_mode") or "").strip(),
        "research_query": str(row.get("research_query") or "").strip(),
        "report_order": intish(row.get("report_order")) or None,
        "comment_heading": str(row.get("comment_heading") or "").strip(),
        "topic_comment_heading": str(row.get("topic_comment_heading") or "").strip(),
        "ai_formal_include": (
            row.get("ai_formal_include")
            if isinstance(row.get("ai_formal_include"), bool)
            else truthy(row.get("ai_formal_include"))
            if str(row.get("ai_formal_include") or "").strip()
            else None
        ),
        "ai_semantic_quality": str(row.get("ai_semantic_quality") or "").strip(),
        "ai_formal_reason": str(row.get("ai_formal_reason") or "").strip(),
        "ai_comment_heading": str(row.get("ai_comment_heading") or "").strip(),
        "interpretive_verified": truthy(row.get("interpretive_verified")),
        "spread_count_display": str(row.get("spread_count_display") or row.get("read_count_display") or "").strip(),
        "comment_id": str(row.get("comment_id") or row.get("cid") or "").strip(),
        "reply_id": str(row.get("reply_id") or "").strip(),
        "parent_comment_id": str(row.get("parent_comment_id") or "").strip(),
        "raw_file": str(row.get("raw_file") or "").strip(),
        "raw_id": str(row.get("raw_id") or "").strip(),
        "origin": str(row.get("origin") or "").strip(),
        "data_origin": str(row.get("data_origin") or "").strip(),
        "in_monitoring_window": row.get("in_monitoring_window") is not False,
        "original_title": str(row.get("original_title") or row.get("title_original") or "").strip(),
        "original_content": str(row.get("original_content") or row.get("content_original") or "").strip(),
        "summary": str(row.get("summary") or "").strip(),
        "excerpt": str(row.get("excerpt") or "").strip(),
        "title_cn": str(row.get("title_cn") or row.get("translation_title_cn") or "").strip(),
        "translation_title_cn": str(row.get("translation_title_cn") or "").strip(),
        "content_cn": str(row.get("content_cn") or "").strip(),
        "translation": str(row.get("translation") or "").strip(),
        "translation_cn": str(row.get("translation_cn") or "").strip(),
        "summary_cn": str(row.get("summary_cn") or "").strip(),
        "interpretive_summary_cn": str(row.get("interpretive_summary_cn") or "").strip(),
        "overseas_category": str(row.get("overseas_category") or row.get("category") or "").strip(),
        "speaker": str(row.get("speaker") or "").strip(),
        "expert": str(row.get("expert") or "").strip(),
        "expert_name": str(row.get("expert_name") or "").strip(),
        "quoted_person": str(row.get("quoted_person") or "").strip(),
        "author_name": str(row.get("author_name") or "").strip(),
        "attribution": str(row.get("attribution") or "").strip(),
        "attribution_status": str(row.get("attribution_status") or "").strip(),
        "topic_hits": list(row.get("topic_hits") or []),
        "quality_flags": [],
    }
    sample.update(sentiment_metadata(row))
    if len(sample["content"]) < 10:
        sample["quality_flags"].append("short_text")
    if platform in {"bing_news", "bing"} or "bing.com" in domain_from_url(sample["url"]):
        sample["quality_flags"].append("aggregated_search_link")
    if sample.get("url") and domain_from_url(sample["url"]) in {"www.bing.com", "bing.com", "cn.bing.com"}:
        sample["quality_flags"].append("not_original_source_url")
    if re.search(r"早报|合集|汇总|一图读懂|每日", sample["title"]):
        sample["quality_flags"].append("digest")
    agenda_markers = ["人工智能", "外贸", "贸易强国", "碳达峰", "国民健康"]
    if sum(1 for marker in agenda_markers if marker in f"{sample['title']} {sample['content']}") >= 3:
        sample["quality_flags"].append("full_agenda_digest")
    sample.update(viewpoint_attribution(sample))
    return sample


def load_samples(path: str | None, topics: list[str]) -> list[dict[str, Any]]:
    if not path:
        return []
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(path)
    if source.suffix.lower() == ".json":
        data = json.loads(source.read_text(encoding="utf-8"))
        rows = data if isinstance(data, list) else data.get("samples", [])
    elif source.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(source, data_only=True, read_only=True)
        rows = []
        for worksheet in workbook.worksheets:
            values = list(worksheet.iter_rows(values_only=True))
            header_index = next(
                (idx for idx, row in enumerate(values) if sum(1 for value in row if str(value or "").strip()) >= 2),
                None,
            )
            if header_index is None:
                continue
            headers = [str(value or "").strip() for value in values[header_index]]
            for value_row in values[header_index + 1 :]:
                row = {headers[col]: value for col, value in enumerate(value_row) if col < len(headers) and headers[col]}
                if any(value not in (None, "") for value in row.values()):
                    row.setdefault("source_sheet", worksheet.title)
                    rows.append(row)
        workbook.close()
    else:
        with source.open("r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
    samples = []
    for idx, row in enumerate(rows, 1):
        if isinstance(row, dict):
            sample = normalize_sample(row, topics, idx)
            if sample:
                samples.append(sample)
    return samples


def load_comment_handoffs(paths: list[str] | None, topics: list[str]) -> list[dict[str, Any]]:
    """Load the strict, traceable comment evidence emitted by sentiment stage."""
    samples: list[dict[str, Any]] = []
    for path in paths or []:
        loaded = load_samples(path, topics)
        invalid = [row for row in loaded if not is_actual_platform_comment(row)]
        if invalid:
            raise ValueError(
                f"评论报告交接文件存在{len(invalid)}条不可追溯记录：{path}。"
                "必须包含评论ID、原链接、is_comment=true、quote_verified=true和"
                "evidence_mode=verbatim_public_comment。"
            )
        samples.extend(loaded)
    return samples


def load_sentiment_summary(path: str | None) -> dict[str, Any]:
    if not path:
        return {}
    source = Path(path).expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(path)
    payload = json.loads(source.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Sentiment summary must be a JSON object")
    return payload


def load_analysis_bundle(path: str | None) -> dict[str, Any]:
    if not path:
        return {}
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(path)
    data = json.loads(source.read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else {}


def analysis_bundle_rows(bundle: dict[str, Any]) -> list[dict[str, Any]]:
    rows = [dict(row) for row in bundle.get("samples", []) if isinstance(row, dict)]
    for row in rows:
        row.setdefault("data_origin", "public_web_supplement")
    comments = (bundle.get("comments") or {}).get("selected") or bundle.get("comments") or []
    if isinstance(comments, list):
        for row in comments:
            if not isinstance(row, dict):
                continue
            item = dict(row)
            item.setdefault("source_type", "netizen_comment")
            item.setdefault("is_comment", True)
            item.setdefault("data_origin", "public_web_supplement")
            rows.append(item)
    selected_candidate_ids = {
        str(evidence.get("candidate_id") or "").strip()
        for item in (bundle.get("viewpoints") or {}).get("by_topic", [])
        if isinstance(item, dict)
        for cluster in item.get("clusters") or []
        if isinstance(cluster, dict)
        for evidence in cluster.get("evidence") or []
        if isinstance(evidence, dict) and str(evidence.get("candidate_id") or "").strip()
    }
    research = ((bundle.get("research_audit") or {}).get("domestic_media_research") or {})
    for pool in research.get("candidate_pool_by_topic") or []:
        if not isinstance(pool, dict):
            continue
        topic = str(pool.get("topic") or "")
        for candidate in pool.get("candidates") or []:
            if not isinstance(candidate, dict) or candidate.get("decision") != "eligible":
                continue
            candidate_id = str(candidate.get("candidate_id") or "").strip()
            if candidate_id and candidate_id in selected_candidate_ids:
                continue
            item = dict(candidate)
            item.setdefault("id", candidate_id)
            item.setdefault("topic", topic)
            item.setdefault("source_type", "mainstream_media")
            item.setdefault("content", item.get("content_summary") or item.get("source_excerpt") or "")
            item.setdefault("data_origin", "public_web_candidate_pool")
            item.setdefault("candidate_pool_status", "eligible_not_selected_for_formal")
            rows.append(item)
    for item in (bundle.get("viewpoints") or {}).get("by_topic", []):
        if not isinstance(item, dict):
            continue
        topic = str(item.get("topic") or "")
        for cluster in item.get("clusters") or []:
            if not isinstance(cluster, dict):
                continue
            for evidence in cluster.get("evidence") or []:
                if not isinstance(evidence, dict):
                    continue
                row = dict(evidence)
                row.setdefault("topic", topic)
                row.setdefault("source_type", "mainstream_media")
                row.setdefault("data_origin", "public_web_supplement")
                rows.append(row)
    return rows


def read_sentiment_results(path: str | None) -> list[dict[str, Any]]:
    if not path:
        return []
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(path)
    if source.suffix.lower() == ".json":
        data = json.loads(source.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return [row for row in data if isinstance(row, dict)]
        for key in ["results", "predictions", "samples", "rows"]:
            rows = data.get(key)
            if isinstance(rows, list):
                return [row for row in rows if isinstance(row, dict)]
        return []
    with source.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def apply_sentiment_results(samples: list[dict[str, Any]], path: str | None) -> tuple[int, list[str]]:
    rows = read_sentiment_results(path)
    if not rows:
        return 0, (["未取得可回灌的情感分析结果。"] if path else [])
    sample_by_id = {str(sample.get("id") or ""): sample for sample in samples if sample.get("id")}
    matched = 0
    unmatched = 0
    for row in rows:
        sample_id = first(row, ["sample_id", "source_id", "id", "raw_id"])
        if not sample_id:
            source_value = str(row.get("source") or "").strip()
            if source_value in sample_by_id:
                sample_id = source_value
        sample = sample_by_id.get(sample_id)
        if not sample or not is_comment(sample):
            explicit_denominator = row.get("in_sentiment_denominator")
            explained_exclusion = (
                explicit_denominator not in {None, ""}
                and not truthy(explicit_denominator)
                and not truthy(row.get("needs_review"))
                and bool(str(row.get("exclusion_reason") or "").strip())
            )
            if explained_exclusion:
                continue
            unmatched += 1
            continue
        needs_review = truthy(row.get("needs_review"))
        explicit_denominator = row.get("in_sentiment_denominator")
        in_denominator = truthy(explicit_denominator) if explicit_denominator not in {None, ""} else True
        label = normalize_sentiment_label(row.get("label") or row.get("predicted_label") or row.get("sentiment"))
        label_source = str(row.get("label_source") or row.get("sentiment_source") or "").strip().lower()
        if not label_source:
            label_source = "active_learning_classifier" if row.get("predicted_label") else "ai_reviewed"
        if needs_review:
            sample.update(
                {
                    "sentiment": "unknown",
                    "sentiment_source": label_source,
                    "sentiment_status": "unprocessed",
                    "in_sentiment_denominator": False,
                    "sentiment_confidence": row.get("prediction_confidence") or row.get("confidence") or "",
                    "sentiment_exclusion_reason": "模型低置信度，仍需AI或人工复核",
                }
            )
        elif not in_denominator:
            sample.update(
                {
                    "sentiment": "unknown",
                    "sentiment_source": label_source,
                    "sentiment_status": "excluded",
                    "in_sentiment_denominator": False,
                    "sentiment_confidence": row.get("prediction_confidence") or row.get("confidence") or "",
                    "sentiment_exclusion_reason": str(row.get("exclusion_reason") or "无关、低信息、垃圾文本或语境不足"),
                }
            )
        elif label in {"positive", "neutral", "negative"}:
            sample.update(
                {
                    "sentiment": label,
                    "sentiment_source": label_source,
                    "sentiment_status": "classified",
                    "in_sentiment_denominator": True,
                    "sentiment_confidence": row.get("prediction_confidence") or row.get("confidence") or "",
                    "sentiment_exclusion_reason": "",
                }
            )
        else:
            sample.update(sentiment_metadata({}))
        matched += 1
    gaps = []
    if unmatched:
        gaps.append(f"情感结果中有{unmatched}条无法按 sample_id 对应到网民评论，未予回灌。")
    return matched, gaps


def reset_inherited_sentiment(samples: list[dict[str, Any]]) -> None:
    """Prevent labels embedded in sample inputs from masquerading as this run's analysis."""

    for sample in samples:
        inherited = {
            "sentiment": sample.get("sentiment"),
            "sentiment_source": sample.get("sentiment_source"),
            "sentiment_status": sample.get("sentiment_status"),
            "in_sentiment_denominator": sample.get("in_sentiment_denominator"),
            "sentiment_confidence": sample.get("sentiment_confidence"),
            "sentiment_exclusion_reason": sample.get("sentiment_exclusion_reason"),
        }
        if any(value not in {None, "", False, "unknown", "unprocessed"} for value in inherited.values()):
            sample["inherited_sentiment_reference"] = inherited
        sample.update(sentiment_metadata({}))


def strip_html(value: str) -> str:
    value = re.sub(r"<[^>]+>", " ", value or "")
    return html.unescape(re.sub(r"\s+", " ", value)).strip()


def fetch_bing_news(query: str, limit: int, timeout: int = 15) -> list[dict[str, Any]]:
    encoded = urllib.parse.urlencode({"q": query, "format": "rss"})
    url = f"https://www.bing.com/news/search?{encoded}"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        xml_data = resp.read()
    root = ET.fromstring(xml_data)
    rows: list[dict[str, Any]] = []
    for item in root.findall(".//item")[:limit]:
        title = item.findtext("title") or ""
        link = item.findtext("link") or ""
        desc = strip_html(item.findtext("description") or "")
        pub_date = item.findtext("pubDate") or ""
        source = item.findtext("{*}Source") or item.findtext("{*}source") or domain_from_url(link)
        rows.append(
            {
                "platform": "bing_news",
                "source_type": "media",
                "source": source,
                "title": strip_html(title),
                "content": desc or strip_html(title),
                "url": link,
                "published_at": pub_date,
                "spread_count": 0,
                "comment_count": 0,
                "is_comment": "false",
                "raw_query": query,
            }
        )
    return rows


def fetch_bing_web(query: str, limit: int, timeout: int = 15) -> list[dict[str, Any]]:
    encoded = urllib.parse.urlencode({"q": query})
    url = f"https://www.bing.com/search?{encoded}"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        html_text = resp.read().decode("utf-8", errors="replace")
    rows: list[dict[str, Any]] = []
    blocks = re.findall(r'<li class="b_algo".*?</li>', html_text, flags=re.S)
    for block in blocks[:limit]:
        link_match = re.search(r'<h2[^>]*>\s*<a href="([^"]+)"[^>]*>(.*?)</a>', block, flags=re.S)
        if not link_match:
            continue
        link = html.unescape(link_match.group(1))
        title = strip_html(link_match.group(2))
        desc_match = re.search(r"<p[^>]*>(.*?)</p>", block, flags=re.S)
        desc = strip_html(desc_match.group(1)) if desc_match else ""
        if not title or not link.startswith("http"):
            continue
        rows.append(
            {
                "platform": "web",
                "source_type": "media",
                "source": domain_from_url(link),
                "title": title,
                "content": desc or title,
                "url": link,
                "published_at": "",
                "spread_count": 0,
                "comment_count": 0,
                "is_comment": "false",
                "raw_query": query,
            }
        )
    return rows


def sample_relevance(row: dict[str, Any], topic: str) -> int:
    text = f"{row.get('title', '')} {row.get('content', '')}"
    terms = [topic_display(topic)] + extract_terms(topic)
    primary_terms = [term for term in extract_terms(topic) if term not in GENERIC_TERMS and len(term) >= 3]
    if primary_terms and not any(term in text for term in primary_terms):
        return 0
    score = sum(2 for term in terms[:4] if term and term in text)
    score += sum(1 for term in terms[4:] if term and term in text)
    if "国务院常务会议" in text:
        score += 2
    return score


def collect_public_web_samples(topics: list[str], per_topic: int, timeout: int) -> tuple[list[dict[str, Any]], list[str]]:
    rows: list[dict[str, Any]] = []
    gaps: list[str] = []
    seen_keys: set[tuple[str, str]] = set()
    for topic in topics:
        topic_rows = 0
        display = topic_display(topic)
        terms = extract_terms(topic)
        queries = [
            f"国务院常务会议 {display}",
            f"李强 国务院常务会议 {display}",
            f"国务院常务会议 {topic}",
        ]
        queries.extend(f"国务院常务会议 {term}" for term in terms[:4])
        candidates: list[dict[str, Any]] = []
        for query in list(dict.fromkeys(queries)):
            try:
                for row in fetch_bing_news(query, max(2, per_topic), timeout):
                    url = row.get("url") or ""
                    key = (url, topic)
                    if url and key in seen_keys:
                        continue
                    seen_keys.add(key)
                    row["topic"] = topic
                    row["id"] = f"web_{len(rows)+1:04d}"
                    row["_relevance"] = sample_relevance(row, topic)
                    candidates.append(row)
            except Exception as exc:
                gaps.append(f"公开新闻检索失败：{query}（{type(exc).__name__}: {exc}）")
            try:
                for row in fetch_bing_web(query, max(3, per_topic), timeout):
                    url = row.get("url") or ""
                    key = (url, topic)
                    if url and key in seen_keys:
                        continue
                    seen_keys.add(key)
                    row["topic"] = topic
                    row["id"] = f"web_{len(rows)+1:04d}"
                    row["_relevance"] = sample_relevance(row, topic)
                    candidates.append(row)
            except Exception as exc:
                gaps.append(f"公开网页检索失败：{query}（{type(exc).__name__}: {exc}）")
        candidates.sort(key=lambda row: row.get("_relevance", 0), reverse=True)
        for row in candidates:
            if row.get("_relevance", 0) <= 0:
                continue
            rows.append(row)
            topic_rows += 1
            if topic_rows >= per_topic:
                break
        if topic_rows == 0:
            gaps.append(f"公开新闻检索未获得可归入“{display}”的样本。")
    normalized = []
    for idx, row in enumerate(rows, 1):
        sample = normalize_sample(row, topics, idx)
        if sample:
            sample["quality_flags"].append("public_web_sample")
            normalized.append(sample)
    return normalized, gaps


def dedupe_samples(samples: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen = set()
    output = []
    for sample in samples:
        if is_comment(sample):
            # Comments sharing one article URL or identical wording are distinct
            # voices.  Deduplicate only their stable comment/reply identity.
            identity = (
                sample.get("comment_id")
                or sample.get("reply_id")
                or sample.get("raw_id")
                or sample.get("id")
            )
            key = ("comment", sample.get("platform"), identity)
        else:
            key = (sample.get("topic"), sample.get("url")) if sample.get("url") else (sample.get("topic"), sample.get("platform"), sample.get("source"), sample.get("title"), sample.get("content", "")[:80])
        if key in seen:
            sample["quality_flags"].append("duplicate")
            continue
        seen.add(key)
        output.append(sample)
    return output


def expand_cross_topic_samples(samples: list[dict[str, Any]], topics: list[str], limit_per_topic: int = 3) -> list[dict[str, Any]]:
    expanded = list(samples)
    counts = Counter(str(sample.get("topic") or "") for sample in samples)
    for topic in topics:
        if counts.get(topic, 0) > 0:
            continue
        added = 0
        for sample in samples:
            text = f"{sample.get('title', '')} {sample.get('content', '')}"
            if sample_relevance(sample, topic) <= 0:
                continue
            clone = dict(sample)
            clone["id"] = f"{sample.get('id', 'sample')}_topic_{len(expanded)+1}"
            clone["topic"] = topic
            flags = list(clone.get("quality_flags") or [])
            flags.append("cross_topic_source")
            clone["quality_flags"] = list(dict.fromkeys(flags))
            clone["content"] = clone.get("content") or text
            expanded.append(clone)
            added += 1
            if added >= limit_per_topic:
                break
    return expanded


def is_comment(sample: dict[str, Any]) -> bool:
    return bool(sample.get("is_comment")) or sample.get("source_type") in {"netizen", "netizen_comment", "comment", "overseas_public_discussion", "overseas_netizen"}


def is_actual_platform_comment(sample: dict[str, Any]) -> bool:
    """Accept comment/reply rows only; public posts and summaries are not comments."""
    if not is_comment(sample) or not str(sample.get("url") or "").strip():
        return False
    mode = str(sample.get("evidence_mode") or "").strip().lower()
    if mode in {"public_discussion_summary", "verbatim_public_post", "public_post", "post"}:
        return False
    raw_file = str(sample.get("raw_file") or "").lower()
    has_comment_identity = bool(
        sample.get("comment_id")
        or sample.get("reply_id")
        or sample.get("parent_comment_id")
        or any(marker in raw_file for marker in ("comment", "reply", "评论", "回复"))
    )
    return bool(
        sample.get("quote_verified")
        and mode in {"verbatim_public_comment", "verbatim_comment", "platform_comment", "platform_reply"}
        and has_comment_identity
    )


def is_overseas(sample: dict[str, Any]) -> bool:
    region = str(sample.get("region") or "").strip().lower()
    source_type = str(sample.get("source_type") or "").strip().lower()
    sample_id = str(sample.get("id") or "").strip().lower()
    return bool(
        region in {"overseas", "foreign", "境外"}
        or source_type.startswith("overseas")
        or source_type in {"境外媒体", "境外自媒体"}
        or sample_id.startswith("system-overseas-")
    )


def source_bucket(sample: dict[str, Any]) -> str:
    if is_overseas(sample) and is_comment(sample):
        return "overseas_comments"
    if is_overseas(sample):
        return "overseas_media"
    if is_comment(sample):
        return "comments"
    source_type = str(sample.get("source_type") or "unknown")
    if source_type in {"central_media", "mainstream_media", "commercial_media", "official_media", "media", "news"}:
        return "domestic_media"
    return "self_media"


def source_label(value: str) -> str:
    return SOURCE_TYPE_LABELS.get(value, value or "未识别")


def platform_label(value: str) -> str:
    return PLATFORM_LABELS.get(value, value or "未识别")


def top_items(items: list[dict[str, Any]], limit: int = 5) -> list[dict[str, Any]]:
    return sorted(items, key=lambda x: (int(x.get("spread_count") or 0), int(x.get("comment_count") or 0), x.get("published_at") or ""), reverse=True)[:limit]


def sentiment_counts(items: list[dict[str, Any]]) -> dict[str, int]:
    counter = Counter(
        str(x.get("sentiment") or "unknown")
        if truthy(x.get("in_sentiment_denominator"))
        else "unknown"
        for x in items
    )
    return {k: counter.get(k, 0) for k in ["positive", "neutral", "negative", "unknown"]}


def sentiment_text(counts: dict[str, int]) -> str:
    return f"正面{counts.get('positive', 0)}条、中性{counts.get('neutral', 0)}条、负面{counts.get('negative', 0)}条、未判定{counts.get('unknown', 0)}条"


def viewpoint_attribution(sample: dict[str, Any]) -> dict[str, Any]:
    explicit = str(sample.get("attribution") or "").strip()
    explicit_status = str(sample.get("attribution_status") or "").strip()
    if explicit:
        return {
            "attribution": explicit,
            "attribution_status": explicit_status or "named_person",
            "needs_attribution_review": explicit_status in {"missing", "anonymous_expert"},
        }
    for key in ("speaker", "expert", "expert_name", "quoted_person", "author_name"):
        value = str(sample.get(key) or "").strip()
        if value:
            return {"attribution": value, "attribution_status": "named_person", "needs_attribution_review": False}
    text = f"{sample.get('title', '')} {sample.get('content', '')}"
    match = re.search(r"([\u4e00-\u9fff]{2,4})(?:认为|表示|指出|建议|强调|称)", text)
    generic_attributions = {"专家", "媒体", "学者", "业内", "机构", "记者", "有关负责人", "舆论", "报道", "文章", "消息", "会议"}
    if match and not any(term in match.group(1) for term in generic_attributions):
        return {"attribution": match.group(1), "attribution_status": "named_person", "needs_attribution_review": False}
    source = str(sample.get("source") or sample.get("platform") or "").strip()
    if source:
        return {"attribution": source, "attribution_status": "media_only", "needs_attribution_review": False}
    return {"attribution": "来源主体待核", "attribution_status": "missing", "needs_attribution_review": True}


def row_sample(sample: dict[str, Any]) -> dict[str, Any]:
    attribution = viewpoint_attribution(sample)
    return {
        "id": sample.get("id"),
        "sample_id": sample.get("id"),
        "topic": sample.get("topic"),
        "source": sample.get("source"),
        "source_type": source_label(str(sample.get("source_type") or "")),
        "platform": platform_label(str(sample.get("platform") or "")),
        "region": sample.get("region"),
        "is_comment": sample.get("is_comment", False),
        "origin": sample.get("origin") or "",
        "data_origin": sample.get("data_origin") or "",
        "title": sample.get("title"),
        "content": sample.get("content"),
        "url": sample.get("url"),
        "published_at": sample.get("published_at"),
        "spread_count": sample.get("spread_count"),
        "comment_count": sample.get("comment_count"),
        "sentiment": sample.get("sentiment"),
        "sentiment_source": sample.get("sentiment_source"),
        "sentiment_status": sample.get("sentiment_status"),
        "in_sentiment_denominator": sample.get("in_sentiment_denominator"),
        "sentiment_confidence": sample.get("sentiment_confidence"),
        "sentiment_exclusion_reason": sample.get("sentiment_exclusion_reason"),
        "quote_verified": sample.get("quote_verified", False),
        "evidence_mode": sample.get("evidence_mode", ""),
        "research_query": sample.get("research_query", ""),
        "candidate_id": sample.get("candidate_id", ""),
        "candidate_pool_status": sample.get("candidate_pool_status", ""),
        "discovery_route": sample.get("discovery_route", ""),
        "selection_reason": sample.get("selection_reason", ""),
        "decision_reason": sample.get("decision_reason", ""),
        "report_order": sample.get("report_order"),
        "comment_heading": sample.get("comment_heading", ""),
        "topic_comment_heading": sample.get("topic_comment_heading", ""),
        "ai_formal_include": sample.get("ai_formal_include"),
        "ai_semantic_quality": sample.get("ai_semantic_quality", ""),
        "ai_formal_reason": sample.get("ai_formal_reason", ""),
        "ai_comment_heading": sample.get("ai_comment_heading", ""),
        "interpretive_verified": sample.get("interpretive_verified", False),
        "spread_count_display": sample.get("spread_count_display", ""),
        "comment_id": sample.get("comment_id", ""),
        "reply_id": sample.get("reply_id", ""),
        "parent_comment_id": sample.get("parent_comment_id", ""),
        "raw_file": sample.get("raw_file", ""),
        "raw_id": sample.get("raw_id", ""),
        "in_monitoring_window": sample.get("in_monitoring_window", True),
        "original_title": sample.get("original_title") or sample.get("title_original") or "",
        "original_content": sample.get("original_content") or sample.get("content_original") or "",
        "summary": sample.get("summary") or "",
        "excerpt": sample.get("excerpt") or "",
        "title_cn": sample.get("title_cn") or sample.get("translation_title_cn") or "",
        "content_cn": sample.get("content_cn") or "",
        "translation": sample.get("translation") or "",
        "translation_cn": sample.get("translation_cn") or "",
        "summary_cn": sample.get("summary_cn") or "",
        "interpretive_summary_cn": sample.get("interpretive_summary_cn") or "",
        "overseas_category": sample.get("overseas_category") or sample.get("category") or "",
        **attribution,
    }


def build_topic_stats(topics: list[str], samples: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for topic in topics:
        items = [x for x in samples if x.get("topic") == topic]
        comment_items = [x for x in items if is_comment(x)]
        rows.append(
            {
                "topic": topic,
                "display": topic_display(topic),
                "total_samples": len(items),
                "spread_count": sum(int(x.get("spread_count") or 0) for x in items),
                "domestic_media": sum(1 for x in items if source_bucket(x) == "domestic_media"),
                "self_media": sum(1 for x in items if source_bucket(x) == "self_media"),
                "comments": sum(1 for x in items if source_bucket(x) == "comments"),
                "overseas_media": sum(1 for x in items if source_bucket(x) == "overseas_media"),
                "sentiment": sentiment_counts(comment_items),
                "sentiment_authority": "large_scale_sentiment_analysis",
                "top_samples": [row_sample(x) for x in top_items(items, 3)],
            }
        )
    return rows


def extract_viewpoint(sample: dict[str, Any]) -> str:
    text = f"{sample.get('title', '')}。{sample.get('content', '')}"
    if any(word in text for word in POSITIVE_WORDS):
        return "认为相关部署有利于稳定预期、完善政策体系或推动行业发展"
    if any(word in text for word in NEGATIVE_WORDS):
        return "关注政策落地、配套机制、治理风险或企业实际压力"
    if any(word in text for word in ["建议", "应当", "需要", "期待", "希望"]):
        return "建议进一步细化执行安排并回应市场和公众关切"
    return "以转述会议部署和政策信号为主"


def build_viewpoints(topics: list[str], samples: list[dict[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {"by_topic": []}
    for topic in topics:
        items = [
            x
            for x in samples
            if x.get("topic") == topic
            and source_bucket(x) in {"domestic_media", "self_media"}
            and not {"duplicate", "short_text", "aggregated_search_link", "full_agenda_digest"}.intersection(set(x.get("quality_flags", [])))
        ]
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        # Native fallback must not silently turn a saturated candidate pool into
        # a short list.  Every non-duplicate eligible domestic sample is
        # grouped; ranking controls order only, never inclusion.
        for item in top_items(items, len(items)):
            grouped[extract_viewpoint(item)].append(item)
        result["by_topic"].append(
            {
                "topic": topic,
                "clusters": [
                    {"summary": summary, "evidence": [row_sample(x) for x in evidence]}
                    for summary, evidence in grouped.items()
                ],
            }
        )
    return result


def sanitize_domestic_viewpoints(viewpoints: dict[str, Any], samples: list[dict[str, Any]]) -> dict[str, Any]:
    """Keep the domestic-opinion section strictly domestic, including imported bundles."""
    by_url: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for sample in samples:
        url = str(sample.get("url") or "").strip()
        if url:
            by_url[url].append(sample)

    cleaned: dict[str, Any] = {"by_topic": []}
    for item in viewpoints.get("by_topic") or []:
        if not isinstance(item, dict):
            continue
        clusters = []
        for cluster in item.get("clusters") or []:
            if not isinstance(cluster, dict):
                continue
            evidence_rows = []
            for evidence in cluster.get("evidence") or []:
                if not isinstance(evidence, dict):
                    continue
                url = str(evidence.get("url") or "").strip()
                matches = by_url.get(url, []) if url else []
                if is_overseas(evidence) or (matches and all(is_overseas(row) for row in matches)):
                    continue
                evidence_rows.append(evidence)
            if evidence_rows:
                cleaned_cluster = dict(cluster)
                cleaned_cluster["evidence"] = evidence_rows
                clusters.append(cleaned_cluster)
        cleaned_item = dict(item)
        cleaned_item["clusters"] = clusters
        cleaned["by_topic"].append(cleaned_item)
    return cleaned


LOW_VALUE_COMMENT_PATTERNS = [
    r"^[\[\]赞666888\s。！!]+$",
    r"^(支持|点赞|转发微博|了解|说得好|很好|加油)[。！!]*$",
]
IRRELEVANT_COMMENT_TERMS = [
    "股票", "基金", "本金", "亏光", "割肉", "成交量", "科大讯飞", "中科曙光",
    "蓝色光标", "华胜天成", "润泽科技", "起飞", "牛逼",
]


def comment_is_low_value(text: str) -> bool:
    compact = re.sub(r"\s+", "", text or "")
    if not compact:
        return True
    if any(re.search(pattern, compact) for pattern in LOW_VALUE_COMMENT_PATTERNS):
        return True
    if any(term in compact for term in IRRELEVANT_COMMENT_TERMS):
        return True
    return False


def comment_rank_score(sample: dict[str, Any]) -> int:
    text = str(sample.get("content") or sample.get("title") or "")
    topic = str(sample.get("topic") or "")
    score = int(sample.get("spread_count") or 0) + int(sample.get("comment_count") or 0) * 5
    score += min(len(text), 120)
    if topic_score("", text, topic) > 0:
        score += 200
    if sample.get("sentiment") in {"positive", "negative"}:
        score += 80
    return score


def select_comments(samples: list[dict[str, Any]], limit: int = 0, *, overseas: bool = False) -> list[dict[str, Any]]:
    comments_by_topic: dict[str, list[dict[str, Any]]] = defaultdict(list)
    seen_comment_ids = set()
    for item in samples:
        if not is_actual_platform_comment(item):
            continue
        if item.get("sentiment_status") == "excluded" or item.get("in_sentiment_denominator") is False:
            continue
        if is_overseas(item) != overseas:
            continue
        text = str(item.get("content") or item.get("title") or "").strip()
        # A reviewed short stance such as “支持” or “反对” is valid evidence.
        # Apply deterministic low-value filtering only before denominator review.
        if comment_is_low_value(text) and not truthy(item.get("in_sentiment_denominator")):
            item["quality_flags"].append("short_comment")
            continue
        key = (
            item.get("platform"),
            item.get("comment_id") or item.get("reply_id") or item.get("raw_id") or item.get("id"),
        )
        if key in seen_comment_ids:
            continue
        seen_comment_ids.add(key)
        comments_by_topic[str(item.get("topic") or "未分类")].append(item)
    for rows in comments_by_topic.values():
        rows.sort(key=comment_rank_score, reverse=True)
    selected: list[dict[str, Any]] = []
    while any(comments_by_topic.values()) and (limit <= 0 or len(selected) < limit):
        for topic in list(comments_by_topic):
            rows = comments_by_topic[topic]
            if not rows:
                continue
            selected.append(rows.pop(0))
            if limit > 0 and len(selected) >= limit:
                break
    return [row_sample(x) for x in selected]


def build_hotwords(topics: list[str], samples: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    seen = set()
    for topic in topics:
        display = topic_display(topic)
        text = " ".join(f"{x.get('title', '')} {x.get('content', '')}" for x in samples if x.get("topic") == topic)
        for word in extract_terms(topic):
            if word == display or display.startswith(word) and len(word) > 8:
                continue
            if len(word) > 10 or word in GENERIC_TERMS or word in {"情况汇报", "有关工作", "行动方案"}:
                continue
            count = text.count(word) if text else 0
            sample_count = sum(1 for x in samples if x.get("topic") == topic and word in f"{x.get('title', '')}{x.get('content', '')}")
            if count <= 0 and word in topic:
                count = 1
            if count <= 0:
                continue
            key = (topic, word)
            if key in seen:
                continue
            seen.add(key)
            example = ""
            for item in samples:
                if item.get("topic") == topic and word in f"{item.get('title', '')}{item.get('content', '')}":
                    example = str(item.get("title") or item.get("content") or "")[:80]
                    break
            rows.append({"topic": topic, "word": word, "count": count, "sample_count": sample_count, "example": example})
    return sorted(rows, key=lambda x: (x["sample_count"], x["count"]), reverse=True)


def hotword_example(item: dict[str, Any]) -> str:
    """Return a report-safe example for both legacy and reviewed hotword schemas."""
    value = item.get("example")
    if not value:
        titles = list(item.get("sample_titles") or [])
        value = titles[0] if titles else ""
    return str(value or "")[:80]


def build_evidence_hotwords(
    topics: list[str],
    samples: list[dict[str, Any]],
    system_candidates: list[str] | None = None,
) -> list[dict[str, Any]]:
    """Build hotwords only from traceable domestic media/self-media and real comments."""
    from cwh_hotword_pipeline import deterministic_selection, generate_candidates, score_hotword_evidence

    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    evidence_universe = []
    for sample in samples:
        bucket = source_bucket(sample)
        if bucket not in {"domestic_media", "self_media", "comments"}:
            continue
        if bucket == "comments" and not is_actual_platform_comment(sample):
            continue
        evidence_universe.append(sample)

    topic_aliases = [list(dict.fromkeys([topic_display(topic), *extract_terms(topic)])) for topic in topics]
    documents = [
        {
            "title": sample.get("title") or "",
            "content": sample.get("content") or "",
            "source": sample.get("source") or sample.get("platform") or "未知来源",
            "account": sample.get("source") or sample.get("platform") or "未知来源",
            "url": sample.get("url") or "",
        }
        for sample in evidence_universe
    ]
    generated = generate_candidates(documents, [topic_display(topic) for topic in topics], topic_aliases)
    selected = deterministic_selection(generated, 65)
    reviewed_candidates: dict[str, dict[str, Any]] = {}
    workbook_candidates: list[str] = []
    raw_candidates: Any = system_candidates or []
    if isinstance(raw_candidates, dict):
        raw_candidates = raw_candidates.get("selected") or []
    for raw in raw_candidates:
        if isinstance(raw, dict):
            word = str(raw.get("term") or raw.get("word") or "").strip()
            if word:
                reviewed_candidates[word] = raw
        else:
            word = str(raw or "").strip()
            if word:
                workbook_candidates.append(word)
    agenda_candidates = list(dict.fromkeys(term for topic in topics for term in extract_terms(topic) if term))
    preferred_candidates = list(dict.fromkeys([*reviewed_candidates, *workbook_candidates, *agenda_candidates]))
    generated_candidates = [str(item.get("term") or "") for item in selected]
    candidate_words = [*preferred_candidates, *generated_candidates]
    comments_available = any(source_bucket(sample) == "comments" for sample in evidence_universe)
    fragment_blacklist = {
        "国务", "务院", "院常", "常务", "常务会", "会议", "召开", "主持", "研究", "听取", "审议", "通过",
        "有关", "情况", "工作", "汇报", "部署", "进一步", "李强", "国常会",
    }
    fragment_substrings = {"国务", "务院", "院常", "常务", "会议", "召开", "主持", "李强", "国常会"}

    for topic in topics:
        evidence_rows = []
        for sample in evidence_universe:
            if sample.get("topic") != topic and topic not in (sample.get("topic_hits") or []):
                continue
            evidence_rows.append(sample)
        for word in dict.fromkeys(candidate_words):
            if hotword_topic(word, topics) != topic:
                continue
            if (
                len(word) > 10
                or word in GENERIC_TERMS
                or word in fragment_blacklist
                or any(fragment in word for fragment in fragment_substrings)
                or word in {"情况汇报", "有关工作", "行动方案"}
                or word.startswith(("究", "取", "议通", "一步"))
                or word.endswith(("有关", "情况", "汇报", "利用有", "建设情", "进展情"))
                or bool(re.search(r"\d", word))
                or (
                    word not in preferred_candidates
                    and len(word) <= 4
                    and any(word in preferred and len(preferred) > len(word) for preferred in preferred_candidates)
                )
                or (
                    word not in preferred_candidates
                    and any(token in word for token in ("部署", "听取", "审议", "通过", "进一步"))
                )
                or (
                    word not in preferred_candidates
                    and word.endswith(("救", "工", "建", "利", "用", "进", "情", "汇"))
                )
            ):
                continue
            aliases = [word]
            reviewed = reviewed_candidates.get(word) or {}
            aliases.extend(str(alias).strip() for alias in (reviewed.get("aliases") or []) if str(alias).strip())
            matched = [
                sample
                for sample in evidence_rows
                if any(alias in f"{sample.get('title', '')}{sample.get('content', '')}" for alias in aliases)
            ]
            if not matched:
                continue
            key = (topic, word)
            if key in seen:
                continue
            seen.add(key)
            media_rows = [sample for sample in matched if source_bucket(sample) in {"domestic_media", "self_media"}]
            comment_rows = [sample for sample in matched if source_bucket(sample) == "comments"]
            document_count = len(
                {
                    canonical_source_url(str(sample.get("url") or ""))
                    or str(sample.get("id") or "")
                    or f"{sample.get('source', '')}|{sample.get('title', '')}|{sample.get('content', '')}"
                    for sample in matched
                }
            )
            is_topic_anchor = word in agenda_candidates
            if document_count < 2 and not (is_topic_anchor and document_count >= 1):
                continue
            source_count = len({str(sample.get("source") or sample.get("platform") or "") for sample in matched if sample.get("source") or sample.get("platform")})
            occurrence_count = sum(
                sum(f"{sample.get('title', '')}{sample.get('content', '')}".count(alias) for alias in aliases)
                for sample in matched
            )
            title_hits = sum(
                1
                for sample in matched
                if any(alias in str(sample.get("title") or "") for alias in aliases)
            )
            scoring_row = {
                "term": word,
                "topic_hits": [topics.index(topic) + 1],
                "document_count": document_count,
                "title_hits": title_hits,
                "source_count": source_count,
                "comment_count": len(comment_rows),
            }
            score, score_components = score_hotword_evidence(
                scoring_row,
                comments_available=comments_available,
                topic_balance=10.0,
                ai_representativeness=reviewed.get("ai_representativeness"),
            )
            evidence = [
                {
                    "source": sample.get("source") or sample.get("platform") or "",
                    "title": sample.get("title") or "",
                    "content": str(sample.get("content") or "")[:160],
                    "url": sample.get("url") or "",
                    "evidence_type": "netizen_comment" if source_bucket(sample) == "comments" else "media_or_self_media",
                }
                for sample in matched[:8]
            ]
            rows.append(
                {
                    "topic": topic,
                    "word": word,
                    "count": occurrence_count,
                    "sample_count": document_count,
                    "media_sample_count": len(media_rows),
                    "comment_sample_count": len(comment_rows),
                    "source_count": source_count,
                    "title_hits": title_hits,
                    "score": score,
                    "display_weight": score,
                    "score_components": score_components,
                    "ai_representativeness": score_components.get("ai_semantic_representativeness", 0),
                    "is_topic_anchor": is_topic_anchor,
                    "aliases": list(dict.fromkeys(aliases)),
                    "selection_method": "ai_semantic_review_with_evidence" if reviewed else "evidence_weighted_fallback",
                    "evidence": evidence,
                    "evidence_basis": "domestic_media_self_media_and_verified_comments",
                    "example": str((matched[0].get("title") or matched[0].get("content") or ""))[:80],
                }
            )
    ranked = sorted(
        rows,
        key=lambda item: (item["score"], item["source_count"], item["sample_count"], len(item["word"])),
        reverse=True,
    )
    selected_rows: list[dict[str, Any]] = []
    topic_counts: Counter[str] = Counter()
    for item in ranked:
        if topic_counts[item["topic"]] >= 8:
            continue
        if any(
            item["topic"] == kept["topic"]
            and (item["word"] in kept["word"] or kept["word"] in item["word"])
            for kept in selected_rows
        ):
            continue
        selected_rows.append(item)
        topic_counts[item["topic"]] += 1
    return selected_rows


def hotword_topic(word: str, topics: list[str]) -> str:
    domain_groups = [
        (["防汛", "抗洪", "救灾"], ["防汛", "抗洪", "救灾", "监测", "预报", "风险", "隐患", "河段", "水库", "水利", "人员转移", "应急", "物资", "防灾减灾", "生命财产"]),
        (["数字中国"], ["数字", "数智", "通信网", "算力", "数据", "科技攻关", "新型安全"]),
        (["流通体系", "物流"], ["流通", "物流", "仓储", "多式联运", "供应链", "骨干网络"]),
        (["支柱产业", "新兴产业"], ["支柱产业", "基础研究", "软硬件", "技术路线", "应用牵引", "技术迭代", "生态完善", "监管模式"]),
        (["自然资源", "国土空间"], ["自然资源", "国土空间", "主体功能区", "确权登记", "城市群", "都市圈", "资产管理"]),
        (["全民健身", "健身计划"], ["全民健身", "健身", "体育", "场地设施", "健康产业", "群众赛事"]),
    ]
    best_topic = topics[0] if topics else ""
    best_score = -1
    for topic in topics:
        display = topic_display(topic)
        score = 0
        if word in display or display in word:
            score += 20
        score += max((len(term) for term in extract_terms(topic) if term and term in word), default=0)
        for topic_hints, word_hints in domain_groups:
            if any(hint in display for hint in topic_hints) and any(hint in word for hint in word_hints):
                score += 10
        if score > best_score:
            best_score = score
            best_topic = topic
    return best_topic


def overseas_category(sample: dict[str, Any]) -> str:
    text = f"{sample.get('title', '')}{sample.get('content', '')}"
    if any(word in text for word in HYPE_WORDS):
        return "借题炒作/风险解读"
    if any(word in text for word in ["认为", "分析", "指出", "解读", "关注"]):
        return "一般解读"
    return "转载报道"


def build_overseas(samples: list[dict[str, Any]]) -> dict[str, Any]:
    groups: dict[str, list[dict[str, Any]]] = {"转载报道": [], "一般解读": [], "借题炒作/风险解读": []}
    public_comments: list[dict[str, Any]] = []
    for sample in samples:
        if not is_overseas(sample):
            continue
        if is_comment(sample):
            public_comments.append(row_sample(sample))
            continue
        groups[overseas_category(sample)].append(row_sample(sample))
    return {
        "summary": {key: len(value) for key, value in groups.items()},
        "groups": groups,
        "public_comments": public_comments,
        "needs_manual_review": True,
    }


STRICT_OVERSEAS_COMMENT_EXCLUSIONS = (
    r"(?:推翻|打倒|消灭)(?:中共|共产党|中国政府)",
    r"(?:中共|共产党|中国政府).{0,18}(?:邪恶|暴政|独裁政权|应当垮台|必须下台)",
    r"(?:CCP|Chinese Communist Party).{0,30}(?:evil|regime must fall|overthrow|dictatorship)",
    r"(?:中共|CCP).{0,24}(?:忽视农村发展|neglected the development of the country side)",
)


def disallowed_overseas_comment(sample: dict[str, Any]) -> bool:
    text = " ".join(
        str(sample.get(key) or "")
        for key in ("content", "title", "translation", "translation_cn", "summary_cn")
    )
    return any(re.search(pattern, text, re.IGNORECASE) for pattern in STRICT_OVERSEAS_COMMENT_EXCLUSIONS)


def overseas_report_category_v2(sample: dict[str, Any]) -> str:
    category_aliases = {
        "转载报道": "事实性报道",
        "事实报道": "事实性报道",
        "事实性报道": "事实性报道",
        "一般解读": "解读性报道",
        "评论解读": "解读性报道",
        "解读性报道": "解读性报道",
        "风险叙事": "借题炒作/风险解读",
        "借题炒作/风险叙事": "借题炒作/风险解读",
        "风险解读": "借题炒作/风险解读",
        "借题炒作/风险解读": "借题炒作/风险解读",
    }
    explicit = ""
    for field in (
        "ai_report_category",
        "reviewed_report_category",
        "overseas_category",
        "category",
    ):
        value = str(sample.get(field) or "").strip()
        if value in category_aliases:
            explicit = category_aliases[value]
            break
    if explicit:
        return explicit
    title = str(sample.get("title") or "")
    text = " ".join(
        str(sample.get(key) or "")
        for key in ("title", "content", "summary", "summary_cn", "interpretive_summary_cn", "translation_cn")
    )
    risk_narrative_markers = ("借题炒作", "失衡", "危机", "威胁论", "抹黑", "质疑", "争议", "批评", "风险叙事")
    if any(word in text for word in risk_narrative_markers):
        return "借题炒作/风险解读"
    factual_title_markers = ("国务院常务会议", "国常会", "李强主持", "部署", "研究", "审议通过", "听取")
    analytical_title_markers = ("分析", "解读", "评论", "失衡", "危机", "质疑", "影响", "挑战", "前景")
    if any(marker in title for marker in factual_title_markers) and not any(marker in title for marker in analytical_title_markers):
        return "事实性报道"
    interpretive_markers = ("认为", "分析", "指出", "解读", "评论", "意味着", "影响", "挑战", "风险", "前景")
    if truthy(sample.get("interpretive_verified")) or any(word in text for word in interpretive_markers):
        return "解读性报道"
    return "事实性报道"


def build_overseas_v2(samples: list[dict[str, Any]]) -> dict[str, Any]:
    groups: dict[str, list[dict[str, Any]]] = {
        "事实性报道": [],
        "解读性报道": [],
        "借题炒作/风险解读": [],
    }
    public_comments: list[dict[str, Any]] = []
    excluded_comments: list[dict[str, Any]] = []
    for sample in samples:
        if not is_overseas(sample):
            continue
        if is_comment(sample):
            if disallowed_overseas_comment(sample):
                excluded_comments.append({**row_sample(sample), "exclusion_reason": "obvious_anti_party_or_anti_government_content"})
            else:
                public_comments.append(row_sample(sample))
            continue
        category = overseas_report_category_v2(sample)
        groups[category].append({**row_sample(sample), "overseas_category": category})
    return {
        "summary": {key: len(value) for key, value in groups.items()},
        "groups": groups,
        "public_comments": public_comments,
        "excluded_public_comments": excluded_comments,
        "needs_manual_review": True,
    }


def build_system_reviewed_overseas(
    system_data: dict[str, Any],
    samples: list[dict[str, Any]],
) -> dict[str, Any]:
    """Build formal overseas evidence only from the workbook's reviewed appendix universe."""
    groups: dict[str, list[dict[str, Any]]] = {
        "事实性报道": [],
        "解读性报道": [],
        "借题炒作/风险解读": [],
    }
    by_url = {
        canonical_source_url(str(sample.get("url") or "")): sample
        for sample in samples
        if canonical_source_url(str(sample.get("url") or ""))
    }
    by_source_title = {
        (
            re.sub(r"\s+", "", str(sample.get("source") or "")).casefold(),
            re.sub(r"\s+", "", str(sample.get("title") or "")).casefold(),
        ): sample
        for sample in samples
        if sample.get("source") and sample.get("title")
    }
    seen: set[str] = set()
    for reviewed in system_data.get("overseas_reports") or []:
        url_key = canonical_source_url(str(reviewed.get("url") or ""))
        source_title_key = (
            re.sub(r"\s+", "", str(reviewed.get("source") or "")).casefold(),
            re.sub(r"\s+", "", str(reviewed.get("title") or "")).casefold(),
        )
        sample = by_url.get(url_key) or by_source_title.get(source_title_key) or {}
        dedupe_key = url_key or "|".join(source_title_key)
        if not dedupe_key or dedupe_key in seen:
            continue
        seen.add(dedupe_key)
        explicit_category = str(
            reviewed.get("ai_report_category")
            or reviewed.get("reviewed_report_category")
            or sample.get("ai_report_category")
            or sample.get("reviewed_report_category")
            or ""
        ).strip()
        if not explicit_category:
            raise ValueError(
                "标准总表境外记录缺少逐条AI报道类型，不能静默默认为事实性报道："
                f"{reviewed.get('source', '')}《{reviewed.get('title', '')}》"
            )
        category = overseas_report_category_v2({"ai_report_category": explicit_category})
        item = {
            "source": reviewed.get("source") or sample.get("source") or "",
            "source_cn_simplified": reviewed.get("source_cn_simplified") or sample.get("source_cn_simplified") or "",
            "published_at": reviewed.get("published_at") or sample.get("published_at") or "",
            "title": reviewed.get("title") or sample.get("title") or "",
            "title_cn": (
                reviewed.get("title_cn_simplified")
                or sample.get("title_cn_simplified")
                or sample.get("title_cn")
                or sample.get("translation_title_cn")
                or reviewed.get("title")
                or ""
            ),
            "title_cn_simplified": reviewed.get("title_cn_simplified") or sample.get("title_cn_simplified") or "",
            "url": reviewed.get("url") or sample.get("url") or "",
            "topic_hits": list(reviewed.get("topic_hits") or sample.get("topic_hits") or []),
            "summary_cn": (
                reviewed.get("summary_cn_simplified")
                or sample.get("summary_cn_simplified")
                or sample.get("summary_cn")
                or sample.get("interpretive_summary_cn")
                or ""
            ),
            "summary_cn_simplified": reviewed.get("summary_cn_simplified") or sample.get("summary_cn_simplified") or "",
            "ai_report_category": explicit_category,
            "classification_reason": reviewed.get("classification_reason") or sample.get("classification_reason") or "标准总表境外AI审核后保留记录",
            "classification_confidence": reviewed.get("classification_confidence") or sample.get("classification_confidence"),
            "interpretive_verified": bool(reviewed.get("interpretive_verified") or sample.get("interpretive_verified")),
            "formal_include": True,
            "meeting_relevance": True,
            "review_status": "workbook_ai_review_complete",
            "authority": "monitoring_system_reviewed_overseas_appendix",
            "overseas_category": category,
            "category": category,
        }
        groups[category].append(item)

    reviewed_comments = [
        row_sample(sample)
        for sample in samples
        if is_overseas(sample)
        and is_comment(sample)
        and sample.get("formal_include") is True
        and sample.get("meeting_relevance") is True
        and not disallowed_overseas_comment(sample)
    ]
    return {
        "summary": {key: len(value) for key, value in groups.items()},
        "groups": groups,
        "public_comments": reviewed_comments,
        "excluded_public_comments": [],
        "needs_manual_review": False,
        "authority": "monitoring_system_reviewed_overseas_appendix",
    }


def reconcile_merged_overseas(data: dict[str, Any]) -> None:
    """Rebuild mutually exclusive overseas groups after supplement enrichment."""
    overseas = data.setdefault("overseas", {})
    appendix = data.setdefault("appendices", {}).setdefault("overseas_reports", [])
    unique_rows: list[dict[str, Any]] = []
    by_key: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(appendix):
        key = canonical_source_url(str(row.get("url") or "")) or f"row:{index}:{row.get('source')}:{row.get('title')}"
        if key in by_key:
            by_key[key].update({field: value for field, value in row.items() if value not in (None, "", [], {})})
            continue
        clone = dict(row)
        by_key[key] = clone
        unique_rows.append(clone)
    groups: dict[str, list[dict[str, Any]]] = {
        "事实性报道": [],
        "解读性报道": [],
        "借题炒作/风险解读": [],
    }
    for row in unique_rows:
        category = overseas_report_category_v2(row)
        row["category"] = category
        row["overseas_category"] = category
        groups[category].append(row)
    appendix[:] = unique_rows

    accepted_comments = []
    excluded_comments = list(overseas.get("excluded_public_comments") or [])
    for row in overseas.get("public_comments") or []:
        if disallowed_overseas_comment(row):
            excluded_comments.append({**row, "exclusion_reason": "obvious_anti_party_or_anti_government_content"})
        else:
            accepted_comments.append(row)
    overseas["groups"] = groups
    overseas["summary"] = {key: len(value) for key, value in groups.items()}
    overseas["public_comments"] = accepted_comments
    overseas["excluded_public_comments"] = excluded_comments


def date_key(value: str) -> str:
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", value or "")
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    match = re.search(r"(\d{1,2})\s+([A-Za-z]{3})\s+(\d{4})", value or "")
    if match:
        month_map = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}
        return f"{int(match.group(3)):04d}-{month_map.get(match.group(2), 1):02d}-{int(match.group(1)):02d}"
    return "未标注日期"


def filter_samples_by_monitoring_period(
    samples: list[dict[str, Any]],
    start: str,
    end: str,
) -> tuple[list[dict[str, Any]], int]:
    if not start and not end:
        return samples, 0
    kept: list[dict[str, Any]] = []
    dropped = 0
    for sample in samples:
        key = date_key(str(sample.get("published_at") or ""))
        if key == "未标注日期":
            kept.append(sample)
            continue
        if start and key < start:
            dropped += 1
            continue
        if end and key > end:
            dropped += 1
            continue
        kept.append(sample)
    return kept, dropped


def reconcile_collection_gaps(gaps: list[str], samples: list[dict[str, Any]], topics: list[str]) -> list[str]:
    covered = {str(sample.get("topic") or "") for sample in samples}
    output = []
    for gap in gaps:
        skip = False
        for topic in topics:
            display = topic_display(topic)
            if topic in covered and f"“{display}”" in gap and "未获得可归入" in gap:
                skip = True
                break
        if not skip:
            output.append(gap)
    return output


def build_statistics(samples: list[dict[str, Any]], topics: list[str]) -> dict[str, Any]:
    return {
        "total_samples": len(samples),
        "total_spread": sum(int(x.get("spread_count") or 0) for x in samples),
        "total_comments": sum(int(x.get("comment_count") or 0) for x in samples),
        "by_platform": dict(Counter(platform_label(str(x.get("platform") or "unknown")) for x in samples)),
        "by_source_bucket": dict(Counter(source_bucket(x) for x in samples)),
        "by_source_type": dict(Counter(source_label(str(x.get("source_type") or "unknown")) for x in samples)),
        "by_region": dict(Counter(str(x.get("region") or "domestic") for x in samples)),
        "by_topic": dict(Counter(str(x.get("topic") or "未归类") for x in samples)),
        "by_sentiment": sentiment_counts([x for x in samples if is_comment(x)]),
        "by_date": dict(Counter(date_key(str(x.get("published_at") or "")) for x in samples)),
        "topic_order": topics,
    }


def collection_limits(profile: str, max_posts: int, max_comments: int) -> tuple[int, int]:
    defaults = {
        "demo": (5, 10),
        "formal": (200, 100),
        "deep": (500, 200),
    }
    default_posts, default_comments = defaults.get(profile, defaults["demo"])
    return max_posts or default_posts, max_comments or default_comments


def collection_targets(args: argparse.Namespace, topics: list[str]) -> dict[str, Any]:
    topic_count = max(1, len(topics))
    return {
        "target_total_samples": args.target_samples,
        "target_comments": args.target_comments,
        "target_overseas": args.target_overseas,
        "target_comments_per_topic": args.target_comments_per_topic,
        "target_samples_per_topic": max(1, args.target_samples // topic_count),
        "max_low_quality_ratio": args.max_low_quality_ratio,
        "max_collection_rounds": args.max_collection_rounds,
    }


def mediaspider_queries(topic: str) -> list[str]:
    queries: list[str] = []
    base_terms = [topic_display(topic), *extract_terms(topic)]
    for term in base_terms:
        term = str(term).strip()
        if not term:
            continue
        queries.extend(
            [
                f"国务院常务会议 {term}",
                f"{term} 国务院",
                f"{term} 政策",
                f"{term} 评论",
            ]
        )
    return list(dict.fromkeys(q for q in queries if q))[:8]


def mediaspider_task_queries(topic: str, collection_profile: str, max_posts: int) -> list[str]:
    queries = mediaspider_queries(topic)
    if collection_profile == "demo":
        return queries[:1]
    return queries


def agent_reach_queries(topic: str) -> list[str]:
    queries: list[str] = []
    for term in [topic_display(topic), *extract_terms(topic)]:
        term = str(term).strip()
        if not term:
            continue
        queries.extend(
            [
                f"国务院常务会议 {term}",
                f"{term} 大家怎么看",
                f"{term} 评论",
            ]
        )
    return list(dict.fromkeys(queries))[:6]


def generate_agent_reach_tasks(topics: list[str], out_dir: Path, platforms: list[str]) -> dict[str, Any]:
    task_dir = out_dir / "agent_reach_tasks"
    task_dir.mkdir(parents=True, exist_ok=True)
    tasks: list[dict[str, Any]] = []
    for topic_index, topic in enumerate(topics, 1):
        for query in agent_reach_queries(topic):
            if "xhs" in platforms or "xiaohongshu" in platforms:
                tasks.append(
                    {
                        "topic": topic,
                        "platform": "xiaohongshu",
                        "query": query,
                        "purpose": "netizen_quote_candidates",
                        "expected_output": "posts plus representative original user wording; comments require reading selected note URLs/IDs with login-capable backend",
                        "command": f'opencli xiaohongshu search "{query}" -f yaml',
                        "argv": ["opencli", "xiaohongshu", "search", query, "-f", "yaml"],
                    }
                )
            if "bili" in platforms or "bilibili" in platforms:
                tasks.append(
                    {
                        "topic": topic,
                        "platform": "bilibili",
                        "query": query,
                        "purpose": "netizen_quote_candidates",
                        "expected_output": "video discussion leads and high-signal public comments when supported by backend",
                        "command": f'bili search "{query}" --type video -n 20',
                        "argv": ["bili", "search", query, "--type", "video", "-n", "20"],
                    }
                )
            if "v2ex" in platforms:
                tasks.append(
                    {
                        "topic": topic,
                        "platform": "v2ex",
                        "query": query,
                        "purpose": "netizen_quote_candidates",
                        "expected_output": "public forum replies suitable as traceable original wording when topic match exists",
                        "command": f'mcporter call "exa.web_search_exa(query: \\"site:v2ex.com/t {query}\\", numResults: 10)"',
                        "argv": ["mcporter", "call", f'exa.web_search_exa(query: "site:v2ex.com/t {query}", numResults: 10)'],
                    }
                )
    manifest = {
        "task_count": len(tasks),
        "topics": topics,
        "platforms": platforms,
        "note": "agent-reach is used for lead discovery and traceable representative netizen quote candidates; MediaSpider remains the main bulk comment collector for volume and platform statistics.",
        "tasks": tasks,
    }
    manifest_path = task_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"task_dir": str(task_dir), "manifest": str(manifest_path), "task_count": len(tasks), "tasks": tasks}


def run_agent_reach_and_ingest(
    args: argparse.Namespace,
    agent_reach_tasks: dict[str, Any] | None,
    topics: list[str],
    out_dir: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any], list[str]]:
    if not agent_reach_tasks:
        return [], {}, ["未生成 agent-reach 任务，无法执行网友原话线索采集。"]

    script_dir = Path(__file__).resolve().parent
    task_dir = out_dir / "agent_reach_tasks"
    raw_dir = task_dir / "raw"
    batch_out = task_dir / "batch_run.json"
    batch_cmd = [
        sys.executable,
        str(script_dir / "run_agent_reach_batch.py"),
        "--manifest",
        str(agent_reach_tasks["manifest"]),
        "--limit",
        str(args.agent_reach_limit),
        "--out-dir",
        str(raw_dir),
        "--out",
        str(batch_out),
        "--run",
    ]
    if args.opencli_profile:
        batch_cmd.extend(["--opencli-profile", args.opencli_profile])
    if args.agent_reach_platforms:
        batch_cmd.extend(["--platforms", args.agent_reach_platforms])
    run_proc = subprocess.run(batch_cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", check=False)

    ingest_out = out_dir / "agent_reach_samples.json"
    ingest_cmd = [
        sys.executable,
        str(script_dir / "ingest_agent_reach_outputs.py"),
        str(raw_dir),
        "--output",
        str(ingest_out),
        "--topics",
        json.dumps(topics, ensure_ascii=False),
    ]
    ingest_proc = subprocess.run(ingest_cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", check=False)

    samples = load_samples(str(ingest_out), topics) if ingest_out.exists() else []
    status = {
        "enabled": True,
        "batch_command": batch_cmd,
        "batch_exit_code": run_proc.returncode,
        "batch_out": str(batch_out),
        "batch_stdout_tail": run_proc.stdout[-3000:],
        "batch_stderr_tail": run_proc.stderr[-3000:],
        "ingest_command": ingest_cmd,
        "ingest_exit_code": ingest_proc.returncode,
        "ingest_out": str(ingest_out),
        "normalized_samples": len(samples),
    }
    gaps: list[str] = []
    if run_proc.returncode != 0:
        gaps.append(f"agent-reach 线索采集未完全成功，退出码{run_proc.returncode}；详见 {batch_out}。")
    if ingest_proc.returncode != 0:
        gaps.append(f"agent-reach 输出归一化失败，退出码{ingest_proc.returncode}；详见 {ingest_out}。")
    if run_proc.returncode == 0 and not samples:
        gaps.append("agent-reach 已执行但未归一化出可用网友原话样本，需检查后端登录态、输出格式或平台命中。")
    return samples, status, gaps


def run_foreign_mediaspider_and_ingest(
    args: argparse.Namespace,
    agenda: str,
    topics: list[str],
    out_dir: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any], list[str]]:
    """Run the server-safe MediaSpider foreign branch and retain strict evidence only."""
    script_dir = Path(__file__).resolve().parent
    collector_dir = out_dir / "foreign_mediaspider"
    collector_dir.mkdir(parents=True, exist_ok=True)
    plan_path = collector_dir / "research_plan.json"
    if args.system_workbook:
        from build_research_plan import build_plan

        plan = build_plan(args.system_workbook, agenda)
    else:
        plan = {
            "monitoring_period": {"start": args.monitor_start or "", "end": args.monitor_end or ""},
            "topics": [
                {
                    "topic": topic,
                    "queries": {"overseas": [f"{args.monitor_start or ''} China State Council executive meeting {topic}".strip()]},
                }
                for topic in topics
            ],
        }
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")

    command = [
        sys.executable,
        str(script_dir / "run_foreign_mediaspider.py"),
        str(plan_path),
        "--output-dir",
        str(collector_dir),
        "--supervisor",
        args.foreign_mediaspider_supervisor or args.mediaspider_supervisor,
        "--platforms",
        args.foreign_mediaspider_platforms,
        "--foreign-mode",
        "fallback-only",
        "--limit",
        str(args.foreign_mediaspider_limit),
        "--run",
    ]
    if args.foreign_media_home or args.media_home:
        command.extend(["--media-home", args.foreign_media_home or args.media_home])

    gaps: list[str] = []
    creationflags = getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0) if os.name == "nt" else 0
    process = subprocess.Popen(
        command,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        creationflags=creationflags,
        start_new_session=os.name != "nt",
    )
    try:
        stdout, stderr = process.communicate(
            timeout=args.foreign_mediaspider_timeout_sec if args.foreign_mediaspider_timeout_sec > 0 else None
        )
        timed_out = False
    except subprocess.TimeoutExpired:
        timed_out = True
        if os.name == "nt":
            subprocess.run(
                ["taskkill", "/PID", str(process.pid), "/T", "/F"],
                capture_output=True,
                text=True,
                check=False,
            )
        else:
            try:
                os.killpg(process.pid, 9)
            except ProcessLookupError:
                pass
        stdout, stderr = process.communicate()
        gaps.append(f"MediaSpider境外补证运行超时（{args.foreign_mediaspider_timeout_sec}秒）。")

    samples_path = collector_dir / "foreign_mediaspider_samples.json"
    normalized = load_samples(str(samples_path), topics) if samples_path.exists() else []
    accepted: list[dict[str, Any]] = []
    for sample in normalized:
        topic = str(sample.get("topic") or "")
        if not sample.get("url") or topic not in topics:
            continue
        if sample_relevance(sample, topic) <= 0:
            continue
        accepted.append(sample)

    status = {
        "enabled": True,
        "command": command,
        "exit_code": process.returncode if process.returncode is not None else -1,
        "timed_out": timed_out,
        "stdout_tail": stdout[-3000:],
        "stderr_tail": stderr[-3000:],
        "samples_path": str(samples_path),
        "normalized_samples": len(normalized),
        "accepted_samples": len(accepted),
        "platforms": [item.strip() for item in args.foreign_mediaspider_platforms.split(",") if item.strip()],
    }
    if not timed_out and process.returncode != 0:
        gaps.append(f"MediaSpider境外补证未完全成功，退出码{process.returncode}；详见{collector_dir}。")
    if not accepted:
        gaps.append("MediaSpider境外补证未取得通过主题相关性和链接校验的新增样本。")
    return accepted, status, gaps


def generate_mediacrawler_tasks(
    agenda: str,
    topics: list[str],
    out_dir: Path,
    platforms: list[str],
    max_posts: int,
    max_comments: int,
    collection_profile: str,
) -> dict[str, Any]:
    task_dir = out_dir / "mediacrawler_tasks"
    task_dir.mkdir(parents=True, exist_ok=True)
    resolved_posts, resolved_comments = collection_limits(collection_profile, max_posts, max_comments)
    tasks = []
    for topic_index, topic in enumerate(topics, 1):
        for platform in platforms:
            task = {
                "platform": platform,
                "crawler_type": "search",
                "login_type": "qrcode",
                "keywords": mediaspider_task_queries(topic, collection_profile, max_posts),
                "get_comment": True,
                "get_sub_comment": False,
                "max_posts": resolved_posts,
                "max_comments_per_post": resolved_comments,
                "max_concurrency": 1,
                "save_data_option": "jsonl",
                "headless": False,
                "cwh_topic": topic,
                "cwh_collection_profile": collection_profile,
            }
            path = task_dir / f"{topic_index:02d}_{platform}_{safe_name(topic_display(topic))}.json"
            path.write_text(json.dumps(task, ensure_ascii=False, indent=2), encoding="utf-8")
            tasks.append({"topic": topic, "platform": platform, "path": str(path), "keywords": task["keywords"]})
    manifest = {
        "agenda": agenda,
        "topics": topics,
        "platforms": platforms,
        "collection_profile": collection_profile,
        "max_posts": resolved_posts,
        "max_comments_per_post": resolved_comments,
        "estimated_post_cap": len(tasks) * resolved_posts,
        "estimated_comment_cap": len(tasks) * resolved_posts * resolved_comments,
        "task_count": len(tasks),
        "tasks": tasks,
    }
    manifest_path = task_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"task_dir": str(task_dir), "manifest": str(manifest_path), "task_count": len(tasks), "tasks": tasks}


def run_mediacrawler_and_ingest(
    args: argparse.Namespace,
    generated_tasks: dict[str, Any] | None,
    topics: list[str],
    out_dir: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any], list[str]]:
    if not generated_tasks:
        return [], {}, ["未生成 MediaSpider 任务，无法启动真实采集。"]

    script_dir = Path(__file__).resolve().parent
    batch_out = out_dir / "mediacrawler_tasks" / "batch_run.json"
    batch_cmd = [
        sys.executable,
        str(script_dir / "run_mediacrawler_batch.py"),
        "--manifest",
        str(generated_tasks["manifest"]),
        "--supervisor",
        args.mediaspider_supervisor,
        "--python",
        sys.executable,
        "--limit",
        str(args.mediaspider_limit),
        "--out",
        str(batch_out),
        "--run",
    ]
    if args.platforms:
        batch_cmd.extend(["--platforms", args.platforms])
    if args.media_home:
        batch_cmd.extend(["--media-home", args.media_home])
    if args.keep_proxy:
        batch_cmd.append("--keep-proxy")
    if args.mediaspider_task_timeout_sec > 0:
        batch_cmd.extend(["--task-timeout-sec", str(args.mediaspider_task_timeout_sec)])

    run_proc = subprocess.run(
        batch_cmd,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )

    ingest_out = out_dir / "mediaspider_samples.json"
    ingest_cmd = [
        sys.executable,
        str(script_dir / "ingest_mediacrawler_outputs.py"),
        str(out_dir / "mediacrawler_tasks"),
        "--output",
        str(ingest_out),
    ]
    ingest_proc = subprocess.run(
        ingest_cmd,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )

    samples = load_samples(str(ingest_out), topics) if ingest_out.exists() else []
    status = {
        "enabled": True,
        "batch_command": batch_cmd,
        "batch_exit_code": run_proc.returncode,
        "batch_out": str(batch_out),
        "batch_stdout_tail": run_proc.stdout[-3000:],
        "batch_stderr_tail": run_proc.stderr[-3000:],
        "ingest_command": ingest_cmd,
        "ingest_exit_code": ingest_proc.returncode,
        "ingest_out": str(ingest_out),
        "ingest_stdout_tail": ingest_proc.stdout[-3000:],
        "ingest_stderr_tail": ingest_proc.stderr[-3000:],
        "normalized_samples": len(samples),
    }
    gaps: list[str] = []
    if run_proc.returncode != 0:
        gaps.append(f"MediaSpider 真实采集未完全成功，退出码{run_proc.returncode}；详见 {batch_out}。")
        diagnostic_text = f"{run_proc.stdout}\n{run_proc.stderr}".lower()
        if "9222" in diagnostic_text or "cdp" in diagnostic_text or "connect" in diagnostic_text:
            gaps.append("MediaSpider 采集疑似卡在浏览器/CDP连接：请确认采集浏览器已开放 9222 端口，或按 MediaSpider supervisor 日志调整浏览器启动方式。")
    if ingest_proc.returncode != 0:
        gaps.append(f"MediaSpider 输出归一化失败，退出码{ingest_proc.returncode}；详见 {ingest_out}。")
    if run_proc.returncode == 0 and not samples:
        gaps.append("MediaSpider 任务运行完成但未归一化出有效样本，需检查 raw 输出目录、登录态或关键词命中。")
    return samples, status, gaps


def escape_svg(text: Any) -> str:
    return html.escape(str(text))


def write_bar_svg(path: Path, title: str, values: dict[str, int], color: str = "#2f6f73") -> None:
    width = 860
    row_h = 34
    if not values:
        values = {"暂无可统计样本": 0}
    height = max(180, 58 + row_h * len(values))
    max_value = max(values.values(), default=1) or 1
    rows = []
    y = 50
    for name, value in values.items():
        bar_w = int((value / max_value) * 560) if max_value else 0
        rows.append(f'<text x="24" y="{y + 18}" font-size="14">{escape_svg(name)}</text>')
        rows.append(f'<rect x="180" y="{y}" width="{bar_w}" height="22" fill="{color}"/>')
        rows.append(f'<text x="{190 + bar_w}" y="{y + 17}" font-size="13">{value}</text>')
        y += row_h
    svg = f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="#fff"/><text x="24" y="28" font-size="18" font-weight="700">{escape_svg(title)}</text>{"".join(rows)}</svg>'
    path.write_text(svg, encoding="utf-8")


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    keys = list(dict.fromkeys(key for row in rows for key in row.keys()))
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=keys)
        writer.writeheader()
        writer.writerows(rows)


def write_sentiment_handoff(data: dict[str, Any], out_dir: Path) -> dict[str, str]:
    sentiment_dir = out_dir / "sentiment_analysis"
    sentiment_dir.mkdir(parents=True, exist_ok=True)
    input_path = sentiment_dir / "sentiment_input.csv"
    schema_path = sentiment_dir / "label_schema.md"
    comments = [sample for sample in data.get("samples", []) if is_comment(sample)]
    rows = [
        {
            "sample_id": sample.get("id") or "",
            "text": sample.get("content") or sample.get("title") or "",
            "topic": sample.get("topic") or "",
            "platform": sample.get("platform") or "",
            "source": sample.get("source") or "",
            "published_at": sample.get("published_at") or "",
            "url": sample.get("url") or "",
            "spread_count": sample.get("spread_count") or 0,
            "comment_count": sample.get("comment_count") or 0,
        }
        for sample in comments
    ]
    write_csv(input_path, rows)
    schema_path.write_text(
        """# CWH网民评论情感标签

正式统计对象仅为与本次会议议题相关、具有可判定语境的网民评论。

| 标签 | 口径 |
|---|---|
| positive | 明确支持、认可、点赞、期待政策或议题进展。 |
| neutral | 与议题相关，以事实陈述、提问、解释或平衡分析为主，无明确正负态度。 |
| negative | 明确质疑、反对、不满、担忧或批评政策及其影响。 |

无关内容、广告垃圾、纯表情、纯转发、语境不足和无法判断的文本不强行归入中性；设置 `in_sentiment_denominator=false` 并填写 `exclusion_reason`。低置信度、讽刺、矛盾线索和边界样本进入复核，不直接作为最终标签。
""",
        encoding="utf-8",
    )
    return {"input": str(input_path), "label_schema": str(schema_path)}


def parse_ymd(value: str) -> datetime | None:
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", value or "")
    if not match:
        return None
    return datetime(int(match.group(1)), int(match.group(2)), int(match.group(3)))


def workbook_dates(data: dict[str, Any]) -> list[str]:
    monitor = data.get("collection", {}).get("monitoring_period") or {}
    start = parse_ymd(str(monitor.get("start") or ""))
    end = parse_ymd(str(monitor.get("end") or ""))
    if start and end and end >= start:
        days = []
        current = start
        while current <= end:
            days.append(current.strftime("%Y-%m-%d"))
            current += timedelta(days=1)
        return days
    dates = sorted(k for k in (data.get("statistics", {}).get("by_date") or {}) if k != "未标注日期")
    if dates:
        return dates
    meeting_date = parse_ymd(str(data.get("meeting", {}).get("date") or ""))
    if meeting_date:
        return [(meeting_date + timedelta(days=i)).strftime("%Y-%m-%d") for i in range(4)]
    return [datetime.now().strftime("%Y-%m-%d")]


def workbook_platform_metrics(samples: list[dict[str, Any]], dates: list[str]) -> dict[str, dict[str, int]]:
    metrics: dict[str, dict[str, int]] = {
        day: {
            "wechat_articles": 0,
            "wechat_reads": 0,
            "wechat_comments": 0,
            "weibo": 0,
            "video_account": 0,
            "domestic_media": 0,
            "overseas_media": 0,
            "douyin": 0,
            "kuaishou": 0,
            "bili": 0,
            "xhs": 0,
            "forum": 0,
            "other_new_media": 0,
        }
        for day in dates
    }
    if not metrics:
        return metrics
    fallback_day = dates[0]
    for sample in samples:
        day = date_key(str(sample.get("published_at") or ""))
        if day not in metrics:
            day = fallback_day if day == "未标注日期" else day
        if day not in metrics:
            continue
        platform = str(sample.get("platform") or "").lower()
        url = str(sample.get("url") or "").lower()
        bucket = source_bucket(sample)
        spread = int(sample.get("spread_count") or 0)
        comments = int(sample.get("comment_count") or 0)
        row = metrics[day]
        if platform in {"wechat", "weixin"} or "mp.weixin.qq.com" in url:
            row["wechat_articles"] += 1
            row["wechat_reads"] += spread
            row["wechat_comments"] += comments
        elif platform in {"weibo", "wb"}:
            row["weibo"] += 1
        elif platform in {"shipinhao", "video_account", "channels"}:
            row["video_account"] += 1
        elif platform in {"douyin", "dy"}:
            row["douyin"] += 1
        elif platform in {"kuaishou", "ks"}:
            row["kuaishou"] += 1
        elif platform in {"bili", "bilibili"}:
            row["bili"] += 1
        elif platform in {"xhs", "xiaohongshu"}:
            row["xhs"] += 1
        elif platform in {"tieba", "v2ex", "forum"}:
            row["forum"] += 1
        elif bucket == "domestic_media":
            row["domestic_media"] += 1
        elif bucket == "overseas_media":
            row["overseas_media"] += 1
        elif bucket == "self_media":
            row["other_new_media"] += 1
    return metrics


def workbook_ratio(counts: dict[str, int], key: str) -> float:
    total = sum(int(counts.get(name, 0) or 0) for name in ["positive", "neutral", "negative"])
    if total <= 0:
        return 0
    return round(int(counts.get(key, 0) or 0) / total, 4)


def workbook_sheet_title(value: str) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "", value).strip()
    return cleaned[:31] or "Sheet"


def write_cwh_data_workbook(data: dict[str, Any], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    samples = data.get("samples") or []
    topics = data.get("meeting", {}).get("topics") or []
    dates = workbook_dates(data)
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    sub_fill = PatternFill("solid", fgColor="EAF2F8")
    title_font = Font(bold=True, size=13)
    header_font = Font(bold=True)

    def style_sheet(ws, widths: dict[int, int] | None = None) -> None:
        ws.freeze_panes = "A3"
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if widths:
            for col, width in widths.items():
                ws.column_dimensions[get_column_letter(col)].width = width

    ws = wb.create_sheet("关键词")
    ws.append(["序号", "标题", "关键词", "排除词", "ID"])
    system_data = data.get("system_data") or {}
    system_master = system_data.get("master_event") or {}
    system_topic_rows = system_data.get("topics") or []
    master_keywords = str(system_master.get("keywords") or "") or "\n".join(
        dict.fromkeys(query for topic in topics for query in keywords_for(topic))
    )
    ws.append([1, "总事件", master_keywords, system_master.get("exclude_terms", ""), system_master.get("event_id") or 200000000])
    ws.append([])
    ws.append(["序号", "标题", "子事件", "排除词", "ID"])
    for idx, topic in enumerate(topics, 1):
        system_topic = next((row for row in system_topic_rows if int(row.get("index", 0) or 0) == idx), {})
        ws.append([
            idx,
            topic_display(topic),
            system_topic.get("keywords") or "\n".join(keywords_for(topic)),
            system_topic.get("exclude_terms", ""),
            system_topic.get("event_id") or 200000000 + idx,
        ])
    for cell in ws[1] + ws[4]:
        cell.fill = header_fill
        cell.font = header_font
    style_sheet(ws, {1: 8, 2: 34, 3: 70, 4: 18, 5: 14})

    def write_trend_sheet(
        ws,
        title: str,
        rows_samples: list[dict[str, Any]],
        system_event: dict[str, Any] | None = None,
        is_total_event: bool = True,
    ) -> None:
        metrics = workbook_platform_metrics(rows_samples, dates)
        if system_event:
            metrics = {}
            for item in system_event.get("daily") or []:
                day = str(item.get("date") or "")
                details = dict(item.get("details") or {})
                metrics[day] = {
                    "wechat_articles": details.get("wechat_articles", 0),
                    "wechat_reads": details.get("wechat_reads", 0),
                    "wechat_comments": details.get("wechat_comments", 0),
                    "weibo": details.get("weibo", 0),
                    "domestic_news": details.get("domestic_news", item.get("domestic_mainstream", 0)),
                    "domestic_app": details.get("domestic_app", 0),
                    "domestic_forum": details.get("domestic_forum", 0),
                    "other_video": details.get("other_video", 0),
                    "overseas_news": details.get("overseas_news", item.get("overseas_media", 0)),
                    "twitter": details.get("twitter", 0),
                    "overseas_other": details.get("overseas_other", 0),
                    "video_account": details.get("video_account", item.get("video_account", 0)),
                    "douyin": details.get("douyin", 0),
                }
        ws["A2"] = title
        ws["A2"].font = title_font
        summary_headers = ["日期", domestic_media_label(data), "境外媒体", "微信公众号", "新浪微博", "视频号", "新媒体", "信息传播量"] if is_total_event else ["日期", domestic_media_label(data), "境外媒体", "新媒体", "信息传播量"]
        detail_headers = ["日期", "公众文章", "公号-在看量", "公号-精选评论量", "新浪微博", "境内新闻", "境内APP", "境内论坛", "其他视频", "境外新闻", "推特", "境外其他", "视频号发文", "抖音"]
        for col, value in enumerate(summary_headers, 1):
            ws.cell(3, col, value)
        for col, value in enumerate(detail_headers, 10):
            ws.cell(3, col, value)
        for cell in ws[3]:
            if cell.value:
                cell.fill = header_fill
                cell.font = header_font
        for i, day in enumerate(dates, 4):
            row = metrics.get(day, {})
            authoritative_day = next((item for item in (system_event or {}).get("daily", []) if str(item.get("date") or "") == day), {})
            detail_values = [
                day,
                row.get("wechat_articles", 0),
                row.get("wechat_reads", 0),
                row.get("wechat_comments", 0),
                row.get("weibo", 0),
                row.get("domestic_news", row.get("domestic_media", 0)),
                row.get("domestic_app", 0),
                row.get("domestic_forum", row.get("forum", 0)),
                row.get("other_video", row.get("bili", 0) + row.get("xhs", 0) + row.get("kuaishou", 0)),
                row.get("overseas_news", row.get("overseas_media", 0)),
                row.get("twitter", 0),
                row.get("overseas_other", 0),
                row.get("video_account", 0),
                row.get("douyin", 0),
            ]
            for col, value in enumerate(detail_values, 10):
                ws.cell(i, col, value)
            if authoritative_day:
                ws.cell(i, 1, day)
                ws.cell(i, 2, int(authoritative_day.get("domestic_mainstream", 0) or 0))
                ws.cell(i, 3, int(authoritative_day.get("overseas_media", 0) or 0))
                if is_total_event:
                    ws.cell(i, 4, int(authoritative_day.get("wechat_public", 0) or 0))
                    ws.cell(i, 5, int(authoritative_day.get("weibo", 0) or 0))
                    ws.cell(i, 6, int(authoritative_day.get("video_account", 0) or 0))
                    ws.cell(i, 7, int(authoritative_day.get("new_media", 0) or 0))
                    ws.cell(i, 8, int(authoritative_day.get("total_spread", 0) or 0))
                else:
                    ws.cell(i, 4, int(authoritative_day.get("new_media", 0) or 0))
                    ws.cell(i, 5, int(authoritative_day.get("total_spread", 0) or 0))
            else:
                ws.cell(i, 1, f"=J{i}")
                ws.cell(i, 2, f"=O{i}")
                ws.cell(i, 3, f"=S{i}")
                if is_total_event:
                    ws.cell(i, 4, f"=K{i}+L{i}+M{i}")
                    ws.cell(i, 5, f"=N{i}")
                    ws.cell(i, 6, f"=V{i}")
                    ws.cell(i, 7, f"=P{i}+Q{i}+R{i}+T{i}+U{i}+W{i}")
                    ws.cell(i, 8, f"=B{i}+C{i}+D{i}+E{i}+F{i}+G{i}")
                else:
                    ws.cell(i, 4, f"=K{i}+L{i}+M{i}+N{i}+P{i}+Q{i}+R{i}+T{i}+U{i}+V{i}+W{i}")
                    ws.cell(i, 5, f"=B{i}+C{i}+D{i}")
        total_row = max(12, len(dates) + 5)
        summary_end_col = 8 if is_total_event else 5
        system_totals = (system_event or {}).get("totals") or {}
        total_keys = ["domestic_mainstream", "overseas_media", "wechat_public", "weibo", "video_account", "new_media", "total_spread"] if is_total_event else ["domestic_mainstream", "overseas_media", "new_media", "total_spread"]
        for offset, col in enumerate(range(2, summary_end_col + 1)):
            key = total_keys[offset]
            if system_totals:
                ws.cell(total_row, col, int(system_totals.get(key, 0) or 0))
            else:
                ws.cell(total_row, col, f"=SUM({get_column_letter(col)}4:{get_column_letter(col)}{3 + len(dates)})")
        for col in range(11, 24):
            ws.cell(total_row, col, f"=SUM({get_column_letter(col)}4:{get_column_letter(col)}{3 + len(dates)})")
        ws.cell(total_row, 1, "合计")
        ws.cell(total_row, 1).font = header_font
        style_sheet(ws, {1: 14, 2: 14, 3: 12, 4: 14, 5: 12, 6: 10, 7: 12, 8: 12, 10: 14})

    write_trend_sheet(wb.create_sheet("总事件"), f"{data.get('meeting', {}).get('date', '')}国务院常务会议总事件舆情走势", samples, system_data.get("overall"), True)
    for idx, topic in enumerate(topics, 1):
        topic_samples = [x for x in samples if x.get("topic") == topic]
        system_event = next((row for row in system_data.get("subevents", []) if int(row.get("index", 0) or 0) == idx), None)
        write_trend_sheet(wb.create_sheet(workbook_sheet_title(f"子事件{idx}")), f"子事件{idx}-{topic_display(topic)}", topic_samples, system_event, False)

    ws = wb.create_sheet("子事件数据汇总")
    ws.append(["序号", "标题", domestic_media_label(data), "新媒体", "境外媒体", "总量", "网民情感", "", "", "", "", "总量/万"])
    ws.append(["", "", "", "", "", "", "正面", "中立", "负面", "合计", "", ""])
    for cell in ws[1] + ws[2]:
        cell.fill = header_fill
        cell.font = header_font
    topic_stats = data.get("topic_stats") or []
    for idx, row in enumerate(topic_stats, 1):
        sentiment = row.get("sentiment") or {}
        domestic_media = int(row.get("domestic_media", 0) or 0)
        new_media = int(row.get("self_media", 0) or 0)
        overseas_media = int(row.get("overseas_media", 0) or 0)
        total = int(row.get("spread_count", row.get("total_samples", domestic_media + new_media + overseas_media)) or 0)
        sentiment_ready = bool(row.get("sentiment_formal_ready"))
        excel_row = idx + 2
        ws.append([
            idx,
            row.get("display") or topic_display(row.get("topic", "")),
            domestic_media,
            new_media,
            overseas_media,
            total,
            workbook_ratio(sentiment, "positive") if sentiment_ready else "",
            workbook_ratio(sentiment, "neutral") if sentiment_ready else "",
            workbook_ratio(sentiment, "negative") if sentiment_ready else "",
            f"=SUM(G{excel_row}:I{excel_row})" if sentiment_ready else "",
            "",
            f"=F{excel_row}/10000",
        ])
    style_sheet(ws, {1: 8, 2: 44, 3: 14, 4: 12, 5: 12, 6: 12, 7: 10, 8: 10, 9: 10, 12: 12})

    ws = wb.create_sheet("外媒报道列表")
    ws.append(["", "", "", "", "", "", "", "涉及子事件（子事件境外数据从这里数）"])
    ws.append(["", "序号", "来源", "报道日期", "超链接标题", "标题", "发布地址", *[str(i) for i in range(1, len(topics) + 1)]])
    for cell in ws[2]:
        if cell.value:
            cell.fill = header_fill
            cell.font = header_font
    overseas_rows = [x for x in samples if is_overseas(x)]
    for idx, sample in enumerate(top_items(overseas_rows, 80), 1):
        excel_row = idx + 2
        title = sample.get("title") or sample.get("content") or ""
        url = sample.get("url") or ""
        flags = []
        text = f"{title} {sample.get('content', '')}"
        for topic_idx, topic in enumerate(topics, 1):
            flags.append(topic_idx if sample.get("topic") == topic or topic_score(str(title), text, topic) > 0 else "")
        ws.append(["", idx, sample.get("source"), date_key(str(sample.get("published_at") or "")), f"=HYPERLINK(G{excel_row},F{excel_row})" if url else title, title, url, *flags])
    style_sheet(ws, {2: 8, 3: 22, 4: 14, 5: 18, 6: 52, 7: 60})

    ws = wb.create_sheet("词云")
    for idx, item in enumerate(data.get("hotwords") or [], 1):
        ws.cell(idx + 1, 1, item.get("word"))
    style_sheet(ws, {1: 28})

    ws = wb.create_sheet("公众TOP")
    ws.append(["", "微信公号文章阅读量TOP", "", "", "", "", "制作超链", ""])
    ws.append(["", "序号", "账号", "标题", "阅读量", "在看量", "原文标题", "链接"])
    for cell in ws[2]:
        if cell.value:
            cell.fill = header_fill
            cell.font = header_font
    wechat_rows = data.get("appendices", {}).get("wechat_top") or [
        row_sample(x)
        for x in top_items(
            [x for x in samples if str(x.get("platform") or "").lower() in {"wechat", "weixin"} or "mp.weixin.qq.com" in str(x.get("url") or "").lower()],
            10,
        )
    ]
    for idx, row in enumerate(wechat_rows[:10], 1):
        excel_row = idx + 2
        url = row.get("url") or ""
        title = row.get("title") or row.get("content") or ""
        ws.append(["", idx, row.get("source"), f"=HYPERLINK(H{excel_row},G{excel_row})" if url else title, row.get("spread_count") or 0, row.get("comment_count") or 0, title, url])
    style_sheet(ws, {2: 8, 3: 22, 4: 28, 5: 12, 6: 10, 7: 52, 8: 60})

    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 2)):
            for cell in row:
                if cell.value:
                    cell.font = Font(bold=True)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, float) and 0 <= cell.value <= 1:
                    cell.number_format = "0%"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def md_table(headers: list[str], rows: list[list[Any]], empty_label: str = "本批次未获得可核验样本") -> str:
    if not rows:
        rows = [[empty_label for _ in headers]]
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(str(x).replace("\n", " ") for x in row) + " |")
    return "\n".join(lines)


def sample_phrase(sample: dict[str, Any]) -> str:
    title = str(sample.get("title") or sample.get("content") or "")[:60]
    source = str(sample.get("source") or "")
    return f"{source}《{title}》" if title else source


def render_report(data: dict[str, Any]) -> str:
    meeting = data["meeting"]
    stats = data["statistics"]
    topic_stats = data["topic_stats"]
    samples = data["samples"]
    selected_comments = data["comments"]["selected"]
    overseas = data["overseas"]
    gaps = data["audit"]["data_gaps"]
    media_samples = [x for x in samples if source_bucket(x) == "domestic_media"]
    self_media_samples = [x for x in samples if source_bucket(x) == "self_media"]
    title_date = f"{int(meeting['date'][5:7])}月{int(meeting['date'][8:10])}日" if re.match(r"\d{4}-\d{2}-\d{2}", meeting["date"]) else meeting["date"]
    date_counts = {k: v for k, v in stats["by_date"].items() if k != "未标注日期"}
    peak_text = ""
    if date_counts:
        peak_date, peak_count = max(date_counts.items(), key=lambda item: item[1])
        peak_text = f"，传播峰值出现在{peak_date}，当日系统传播量为{peak_count}条" if stats.get("authority") == "monitoring_system" else f"，样本传播峰值出现在{peak_date}，当日归集{peak_count}条"
    if stats.get("authority") == "monitoring_system":
        generation_note = "本报告由 CWH 自动化流程依据部门舆情监测系统汇总数据和已接入明细生成；系统未导出的明细统一列入文末审核清单。"
        overall_note = f"监测系统给出相关信息传播总量{stats['total_spread']}条；当前接入可追溯证据明细{len(samples)}条、网民评论明细{stats['total_comments']}条。其中，{domestic_media_label(data)}传播量{stats['by_source_bucket'].get('domestic_media', 0)}条，新媒体传播量{stats['by_source_bucket'].get('self_media', 0)}条，境外媒体传播量{stats['by_source_bucket'].get('overseas_media', 0)}条。"
    else:
        generation_note = "本报告由 CWH 自动化流程根据公开检索、已接入样本和结构化分析结果生成；采集不足或需人工判断的部分统一列入文末审核清单。"
        overall_note = f"当前批次共归集相关样本{stats['total_samples']}条，累计可读传播量/热度值{stats['total_spread']}，评论量{stats['total_comments']}。其中，境内媒体样本{stats['by_source_bucket'].get('domestic_media', 0)}条，自媒体及社交平台样本{stats['by_source_bucket'].get('self_media', 0)}条，网民评论样本{stats['by_source_bucket'].get('comments', 0)}条，境外媒体样本{stats['by_source_bucket'].get('overseas_media', 0)}条。"

    lines = [
        f"# {title_date}国务院常务会议舆情情况报告",
        "",
        opening_paragraph(meeting, title_date, "会议议题包括：" + "、".join(topic_display(x) for x in meeting["topics"])) + generation_note,
        "",
        "## 一、整体传播情况",
        "",
        overall_note,
        "",
        "平台分布方面：" + ("、".join(f"{k}{v}条" for k, v in stats["by_platform"].items()) if stats["by_platform"] else "当前批次尚无平台样本") + "。传播时间上：" + ("、".join(f"{k}{v}条" for k, v in stats["by_date"].items()) if stats["by_date"] else "当前批次尚未形成时间序列") + peak_text + "。",
        "",
        "## 二、子议题传播情况",
        "",
        md_table(
            ["子议题", "样本数", "传播量/热度", "境内媒体", "自媒体", "网民评论", "境外媒体", "情感倾向"],
            [
                [
                    row["display"],
                    row["total_samples"],
                    row["spread_count"],
                    row["domestic_media"],
                    row["self_media"],
                    row["comments"],
                    row["overseas_media"],
                    sentiment_text(row["sentiment"]),
                ]
                for row in topic_stats
            ],
        ),
        "",
        "## 三、境内媒体及自媒体观点",
        "",
    ]
    for item in data["viewpoints"]["by_topic"]:
        lines.append(f"**{topic_display(item['topic'])}**")
        if not item["clusters"]:
            lines.append("当前批次未采到可归纳的媒体/自媒体样本。")
        for cluster in item["clusters"]:
            evidence = "；".join(sample_phrase(x) for x in cluster["evidence"])
            lines.append(f"- {cluster['summary']}。代表样本：{evidence}。")
        lines.append("")

    lines.extend(
        [
            "## 四、网民评论情况",
            "",
            f"当前批次筛选出可读评论样本{len(selected_comments)}条，情感分布为{sentiment_text(data['comments']['sentiment'])}。代表性评论如下：",
            "",
        ]
    )
    if selected_comments:
        for comment in selected_comments[:8]:
            lines.append(f"- {comment['source']}（{comment['platform']}）：{comment['content']}")
    else:
        lines.append("- 公开网页检索暂未获得可直接引用的网民评论；需继续接入微博、抖音、小红书、B站等评论采集结果。")

    lines.extend(["", "## 五、候选热词及关注点", ""])
    hotword_rows = [
        [
            x.get("word", ""),
            topic_display(str(x.get("topic") or "")),
            x.get("count", 0),
            x.get("sample_count", 0),
            hotword_example(x),
        ]
        for x in data.get("hotwords", [])[:12]
    ]
    lines.append(md_table(["候选热词", "所属议题", "出现次数", "来源样本数", "示例"], hotword_rows, "待真实样本确认"))

    lines.extend(
        [
            "",
            "## 六、境外关注情况",
            "",
            f"当前批次识别境外媒体样本{sum(overseas['summary'].values())}条，其中转载报道{overseas['summary'].get('转载报道', 0)}条，一般解读{overseas['summary'].get('一般解读', 0)}条，借题炒作/风险解读{overseas['summary'].get('借题炒作/风险解读', 0)}条。境外分类为自动预分类，正式成稿前需人工复核。",
            "",
        ]
    )
    for category, rows in overseas["groups"].items():
        lines.append(f"**{category}**")
        if rows:
            for row in rows[:5]:
                lines.append(f"- {row['source']}《{row['title']}》")
        else:
            lines.append("- 本批次未识别该类境外报道。")
        lines.append("")

    lines.extend(
        [
            "## 七、附录及人工审核清单",
            "",
            "### 附录一：境外媒体报道列表",
            "",
            md_table(["来源", "标题", "分类", "链接"], [[row["source"], row["title"], category, row["url"]] for category, rows in overseas["groups"].items() for row in rows[:10]]),
            "",
            "### 附录二：传播量较大的媒体/自媒体样本",
            "",
            md_table(["来源", "平台", "标题", "传播量/热度", "链接"], [[x["source"], platform_label(x["platform"]), x["title"], x["spread_count"], x["url"]] for x in top_items(media_samples + self_media_samples, 10)]),
            "",
            "### 附录三：人工审核清单",
            "",
        ]
    )
    for item in data["audit"]["review_items"]:
        lines.append(f"- {item}")
    if gaps:
        lines.extend(["", "### 当前数据接入状态", ""])
        for gap in gaps:
            lines.append(f"- {gap}")
    return "\n".join(lines) + "\n"


def build_expansion_plan(
    samples: list[dict[str, Any]],
    topics: list[str],
    targets: dict[str, Any],
    generated_tasks: dict[str, Any] | None,
    agent_reach_tasks: dict[str, Any] | None,
) -> dict[str, Any]:
    total = len(samples)
    comments = [x for x in samples if is_comment(x)]
    overseas = [x for x in samples if is_overseas(x)]
    low_quality = [x for x in samples if x.get("quality_flags")]
    by_topic = Counter(str(x.get("topic") or "") for x in samples)
    comments_by_topic = Counter(str(x.get("topic") or "") for x in comments)
    actions: list[dict[str, Any]] = []

    if total < int(targets["target_total_samples"]):
        actions.append(
            {
                "priority": "P0",
                "action": "continue_mediaspider_bulk_collection",
                "reason": f"总样本{total}条，低于目标{targets['target_total_samples']}条。",
                "task_manifest": generated_tasks.get("manifest") if generated_tasks else "",
            }
        )
    if len(comments) < int(targets["target_comments"]):
        actions.append(
            {
                "priority": "P0",
                "action": "deepen_platform_comments",
                "reason": f"网民评论{len(comments)}条，低于目标{targets['target_comments']}条。",
                "preferred_platforms": ["wb", "douyin", "xhs", "bili"],
            }
        )
    weak_topics = [
        topic
        for topic in topics
        if comments_by_topic.get(topic, 0) < int(targets["target_comments_per_topic"])
        or by_topic.get(topic, 0) < int(targets["target_samples_per_topic"])
    ]
    if weak_topics:
        actions.append(
            {
                "priority": "P0",
                "action": "expand_weak_topics",
                "reason": "部分子议题样本或评论不足。",
                "topics": [topic_display(x) for x in weak_topics],
            }
        )
    if len(overseas) < int(targets["target_overseas"]):
        actions.append(
            {
                "priority": "P1",
                "action": "expand_overseas_sources",
                "reason": f"境外媒体样本{len(overseas)}条，低于目标{targets['target_overseas']}条。",
                "suggested_sources": ["联合早报", "南华早报", "香港电台", "中评社", "澳门新华澳报", "路透", "彭博"],
            }
        )
    low_quality_ratio = len(low_quality) / total if total else 0
    if total and low_quality_ratio > float(targets["max_low_quality_ratio"]):
        actions.append(
            {
                "priority": "P0",
                "action": "replace_low_quality_sources",
                "reason": f"低质量样本占比{low_quality_ratio:.0%}，高于阈值{float(targets['max_low_quality_ratio']):.0%}。",
                "avoid": ["bing_news", "aggregated_search_link"],
            }
        )
    if agent_reach_tasks:
        actions.append(
            {
                "priority": "P1",
                "action": "run_agent_reach_lead_discovery",
                "reason": "用 agent-reach 补充公开讨论线索，再反哺 MediaSpider 关键词和典型样本。",
                "task_manifest": agent_reach_tasks.get("manifest", ""),
            }
        )
    return {
        "targets": targets,
        "current": {
            "total_samples": total,
            "comments": len(comments),
            "overseas": len(overseas),
            "low_quality_ratio": round(low_quality_ratio, 4) if total else 0,
            "by_topic": dict(by_topic),
            "comments_by_topic": dict(comments_by_topic),
        },
        "actions": actions,
        "next_round_recommended": bool(actions),
    }


def build_audit(
    samples: list[dict[str, Any]],
    had_input_samples: bool,
    generated_tasks: dict[str, Any] | None,
    agent_reach_tasks: dict[str, Any] | None,
    collection_gaps: list[str],
    targets: dict[str, Any],
    expansion_plan: dict[str, Any],
) -> dict[str, Any]:
    gaps = list(collection_gaps)
    if not had_input_samples and not samples:
        gaps.append("当前未输入样本文件，公开检索也未获得可用样本；报告只能生成结构和采集任务包，不能形成传播判断。")
    if generated_tasks:
        task_note = "已生成 MediaSpider 兼容采集任务；真实运行通常需要平台登录、二维码、CDP端口或验证码处理。"
        manifest = generated_tasks.get("manifest")
        if manifest:
            task_note += f"任务清单：{manifest}。"
        gaps.append(task_note)
    if samples and not any(is_overseas(x) for x in samples):
        gaps.append("境外媒体样本不足，需继续接入境外新闻/RSS/港澳台媒体检索。")
    if not any(is_comment(x) for x in samples):
        gaps.append("网民评论样本不足，需接入微博、抖音、小红书、B站等评论采集。")
    low_quality_raw = [x for x in samples if x.get("quality_flags")]
    low_quality = [row_sample(x) for x in low_quality_raw]
    low_quality_ratio = round(len(low_quality_raw) / len(samples), 4) if samples else 0
    comment_samples = [x for x in samples if is_comment(x)]
    sentiment_status_counts = Counter(str(x.get("sentiment_status") or "unprocessed") for x in comment_samples)
    sentiment_denominator = sum(1 for x in comment_samples if truthy(x.get("in_sentiment_denominator")))
    acceptance_blockers = []
    if samples and low_quality_ratio >= float(targets["max_low_quality_ratio"]):
        acceptance_blockers.append(f"低质量样本占比{low_quality_ratio:.0%}，超过正式交付阈值{float(targets['max_low_quality_ratio']):.0%}。")
    if len(samples) < int(targets["target_total_samples"]):
        acceptance_blockers.append(f"总样本{len(samples)}条，低于正式交付目标{targets['target_total_samples']}条。")
    if sum(1 for x in samples if is_comment(x)) < int(targets["target_comments"]):
        acceptance_blockers.append(f"网民评论样本不足，低于正式交付目标{targets['target_comments']}条。")
    if sum(1 for x in samples if is_overseas(x)) < int(targets["target_overseas"]):
        acceptance_blockers.append(f"境外媒体样本不足，低于正式交付目标{targets['target_overseas']}条。")
    if samples and all("public_web_sample" in set(x.get("quality_flags") or []) for x in samples):
        acceptance_blockers.append("当前有效样本全部来自公开搜索兜底，未接入正式平台采集或内部监测数据。")
    if samples and not any(is_comment(x) for x in samples):
        acceptance_blockers.append("缺少可核验网民原始评论，不能支撑正式报告的网民评论模块。")
    if samples and not any(is_overseas(x) for x in samples):
        acceptance_blockers.append("缺少境外媒体样本，不能支撑正式报告的境外关注模块。")
    if comment_samples and sentiment_status_counts.get("unprocessed", 0):
        acceptance_blockers.append(
            f"有{sentiment_status_counts['unprocessed']}条网民评论尚未经过大规模情感分析流程或仍待复核，不能计算正式情感比例。"
        )
    if comment_samples and sentiment_denominator <= 0:
        acceptance_blockers.append("网民评论情感有效分母为0，不能生成子议题正面、中性、负面比例。")
    return {
        "data_gaps": gaps,
        "low_quality_samples": low_quality,
        "quality_summary": {
            "total_samples": len(samples),
            "low_quality_samples": len(low_quality_raw),
            "low_quality_ratio": low_quality_ratio,
            "flag_counts": dict(Counter(flag for sample in samples for flag in sample.get("quality_flags", []))),
            "sentiment_status_counts": dict(sentiment_status_counts),
            "sentiment_denominator": sentiment_denominator,
        },
        "collection_targets": targets,
        "expansion_plan": expansion_plan,
        "acceptance": {
            "ready_for_formal_delivery": not acceptance_blockers,
            "blockers": acceptance_blockers,
        },
        "review_items": [
            "复核子议题归类是否准确，尤其是标题泛化或新闻合集类样本。",
            "复核候选热词，删除过泛词，保留能解释公众关注点的词。",
            "复核网民评论是否真实、可引用、无敏感风险，不要过度润色原话。",
            "确认情感结果由 large-scale-sentiment-analysis 的AI语义标注、分类器和复核流程生成，并保留分母与排除原因。",
            "复核境外样本的事实性报道、解读性报道、借题炒作/风险解读分类及中文译文。",
            "正式交付前用内部传播系统或真实采集结果替换公开检索/小样本结果。",
        ],
    }


def build_appendices(samples: list[dict[str, Any]], overseas: dict[str, Any], selected_comments: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    media = [row_sample(x) for x in top_items([x for x in samples if source_bucket(x) in {"domestic_media", "self_media"}], 30)]
    wechat_candidates = top_items(
        [x for x in samples if str(x.get("platform")).lower() in {"wechat", "weixin"}],
        len(samples),
    )
    wechat = []
    seen_wechat_sources: set[str] = set()
    for sample in wechat_candidates:
        source_key = re.sub(
            r"[^0-9a-z\u3400-\u9fff]+",
            "",
            unicodedata.normalize("NFKC", str(sample.get("source") or "")).lower(),
        )
        source_key = source_key or canonical_source_url(str(sample.get("url") or "")) or str(sample.get("id") or "")
        if source_key in seen_wechat_sources:
            continue
        seen_wechat_sources.add(source_key)
        wechat.append(row_sample(sample))
        if len(wechat) >= 20:
            break
    overseas_rows = []
    for category, rows in overseas["groups"].items():
        for row in rows:
            item = dict(row)
            item["category"] = category
            overseas_rows.append(item)
    return {"media_top": media, "wechat_top": wechat, "selected_comments": selected_comments, "overseas_reports": overseas_rows}


def system_ratio_counts(ratios: dict[str, Any], scale: int = 10000) -> dict[str, int]:
    counts = {
        name: max(0, round(float(ratios.get(name, 0) or 0) * scale))
        for name in ["positive", "neutral", "negative"]
    }
    delta = scale - sum(counts.values())
    if delta and any(counts.values()):
        counts["neutral"] = max(0, counts["neutral"] + delta)
    return counts


def apply_system_authority(
    system_data: dict[str, Any],
    samples: list[dict[str, Any]],
    statistics: dict[str, Any],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    overall = system_data.get("overall") or {}
    totals = overall.get("totals") or {}
    daily = overall.get("daily") or []
    domestic = int(totals.get("domestic_mainstream", 0) or 0)
    overseas = int(totals.get("overseas_media", 0) or 0)
    wechat = int(totals.get("wechat_public", 0) or 0)
    weibo = int(totals.get("weibo", 0) or 0)
    video = int(totals.get("video_account", 0) or 0)
    new_media = int(totals.get("new_media", 0) or 0)
    total_spread = int(totals.get("total_spread", 0) or 0)
    authoritative = dict(statistics)
    authoritative.update(
        {
            "total_spread": total_spread,
            "total_comments": sum(1 for item in samples if is_comment(item)),
            "by_platform": {
                "wechat": wechat,
                "weibo": weibo,
                "video_account": video,
                "new_media": new_media,
            },
            "by_source_bucket": {
                "domestic_media": domestic,
                "overseas_media": overseas,
                "self_media": wechat + weibo + video + new_media,
                "comments": sum(1 for item in samples if is_comment(item)),
            },
            "by_region": {"domestic": max(0, total_spread - overseas), "overseas": overseas},
            "by_date": {str(item.get("date")): int(item.get("total_spread", 0) or 0) for item in daily if item.get("date")},
            "authority": "monitoring_system",
        }
    )
    authoritative["source_bucket_labels"] = dict(system_data.get("channel_labels") or {})
    topic_stats: list[dict[str, Any]] = []
    for item in system_data.get("subevents") or []:
        title = str(item.get("title") or "")
        topic_samples = [
            sample
            for sample in samples
            if sample.get("topic") == title or title in (sample.get("topic_hits") or [])
        ]
        topic_stats.append(
            {
                "topic": title,
                "display": topic_display(title),
                "total_samples": int(item.get("total_spread", 0) or 0),
                "spread_count": int(item.get("total_spread", 0) or 0),
                "comment_count": sum(int(sample.get("comment_count", 0) or 0) for sample in topic_samples),
                "domestic_media": int(item.get("domestic_mainstream", 0) or 0),
                "self_media": int(item.get("new_media", 0) or 0),
                "comments": sum(1 for sample in topic_samples if is_comment(sample)),
                "overseas_media": int(item.get("overseas_media", 0) or 0),
                # The workbook's sentiment is retained only as a reference field. Formal
                # ratios must come from audited comment-level analysis results.
                "sentiment": sentiment_counts([sample for sample in topic_samples if is_comment(sample)]),
                "system_sentiment_reference": item.get("sentiment") or {},
                "sentiment_authority": "large_scale_sentiment_analysis",
                "top_samples": [row_sample(sample) for sample in top_items(topic_samples, 3)],
                "authority": "monitoring_system",
            }
        )
    return authoritative, topic_stats


def apply_sentiment_summary(
    topic_stats: list[dict[str, Any]],
    summary: dict[str, Any],
) -> list[dict[str, Any]]:
    """Attach the workbook denominator gate without hiding observed comments."""
    summary_rows = list(summary.get("topics") or [])
    for row in topic_stats:
        row_key = canonical_topic_key(row.get("topic"))
        matched = next(
            (
                item
                for item in summary_rows
                if canonical_topic_key(item.get("title")) == row_key
                or (
                    row_key
                    and (
                        row_key in canonical_topic_key(item.get("title"))
                        or canonical_topic_key(item.get("title")) in row_key
                    )
                )
            ),
            None,
        )
        if not matched:
            row["sentiment_formal_ready"] = False
            row["sentiment_sample_status"] = "pending"
            row["sentiment_denominator"] = 0
            continue
        row["sentiment_formal_ready"] = matched.get("status") == "ready"
        row["sentiment_sample_status"] = str(matched.get("status") or "pending")
        row["sentiment_denominator"] = int(matched.get("denominator") or 0)
        if isinstance(matched.get("counts"), dict):
            row["sentiment"] = {
                label: int((matched.get("counts") or {}).get(label, 0) or 0)
                for label in ["positive", "neutral", "negative"]
            }
    return topic_stats


def build_system_audit(
    system_data: dict[str, Any],
    samples: list[dict[str, Any]],
    collection_gaps: list[str],
    viewpoints: dict[str, Any],
    analysis_metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    gaps = list(collection_gaps)
    blockers: list[str] = []
    research_budget_gaps = list((analysis_metadata or {}).get('research_budget_gaps') or [])
    gaps.extend(research_budget_gaps)
    blockers.extend(research_budget_gaps)
    validation = system_data.get("validation") or []
    for item in validation:
        message = str(item.get("message") or item)
        gaps.append(message)
        if item.get("severity") == "error":
            blockers.append(message)
    total_spread = int(((system_data.get("overall") or {}).get("totals") or {}).get("total_spread", 0) or 0)
    subevents = system_data.get("subevents") or []
    topics = [str(item.get("title") or "") for item in subevents]
    comments = [sample for sample in samples if is_comment(sample)]
    traceable_comments = [sample for sample in comments if is_actual_platform_comment(sample)]
    formally_reviewed_comments = [
        sample for sample in traceable_comments
        if isinstance(sample.get("ai_formal_include"), bool)
        and str(sample.get("ai_semantic_quality") or "").strip()
        and str(sample.get("ai_formal_reason") or "").strip()
        and (
            sample.get("ai_formal_include") is False
            or str(sample.get("topic_comment_heading") or sample.get("comment_heading") or sample.get("ai_comment_heading") or "").strip()
        )
    ]
    domestic_detail = [sample for sample in samples if source_bucket(sample) == "domestic_media"]
    viewpoint_details = [sample for sample in samples if source_bucket(sample) in {"domestic_media", "self_media"}]
    analyzed_comments = [
        sample for sample in comments
        if sample.get("sentiment_status") == "classified"
        and sample.get("in_sentiment_denominator")
        and sample.get("sentiment_source") in VERIFIED_SENTIMENT_SOURCES
    ]
    system_sentiment_final = bool(comments) and len(analyzed_comments) == len(
        [sample for sample in comments if sample.get("sentiment_status") != "excluded"]
    )
    if total_spread <= 0:
        blockers.append("系统数据包未提供总事件传播总量。")
    if not subevents or any(int(item.get("total_spread", 0) or 0) <= 0 for item in subevents):
        blockers.append("系统数据包缺少一个或多个子事件的传播总量。")
    if not traceable_comments:
        blockers.append("尚未取得可追溯的公开网民原话；必须完成批准的境内平台评论采集和逐条AI审核后才能生成正式报告。")
    elif len(formally_reviewed_comments) != len(traceable_comments):
        blockers.append("可追溯网民评论尚未全部完成正式引用AI审核；每条必须明确是否适合引用、语义质量、审核理由，拟引用评论还须填写态度型小标题。")
    if not viewpoint_details:
        blockers.append("尚未取得可追溯的境内媒体、自媒体或专家公开证据。")
    elif not domestic_detail:
        gaps.append("系统数据包未包含完整境内主流媒体文章明细；当前定性分析仅可使用经核验的公开网络补证。")
    missing_viewpoint_topics = [topic_display(topic) for topic in topics if not any(sample.get("topic") == topic or topic in (sample.get("topic_hits") or []) for sample in viewpoint_details)]
    if missing_viewpoint_topics:
        blockers.append("以下子议题缺少可追溯的媒体/自媒体观点证据：" + "、".join(missing_viewpoint_topics) + "。")
    missing_comment_topics = [topic_display(topic) for topic in topics if not any(sample.get("topic") == topic or topic in (sample.get("topic_hits") or []) for sample in traceable_comments)]
    thin_viewpoint_clusters: list[str] = []
    viewpoint_density_exceptions: list[dict[str, Any]] = []
    for item in viewpoints.get("by_topic") or []:
        topic_name = topic_display(str(item.get("topic") or ""))
        for cluster in item.get("clusters") or []:
            density = cluster_density_result(cluster)
            if not density["passed"]:
                thin_viewpoint_clusters.append(f"{topic_name}：{cluster.get('summary') or '未命名观点'}")
            elif density["accepted_exception"]:
                viewpoint_density_exceptions.append({
                    "topic": topic_name,
                    "summary": cluster.get("summary"),
                    **density,
                })
    if thin_viewpoint_clusters:
        blockers.append(
            "以下境内观点簇未达到正式成稿的来源与论述密度门禁："
            + "、".join(thin_viewpoint_clusters)
            + "。"
        )
    verified_comment_groups: Counter[str] = Counter(
        str(sample.get("system_topic") or sample.get("topic") or sample.get("subtopic") or "")
        for sample in traceable_comments
        if sample.get("ai_formal_include") is True
    )
    thin_comment_groups = [name for name, count in verified_comment_groups.items() if name and count < 2]
    inconsistent_comment_topics: list[str] = []
    for topic_name in verified_comment_groups:
        rows = [
            sample for sample in traceable_comments
            if str(sample.get("system_topic") or sample.get("topic") or sample.get("subtopic") or "") == topic_name
            and sample.get("ai_formal_include") is True
        ]
        common = {
            str(row.get("topic_comment_heading") or "").strip()
            for row in rows if str(row.get("topic_comment_heading") or "").strip()
        }
        per_row = {
            str(row.get("comment_heading") or row.get("ai_comment_heading") or "").strip()
            for row in rows if str(row.get("comment_heading") or row.get("ai_comment_heading") or "").strip()
        }
        if len(common) > 1 or (not common and len(per_row) > 1):
            inconsistent_comment_topics.append(topic_name)
    if inconsistent_comment_topics:
        blockers.append(
            "同一子议题的评论存在多个并列小标题；须由AI填写唯一topic_comment_heading后再成稿："
            + "、".join(inconsistent_comment_topics) + "。"
        )
    viewpoint_quality_issues = domestic_viewpoint_quality_issues({
        "viewpoints": viewpoints, "metadata": analysis_metadata or {},
    })
    viewpoint_quality_errors = [
        item for item in viewpoint_quality_issues
        if item.get("severity") == "error" or item.get("strict_severity") == "error"
        or item.get("code") == "domestic_interpretation_gap"
    ]
    blockers.extend(str(item.get("message") or "") for item in viewpoint_quality_errors)
    if comments and not system_sentiment_final:
        blockers.append("评论明细尚未完成自有情感分析；需运行 large-scale-sentiment-analysis 并按 sample_id 回灌，系统原情感标签不进入正式统计。")
    actions = []
    if not traceable_comments:
        actions.append({"priority": "P0", "action": "collect_domestic_platform_comments", "reason": "运行批准的境内平台评论采集，保留评论ID、原帖链接、时间和原始文件，再进入逐条AI审核与情感阶段。"})
    elif missing_comment_topics:
        actions.append({"priority": "P1", "action": "research_public_comment_evidence", "reason": "继续补充缺失子议题的公开讨论原话：" + "、".join(missing_comment_topics) + "。"})
    if not viewpoint_details:
        actions.append({"priority": "P0", "action": "research_public_media_evidence", "reason": "自动检索媒体、专家和自媒体公开解读，用于观点聚类；不得改写系统统计量。"})
    elif missing_viewpoint_topics:
        actions.append({"priority": "P0", "action": "research_public_media_evidence", "reason": "自动补充缺失子议题的公开媒体/自媒体证据：" + "、".join(missing_viewpoint_topics) + "。"})
    if thin_viewpoint_clusters:
        actions.append({
            "priority": "P1",
            "action": "deepen_viewpoint_evidence",
            "reason": f"{len(thin_viewpoint_clusters)}个观点簇未通过共享成文密度门禁；补充真实证据或可审核的thin_cluster_exception，不得堆字或重复归因。",
        })
    if viewpoint_quality_issues:
        actions.append({
            "priority": "P0" if viewpoint_quality_errors else "P1",
            "action": "revise_domestic_viewpoint_prose",
            "reason": (
                f"境内媒体自媒体成文发现{len(viewpoint_quality_issues)}项来源复用、归因或原话密度问题；"
                "应回到证据原文合并重复声音、修正媒体主体写法并补足过短表述。"
            ),
        })
    if thin_comment_groups:
        actions.append({
            "priority": "P1",
            "action": "deepen_public_comment_evidence",
            "reason": "以下评论组仅有1条可核验原话，建议补至2至3条：" + "、".join(thin_comment_groups) + "。",
        })
    if not system_data.get("wechat_top"):
        actions.append({"priority": "P1", "action": "request_system_wechat_top", "reason": "补导公众号阅读量和在看量榜单。"})
    if not system_data.get("overseas_reports"):
        actions.append({"priority": "P1", "action": "request_system_overseas_list", "reason": "补导境外报道明细及子事件命中标签。"})
    gaps.extend(item for item in blockers if item not in gaps)
    return {
        "data_gaps": gaps,
        "low_quality_samples": [],
        "quality_summary": {
            "data_authority": "monitoring_system",
            "system_total_spread": total_spread,
            "system_subevents": len(subevents),
            "detail_samples": len(samples),
            "comment_details": len(comments),
            "traceable_comment_details": len(traceable_comments),
            "formally_reviewed_comment_details": len(formally_reviewed_comments),
            "domestic_media_details": len(domestic_detail),
            "system_sentiment_final": system_sentiment_final,
            "sentiment_authority": "large_scale_sentiment_analysis",
            "analyzed_comment_details": len(analyzed_comments),
            "missing_viewpoint_topics": missing_viewpoint_topics,
            "missing_comment_topics": missing_comment_topics,
            "thin_viewpoint_clusters": thin_viewpoint_clusters,
            "viewpoint_density_exceptions": viewpoint_density_exceptions,
            "thin_comment_groups": thin_comment_groups,
            "domestic_viewpoint_quality_issues": viewpoint_quality_issues,
        },
        "collection_targets": {"mode": "system_totals_plus_public_evidence", "fixed_sample_targets_disabled": True},
        "expansion_plan": {"mode": "automatic_public_evidence_research", "actions": actions, "next_round_recommended": bool(actions)},
        "acceptance": {"ready_for_formal_delivery": not blockers, "blockers": list(dict.fromkeys(blockers))},
        "review_items": [
            "核对系统导出的监测起止时间、事件ID和统计口径。",
            "确认总事件全局去重、子事件允许重复命中的口径未被改写。",
            "复核媒体、专家和自媒体公开证据对应的观点聚类，不得把公开样本数写成系统传播量。",
            "复核网民评论原话、平台、时间和链接；无法追溯的内容只能作为观点概括，不得加引号。",
            "正式情感比例只使用 large-scale-sentiment-analysis 回灌结果；系统原标签仅作原始字段留存。",
        ],
    }


def foreign_collection_audit_succeeded(evidence: dict[str, Any]) -> bool:
    if not evidence or evidence.get("dry_run") is not False:
        return False
    if evidence.get("collection_completed") is True:
        return True
    return int(evidence.get("command_exit_code", 0) or 0) == 0


def formal_delivery_ready(data: dict[str, Any]) -> bool:
    return bool(
        ((data.get("audit") or {}).get("acceptance") or {}).get("ready_for_formal_delivery")
    )


def write_outputs(data: dict[str, Any], out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    chart_dir = out_dir / "charts"
    appendix_dir = out_dir / "appendices"
    chart_dir.mkdir(exist_ok=True)
    appendix_dir.mkdir(exist_ok=True)
    data["artifacts"]["sentiment_analysis"] = write_sentiment_handoff(data, out_dir)

    data_path = out_dir / "report_data.json"
    audit_path = out_dir / "cwh_audit.json"
    data["artifacts"]["report_data"] = str(data_path)
    data["artifacts"]["audit"] = str(audit_path)
    from cwh_available_delivery import available_delivery
    deliver_available = available_delivery(data.get("analysis_bundle") or {})
    if not formal_delivery_ready(data) and not deliver_available:
        data["artifacts"]["formal_report_status"] = "blocked_by_acceptance_gate"
        data_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        audit_path.write_text(json.dumps(data["audit"], ensure_ascii=False, indent=2), encoding="utf-8")
        return

    if not formal_delivery_ready(data):
        data["delivery_class"] = "available_with_gaps"
        data["artifacts"]["formal_report_status"] = "generated_with_evidence_gaps"

    report = render_report(data)
    report_path = out_dir / "cwh_report.md"
    report_path.write_text(report, encoding="utf-8")
    data["artifacts"]["report"] = str(report_path)
    data_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    audit_path.write_text(json.dumps(data["audit"], ensure_ascii=False, indent=2), encoding="utf-8")

    write_bar_svg(chart_dir / "platform_distribution.svg", "平台样本分布", data["statistics"]["by_platform"])
    write_bar_svg(chart_dir / "topic_distribution.svg", "子议题样本分布", {topic_display(k): v for k, v in data["statistics"]["by_topic"].items()}, "#7b4f2c")
    write_bar_svg(chart_dir / "hotword_distribution.svg", "候选热词分布", {x["word"]: int(x["count"]) for x in data["hotwords"][:12]}, "#5a6f2f")
    ordered_dates = dict(sorted(data["statistics"]["by_date"].items()))
    write_bar_svg(chart_dir / "trend_distribution.svg", "传播走势", ordered_dates, "#345995")
    data["artifacts"]["charts"] = {
        "platform_distribution": str(chart_dir / "platform_distribution.svg"),
        "topic_distribution": str(chart_dir / "topic_distribution.svg"),
        "hotword_distribution": str(chart_dir / "hotword_distribution.svg"),
        "trend_distribution": str(chart_dir / "trend_distribution.svg"),
    }
    for name, rows in data["appendices"].items():
        path = appendix_dir / f"{name}.csv"
        write_csv(path, rows)
        data["artifacts"].setdefault("appendices", {})[name] = str(path)
    try:
        workbook_path = out_dir / "cwh_data_workbook.xlsx"
        write_cwh_data_workbook(data, workbook_path)
        data["artifacts"]["data_workbook"] = str(workbook_path)
    except Exception as exc:
        data.setdefault("audit", {}).setdefault("data_gaps", []).append(f"数据工作簿生成失败：{type(exc).__name__}: {exc}")
    try:
        from formalize_cwh_report import formalize_report

        formalize_report(data, out_dir)
    except Exception as exc:
        data.setdefault("audit", {}).setdefault("data_gaps", []).append(f"正式稿/Word渲染失败：{type(exc).__name__}: {exc}")
    if deliver_available and not formal_delivery_ready(data):
        # Formatting can add a late quality warning. Label the dashboard from
        # the final audit, not only the audit observed before generation.
        data["delivery_class"] = "available_with_gaps"
        data["artifacts"]["formal_report_status"] = "generated_with_evidence_gaps"
    try:
        from generate_dashboard import generate_dashboard

        dashboard_path = generate_dashboard(data, out_dir / "cwh_dashboard.html")
        data["artifacts"]["dashboard"] = str(dashboard_path)
    except Exception as exc:
        data.setdefault("audit", {}).setdefault("data_gaps", []).append(f"HTML看板生成失败：{type(exc).__name__}: {exc}")
    data_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    audit_path.write_text(json.dumps(data["audit"], ensure_ascii=False, indent=2), encoding="utf-8")


def orchestrate(args: argparse.Namespace) -> dict[str, Any]:
    agenda = read_text(args.input_file, args.input) or "本次国务院常务会议"
    topics = infer_topics(agenda)
    out_dir = Path(args.out_dir)
    system_data: dict[str, Any] = {}
    if args.system_workbook:
        system_data = ingest_workbook(args.system_workbook)
        system_topics = [str(item.get("title") or "").strip() for item in system_data.get("topics") or []]
        if system_topics:
            topics = system_topics
        system_period = system_data.get("monitoring_period") or {}
        args.monitor_start = args.monitor_start or str(system_period.get("start") or "")
        args.monitor_end = args.monitor_end or str(system_period.get("end") or "")
    system_mode = bool(system_data) or args.data_mode == "system"
    if system_mode and not str(getattr(args, "hotword_audit", "") or "").strip():
        raise ValueError(
            "正式系统数据流程必须提供AI审核完成的--hotword-audit；"
            "禁止由报告阶段使用规则候选词生成正式热词"
        )
    analysis_bundle = load_analysis_bundle(args.analysis_bundle)
    sentiment_summary = load_sentiment_summary(getattr(args, "sentiment_summary", ""))
    input_samples = load_samples(args.samples, topics)
    comment_handoff_files = list(getattr(args, "comment_handoff", []) or [])
    comment_handoff_samples = load_comment_handoffs(comment_handoff_files, topics)
    if args.system_details:
        input_samples.extend(load_samples(args.system_details, topics))
    system_samples: list[dict[str, Any]] = []
    for idx, row in enumerate(system_data.get("evidence_samples") or [], 1):
        sample = normalize_sample(row, topics, idx)
        if sample:
            sample["topic_hits"] = row.get("topic_hits") or []
            system_samples.append(sample)
    analysis_samples: list[dict[str, Any]] = []
    for idx, row in enumerate(analysis_bundle_rows(analysis_bundle), 1):
        sample = normalize_sample(row, topics, 100000 + idx)
        if sample:
            analysis_samples.append(sample)
    web_samples: list[dict[str, Any]] = []
    mediaspider_samples: list[dict[str, Any]] = []
    foreign_mediaspider_samples: list[dict[str, Any]] = []
    agent_reach_samples: list[dict[str, Any]] = []
    mediaspider_run: dict[str, Any] = {}
    foreign_mediaspider_run: dict[str, Any] = {}
    agent_reach_run: dict[str, Any] = {}
    agent_reach_tasks = None
    collection_gaps: list[str] = []
    supplemental_collection = (
        args.supplemental_collection
        or args.run_mediaspider
        or args.run_foreign_mediaspider
        or args.run_agent_reach
    )
    if args.auto_web and (not system_mode or supplemental_collection) and (not input_samples or args.web_always):
        web_samples, collection_gaps = collect_public_web_samples(topics, args.web_limit, args.web_timeout)

    generated_tasks = None
    if not args.no_tasks and (not system_mode or supplemental_collection):
        platforms = [x.strip() for x in args.platforms.split(",") if x.strip()]
        generated_tasks = generate_mediacrawler_tasks(
            agenda,
            topics,
            out_dir,
            platforms,
            args.max_posts,
            args.max_comments,
            args.collection_profile,
        )
        if args.agent_reach_plan:
            agent_reach_platforms = [x.strip() for x in args.agent_reach_platforms.split(",") if x.strip()]
            agent_reach_tasks = generate_agent_reach_tasks(topics, out_dir, agent_reach_platforms)

    if args.run_agent_reach:
        agent_reach_samples, agent_reach_run, agent_reach_gaps = run_agent_reach_and_ingest(args, agent_reach_tasks, topics, out_dir)
        collection_gaps.extend(agent_reach_gaps)

    if args.run_mediaspider:
        mediaspider_samples, mediaspider_run, mediaspider_gaps = run_mediacrawler_and_ingest(args, generated_tasks, topics, out_dir)
        collection_gaps.extend(mediaspider_gaps)

    if args.run_foreign_mediaspider:
        foreign_mediaspider_samples, foreign_mediaspider_run, foreign_mediaspider_gaps = run_foreign_mediaspider_and_ingest(
            args,
            agenda,
            topics,
            out_dir,
        )
        collection_gaps.extend(foreign_mediaspider_gaps)

    samples = dedupe_samples(
        input_samples
        + comment_handoff_samples
        + system_samples
        + analysis_samples
        + web_samples
        + mediaspider_samples
        + foreign_mediaspider_samples
        + agent_reach_samples
    )
    samples, dropped_by_period = filter_samples_by_monitoring_period(samples, args.monitor_start, args.monitor_end)
    if dropped_by_period:
        collection_gaps.append(f"已按监测窗口过滤样本{dropped_by_period}条，避免后续信息倒灌。")

    reset_inherited_sentiment(samples)
    sentiment_result_count, sentiment_gaps = apply_sentiment_results(samples, args.sentiment_results)
    collection_gaps.extend(sentiment_gaps)

    collection_gaps = reconcile_collection_gaps(collection_gaps, samples, topics)

    statistics = build_statistics(samples, topics)
    topic_stats = build_topic_stats(topics, samples)
    if system_data:
        statistics, topic_stats = apply_system_authority(system_data, samples, statistics)
    if sentiment_summary:
        topic_stats = apply_sentiment_summary(topic_stats, sentiment_summary)
    viewpoints = build_viewpoints(topics, samples)
    if (analysis_bundle.get("viewpoints") or {}).get("by_topic"):
        viewpoints = analysis_bundle["viewpoints"]
    viewpoints = sanitize_domestic_viewpoints(viewpoints, samples)
    selected_comments = select_comments(samples, args.comment_limit)
    reviewed_hotwords = analysis_bundle.get("hotwords") or []
    hotwords = (
        []
        if str(getattr(args, "hotword_audit", "") or "").strip()
        else build_evidence_hotwords(topics, samples, reviewed_hotwords)
    )
    overseas = (
        build_system_reviewed_overseas(system_data, samples)
        if system_data and system_data.get("overseas_reports")
        else build_overseas_v2(samples)
    )
    appendices = build_appendices(samples, overseas, selected_comments)
    targets = collection_targets(args, topics)
    if system_data:
        audit = build_system_audit(system_data, samples, collection_gaps, viewpoints,
                                   analysis_bundle.get("metadata"))
        if not (analysis_bundle.get("viewpoints") or {}).get("by_topic"):
            message = "正式系统报告缺少经过公开网络研究和AI审核的analysis_bundle.json。"
            audit.setdefault("data_gaps", []).append(message)
            acceptance = audit.setdefault("acceptance", {})
            acceptance.setdefault("blockers", []).append(message)
            acceptance["blockers"] = list(dict.fromkeys(acceptance["blockers"]))
            acceptance["ready_for_formal_delivery"] = False
            audit.setdefault("expansion_plan", {}).setdefault("actions", []).append(
                {
                    "priority": "P0",
                    "action": "research_public_media_evidence",
                    "reason": "按全部系统子议题完成境内媒体、专家和自媒体公开证据研究，并生成analysis_bundle.json。",
                }
            )
            audit["expansion_plan"]["next_round_recommended"] = True
        expansion_plan = audit["expansion_plan"]
    else:
        expansion_plan = build_expansion_plan(samples, topics, targets, generated_tasks, agent_reach_tasks)
        audit = build_audit(samples, bool(input_samples), generated_tasks, agent_reach_tasks, collection_gaps, targets, expansion_plan)
    corpus_reviews = (((analysis_bundle.get("research_audit") or {}).get("domestic_media_research") or {}).get("public_article_corpus_review") or {}).get("topic_reviews") or []
    audit.setdefault("quality_summary", {})["raw_corpus_review_coverage"] = [
        {"topic": row.get("topic"), "reviewed_documents": len(set(row.get("reviewed_record_ids") or [])),
         "deferred_documents": len(set(row.get("deferred_record_ids") or [])), "deferral_reason": row.get("deferral_reason") or ""}
        for row in corpus_reviews if isinstance(row, dict)
    ]
    data = {
        "version": SCHEMA_VERSION,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "meeting": {
            "date": args.monitor_start or normalize_date(agenda),
            "agenda": agenda,
            "topics": topics,
            "topic_keywords": {topic: extract_terms(topic) for topic in topics},
            "search_queries": {topic: keywords_for(topic) for topic in topics},
        },
        "collection": {
            "data_mode": "system" if system_mode else "crawler_or_samples",
            "system_workbook": args.system_workbook or "",
            "system_details_file": args.system_details or "",
            "system_evidence_sample_count": len(system_samples),
            "analysis_bundle_file": args.analysis_bundle or "",
            "analysis_sample_count": len(analysis_samples),
            "input_sample_file": args.samples or "",
            "sample_count": len(samples),
            "input_sample_count": len(input_samples),
            "comment_handoff_files": comment_handoff_files,
            "comment_handoff_count": len(comment_handoff_samples),
            "public_web_sample_count": len(web_samples),
            "mediaspider_sample_count": len(mediaspider_samples),
            "foreign_mediaspider_sample_count": len(foreign_mediaspider_samples),
            "agent_reach_sample_count": len(agent_reach_samples),
            "sentiment_results_file": args.sentiment_results or "",
            "sentiment_result_count": sentiment_result_count,
            "mediacrawler_tasks": generated_tasks,
            "agent_reach_tasks": agent_reach_tasks,
            "mediaspider_run": mediaspider_run,
            "foreign_mediaspider_run": foreign_mediaspider_run,
            "agent_reach_run": agent_reach_run,
            "collection_profile": args.collection_profile,
            "collection_targets": targets,
            "collection_gaps": collection_gaps,
            "monitoring_period": {"start": args.monitor_start or "", "end": args.monitor_end or ""},
            "data_authority_note": "正式传播数据以监测系统导出为权威口径；MediaSpider和agent-reach仅可作为研发验证或缺失证据补充，不得替代系统传播总量。" if system_mode else "当前未接入监测系统工作簿。",
            "data_scope_note": "正式传播量以监测系统导出为主；正式情感比例只使用large-scale-sentiment-analysis对原始评论逐条分析后的回灌结果，系统原情感标签不进入分母。" if system_mode else "兼容演示路径可使用样本或MediaSpider；该路径不得冒充部门正式监测数据。",
        },
        "system_data": system_data,
        "analysis_bundle": analysis_bundle,
        "sentiment_summary": sentiment_summary,
        "statistics": statistics,
        "topic_stats": topic_stats,
        "viewpoints": viewpoints,
        "comments": {"selected": selected_comments, "sentiment": sentiment_counts([x for x in samples if is_comment(x) and not is_overseas(x)])},
        "hotwords": hotwords,
        "overseas": overseas,
        "appendices": appendices,
        "audit": audit,
        "samples": samples,
        "artifacts": {
            **({"wordcloud_image": str(Path(args.wordcloud_image).resolve())} if args.wordcloud_image else {}),
            **({"trend_chart_image": str(Path(args.trend_chart_image).resolve())} if getattr(args, "trend_chart_image", "") else {}),
            **({"topic_chart_image": str(Path(args.topic_chart_image).resolve())} if getattr(args, "topic_chart_image", "") else {}),
        },
    }
    overseas_supplements = str(getattr(args, "overseas_supplements", "") or "").strip()
    if overseas_supplements:
        from merge_public_supplements import load_json as load_supplement_json
        from merge_public_supplements import merge as merge_public_supplements

        supplement_path = Path(overseas_supplements).resolve()
        data = merge_public_supplements(data, load_supplement_json(supplement_path))
        reconcile_merged_overseas(data)
        data.setdefault("collection", {})["overseas_supplements_file"] = str(supplement_path)
    foreign_collection_audit_path = str(getattr(args, "foreign_collection_audit", "") or "").strip()
    supplied_foreign_audit: dict[str, Any] = {}
    if foreign_collection_audit_path:
        audit_path = Path(foreign_collection_audit_path).resolve()
        supplied_foreign_audit = json.loads(audit_path.read_text(encoding="utf-8-sig"))
        data.setdefault("collection", {})["foreign_collection_audit_file"] = str(audit_path)
    if system_mode:
        foreign_run_evidence = supplied_foreign_audit or foreign_mediaspider_run
        foreign_attempted = bool(
            args.run_foreign_mediaspider
            or (foreign_run_evidence and foreign_run_evidence.get("dry_run") is False)
        )
        foreign_succeeded = bool(foreign_attempted and foreign_collection_audit_succeeded(foreign_run_evidence))
        data.setdefault("collection", {})["foreign_public_discussion_collection_attempted"] = foreign_attempted
        data["collection"]["foreign_public_discussion_collection_succeeded"] = foreign_succeeded
        message = ""
        if not foreign_attempted:
            message = "尚未实际运行境外公开讨论采集；不得以境外媒体报道列表代替境外网民评论采集，也不得据此写‘暂无评论性观点’。"
        elif not foreign_succeeded:
            message = "境外公开讨论采集已尝试但未成功完成；需处理平台登录、风控、网络或采集器故障后再生成正式报告。"
        if message:
            audit = data.setdefault("audit", {})
            audit.setdefault("data_gaps", []).append(message)
            acceptance = audit.setdefault("acceptance", {})
            acceptance.setdefault("blockers", []).append(message)
            acceptance["blockers"] = list(dict.fromkeys(acceptance["blockers"]))
            acceptance["ready_for_formal_delivery"] = False
    if args.hotword_audit:
        from formalize_cwh_report import apply_hotword_audit

        apply_hotword_audit(data, Path(args.hotword_audit))
    for notice in (sentiment_summary.get('notice'), supplied_foreign_audit.get('notice')):
        if notice:
            audit = data.setdefault('audit', {})
            audit.setdefault('data_gaps', []).append(notice)
            acceptance = audit.setdefault('acceptance', {})
            acceptance['ready_for_formal_delivery'] = False
            acceptance.setdefault('blockers', []).append(notice)
    write_outputs(data, out_dir)
    return data


def main() -> None:
    parser = argparse.ArgumentParser(description="One-click CWH public-opinion report demo orchestrator.")
    parser.add_argument("--input", default="", help="Meeting agenda text.")
    parser.add_argument("--input-file", help="Read agenda text from UTF-8 file.")
    parser.add_argument("--samples", help="CSV/JSON normalized or raw sample table.")
    parser.add_argument(
        "--comment-handoff",
        action="append",
        default=[],
        help="Strict report_comment_handoff.csv from the sentiment stage; may be repeated.",
    )
    parser.add_argument("--system-workbook", default="", help="Authoritative monitoring-system CWH workbook (.xlsx).")
    parser.add_argument("--system-details", default="", help="Optional system-exported detail CSV/JSON for media articles and netizen comments.")
    parser.add_argument("--analysis-bundle", default="", help="Internal AI/research evidence bundle generated automatically by the skill.")
    parser.add_argument(
        "--overseas-supplements",
        default="",
        help="Reviewed JSON bundle with media/comments keys. Merged without changing monitoring-system totals.",
    )
    parser.add_argument(
        "--foreign-collection-audit",
        default="",
        help="Completed non-dry-run foreign_mediaspider_run.json proving that overseas public-discussion collection was actually attempted.",
    )
    parser.add_argument("--data-mode", choices=["auto", "system", "crawler"], default="auto", help="Formal data authority mode; auto selects system when --system-workbook is supplied.")
    parser.add_argument("--supplemental-collection", action="store_true", help="Allow crawler/search supplements without replacing system totals.")
    parser.add_argument("--sentiment-results", default="", help="CSV/JSON sentiment results keyed by sample_id from large-scale-sentiment-analysis.")
    parser.add_argument(
        "--sentiment-summary",
        default="",
        help="sentiment_workbook_summary.json; controls formal percentage readiness in Excel/dashboard/report.",
    )
    parser.add_argument(
        "--wordcloud-image",
        default="",
        help="Optional reviewed word-cloud PNG. It replaces only the report/dashboard word cloud and preserves system trend/topic charts.",
    )
    parser.add_argument("--trend-chart-image", default="", help="Optional reviewed monitoring-workbook trend chart PNG.")
    parser.add_argument("--topic-chart-image", default="", help="Optional reviewed monitoring-workbook subtopic chart PNG.")
    parser.add_argument(
        "--hotword-audit",
        default="",
        help="AI-reviewed hotword_audit.json; required for formal system-workbook reports.",
    )
    parser.add_argument("--out-dir", default="outputs/cwh_orchestrator_demo", help="Output directory.")
    parser.add_argument("--no-tasks", action="store_true", help="Do not generate MediaSpider task JSONs.")
    parser.add_argument("--platforms", default="wb,dy,ks,bili,zhihu", help="Domestic MediaSpider platforms; the current project profile excludes xhs, tieba and foreign platforms.")
    parser.add_argument("--collection-profile", choices=["demo", "formal", "deep"], default="demo", help="MediaSpider collection scale preset.")
    parser.add_argument("--max-posts", type=int, default=0, help="Override max posts per MediaSpider task.")
    parser.add_argument("--max-comments", type=int, default=0, help="Override max comments per post.")
    parser.add_argument("--target-samples", type=int, default=5000, help="Formal-delivery target for total usable samples.")
    parser.add_argument("--target-comments", type=int, default=200, help="Formal-delivery target for verifiable netizen comments.")
    parser.add_argument("--target-comments-per-topic", type=int, default=30, help="Formal-delivery target for comments per subtopic.")
    parser.add_argument("--target-overseas", type=int, default=10, help="Formal-delivery target for overseas media samples.")
    parser.add_argument("--max-low-quality-ratio", type=float, default=0.5, help="Maximum accepted low-quality sample ratio.")
    parser.add_argument("--max-collection-rounds", type=int, default=3, help="Maximum recommended collection expansion rounds.")
    parser.add_argument("--agent-reach-plan", action="store_true", default=True, help="Generate agent-reach lead discovery tasks.")
    parser.add_argument("--no-agent-reach-plan", dest="agent_reach_plan", action="store_false", help="Skip agent-reach lead discovery task generation.")
    parser.add_argument("--agent-reach-platforms", default="bili,v2ex", help="Comma list of domestic agent-reach lead platforms; xhs is excluded by the current project profile.")
    parser.add_argument("--run-agent-reach", action="store_true", help="Actually run agent-reach lead tasks and ingest normalized quote samples.")
    parser.add_argument("--agent-reach-limit", type=int, default=1, help="Maximum agent-reach tasks to run; 0 means all matched tasks.")
    parser.add_argument("--opencli-profile", default="", help="Optional OpenCLI Browser Bridge profile for agent-reach/OpenCLI commands.")
    parser.add_argument("--run-mediaspider", action="store_true", help="Actually run MediaSpider tasks and ingest raw outputs before report generation.")
    parser.add_argument("--mediaspider-supervisor", default="D:/Codex/2026-06-15/ai-1-1-https-drive-weixin/work/mediaspider-supervisor", help="mediaspider-supervisor directory.")
    parser.add_argument("--media-home", default="", help="Optional MediaSpider engine home path.")
    parser.add_argument("--mediaspider-limit", type=int, default=1, help="Maximum MediaSpider tasks to run; 0 means all matched tasks.")
    parser.add_argument("--mediaspider-task-timeout-sec", type=int, default=0, help="Per MediaSpider task timeout in seconds; 0 means no timeout.")
    parser.add_argument("--keep-proxy", action="store_true", help="Keep inherited HTTP/HTTPS proxy variables for MediaSpider.")
    parser.add_argument("--run-foreign-mediaspider", action="store_true", help="Run the server-safe MediaSpider foreign evidence collector.")
    parser.add_argument("--foreign-mediaspider-supervisor", default="", help="Optional MediaSpider supervisor path for the foreign collector.")
    parser.add_argument("--foreign-media-home", default="", help="Optional MediaSpider engine path for the foreign collector.")
    parser.add_argument("--foreign-mediaspider-platforms", default="grounding,youtube", help="Comma-separated server-safe foreign evidence platforms.")
    parser.add_argument("--foreign-mediaspider-limit", type=int, default=8, help="Maximum rows per foreign source and topic.")
    parser.add_argument("--foreign-mediaspider-timeout-sec", type=int, default=300, help="Whole foreign collection timeout in seconds; 0 means no timeout.")
    parser.add_argument("--comment-limit", type=int, default=0, help="Maximum retained comment evidence rows; 0 keeps all qualified rows.")
    parser.add_argument("--auto-web", dest="auto_web", action="store_true", default=True, help="Collect public web/news samples when samples are missing.")
    parser.add_argument("--no-auto-web", dest="auto_web", action="store_false", help="Disable public web/news fallback collection.")
    parser.add_argument("--web-always", action="store_true", help="Run public web/news collection even when sample file is provided.")
    parser.add_argument("--web-limit", type=int, default=5, help="Public web/news samples per topic.")
    parser.add_argument("--web-timeout", type=int, default=15, help="Public web/news request timeout seconds.")
    parser.add_argument("--monitor-start", default="", help="Monitoring start date YYYY-MM-DD.")
    parser.add_argument("--monitor-end", default="", help="Monitoring end date YYYY-MM-DD.")
    args = parser.parse_args()

    data = orchestrate(args)
    print(
        json.dumps(
            {
                "status": "ok" if formal_delivery_ready(data) else "blocked",
                "out_dir": args.out_dir,
                "report": data["artifacts"].get("report"),
                "report_data": data["artifacts"].get("report_data"),
                "dashboard": data["artifacts"].get("dashboard"),
                "samples": data["statistics"]["total_samples"],
                "public_web_samples": data["collection"]["public_web_sample_count"],
                "mediaspider_samples": data["collection"].get("mediaspider_sample_count", 0),
                "foreign_mediaspider_samples": data["collection"].get("foreign_mediaspider_sample_count", 0),
                "mediaspider_run": bool(data["collection"].get("mediaspider_run")),
                "topics": len(data["meeting"]["topics"]),
                "mediacrawler_tasks": data["collection"]["mediacrawler_tasks"]["task_count"] if data["collection"].get("mediacrawler_tasks") else 0,
                "data_gaps": len(data["audit"]["data_gaps"]),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
