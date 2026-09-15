from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any


SUMMARY_KEYS = [
    "date",
    "domestic_mainstream",
    "overseas_media",
    "wechat_public",
    "weibo",
    "video_account",
    "new_media",
    "total_spread",
]

DETAIL_HEADER_MAP = {
    "公众文章": "wechat_articles",
    "公号-在看量": "wechat_reads",
    "公号-精选评论量": "wechat_comments",
    "新浪微博": "weibo",
    "境内新闻": "domestic_news",
    "境内APP": "domestic_app",
    "境内论坛": "domestic_forum",
    "其他视频": "other_video",
    "境外新闻": "overseas_news",
    "推特": "twitter",
    "境外其他": "overseas_other",
    "视频号发文": "video_account",
    "抖音": "douyin",
}


def as_int(value: Any) -> int:
    if value in (None, ""):
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(",", "")
    if text.endswith("+"):
        text = text[:-1]
    try:
        return int(float(text))
    except ValueError:
        return 0


def as_ratio(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().rstrip("%")
        try:
            number = float(text)
        except ValueError:
            return 0.0
        if str(value).strip().endswith("%"):
            number /= 100
    return round(number / 100 if number > 1 else number, 6)


def date_text(value: Any, default_year: int | None = None) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    match = re.search(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    short = re.fullmatch(r"(\d{1,2})[/.\-](\d{1,2})", text)
    if short and default_year:
        return f"{default_year:04d}-{int(short.group(1)):02d}-{int(short.group(2)):02d}"
    return text[:10]


def find_sheet(wb: Any, exact: str) -> Any | None:
    if exact in wb.sheetnames:
        return wb[exact]
    for name in wb.sheetnames:
        if exact in str(name):
            return wb[name]
    return None


def parse_keywords(ws: Any | None) -> dict[str, Any]:
    if ws is None:
        return {"master": {}, "topics": []}
    master = {
        "title": str(ws.cell(2, 2).value or "总事件"),
        "keywords": str(ws.cell(2, 3).value or ""),
        "exclude_terms": str(ws.cell(2, 4).value or ""),
        "event_id": str(ws.cell(2, 5).value or ""),
    }
    header_row = 0
    for row in range(1, min(ws.max_row, 30) + 1):
        if str(ws.cell(row, 3).value or "").strip() == "子事件":
            header_row = row
            break
    topics: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        index = ws.cell(row, 1).value
        title = str(ws.cell(row, 2).value or "").strip()
        if not title:
            continue
        topics.append(
            {
                "index": as_int(index) or len(topics) + 1,
                "title": title,
                "keywords": str(ws.cell(row, 3).value or ""),
                "exclude_terms": str(ws.cell(row, 4).value or ""),
                "event_id": str(ws.cell(row, 5).value or ""),
            }
        )
    return {"master": master, "topics": topics}


def parse_trend_sheet(ws: Any | None) -> dict[str, Any]:
    if ws is None:
        return {"title": "", "daily": [], "totals": {}, "detail_headers": []}
    title = str(ws.cell(2, 1).value or "")
    title_year_match = re.search(r"(20\d{2})", title)
    default_year = int(title_year_match.group(1)) if title_year_match else None
    header_row = 0
    for row in range(1, min(ws.max_row, 12) + 1):
        if str(ws.cell(row, 1).value or "").strip() == "日期":
            header_row = row
            break
    if not header_row:
        return {"title": title, "daily": [], "totals": {}, "detail_headers": []}

    date_columns = [
        col
        for col in range(1, ws.max_column + 1)
        if str(ws.cell(header_row, col).value or "").strip() == "日期"
    ]
    detail_start = date_columns[1] if len(date_columns) > 1 else 0
    if default_year is None and detail_start:
        # Subevent titles often omit the year while the compact left-hand date
        # cells contain only M/D.  Infer the year from the authoritative raw
        # date column on the right before parsing those cached display values.
        for candidate_row in range(header_row + 1, min(ws.max_row, header_row + 32) + 1):
            raw_detail_date = ws.cell(candidate_row, detail_start).value
            if isinstance(raw_detail_date, (datetime, date)):
                default_year = raw_detail_date.year
                break
            match = re.search(r"(20\d{2})", str(raw_detail_date or ""))
            if match:
                default_year = int(match.group(1))
                break
    summary_header_values = [str(ws.cell(header_row, col).value or "").strip() for col in range(1, 9)]
    is_total_event = "微信公众号" in summary_header_values
    detail_columns: dict[int, str] = {}
    if detail_start:
        for col in range(detail_start + 1, ws.max_column + 1):
            label = str(ws.cell(header_row, col).value or "").strip()
            if label in DETAIL_HEADER_MAP:
                detail_columns[col] = DETAIL_HEADER_MAP[label]

    daily: list[dict[str, Any]] = []
    row = header_row + 1
    while row <= ws.max_row:
        raw_date = ws.cell(row, 1).value
        day = date_text(raw_date, default_year)
        if not re.fullmatch(r"20\d{2}-\d{2}-\d{2}", day):
            break
        if is_total_event:
            record = dict(zip(SUMMARY_KEYS, [day, *[as_int(ws.cell(row, col).value) for col in range(2, 9)]]))
        else:
            record = {
                "date": day,
                "domestic_mainstream": as_int(ws.cell(row, 2).value),
                "overseas_media": as_int(ws.cell(row, 3).value),
                "new_media": as_int(ws.cell(row, 4).value),
                "total_spread": as_int(ws.cell(row, 5).value),
                "wechat_public": 0,
                "weibo": 0,
                "video_account": 0,
            }
        details = {key: as_int(ws.cell(row, col).value) for col, key in detail_columns.items()}
        if details:
            record["details"] = details
        daily.append(record)
        row += 1

    totals: dict[str, int] = {}
    for key in SUMMARY_KEYS[1:]:
        totals[key] = sum(as_int(item.get(key)) for item in daily)
    if not is_total_event:
        totals["wechat_public"] = 0
        totals["weibo"] = 0
        totals["video_account"] = 0
    return {"title": title, "daily": daily, "totals": totals, "detail_headers": list(detail_columns.values()),
            "channel_labels": {"domestic_media": str(ws.cell(header_row, 2).value or "").strip()}}


def parse_topic_summary(ws: Any | None) -> list[dict[str, Any]]:
    if ws is None:
        return []
    rows: list[dict[str, Any]] = []
    for row in range(1, ws.max_row + 1):
        index = as_int(ws.cell(row, 1).value)
        title = str(ws.cell(row, 2).value or "").strip()
        if index <= 0 or not title:
            continue
        total = as_int(ws.cell(row, 6).value)
        if total <= 0:
            continue
        rows.append(
            {
                "index": index,
                "title": title,
                "domestic_mainstream": as_int(ws.cell(row, 3).value),
                "new_media": as_int(ws.cell(row, 4).value),
                "overseas_media": as_int(ws.cell(row, 5).value),
                "total_spread": total,
                "sentiment": {
                    "positive": as_ratio(ws.cell(row, 7).value),
                    "neutral": as_ratio(ws.cell(row, 8).value),
                    "negative": as_ratio(ws.cell(row, 9).value),
                },
            }
        )
    return rows


def parse_overseas(ws: Any | None, topic_count: int) -> list[dict[str, Any]]:
    if ws is None:
        return []
    header_row = 0
    for row in range(1, min(ws.max_row, 15) + 1):
        values = {str(ws.cell(row, col).value or "").strip() for col in range(1, min(ws.max_column, 20) + 1)}
        if "来源" in values and "发布地址" in values:
            header_row = row
            break
    headers = {
        str(ws.cell(header_row, col).value or "").strip(): col
        for col in range(1, ws.max_column + 1)
        if str(ws.cell(header_row, col).value or "").strip()
    }
    source_col = headers.get("来源", 3)
    date_col = headers.get("报道日期", 4)
    title_col = headers.get("标题", 6)
    url_col = headers.get("发布地址", 7)
    topic_labels = ["一", "二", "三", "四", "五", "六", "七", "八"][:topic_count]
    topic_columns = [headers.get(label, 8 + index) for index, label in enumerate(topic_labels)]
    rows: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        source = str(ws.cell(row, source_col).value or "").strip()
        title = str(ws.cell(row, title_col).value or "").strip()
        url = str(ws.cell(row, url_col).value or "").strip()
        if not source or not title:
            continue
        hits = [idx for idx, column in enumerate(topic_columns, 1) if ws.cell(row, column).value not in (None, "", 0)]
        category = str(ws.cell(row, headers.get("报道类型", 0)).value or "").strip() if headers.get("报道类型") else ""
        summary_cn = str(ws.cell(row, headers.get("简体中文摘要", 0)).value or "").strip() if headers.get("简体中文摘要") else ""
        rows.append(
            {
                "source": source,
                "published_at": date_text(ws.cell(row, date_col).value),
                "title": title,
                "url": url,
                "topic_hits": hits,
                "ai_report_category": category,
                "classification_reason": (
                    str(ws.cell(row, headers["审核理由"]).value or "").strip()
                    if headers.get("审核理由") else ""
                ),
                "classification_confidence": (
                    ws.cell(row, headers["分类置信度"]).value
                    if headers.get("分类置信度") else None
                ),
                "source_cn_simplified": (
                    str(ws.cell(row, headers["简体中文来源"]).value or "").strip()
                    if headers.get("简体中文来源") else ""
                ),
                "title_cn_simplified": (
                    str(ws.cell(row, headers["简体中文标题"]).value or "").strip()
                    if headers.get("简体中文标题") else ""
                ),
                "summary_cn_simplified": summary_cn,
                "interpretive_verified": category == "解读性报道" and bool(summary_cn),
                "meeting_relevance": True,
                "formal_include": True,
            }
        )
    return rows


def parse_hotwords(ws: Any | None) -> list[str]:
    if ws is None:
        return []
    output = []
    for row in range(1, ws.max_row + 1):
        value = str(ws.cell(row, 1).value or "").strip()
        if value and value not in output:
            output.append(value)
    return output


def parse_wechat_top(ws: Any | None) -> list[dict[str, Any]]:
    if ws is None:
        return []
    header_row = 0
    for row in range(1, min(ws.max_row, 15) + 1):
        values = {str(ws.cell(row, col).value or "").strip() for col in range(1, min(ws.max_column, 12) + 1)}
        if "账号" in values and "阅读量" in values:
            header_row = row
            break
    output: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        source = str(ws.cell(row, 3).value or "").strip()
        title = str(ws.cell(row, 7).value or ws.cell(row, 4).value or "").strip()
        url = str(ws.cell(row, 8).value or "").strip()
        read_value = ws.cell(row, 5).value
        if not source or not title:
            continue
        output.append(
            {
                "source": source,
                "title": title,
                "spread_count": as_int(read_value),
                "spread_count_display": str(read_value or "0").strip(),
                "comment_count": as_int(ws.cell(row, 6).value),
                "url": url,
            }
        )
    return output


def build_evidence_samples(data: dict[str, Any]) -> list[dict[str, Any]]:
    topics = data.get("topics") or []
    samples: list[dict[str, Any]] = []
    for idx, row in enumerate(data.get("overseas_reports") or [], 1):
        hits = row.get("topic_hits") or []
        topic = topics[hits[0] - 1]["title"] if hits and hits[0] <= len(topics) else (topics[0]["title"] if topics else "")
        samples.append(
            {
                "id": f"system-overseas-{idx}",
                "data_origin": "monitoring_system_excel",
                "topic": topic,
                "topic_hits": [topics[i - 1]["title"] for i in hits if 0 < i <= len(topics)],
                "platform": "news",
                "source_type": "overseas_media",
                "region": "overseas",
                **row,
                "content": row.get("title", ""),
                "spread_count": 0,
                "comment_count": 0,
                "is_comment": False,
                "quality_flags": [],
            }
        )
    for idx, row in enumerate(data.get("wechat_top") or [], 1):
        samples.append(
            {
                "id": f"system-wechat-{idx}",
                "data_origin": "monitoring_system_excel",
                "topic": "",
                "platform": "wechat",
                "source_type": "self_media_article",
                "region": "domestic",
                **row,
                "content": row.get("title", ""),
                "published_at": "",
                "is_comment": False,
                "quality_flags": [],
            }
        )
    return samples


def ingest_workbook(path: str | Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    source = Path(path)
    # Some generated workbooks retain a stale worksheet dimension of A1.  In
    # read-only mode openpyxl trusts that metadata and silently hides the real
    # subevent rows.  Normal mode reads the worksheet XML and is therefore the
    # accuracy-safe choice for the small standard workbook used here.
    wb = load_workbook(source, data_only=True, read_only=False)
    keyword_data = parse_keywords(find_sheet(wb, "关键词"))
    topics = keyword_data["topics"]
    overall = parse_trend_sheet(find_sheet(wb, "总事件"))
    topic_summary = parse_topic_summary(find_sheet(wb, "子事件数据汇总"))
    subevents = []
    for idx, topic in enumerate(topics, 1):
        summary = next((item for item in topic_summary if item["index"] == idx), {})
        trend = parse_trend_sheet(find_sheet(wb, f"子事件{idx}"))
        subevents.append({**topic, **summary, "daily": trend.get("daily", []), "totals": trend.get("totals", {}), "trend_totals": trend.get("totals", {})})

    dates = [row["date"] for row in overall.get("daily", []) if row.get("date")]
    result = {
        "schema_version": "1.0",
        "source_type": "monitoring_system_workbook",
        "source_file": str(source.resolve()),
        "master_event": keyword_data["master"],
        "topics": topics,
        "monitoring_period": {"start": min(dates) if dates else "", "end": max(dates) if dates else ""},
        "overall": overall,
        "channel_labels": overall.get("channel_labels") or {},
        "subevents": subevents,
        "overseas_reports": parse_overseas(find_sheet(wb, "外媒报道列表"), len(topics)),
        "hotwords": parse_hotwords(find_sheet(wb, "词云")),
        "wechat_top": parse_wechat_top(find_sheet(wb, "公众TOP")),
        "validation": [],
    }
    wb.close()
    if not result["overall"].get("daily"):
        result["validation"].append({"severity": "error", "code": "missing_overall_daily", "message": "未识别总事件逐日数据。"})
    if not result["subevents"]:
        result["validation"].append({"severity": "error", "code": "missing_subevents", "message": "未识别子事件数据。"})
    missing_daily = [str(item.get("index") or "") for item in result["subevents"] if not item.get("daily")]
    if missing_daily:
        result["validation"].append({
            "severity": "error",
            "code": "missing_subevent_daily",
            "message": f"子事件{'、'.join(missing_daily)}未识别逐日数据。",
        })
    if any(not item.get("total_spread") for item in result["subevents"]):
        result["validation"].append({"severity": "warning", "code": "incomplete_subevent_totals", "message": "部分子事件缺少系统传播总量。"})
    result["evidence_samples"] = build_evidence_samples(result)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Normalize a mentor/system CWH workbook into authoritative JSON.")
    parser.add_argument("workbook", help="Monitoring-system CWH workbook (.xlsx).")
    parser.add_argument("--output", default="outputs/cwh_system_data.json")
    args = parser.parse_args()
    data = ingest_workbook(args.workbook)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": "ok", "output": str(output), "topics": len(data["topics"]), "total_spread": data["overall"].get("totals", {}).get("total_spread", 0)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
