from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import zipfile
from xml.etree import ElementTree as ET
from datetime import date
from functools import lru_cache
from pathlib import Path
from typing import Any

try:
    from opencc import OpenCC
except ImportError:  # pragma: no cover - exercised by deployment dependency checks
    OpenCC = None

from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

from report_rules import (
    cluster_has_named_person,
    cluster_needs_attribution_review,
    enrich_viewpoint_titles,
    formal_sentiment_available,
    formal_sentiment_row_ready,
)
from source_identity import public_source_family
from cwh_writing_rules import chinese_number, opening_paragraph, ordinal_prefix, writing_rules, writing_rules_sha256, domestic_media_label, editorial_profile
from normalize_cwh_analysis import assemble_cluster_details, evidence_sentence
from cwh_docx_template import TEMPLATE, apply_template
from cwh_chart_style import valid_image, write_topic_chart, verify_topic_chart_manifest
from cwh_report_visuals import write_report_charts, verify_report_charts, add_fixed_topic_table, verify_topic_table


@lru_cache(maxsize=1)
def public_source_identity_config() -> dict[str, Any]:
    path = Path(__file__).resolve().parent.parent / "config" / "raw_workbook_mapping.json"
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        return {}


def topic_display(topic: str) -> str:
    value = topic.strip()
    for prefix in ["进一步部署", "听取", "研究", "审议通过", "审议", "部署", "决定"]:
        if value.startswith(prefix):
            value = value[len(prefix) :]
            break
    value = value.removesuffix("有关工作").removesuffix("情况汇报")
    return value.strip(" ，。；、") or topic


def apply_hotword_audit(data: dict[str, Any], audit_path: Path) -> None:
    payload = json.loads(audit_path.read_text(encoding="utf-8"))
    from cwh_semantic_recovery import is_deferred_hotwords
    if is_deferred_hotwords(payload):
        data['hotwords'] = []
        data['hotword_review_gap'] = payload['notice']
        data.setdefault('audit', {}).setdefault('data_gaps', []).append(payload['notice'])
        acceptance = data['audit'].setdefault('acceptance', {})
        acceptance['ready_for_formal_delivery'] = False
        acceptance.setdefault('blockers', []).append(payload['notice'])
        data.setdefault('artifacts', {})['hotword_audit'] = str(audit_path)
        return
    if payload.get("status") != "ai_review_complete":
        raise ValueError(f"词云审计未完成AI审核：{audit_path}")
    if payload.get("method") != "ai_semantic_review_with_evidence":
        raise ValueError(f"词云审计方法不是AI语义审核：{audit_path}")
    if payload.get("second_pass_completed") is not True:
        raise ValueError(f"词云审计缺少第二遍议题覆盖、同义词和噪声复核：{audit_path}")
    selected = list(payload.get("selected") or [])
    if not selected:
        raise ValueError(f"词云审计文件没有selected词条：{audit_path}")
    minimum_term_count = int((payload.get("settings") or {}).get("minimum_term_count") or 36)
    if len(selected) < minimum_term_count:
        raise ValueError(
            f"词云审计仅有{len(selected)}个合格词，少于最低{minimum_term_count}个；"
            "必须让AI从证据中继续补提"
        )
    from cwh_hotword_pipeline import (
        HOTWORD_SEMANTIC_TYPES,
        PROCEDURAL_HOTWORD_MARKERS,
        valid_candidate,
    )

    invalid_terms = []
    for row in selected:
        term = str(row.get("term") or "").strip()
        tier = str(row.get("evidence_tier") or "").strip().lower()
        semantic_type = str(row.get("semantic_type") or "").strip().lower()
        document_count = int(row.get("document_count") or 0)
        if (
            row.get("ai_reviewed") is not True
            or not valid_candidate(term)
            or not str(row.get("selection_reason") or "").strip()
            or tier not in {"core", "supporting"}
            or semantic_type not in HOTWORD_SEMANTIC_TYPES
            or row.get("standalone_topic_label") is not True
            or any(term.startswith(marker) or term.endswith(marker) for marker in PROCEDURAL_HOTWORD_MARKERS)
            or document_count < (2 if tier == "core" else 1)
        ):
            invalid_terms.append(term or "<empty>")
    if invalid_terms:
        raise ValueError("词云审计含未通过AI与证据门禁的词：" + "、".join(invalid_terms[:12]))
    topics = list((data.get("meeting") or {}).get("topics") or [])
    existing_topics = {
        str(row.get("word") or "").strip(): str(row.get("topic") or "").strip()
        for row in data.get("hotwords") or []
        if str(row.get("word") or "").strip()
    }
    topic_corpora = {str(topic): str(topic) for topic in topics}
    for viewpoint in (data.get("viewpoints") or {}).get("by_topic") or []:
        topic_name = str(viewpoint.get("topic") or "")
        if topic_name not in topic_corpora:
            continue
        cluster_text = " ".join(
            f"{cluster.get('summary', '')} {cluster.get('details', '')}"
            for cluster in viewpoint.get("clusters") or []
        )
        topic_corpora[topic_name] += " " + cluster_text
    for sample in data.get("samples") or []:
        topic_name = str(sample.get("topic") or "")
        if topic_name in topic_corpora:
            topic_corpora[topic_name] += f" {sample.get('title', '')} {sample.get('content', '')}"
    normalized = []
    for position, row in enumerate(selected, 1):
        term = str(row.get("term") or "").strip()
        if not term:
            continue
        topic_index = row.get("topic_index")
        if not topic_index:
            hits = list(row.get("topic_hits") or [])
            topic_index = hits[0] if hits else None
        topic = existing_topics.get(term, "")
        try:
            index = int(topic_index)
        except (TypeError, ValueError):
            index = 0
        if 1 <= index <= len(topics):
            topic = str(topics[index - 1])
        if not topic:
            scores = [(corpus.count(term), topic_name) for topic_name, corpus in topic_corpora.items()]
            best_score, best_topic = max(scores, default=(0, ""))
            if best_score > 0:
                topic = best_topic
        sample_titles = list(row.get("sample_titles") or [])
        sample_urls = list(row.get("sample_urls") or [])
        normalized.append(
            {
                "word": term,
                "count": int(row.get("weight") or max(1, len(selected) - position + 1)),
                "rank": int(row.get("rank") or position),
                "topic": topic,
                "sample_count": int(row.get("document_count") or 0),
                "source_count": int(row.get("source_count") or 0),
                "exact_match_count": int(row.get("exact_match_count") or 0),
                "evidence_tier": str(row.get("evidence_tier") or "core"),
                "semantic_type": str(row.get("semantic_type") or ""),
                "standalone_topic_label": row.get("standalone_topic_label") is True,
                "selection_reason": str(row.get("selection_reason") or ""),
                "sample_titles": sample_titles,
                "sample_urls": sample_urls,
                "example": str(sample_titles[0])[:80] if sample_titles else "",
                "selection_method": str(row.get("selection_method") or payload.get("method") or "reviewed_wordcloud"),
                "authority": "reviewed_wordcloud_audit",
            }
        )
    data["hotwords"] = normalized
    shortfall = payload.get("count_shortfall") or {}
    if payload.get("delivery_policy") == "deliver_available_with_gaps" and shortfall.get("notice"):
        audit = data.setdefault("audit", {})
        audit.setdefault("data_gaps", []).append(shortfall["notice"])
        acceptance = audit.setdefault("acceptance", {})
        acceptance.setdefault("blockers", []).append(shortfall["notice"])
        acceptance["blockers"] = list(dict.fromkeys(acceptance["blockers"]))
        acceptance["ready_for_formal_delivery"] = False
    data.setdefault("artifacts", {})["hotword_audit"] = str(audit_path.resolve())


def apply_raw_public_top_audit(data: dict[str, Any]) -> None:
    """Carry a hash-bound raw ranking deficit into all synchronized outputs."""
    value = str((data.get('collection') or {}).get('system_workbook') or '').strip()
    if not value:
        return
    workbook = Path(value)
    sidecar = workbook.parent / 'run' / 'public_top_audit.json'
    if not workbook.is_file() or not sidecar.is_file():
        return
    payload = json.loads(sidecar.read_text(encoding='utf-8'))
    if payload.get('workbook_sha256') != hashlib.sha256(workbook.read_bytes()).hexdigest():
        return
    shortfall = payload.get('evidence_shortfall') or {}
    if not shortfall:
        return
    notice = f"公众号TOP仅取得{shortfall['actual']}个已审核来源，目标{shortfall['required']}个；候选审阅已达上限或耗尽，不代表完整TOP排名。"
    audit = data.setdefault('audit', {})
    audit['raw_public_top_review'] = {'path': str(sidecar.resolve()), 'workbook_sha256': payload['workbook_sha256'],
                                    'evidence_shortfall': shortfall}
    audit['data_gaps'] = list(dict.fromkeys([*(audit.get('data_gaps') or []), notice]))
    acceptance = audit.setdefault('acceptance', {})
    acceptance['blockers'] = list(dict.fromkeys([*(acceptance.get('blockers') or []), notice]))
    acceptance['ready_for_formal_delivery'] = False


def metric(data: dict[str, Any], key: str, default: int = 0) -> int:
    try:
        return int(data.get(key, default) or default)
    except Exception:
        return default


def sentiment_text(row: dict[str, Any]) -> str:
    counts = row.get("sentiment") or {}
    total = sum(int(counts.get(name, 0) or 0) for name in ["positive", "neutral", "negative"])
    if total <= 0:
        return "尚未形成有效情感分母"
    pos = round(int(counts.get("positive", 0)) / total * 100)
    neu = round(int(counts.get("neutral", 0)) / total * 100)
    neg = max(0, 100 - pos - neu)
    return f"正面{pos}%、中性{neu}%、负面{neg}%"


def md_table(headers: list[str], rows: list[list[Any]], empty: str = "未接入可核验数据") -> str:
    if not rows:
        rows = [[empty for _ in headers]]
    def cell(value: Any) -> str:
        return ("" if value is None else str(value)).replace("\r", " ").replace("\n", " ").replace("|", "\\|")
    lines = ["| " + " | ".join(cell(value) for value in headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(cell(value) for value in row) + " |")
    return "\n".join(lines)


def evidence_text(evidence: list[dict[str, Any]]) -> str:
    """Build concrete attributed prose when reviewed cluster details are absent.

    Only independent evidence rows are used.  Reposts and mirror pages stored in
    ``supporting_samples``/``supporting_links`` are audit records and must never
    create a second formal sentence for the same viewpoint.
    """
    parts: list[str] = []
    seen_claims: set[str] = set()
    attribution_verb = re.compile(r"(认为|指出|建议|强调|表示|称|提出|研判|预计|呼吁|担忧|主张)")
    for item in evidence:
        raw_claim = clean_sentence(
            item.get("formal_claim")
            or item.get("claim")
            or item.get("source_excerpt")
            or item.get("summary")
            or item.get("content")
        )
        if not raw_claim:
            continue
        subject = clean_sentence(item.get("attribution") or item.get("source") or item.get("platform"))
        match = attribution_verb.search(raw_claim)
        if subject and subject not in raw_claim:
            if match:
                raw_claim = f"{subject}{raw_claim[match.start():]}"
            else:
                raw_claim = f"{subject}认为，{raw_claim}"
        signature = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "", raw_claim).lower()
        if not signature or signature in seen_claims:
            continue
        seen_claims.add(signature)
        parts.append(raw_claim)
    return "。".join(parts)


def display_cluster_details(cluster: dict[str, Any]) -> str:
    """Render current atomic evidence, never a stale preassembled paragraph."""
    if any(isinstance(row, dict) and row.get('formal_claim') for row in cluster.get('evidence') or []):
        return assemble_cluster_details(cluster)
    return cluster.get('details') or cluster.get('analysis') or ''


def cluster_wording(cluster: dict[str, Any]) -> str:
    summary = clean_sentence(cluster.get("summary"))
    details = clean_sentence(display_cluster_details(cluster))
    warning = attribution_review_marker(cluster)
    if details:
        return f"{summary}。" + '\n\n'.join(cluster_detail_paragraphs(cluster)) + warning
    evidence = evidence_text(cluster.get("evidence") or [])
    if evidence:
        return f"{summary}。{evidence}。{warning}"
    return f"{summary}。{warning}"


ATTRIBUTION_REVIEW_TEXT = "【未识别到具体专家姓名，建议核验原文】"


def attribution_review_marker(cluster: dict[str, Any]) -> str:
    return ATTRIBUTION_REVIEW_TEXT if cluster_needs_attribution_review(cluster) else ""


AGENDA_HEADING_PREFIXES = ("听取", "研究", "审议", "进一步部署", "部署", "决定")
STANCE_HEADING_PREFIXES = tuple(writing_rules()["viewpoint"]["heading_stance_verbs"])


def stance_heading(text: Any) -> str:
    value = clean_sentence(text)
    value = re.sub(r"^(?:舆论|媒体|专家|机构|网民)(?:普遍)?", "", value).strip()
    return value if value.startswith(STANCE_HEADING_PREFIXES) else ""


def comment_stance_heading(text: Any) -> str:
    """Normalize a concrete reviewed comment claim into report heading grammar."""
    value = clean_sentence(text)
    value = re.sub(r"^(?:舆论|媒体|专家|机构|网民)(?:普遍)?", "", value).strip()
    prefixes = STANCE_HEADING_PREFIXES + tuple(writing_rules()['comments']['additional_heading_verbs'])
    if value.startswith(prefixes):
        return value
    # Domestic models often return a valid policy proposition instead of an
    # attribution verb. Preserve the proposition and add the neutral report
    # verb deterministically; generic topic labels still fail this gate.
    if re.search(r"(?:应当|应|需要|需|须|不能|不应|有必要|可以|可)", value):
        return "认为" + value
    return ""


def topic_heading(item: dict[str, Any]) -> str:
    heading = clean_sentence(item.get("heading"))
    if heading and item.get('_reviewed_display_heading') == heading:
        return heading
    reviewed = stance_heading(heading)
    if reviewed:
        return reviewed
    if heading and "尚未形成评论性观点" in heading:
        return heading
    clusters = item.get("clusters") or []
    if clusters:
        summary = clean_sentence(clusters[0].get("summary"))
        summary = re.sub(r"^(舆论|媒体|专家|机构)(普遍)?", "", summary)
        reviewed = stance_heading(summary)
        if reviewed:
            return reviewed
    return topic_display(str(item.get("topic") or heading))


def numbered(index: int) -> str:
    return chinese_number(index + 1)


def clean_sentence(text: Any) -> str:
    return str(text or "").strip().rstrip("。；; ")


FORMAL_REACTION_TOKEN_RE = re.compile(
    r"[\[【](?:赞|点赞|鼓掌|玫瑰|心|爱心|比心|福|点亮平安灯|呲牙|微笑|大笑|笑哭|捂脸|偷笑|调皮|可爱|加油|奋斗|握手|抱拳|祈祷|祈福|合十|狗头|强|支持|庆祝|烟花|礼花|蜡烛|胜利|OK)[\]】]",
    re.IGNORECASE,
)
FORMAL_EMOJI_RE = re.compile(
    "["
    "\U0001F1E6-\U0001F1FF"
    "\U0001F300-\U0001FAFF"
    "\U00002600-\U000027BF"
    "\U0000FE0F"
    "\U0000200D"
    "]+",
    re.UNICODE,
)


def clean_formal_comment(text: Any) -> str:
    """Remove reaction emoji from formal prose without changing evidence rows."""

    value = str(text or "")
    value = FORMAL_REACTION_TOKEN_RE.sub("", value)
    value = FORMAL_EMOJI_RE.sub("", value)
    value = re.sub(r"\s+", " ", value)
    value = re.sub(r"([，。！？；、])\1+", r"\1", value)
    value = re.sub(r'"([^"\n]{1,160})[”"]', r'‘\1’', value)
    value = re.sub(r'“([^”"\n]{1,160})"', r'‘\1’', value)
    value = re.sub(r"(?<=\d)\s*[~～]\s*(?=\d)", "至", value)
    return value.strip(" ，。！？；、;,.!?\t\r\n")


def month_day_text(date_value: str) -> str:
    if len(date_value or "") >= 10:
        try:
            return f"{int(date_value[5:7])}月{int(date_value[8:10])}日"
        except Exception:
            pass
    return date_value or "本次"


def source_bucket_counts(stats: dict[str, Any]) -> dict[str, int]:
    buckets = stats.get("by_source_bucket") or {}
    return {key: int(value or 0) for key, value in buckets.items()}


def platform_count(stats: dict[str, Any], *names: str) -> int:
    platforms = stats.get("by_platform") or {}
    lowered = {str(key).lower(): int(value or 0) for key, value in platforms.items()}
    total = 0
    for name in names:
        total += lowered.get(name.lower(), 0)
    return total


def sentiment_ratio_cells(row: dict[str, Any]) -> list[str]:
    if not formal_sentiment_row_ready(row):
        return ["待分析", "待分析", "待分析"]
    counts = row.get("sentiment") or {}
    total = sum(int(counts.get(name, 0) or 0) for name in ["positive", "neutral", "negative"])
    if total <= 0:
        return ["待分析", "待分析", "待分析"]
    pos = round(int(counts.get("positive", 0)) / total * 100)
    neu = round(int(counts.get("neutral", 0)) / total * 100)
    neg = max(0, 100 - pos - neu)
    return [f"{pos}%", f"{neu}%", f"{neg}%"]


def topic_total(row: dict[str, Any]) -> int:
    if row.get("authority") == "monitoring_system":
        return int(row.get("spread_count", 0) or row.get("total_samples", 0) or 0)
    return (
        int(row.get("domestic_media", 0) or 0)
        + int(row.get("self_media", 0) or 0)
        + int(row.get("comments", 0) or 0)
        + int(row.get("overseas_media", 0) or 0)
    )


def netizen_sentiment_lead(data: dict[str, Any]) -> str:
    counts = (data.get("comments") or {}).get("sentiment") or {}
    positive = int(counts.get("positive", 0) or 0)
    neutral = int(counts.get("neutral", 0) or 0)
    negative = int(counts.get("negative", 0) or 0)
    total = positive + neutral + negative
    if total <= 0:
        aggregate = {"positive": 0, "neutral": 0, "negative": 0}
        for row in data.get("topic_stats") or []:
            sentiment = row.get("sentiment") or {}
            for label in aggregate:
                aggregate[label] += int(sentiment.get(label, 0) or 0)
        positive = aggregate["positive"]
        neutral = aggregate["neutral"]
        negative = aggregate["negative"]
        total = positive + neutral + negative
    if total <= 0:
        return "网民评论尚未完成有效情感分类，暂不判断整体情感倾向。"
    topic_rows = list(data.get("topic_stats") or [])
    has_explicit_gate = any("sentiment_formal_ready" in row for row in topic_rows)
    if has_explicit_gate and not any(bool(row.get("sentiment_formal_ready")) for row in topic_rows):
        if neutral >= positive and neutral >= negative:
            observed = "客观中立评论相对较多"
        elif negative > positive + neutral:
            observed = "批评质疑类评论相对较多"
        else:
            observed = "积极正面评论相对较多"
        return f"本批共取得并复核{total}条有效网民评论，样本尚不足以形成正式情感比例；从本批样本看，{observed}。"
    if negative > positive + neutral:
        return "从已审核的网民评论样本看，负面情绪相对突出，同时存在积极正面和客观中立观点。"
    if neutral >= positive and neutral >= negative:
        return "从已审核的网民评论样本看，情感属性以客观中立为主，积极正面和负面观点并存。"
    return "从已审核的网民评论样本看，情感属性以积极正面和客观中立为主。"


def formal_topic_rows(data: dict[str, Any], include_sentiment: bool | None = None) -> list[list[Any]]:
    if include_sentiment is None:
        include_sentiment = formal_sentiment_available(data)
    rows = []
    for idx, row in enumerate(data.get("topic_stats", []), 1):
        values = [
            idx,
            row.get("topic") or row.get("display") or "",
            row.get("domestic_media", 0),
            int(row.get("self_media", 0) or 0)
            if row.get("authority") == "monitoring_system"
            else int(row.get("self_media", 0) or 0) + int(row.get("comments", 0) or 0),
            row.get("overseas_media", 0),
            topic_total(row),
        ]
        if include_sentiment:
            values.extend(sentiment_ratio_cells(row))
        rows.append(values)
    return rows


def peak_date_text(stats: dict[str, Any]) -> str:
    date_counts = {k: int(v or 0) for k, v in (stats.get("by_date") or {}).items() if k != "未标注日期"}
    if not date_counts:
        return "监测期内"
    peak_date, _ = max(date_counts.items(), key=lambda item: item[1])
    try:
        return f"{int(peak_date[5:7])}月{int(peak_date[8:10])}日"
    except Exception:
        return peak_date


def quote_title(topic: str) -> str:
    return topic.strip("，。；") or topic_display(topic)


def joined_agenda_topics(topics: list[str]) -> str:
    """Join consecutive agenda objects sharing the same governing verb."""

    prefixes = ("进一步部署", "审议通过", "听取", "研究", "审议", "部署", "决定")
    groups: list[tuple[str, list[str]]] = []
    for raw_topic in topics:
        topic = quote_title(raw_topic)
        prefix = next((item for item in prefixes if topic.startswith(item)), "")
        obj = topic[len(prefix) :].strip() if prefix else topic
        if groups and groups[-1][0] == prefix:
            groups[-1][1].append(obj)
        else:
            groups.append((prefix, [obj]))

    def join_objects(objects: list[str]) -> str:
        if len(objects) == 1:
            return objects[0]
        return "、".join(objects[:-1]) + "和" + objects[-1]

    return "，".join(f"{prefix}{join_objects(objects)}" for prefix, objects in groups)


def total_event_paragraphs(data: dict[str, Any]) -> list[str]:
    stats = data.get("statistics", {})
    buckets = source_bucket_counts(stats)
    total = int(stats["total_spread"]) if stats.get("total_spread") is not None else int(stats.get("total_samples") or 0)
    domestic = buckets.get("domestic_media", 0)
    overseas = buckets.get("overseas_media", 0)
    wechat = platform_count(stats, "微信", "wechat", "weixin")
    weibo = platform_count(stats, "微博", "weibo", "wb")
    video = platform_count(stats, "视频号", "video_account", "channels")
    other_new_media = platform_count(stats, "new_media")
    if other_new_media <= 0:
        other_new_media = max(0, buckets.get("self_media", 0) + buckets.get("comments", 0) - wechat - weibo - video)
    total_text = f"{total / 10000:.1f}万" if total >= 10000 else str(total)
    representative = representative_overseas_rows(data)
    overseas_sources = "、".join(formal_overseas_source(row) for row in representative[:3]
                                if row.get("source") and not overseas_source_name_pending(row))
    frames = writing_rules()["propagation"]
    dated_counts = {}
    for key, value in (stats.get("by_date") or {}).items():
        try:
            day = date.fromisoformat(str(key))
            count = int(value or 0)
        except (TypeError, ValueError):
            continue
        if count > 0:
            dated_counts[day.isoformat()] = count
    peak = peak_date_text({"by_date": dated_counts})
    peak_sentence = frames["peak_template"].format(peak_date=peak) if dated_counts else ""
    examples = frames["overseas_examples_template"].format(sources=overseas_sources) if overseas_sources and overseas > 0 else ""
    # Current, explicitly sourced headline only; never borrow a benchmark headline.
    meeting = data.get('meeting') or {}
    focus = str(meeting.get('focus_title') or '').strip()
    focus_sentence = frames['focus_template'].format(focus=focus) if focus and meeting.get('focus_source') else ''
    if not focus_sentence and frames.get('focus_from_current_topics'):
        current_topics = report_topics(data)
        if current_topics:
            focus_sentence = frames['focus_template'].format(focus='”“'.join(current_topics))
    return [
        frames["total_template"].format(total_text=total_text, focus_sentence=focus_sentence, peak_sentence=peak_sentence),
        frames["domestic_template"].format(count=domestic, domestic_label=domestic_media_label(data)),
        frames["new_media_template"].format(wechat=wechat, weibo=weibo, video=video, other=other_new_media),
        frames["overseas_template"].format(count=overseas, examples=examples),
    ]


def report_topics(data: dict[str, Any]) -> list[str]:
    """Use current input order, never historical topic names or fuzzy assignment."""
    values = [row.get('topic') or row.get('display') for row in data.get('topic_stats') or []]
    if not values:
        values = (data.get('meeting') or {}).get('topics') or []
    if not values:
        values = [row.get('topic') for row in (data.get('viewpoints') or {}).get('by_topic') or []]
    return list(dict.fromkeys(str(value).strip() for value in values if isinstance(value, str) and value.strip()))


def report_topic_items(data: dict[str, Any]) -> list[dict[str, Any]]:
    rows = list((data.get('viewpoints') or {}).get('by_topic') or [])
    topics = report_topics(data)
    if not topics:
        return rows
    output, consumed = [], set()
    for topic in topics:
        matches = [(index, row) for index, row in enumerate(rows) if str(row.get('topic') or '').strip() == topic]
        if matches:
            output.extend(row for _, row in matches)
            consumed.update(index for index, _ in matches)
        else:
            output.append({'topic': topic, 'heading': topic, 'clusters': [],
                           'evidence_gap': {'notice': '暂未取得该议题可引用的媒体自媒体观点。'}})
    # Unmatched legacy rows remain visible, not silently reassigned or deleted.
    output.extend(row for index, row in enumerate(rows) if index not in consumed)
    return output


def report_comment_groups(data: dict[str, Any]) -> list[tuple[str, list[dict[str, Any]]]]:
    groups = comment_groups((data.get('comments') or {}).get('selected') or [])
    if not writing_rules()['comments'].get('cover_all_topics'):
        return groups
    topics = report_topics(data)
    if not topics:
        return groups
    output, consumed = [], set()
    for topic in topics:
        matches = [(index, group) for index, group in enumerate(groups)
                   if all(str(row.get('system_topic') or row.get('topic') or row.get('subtopic') or '').strip() == topic
                          for row in group[1])]
        if matches:
            output.extend(group for _, group in matches)
            consumed.update(index for index, _ in matches)
        else:
            output.append((topic, []))
    output.extend(group for index, group in enumerate(groups) if index not in consumed)
    return output


def comment_groups(comments: list[dict[str, Any]]) -> list[tuple[str, list[dict[str, Any]]]]:
    groups: dict[str, list[dict[str, Any]]] = {}
    for item in comments:
        if not is_actual_comment_row(item):
            continue
        if item.get("ai_formal_include") is not True:
            continue
        text = clean_formal_comment(item.get("content") or item.get("title"))
        if not text or not comment_is_substantive(item, text):
            continue
        reviewed_heading = (comment_stance_heading(item.get("topic_comment_heading"))
                            or comment_stance_heading(item.get("comment_heading"))
                            or comment_stance_heading(item.get("ai_comment_heading")))
        if not reviewed_heading:
            continue
        # One system agenda/subtopic is one numbered report item.  Different
        # reactions under the same agenda are evidence dimensions inside that
        # item, never separate “一是/二是” headings.
        topic_key = re.sub(
            r"\s+",
            "",
            str(item.get("system_topic") or item.get("topic") or item.get("subtopic") or "").strip(),
        )
        if not topic_key:
            topic_key = f"heading:{reviewed_heading}"
        groups.setdefault(topic_key, []).append(
            {**item, "content": text, "_reviewed_comment_heading": reviewed_heading}
        )
    output = []
    for _, rows in groups.items():
        deduped = []
        seen = set()
        for row in rows:
            key = re.sub(r"\s+", "", str(row.get("content") or ""))
            if not key or key in seen:
                continue
            seen.add(key)
            deduped.append(row)
        if not deduped:
            continue
        group_headings = [
            comment_stance_heading(row.get("topic_comment_heading"))
            for row in deduped
            if comment_stance_heading(row.get("topic_comment_heading"))
        ]
        unique_row_headings = list(dict.fromkeys(
            row.get("_reviewed_comment_heading") for row in deduped if row.get("_reviewed_comment_heading")
        ))
        # A missing shared heading is not permission to drop approved quotes.
        # Preserve the reviewed individual headings verbatim as coordinated
        # clauses; do not invent a new stance or create extra numbered topics.
        heading = group_headings[0] if group_headings else "；".join(unique_row_headings)
        if heading:
            output.append((heading, deduped))
    output.sort(key=lambda item: min(int(row.get("report_order") or 999) for row in item[1]))
    return output


def comment_lead(data: dict[str, Any], groups: list[tuple[str, list[dict[str, Any]]]]) -> str:
    rules = writing_rules()["comments"]

    def lead_summary(name: str) -> str:
        value = clean_sentence(name)
        clauses = [part.strip() for part in re.split(r"[，；]", value) if part.strip()]
        value = clauses[0] if clauses else value
        value = re.sub(r"^认为", "", value)
        # Length is an authoring target, not permission to cut a Chinese word
        # or qualification from a reviewed clause. Preserve the whole clause.
        return value

    summaries = "、".join(lead_summary(name) for name, rows in groups if rows)
    if not summaries:
        return empty_comment_notice(data)
    if rules.get('use_editorial_lead'):
        return rules['lead_template'].format(stance_summaries=summaries)
    if formal_sentiment_available(data):
        return rules["sentiment_lead_template"].format(
            sentiment_lead=netizen_sentiment_lead(data).rstrip("。"), stance_summaries=summaries,
        )
    return rules["lead_template"].format(stance_summaries=summaries)


def comment_is_substantive(item: dict[str, Any], text: str | None = None) -> bool:
    """Enforce the AI review decision and reject empty reactions as a final guard."""

    value = clean_formal_comment(text if text is not None else item.get("content") or item.get("title"))
    if not comment_is_report_quote_suitable(value):
        return False
    if item.get("ai_formal_include") is not True:
        return False
    quality = str(item.get("ai_semantic_quality") or item.get("semantic_quality") or "").lower()
    return quality in {"substantive", "meaningful", "high", "合格", "有效"}


def comment_is_report_quote_suitable(text: str) -> bool:
    """Keep traceable comments in the dataset but avoid article-like factual dumps as quotes."""
    value = clean_formal_comment(text)
    if not value:
        return False
    if len(value) > writing_rules()['comments']['quote_max_chars']:
        return False
    enumerations = len(re.findall(r"(?:^|[\s；;。])(?:[1-9][.、]|[一二三四五六七八九十]是)", value))
    if enumerations >= 2 and re.search(r"(?:省流|核心(?:调整|变化)|分为|包括)", value):
        return False
    return True


def comment_wording(rows: list[dict[str, Any]]) -> str:
    parts = []
    for item in rows:
        text = clean_sentence(clean_formal_comment(item.get("content") or item.get("title")))
        if not text:
            continue
        if item.get("quote_verified") and item.get("url"):
            parts.append(f"“{text}”")
        else:
            parts.append(f"有网民关注{text}。")
        if len(parts) >= writing_rules()["comments"]["quotes_per_ready_topic"][1]:
            break
    if parts and all(part.startswith("“") for part in parts):
        return "网民称，" + "".join(parts) + "。"
    return " ".join(parts)


def hotword_paragraph(data: dict[str, Any]) -> str:
    if data.get('hotword_review_gap'):
        return str(data['hotword_review_gap'])
    frames = writing_rules()["hotwords"]
    hotwords = data.get("hotwords") or []
    if not hotwords:
        return "从热词分布来看，当前批次样本不足，尚未形成稳定热词分布。"
    chunks = []
    by_topic: dict[str, list[str]] = {}
    for item in hotwords:
        by_topic.setdefault(topic_display(str(item.get("topic") or "")), []).append(str(item.get("word") or ""))
    topic_totals = {
        topic_display(str(row.get("topic") or "")): topic_total(row)
        for row in data.get("topic_stats") or []
    }
    ordered_topics = sorted(by_topic, key=lambda topic: topic_totals.get(topic, 0), reverse=True)
    focus_by_topic: dict[str, str] = {}
    for item in (data.get("viewpoints") or {}).get("by_topic") or []:
        display = topic_display(str(item.get("topic") or ""))
        clusters = item.get("clusters") or []
        if clusters:
            focus = clean_sentence(clusters[0].get("summary") or clusters[0].get("details"))
            focus = re.sub(r"^(?:舆论|媒体|专家|网民)(?:普遍)?", "", focus).strip()
            focus_by_topic[display] = focus
    focus_phrases = frames["focus_phrases"]
    for index, topic in enumerate(ordered_topics):
        words = by_topic[topic]
        selected = [w for w in words if w][:frames["words_per_topic_in_prose"]]
        if not selected:
            continue
        quoted = "、".join(f"“{word}”" for word in selected)
        focus = focus_by_topic.get(topic)
        if focus and focus.startswith(STANCE_HEADING_PREFIXES):
            focus_clause = "，" + frames['stance_focus_subject'] + focus
        else:
            focus_clause = ("，" + focus_phrases[index % len(focus_phrases)] + focus) if focus else ""
        lead = "" if index < 4 else "此外，" if index == 4 else "同时，"
        word_group = frames['word_group_template'].format(quoted=quoted, topic=topic)
        chunks.append(f"{lead}{word_group}{focus_clause}")
    if not chunks:
        return "从热词分布来看，当前批次样本不足，尚未形成稳定热词分布。"
    return frames["opening"] + "。".join(chunks) + "。"


def overseas_report_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    rows = data.get("appendices", {}).get("overseas_reports") or []
    if rows:
        return [row for row in rows if not row.get("is_comment") and row.get("source_type") not in {"overseas_public_discussion", "overseas_netizen"}]
    output = []
    for group_name, group_rows in (data.get("overseas", {}).get("groups") or {}).items():
        output.extend({**row, "_overseas_group": group_name} for row in group_rows)
    return output


def contains_chinese(text: Any) -> bool:
    return bool(re.search(r"[\u3400-\u9fff]", str(text or "")))


_T2S_CONVERTER = OpenCC("t2s") if OpenCC is not None else None


def to_simplified_chinese(text: Any) -> str:
    value = re.sub(r"\s+", " ", str(text or "")).strip()
    if not value:
        return ""
    if _T2S_CONVERTER is None:
        raise RuntimeError(
            "正式外媒中文门禁需要 opencc-python-reimplemented；"
            "请先安装 cwh-report-skill/deploy/with/requirements.txt 中的依赖。"
        )
    return _T2S_CONVERTER.convert(value)


def is_simplified_chinese_text(text: Any) -> bool:
    value = re.sub(r"\s+", " ", str(text or "")).strip()
    return bool(value and contains_chinese(value) and to_simplified_chinese(value) == value)


def first_chinese_text(row: dict[str, Any], fields: tuple[str, ...]) -> str:
    for field in fields:
        value = re.sub(r"\s+", " ", str(row.get(field) or "")).strip()
        if value and contains_chinese(value):
            return value
    return ""


def formal_overseas_title(row: dict[str, Any]) -> str:
    title = first_chinese_text(
        row,
        (
            "title_cn_simplified",
            "formal_title_cn_simplified",
            "title_cn",
            "title_zh",
            "translation_title_cn",
            "translated_title_cn",
            "formal_title_cn",
        ),
    )
    if title:
        return to_simplified_chinese(title)
    raw_title = re.sub(r"\s+", " ", str(row.get("title") or "")).strip()
    return to_simplified_chinese(raw_title) if contains_chinese(raw_title) else ""


def formal_overseas_summary(row: dict[str, Any]) -> str:
    summary = first_chinese_text(
        row,
        (
            "summary_cn_simplified",
            "formal_summary_cn_simplified",
            "summary_cn",
            "translation_cn",
            "formal_summary_cn",
            "interpretive_summary_cn",
            "content_cn",
        ),
    )
    if summary:
        value = clean_sentence(to_simplified_chinese(summary))
    else:
        fallback = first_chinese_text(row, ("summary", "description"))
        value = clean_sentence(to_simplified_chinese(fallback)) if fallback else ""
    value = re.split(
        r"(?:您查看的内容可能不完整|请对本站关闭广告拦截|更多内容访问|热度：\s*加载中)",
        value,
        maxsplit=1,
    )[0].strip()
    # Length is an authoring target, not permission to cut a reviewed condition.
    match = re.match(r'^(?:原文|原分析)(?:认为|指出|表示|称)[，,:：\s]*', value)
    if match and value[match.end():].strip():
        value = value[match.end():].strip()
    return value


FOREIGN_SOURCE_CN = {
    "Wedoany English": "Wedoany英文网",
    "Reuters": "路透社",
    "Bloomberg": "彭博社",
    "South China Morning Post": "香港南华早报",
    "Radio Television Hong Kong": "香港电台",
    "United Daily News": "台湾联合新闻网",
}


def formal_overseas_source(row: dict[str, Any]) -> str:
    source = re.sub(r"\s+", " ", str(row.get("source") or "境外媒体")).strip()
    source_cn = re.sub(
        r"\s+",
        " ",
        str(row.get("source_cn_simplified") or row.get("formal_source_cn_simplified") or row.get("source_cn") or row.get("source_zh") or FOREIGN_SOURCE_CN.get(source, "")),
    ).strip()
    name = to_simplified_chinese(source_cn) if contains_chinese(source_cn) else (
        to_simplified_chinese(source) if contains_chinese(source) else "境外媒体")
    # These exact labels describe geography or an unknown source, not an outlet.
    # This is presentation only: retain raw rows, eligibility and all counts.
    if name in {"香港", "澳门", "台湾", "新加坡", "境外", "海外", "境外媒体", "海外媒体"}:
        return name + "（媒体名称待核）"
    return name


def overseas_source_name_pending(row: dict[str, Any]) -> bool:
    return formal_overseas_source(row).endswith("（媒体名称待核）")


def overseas_report_category(row: dict[str, Any]) -> str:
    category_aliases = {
        "转载报道": "事实性报道",
        "事实报道": "事实性报道",
        "事实性报道": "事实性报道",
        "一般解读": "解读性报道",
        "评论解读": "解读性报道",
        "解读性报道": "解读性报道",
        "风险叙事": "借题炒作/风险解读",
        "借题炒作/风险叙事": "借题炒作/风险解读",
        "风险解读": "借题炒作/风险解读",
        "借题炒作/风险解读": "借题炒作/风险解读",
    }
    for field in (
        "ai_report_category",
        "reviewed_report_category",
        "overseas_category",
        "report_category",
        "analysis_type",
        "narrative_type",
        "category",
        "_overseas_group",
    ):
        value = str(row.get(field) or "").strip()
        if value in category_aliases:
            return category_aliases[value]
    if row.get("interpretive_verified"):
        return "解读性报道"
    return ""


def overseas_report_is_formal_eligible(row: dict[str, Any]) -> bool:
    """Keep relevance review separate from report-type classification."""
    if row.get("formal_include") is not True or row.get("meeting_relevance") is not True:
        return False
    relevance = str(
        row.get("ai_relevance")
        or row.get("reviewed_relevance")
        or row.get("relevance_status")
        or ""
    ).strip().lower()
    return relevance not in {"irrelevant", "unrelated", "exclude", "不相关", "排除"}


def is_actual_comment_row(row: dict[str, Any]) -> bool:
    mode = str(row.get("evidence_mode") or "").strip().lower()
    raw_file = str(row.get("raw_file") or "").lower()
    has_identity = bool(
        row.get("comment_id")
        or row.get("reply_id")
        or row.get("parent_comment_id")
        or any(marker in raw_file for marker in ("comment", "reply", "评论", "回复"))
    )
    return bool(
        row.get("quote_verified")
        and row.get("url")
        and mode in {"verbatim_public_comment", "verbatim_comment", "platform_comment", "platform_reply"}
        and has_identity
    )


def overseas_comment_rows(
    data: dict[str, Any], *, in_window_only: bool = True
) -> list[dict[str, Any]]:
    rows = (data.get("overseas") or {}).get("public_comments") or []
    output = []
    for row in rows:
        if in_window_only and row.get("in_monitoring_window") is False:
            continue
        if row.get("ai_formal_include") is False:
            continue
        all_text = " ".join(str(value or "") for value in row.values())
        if re.search(
            r"推翻(?:现政权|政府|共产党)|颠覆(?:政权|政府)|打倒(?:共产党|中共|政府)|"
            r"(?:中共|共产党|中国政府).{0,12}(?:独裁|邪恶|暴政|垃圾|该死)|"
            r"(?:独裁|邪恶|暴政).{0,12}(?:中共|共产党|中国政府)|"
            r"overthrow\s+(?:the\s+)?(?:regime|government|ccp)|down\s+with\s+(?:the\s+)?ccp|"
            r"(?:evil|dictator(?:ship)?|tyranny).{0,20}(?:ccp|chinese\s+government)|"
            r"(?:ccp|chinese\s+government).{0,20}(?:evil|dictator(?:ship)?|tyranny)|fuck\s+(?:the\s+)?ccp",
            all_text,
            re.IGNORECASE,
        ):
            continue
        chinese_content = clean_formal_comment(
            first_chinese_text(
                row,
                (
                    "translation_cn",
                    "translation",
                    "summary_cn",
                    "formal_summary_cn",
                    "content_cn",
                ),
            )
        )
        if not chinese_content:
            raw_content = clean_formal_comment(row.get("content") or row.get("title"))
            chinese_content = raw_content if contains_chinese(raw_content) else ""
        if chinese_content and comment_is_substantive(row, chinese_content):
            output.append({**row, "_formal_content_cn": chinese_content})
    return output


def overseas_comment_text(row: dict[str, Any]) -> str:
    source = str(row.get("source") or row.get("platform") or "境外平台用户").strip()
    content = str(row.get("_formal_content_cn") or "").strip().rstrip("。！？.!?")
    return f"{source}用户认为，{content}。"


OVERSEAS_COMMENT_THEMES = (
    ("应急救灾与人员安全", ("救灾", "洪水", "洪灾", "防汛", "受灾", "救援", "伤亡", "安全", "平安", "撤离", "生命", "救人", "无人机", "军人", "广西", "不孤单", "团结", "协作", "团队合作", "上下一条心", "效率"), "关注应急响应、救援协作和人员安全，希望相关措施尽快落地"),
    ("基础设施与工程质量", ("基础设施", "排水", "大坝", "堤坝", "工程", "施工", "质量", "豆腐渣", "重建"), "关注防灾基础设施、工程质量和灾后重建，希望提升长期防灾韧性"),
    ("公共治理与政策执行", ("政府", "政策", "执行", "治理", "规划", "监管", "部门", "落实", "效率"), "关注政策执行和公共治理效能，并就具体措施的落实提出意见"),
    ("产业发展与民生改善", ("产业", "经济", "就业", "贸易", "人工智能", "数字", "健康", "健身", "医疗", "民生"), "关注产业发展、公共服务和民生改善带来的实际影响"),
    ("气候变化与灾害风险", ("气候", "全球变暖", "极端天气", "气候变化"), "关注气候变化与极端天气风险，希望提升长期防灾减灾能力"),
    ("信息发布与社会沟通", ("信息", "公开", "透明", "媒体", "报道", "宣传", "沟通", "数据"), "关注信息发布的及时性、透明度和社会沟通效果"),
)


def overseas_comment_platform(row: dict[str, Any]) -> str:
    platform = str(row.get("platform") or "").strip()
    source = str(row.get("source") or "").strip()
    if platform.lower() in {"x", "twitter"} or source.startswith("@") or "x.com/" in str(row.get("url") or ""):
        return "X平台"
    if platform:
        return f"{platform}平台"
    return "境外社交平台"


def overseas_comment_quote_allowed(text: str) -> bool:
    return not re.search(
        r"中共|共产党|CCP|摆拍|宣传|笼络人心|拍一个救灾|正常的国家|邪恶|独裁|暴政|推翻|打倒|颠覆|"
        r"西方建设|豆腐渣|侵略|真厉害|应该很容易|"
        r"overthrow|down\s+with|fuck\s+(?:the\s+)?ccp",
        text,
        re.IGNORECASE,
    )


def overseas_comment_quote_score(text: str) -> int:
    score = 0
    for marker in ("平安", "安全", "准备", "拯救生命", "团结", "协作", "效率", "救援", "建设", "重建"):
        if marker in text:
            score += 2
    if len(text) <= 36:
        score += 1
    return score


def overseas_comment_paragraphs(data: dict[str, Any]) -> list[str]:
    # Formal prose treats approved system and supplemental comments identically.
    # Provenance and monitoring-window status remain visible in structured data/dashboard.
    rows = overseas_comment_rows(data, in_window_only=False)
    if not rows:
        return []
    grouped: dict[str, list[dict[str, Any]]] = {}
    summaries: dict[str, str] = {}
    for row in rows:
        content = str(row.get("_formal_content_cn") or "")
        theme_name = "其他相关议题"
        theme_summary = "期待相关政策安排取得实际成效"
        reviewed_heading = str(row.get("comment_heading") or row.get("ai_comment_heading") or "").strip()
        reviewed_summary = str(row.get("comment_summary") or row.get("ai_comment_summary") or "").strip()
        if reviewed_heading:
            theme_name = reviewed_heading
            theme_summary = reviewed_summary
            grouped.setdefault(theme_name, []).append(row)
            summaries[theme_name] = theme_summary
            continue
        for name, keywords, summary in OVERSEAS_COMMENT_THEMES:
            if any(keyword in content for keyword in keywords):
                theme_name = name
                theme_summary = summary
                break
        grouped.setdefault(theme_name, []).append(row)
        summaries[theme_name] = theme_summary
    paragraphs = ["境外网民观点主要集中在以下方面："]
    for index, (theme_name, theme_rows) in enumerate(grouped.items()):
        quotes = []
        seen_quotes = set()
        ranked_theme_rows = sorted(
            theme_rows,
            key=lambda row: overseas_comment_quote_score(
                clean_sentence(row.get("_formal_content_cn"))
            ),
            reverse=True,
        )
        for row in ranked_theme_rows:
            quote = clean_sentence(row.get("_formal_content_cn"))
            if not quote or quote in seen_quotes or not overseas_comment_quote_allowed(quote):
                continue
            seen_quotes.add(quote)
            if len(quote) > 58:
                quote = quote[:57].rstrip("，、；： ") + "……"
            quotes.append((overseas_comment_platform(row), quote))
            if len(quotes) >= 2:
                break
        if not quotes:
            continue
        quote_clause = ""
        if quotes:
            examples = "；另有".join(f"{platform}网友称“{quote}”" for platform, quote in quotes)
            quote_clause = f"。如{examples}"
        lead = theme_name
        summary = summaries[theme_name]
        if summary:
            lead = f"{lead}，{summary}"
        paragraphs.append(f"{numbered(len(paragraphs) - 1)}是{lead}{quote_clause}。")
    return paragraphs


def overseas_comment_summary(data: dict[str, Any]) -> str:
    """Backward-compatible text view of the reviewed overseas comment groups."""

    return "\n".join(overseas_comment_paragraphs(data))


MAINLAND_FOREIGN_LANGUAGE_OUTLETS = (
    "新华社海外版",
    "新华网海外版",
    "CGTN",
    "人民网英文版",
    "中国经济网英文版",
    "央视网英文版",
    "中国日报英文版",
)


def is_mainland_foreign_language_outlet(row: dict[str, Any]) -> bool:
    source = str(row.get("source") or "")
    return any(name.lower() in source.lower() for name in MAINLAND_FOREIGN_LANGUAGE_OUTLETS)


def overseas_region_priority(row: dict[str, Any]) -> int:
    text = f"{row.get('source', '')}{row.get('title', '')}"
    region_markers = (
        (0, ("\u9999\u6e2f", "\u6e2f\u6fb3")),
        (1, ("\u53f0\u6e7e", "\u4e2d\u8bc4")),
        (2, ("\u65b0\u52a0\u5761", "\u8054\u5408\u65e9\u62a5")),
        (3, ("\u6fb3\u95e8",)),
    )
    for priority, markers in region_markers:
        if any(marker in text for marker in markers):
            return priority
    return 4


def formal_overseas_story_token(row: dict[str, Any]) -> str:
    url = str(row.get("url") or "")
    match = re.search(
        r"(?:NOW[./_-]?|[?&](?:sn|newsid|articleid|storyid|id)=|/)([0-9]{6,})(?:\D|$)",
        url,
        flags=re.I,
    )
    return match.group(1) if match else ""


def formal_overseas_title_key(row: dict[str, Any]) -> str:
    title = formal_overseas_title(row)
    title = re.sub(r"[\W_]+", "", title, flags=re.UNICODE).lower()
    return title.replace("热门内容", "").replace("熱門內容", "")


def dedupe_formal_overseas_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Keep one readable representative per story in the formal partial appendix.

    The monitoring total may still count cross-publisher reprints.  This display
    layer removes same-article mirrors and materially identical syndicated titles
    so the partial appendix does not read like repeated rows.
    """
    selected: list[dict[str, Any]] = []
    story_tokens: set[str] = set()
    title_keys: list[str] = []
    for row in sorted(rows, key=overseas_source_name_pending):
        token = formal_overseas_story_token(row)
        title_key = formal_overseas_title_key(row)
        if token and token in story_tokens:
            continue
        if title_key and any(
            title_key == existing
            or (min(len(title_key), len(existing)) >= 12 and (title_key in existing or existing in title_key))
            for existing in title_keys
        ):
            continue
        selected.append(row)
        if token:
            story_tokens.add(token)
        if title_key:
            title_keys.append(title_key)
    return selected


def appendix_overseas_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for row in overseas_report_rows(data):
        if is_mainland_foreign_language_outlet(row):
            continue
        if not overseas_report_is_formal_eligible(row):
            continue
        title_cn = formal_overseas_title(row)
        category = overseas_report_category(row)
        if not title_cn or not category:
            continue
        rows.append(
            {
                **row,
                "_formal_title_cn": title_cn,
                "_formal_summary_cn": formal_overseas_summary(row),
                "_formal_category": category,
            }
        )
    ordered = sorted(rows, key=lambda row: (str(row.get("published_at") or ""), overseas_region_priority(row)))
    return dedupe_formal_overseas_rows(ordered)


def formal_appendix_overseas_rows(data: dict[str, Any], limit: int = 10) -> list[dict[str, Any]]:
    """Select representative rows for the formal partial appendix; keep full evidence elsewhere."""
    rows = appendix_overseas_rows(data)
    if len(rows) <= limit:
        return rows

    def confidence(row: dict[str, Any]) -> float:
        for field in ("ai_relevance_confidence", "reviewed_confidence", "confidence"):
            try:
                return float(row.get(field))
            except (TypeError, ValueError):
                continue
        return 0.0

    ranked = sorted(rows, key=lambda row: (-confidence(row), str(row.get("published_at") or ""), overseas_region_priority(row)))
    selected: list[dict[str, Any]] = []

    def add_first(predicate: Any) -> None:
        match = next((row for row in ranked if row not in selected and predicate(row)), None)
        if match is not None:
            selected.append(match)

    for category in ("事实性报道", "解读性报道", "借题炒作/风险解读"):
        add_first(lambda row, category=category: row.get("_formal_category") == category)
    for region_priority in range(4):
        add_first(lambda row, region_priority=region_priority: overseas_region_priority(row) == region_priority)
    seen_topics = {str(row.get("topic") or row.get("subtopic") or "") for row in selected}
    for row in ranked:
        topic = str(row.get("topic") or row.get("subtopic") or "")
        if topic and topic not in seen_topics and len(selected) < limit:
            selected.append(row)
            seen_topics.add(topic)
    for row in ranked:
        if len(selected) >= limit:
            break
        if row not in selected:
            selected.append(row)
    return selected[:limit]


def representative_overseas_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    rows = [row for row in appendix_overseas_rows(data) if row.get("_formal_category") == "事实性报道"]
    # Prefer named examples without removing unnamed material from the appendix.
    named_rows = [row for row in rows if not overseas_source_name_pending(row)]
    rows = named_rows or rows
    region_markers = [
        ("新加坡", ("联合早报", "新加坡")),
        ("香港", ("香港", "港澳")),
        ("台湾", ("台湾", "中评", "中央社")),
        ("澳门", ("澳门",)),
    ]
    selected: list[dict[str, Any]] = []
    selected_sources: set[str] = set()

    def add(row: dict[str, Any] | None) -> None:
        if not row:
            return
        source_key = re.sub(r"\W+", "", formal_overseas_source(row)).casefold()
        if source_key in selected_sources:
            return
        selected.append(row)
        selected_sources.add(source_key)

    for _, markers in region_markers:
        match = next(
            (
                row
                for row in rows
                if any(marker in f"{row.get('source', '')}{row.get('title', '')}" for marker in markers)
                and re.sub(r"\W+", "", formal_overseas_source(row)).casefold() not in selected_sources
            ),
            None,
        )
        add(match)
    for row in rows:
        add(row)
    return selected


def overseas_media_body_paragraphs(data: dict[str, Any]) -> list[str]:
    frames = writing_rules()["overseas"]
    appendix_rows = appendix_overseas_rows(data)
    if not appendix_rows:
        return [writing_rules()['overseas']['no_media_evidence_sentence']]

    factual_rows = [row for row in appendix_rows if row.get("_formal_category") == "事实性报道"]
    interpretive_rows = [
        row for row in appendix_rows
        if row.get("_formal_category") in {"解读性报道", "借题炒作/风险解读"}
        and row.get("_formal_summary_cn")
    ]
    paragraphs: list[str] = []

    if factual_rows:
        representative_keys = {
            str(row.get("url") or row.get("title") or "")
            for row in representative_overseas_rows(data)[:frames["max_factual_examples"]]
        }
        ordered_rows = [
            *[row for row in factual_rows if str(row.get("url") or row.get("title") or "") in representative_keys],
            *[row for row in factual_rows if str(row.get("url") or row.get("title") or "") not in representative_keys],
        ]
        examples = "、".join(
            f"{formal_overseas_source(row)}文章《{row.get('_formal_title_cn')}》"
            for row in ordered_rows[:frames["max_factual_examples"]]
        )
        factual_majority = len(factual_rows) > len(appendix_rows) / 2
        lead_key = "factual_lead_template" if factual_majority else "mixed_factual_lead_template"
        ending_key = ("interpretive_transition" if factual_majority else "mixed_interpretive_transition") if interpretive_rows else "no_interpretation_suffix"
        paragraphs.append(frames[lead_key].format(examples=examples) + frames[ending_key])

    if interpretive_rows:
        descriptions = []
        for row in interpretive_rows[:frames["max_interpretive_examples"]]:
            source = formal_overseas_source(row)
            title = str(row.get("_formal_title_cn") or "").strip()
            summary = clean_sentence(row.get("_formal_summary_cn"))
            summary = re.sub(r"^(?:文章|报道|原(?:文)?分析)?(?:认为|指出|称)[，,:：\s]*", "", summary)
            # Reported analysis may belong to a quoted speaker, not the outlet.
            descriptions.append(f"{source}文章《{title}》称，{summary}")
        # The factual paragraph has already introduced the interpretations.
        lead = "" if factual_rows else frames["interpretive_lead"]
        paragraphs.append(lead + "；".join(descriptions) + "。")

    if not paragraphs:
        return ["数据周期内，暂未取得通过报道类型和简体中文门禁的境外报道样本。"]
    return paragraphs


def wechat_top_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    appendices = data.get("appendices") or {}
    rows = appendices.get("wechat_top") or []
    if not rows:
        rows = appendices.get("media_top") or []

    config = public_source_identity_config()

    def normalized_source(row: dict[str, Any]) -> str:
        return public_source_family(
            row.get("source") or row.get("account") or row.get("author") or "",
            config,
        )

    def read_count(row: dict[str, Any]) -> float:
        value = row.get("spread_count")
        if value in (None, ""):
            value = row.get("read_count") or row.get("reads") or row.get("spread_count_display") or 0
        text = str(value).replace(",", "").strip()
        match = re.search(r"\d+(?:\.\d+)?", text)
        if not match:
            return 0.0
        number = float(match.group())
        return number * 10000 if "万" in text else number

    best_by_source: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(rows):
        key = normalized_source(row) or f"__missing_source_{index}"
        current = best_by_source.get(key)
        if current is None or read_count(row) > read_count(current):
            best_by_source[key] = row
    return sorted(best_by_source.values(), key=read_count, reverse=True)


def write_bar_png(path: Path, title: str, values: dict[str, Any]) -> None:
    from PIL import Image, ImageDraw, ImageFont

    def chart_font(size: int):
        for candidate in [
            Path("C:/Windows/Fonts/msyh.ttc"),
            Path("C:/Windows/Fonts/simhei.ttf"),
            Path("C:/Windows/Fonts/simsun.ttc"),
        ]:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size=size)
        return ImageFont.load_default()

    items = [(str(k), int(v or 0)) for k, v in (values or {}).items()]
    if not items:
        items = [("暂无可统计样本", 0)]
    width = 1200
    row_h = 54
    height = 120 + row_h * len(items)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = chart_font(18)
    title_font = chart_font(25)
    max_value = max([v for _, v in items] + [1])
    draw.text((28, 24), title, fill=(20, 20, 20), font=title_font)
    y = 78
    for name, value in items:
        label = name[:34]
        draw.text((28, y + 8), label, fill=(20, 20, 20), font=font)
        bar_w = int((value / max_value) * 700) if max_value else 0
        draw.rectangle((360, y, 360 + bar_w, y + 26), fill=(47, 111, 115))
        draw.text((370 + bar_w, y + 6), str(value), fill=(20, 20, 20), font=font)
        y += row_h
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def write_wordcloud_png(path: Path, title: str, values: dict[str, Any]) -> None:
    from PIL import Image, ImageDraw, ImageFont

    font_path = next(
        (candidate for candidate in [Path("C:/Windows/Fonts/simhei.ttf"), Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simsun.ttc")] if candidate.exists()),
        None,
    )

    def cloud_font(size: int):
        return ImageFont.truetype(str(font_path), size=size) if font_path else ImageFont.load_default()

    items = sorted(((str(key), max(1, int(value or 1))) for key, value in (values or {}).items() if key), key=lambda item: item[1], reverse=True)
    width, height = 1200, 680
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((36, 24), title, fill=(20, 20, 20), font=cloud_font(28))
    if not items:
        draw.text((36, 100), "暂无候选词", fill=(80, 80, 80), font=cloud_font(30))
    else:
        maximum = max(value for _, value in items)
        palette = [(40, 101, 108), (166, 63, 52), (55, 91, 145), (70, 116, 76), (98, 76, 133), (70, 70, 70)]
        x, y, row_height = 36, 95, 0
        for index, (word, value) in enumerate(items[:30]):
            size = 26 + round(28 * value / maximum)
            font = cloud_font(size)
            left, top, right, bottom = draw.textbbox((0, 0), word, font=font)
            word_width, word_height = right - left, bottom - top
            if x + word_width > width - 36:
                x = 36
                y += row_height + 18
                row_height = 0
            if y + word_height > height - 28:
                break
            draw.text((x, y), word, fill=palette[index % len(palette)], font=font)
            x += word_width + 34
            row_height = max(row_height, word_height)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def reviewed_hotword_pipeline_ready(data: dict[str, Any]) -> bool:
    rows = [row for row in data.get("hotwords") or [] if isinstance(row, dict)]
    return bool(rows) and all(
        str(row.get("authority") or "") == "reviewed_wordcloud_audit"
        and str(row.get("selection_method") or "") in {"ai_semantic_review", "manual_review"}
        for row in rows
    )


def reviewed_hotword_image_path(data: dict[str, Any]) -> Path | None:
    audit_value = str((data.get("artifacts") or {}).get("hotword_audit") or "").strip()
    if not audit_value:
        return None
    audit_path = Path(audit_value)
    if not audit_path.exists() or audit_path.stat().st_size <= 0:
        return None
    try:
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError):
        return None
    if (
        str(audit.get("status") or "") != "ai_review_complete"
        or str(audit.get("review_method") or "") not in {"ai_semantic_review", "manual_review"}
        or audit.get("second_pass_completed") is not True
    ):
        return None
    image_value = str(audit.get("image_path") or "").strip()
    if not image_value:
        return None
    image_path = Path(image_value)
    if not image_path.is_absolute():
        image_path = audit_path.parent / image_path
    if not image_path.exists() or image_path.stat().st_size <= 0:
        return None
    return image_path


def ensure_docx_chart_images(data: dict[str, Any], out_dir: Path) -> dict[str, str]:
    chart_dir = out_dir / "charts"
    chart_dir.mkdir(parents=True, exist_ok=True)
    stats = data.get("statistics", {})
    hotwords = {x.get("word", ""): x.get("count", 0) for x in data.get("hotwords", [])[:12]}
    topic_values = [
        (str(row.get("display") or topic_display(row.get("topic", ""))), int(row.get("spread_count", 0) or row.get("total_samples", 0) or 0))
        for row in data.get("topic_stats") or []
    ]
    charts = {
        "platform_distribution": ("平台样本分布", stats.get("by_platform") or {}),
        "topic_distribution": ("子议题传播量", topic_values),
        "trend_distribution": ("传播走势", dict(sorted((stats.get("by_date") or {}).items()))),
        "hotword_distribution": ("候选热词", hotwords),
    }
    output, visual_manifests = write_report_charts(data, chart_dir)
    data.setdefault('audit', {})['report_visual_templates'] = visual_manifests
    # Source charts remain untouched/auditable. They are no longer a styling
    # authority: every report uses the same current-data baseline templates.
    artifacts = data.get("artifacts") or {}
    reviewed_image = reviewed_hotword_image_path(data)
    # Regeneration often follows a copied or restored report directory.  A stale
    # absolute path must not mask the valid reviewed word cloud already copied
    # into that directory.  Try every authoritative candidate before allowing
    # the generic fallback renderer to run.
    wordcloud_candidates = [
        artifacts.get("wordcloud_image"),
        (data.get("collection") or {}).get("wordcloud_image"),
        str(reviewed_image) if reviewed_image is not None else "",
        chart_dir / "hotword_distribution_pipeline.png",
        (artifacts.get("docx_charts") or {}).get("hotword_distribution"),
        (artifacts.get("charts") or {}).get("hotword_distribution"),
        chart_dir / "hotword_distribution_system.png",
    ]
    if data.get('hotword_review_gap'):
        wordcloud_candidates = []
        for owner in (data.setdefault('artifacts', {}), data.get('collection') or {}):
            owner.pop('wordcloud_image', None)
        for key in ('charts', 'docx_charts'):
            (data['artifacts'].get(key) or {}).pop('hotword_distribution', None)
    explicit_wordcloud_value = ""
    invalid_wordcloud_values: list[str] = []
    for candidate_value in wordcloud_candidates:
        candidate_text = str(candidate_value or "").strip()
        if not candidate_text:
            continue
        candidate = Path(candidate_text)
        if candidate.is_file() and candidate.stat().st_size > 0:
            explicit_wordcloud_value = candidate_text
            break
        invalid_wordcloud_values.append(candidate_text)
    if invalid_wordcloud_values:
        gaps = data.setdefault("audit", {}).setdefault("data_gaps", [])
        for invalid_value in dict.fromkeys(invalid_wordcloud_values):
            message = f"指定的词云图片不存在或为空，已继续寻找有效正式词云：{invalid_value}"
            if message not in gaps:
                gaps.append(message)
    explicit_wordcloud_used = False
    if explicit_wordcloud_value:
        explicit_wordcloud = Path(explicit_wordcloud_value)
        if explicit_wordcloud.is_file() and explicit_wordcloud.stat().st_size > 0:
            target = chart_dir / "hotword_distribution_pipeline.png"
            if explicit_wordcloud.resolve() != target.resolve():
                shutil.copyfile(explicit_wordcloud, target)
            output["hotword_distribution"] = str(target)
            data.setdefault("artifacts", {})["wordcloud_image"] = str(target)
            explicit_wordcloud_used = True
    for key, (title, values) in charts.items():
        if key == 'hotword_distribution' and data.get('hotword_review_gap'):
            output.pop(key, None)
            continue
        if key in output:
            continue
        path = chart_dir / f"{key}.png"
        if key == "hotword_distribution":
            write_wordcloud_png(path, title, values)
        elif key == 'topic_distribution':
            data.setdefault('audit', {})['topic_chart_style'] = write_topic_chart(path, values)
        else:
            write_bar_png(path, title, values)
        output[key] = str(path)
    output.pop("source_workbook", None)
    # Monitoring assets remain source artifacts, never override fixed report styling.
    data.setdefault("artifacts", {})["chart_authority"] = (
        'baseline_template_plus_reviewed_wordcloud' if explicit_wordcloud_used
        else 'baseline_template_without_reviewed_wordcloud')
    data.setdefault("artifacts", {}).setdefault("docx_charts", {}).update(output)
    return output


def render_formal_markdown(data: dict[str, Any], out_dir: Path) -> str:
    document_rules = writing_rules()["document"]
    meeting = data["meeting"]
    stats = data["statistics"]
    topics = meeting.get("topics", [])
    month_day = month_day_text(meeting.get("date", ""))
    agenda_topics = joined_agenda_topics(topics) or meeting.get("agenda", "")

    lines: list[str] = [
        "# " + document_rules["title_template"].format(date=month_day),
        "",
        opening_paragraph(meeting, month_day, agenda_topics),
        "",
        "## " + document_rules["fixed_chapters"][0],
        "",
        "### （一）总事件传播情况",
        "",
    ]
    lines.extend(total_event_paragraphs(data))
    lines.extend(["", "### （二）子议题传播情况", ""])
    sentiment_visible = formal_sentiment_available(data)
    topic_headers = ["序号", "标题", domestic_media_label(data), "新媒体", "境外媒体", "总量"]
    topic_headers.extend(["正面", "中立", "负面"])
    lines.append(md_table(topic_headers, [[('' if v == '待分析' else v) for v in row]
                                         for row in formal_topic_rows(data, True)]))

    lines.extend(["", "## " + document_rules["fixed_chapters"][1], "", "### " + document_rules["domestic_subsections"][0], ""])
    for topic_idx, item in enumerate(report_topic_items(data), 1):
        clusters = item.get("clusters") or []
        if not clusters:
            lines.append(f"{topic_idx}.{topic_heading(item)}")
            lines.append((item.get("evidence_gap") or {}).get("notice") or "当前公开样本密度不足，尚未形成可稳定归纳的媒体自媒体观点。")
            lines.append("")
            continue
        if len(clusters) == 1:
            details = clean_sentence(display_cluster_details(clusters[0]))
            body = details or clean_sentence(clusters[0].get("summary"))
            warning = ATTRIBUTION_REVIEW_TEXT if cluster_needs_attribution_review(clusters[0]) else ""
            paragraphs = cluster_detail_paragraphs(clusters[0])
            lines.append(f"{topic_idx}.{topic_heading(item)}。" + ('\n\n'.join(paragraphs) if paragraphs else body + '。') + warning)
            lines.append("")
            continue
        lines.append(f"{topic_idx}.{topic_heading(item)}")
        for idx, cluster in enumerate(clusters, 1):
            prefix = ordinal_prefix(idx)
            lines.append(f"{prefix}{cluster_wording(cluster)}")
        lines.append("")

    comments = data.get("comments", {}).get("selected") or []
    lines.extend(["### " + document_rules["domestic_subsections"][1], ""])
    groups = report_comment_groups(data)
    if groups:
        lines.append(comment_lead(data, groups))
        for idx, (name, rows) in enumerate(groups, 1):
            prefix = ordinal_prefix(idx)
            wording = comment_wording(rows) if rows else writing_rules()['comments']['missing_topic_notice']
            lines.append(f"{prefix}{name}。{wording}")
    else:
        lines.append(empty_comment_notice(data))

    lines.extend(["", "### " + document_rules["domestic_subsections"][2], "", hotword_paragraph(data)])

    appendix_rows = formal_appendix_overseas_rows(data)
    lines.extend(["", "## " + document_rules["fixed_chapters"][2], "", "### " + document_rules["overseas_subsections"][0], ""])
    lines.extend(overseas_media_body_paragraphs(data))
    foreign_comment_paragraphs = overseas_comment_paragraphs(data)
    lines.extend(["", "### " + document_rules["overseas_subsections"][1]])
    if foreign_comment_paragraphs:
        lines.extend(["", *foreign_comment_paragraphs])
    else:
        lines.append(writing_rules()["overseas"]["no_comment_evidence_sentence"])

    lines.extend(["", "## " + document_rules["fixed_chapters"][3], "", "### （一）外媒报道列表（部分）", ""])
    lines.append(
        md_table(
            ["序号", "来源", "报道日期", "超链接标题"],
            [[idx, formal_overseas_source(x), x.get("published_at") or "", x.get("_formal_title_cn")] for idx, x in enumerate(appendix_rows, 1)],
        )
    )
    lines.extend(["", "### （二）传播量较大的公众号文章", "", "微信公号文章阅读量TOP", ""])
    wechat_rows = wechat_top_rows(data)
    lines.append(
        md_table(
            ["序号", "账号", "标题", "阅读量", "在看量"],
            [
                [
                    idx,
                    x.get("source"),
                    x.get("title"),
                    display_metric(x.get("spread_count"), x.get("spread_count_display")),
                    x.get("comment_count"),
                ]
                for idx, x in enumerate(wechat_rows[:10], 1)
            ],
            "未接入公众号阅读量榜单数据",
        )
    )
    return "\n".join(lines) + "\n"


def set_docx_run_font(run: Any, font_name: str, size_pt: int, bold: bool = False) -> None:
    run.font.name = font_name
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    for attr in ["eastAsiaTheme", "asciiTheme", "hAnsiTheme", "cstheme"]:
        r_fonts.attrib.pop(qn(f"w:{attr}"), None)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_docx_table(table: Any, font_name: str, size_pt: int, widths_pt: list[float]) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is not None:
        tbl_w.set(qn("w:w"), str(round(sum(widths_pt) * 20)))
        tbl_w.set(qn("w:type"), "dxa")
    for grid_col, width_pt in zip(table._tbl.tblGrid.gridCol_lst, widths_pt):
        grid_col.set(qn("w:w"), str(round(width_pt * 20)))
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        is_header = tr_pr.find(qn("w:tblHeader")) is not None
        for column_index, cell in enumerate(row.cells):
            width_pt = widths_pt[min(column_index, len(widths_pt) - 1)]
            cell.width = Pt(width_pt)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(round(width_pt * 20)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                # Keep all consecutive header rows with the first data row.
                paragraph.paragraph_format.keep_with_next = is_header
                for run in paragraph.runs:
                    set_docx_run_font(run, font_name, size_pt, is_header)


def shade_cell(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def add_hyperlink(paragraph: Any, text: str, url: str, *, font_name: str, size_pt: int) -> Any:
    if not url:
        run = paragraph.add_run(text)
        set_docx_run_font(run, font_name, size_pt)
        return run
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    for child_run in hyperlink.findall(qn("w:r")):
        r_pr = child_run.find(qn("w:rPr"))
        fonts = OxmlElement("w:rFonts")
        for attr in ["eastAsia", "ascii", "hAnsi"]:
            fonts.set(qn(f"w:{attr}"), font_name)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), str(size_pt * 2))
        size_cs = OxmlElement("w:szCs")
        size_cs.set(qn("w:val"), str(size_pt * 2))
        r_pr.insert(0, fonts)
        r_pr.extend([size, size_cs])
    return hyperlink


def mark_row_as_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def add_table(
    document: Any,
    headers: list[str],
    rows: list[list[Any]],
    *,
    font_name: str = "微软雅黑",
    size_pt: int = 11,
    widths_pt: list[float] | None = None,
) -> Any:
    table = document.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr[idx].text = str(header)
    mark_row_as_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = "" if value is None else str(value)
    style_docx_table(table, font_name, size_pt, widths_pt or [60.0] * len(headers))
    return table


def add_topic_table(document: Any, rows: list[list[Any]], include_sentiment: bool,
                    domestic_label: str = "境内主流媒体") -> Any:
    return add_fixed_topic_table(document, rows, domestic_label)


def add_wechat_top_table(document: Any, rows: list[dict[str, Any]]) -> Any:
    table = document.add_table(rows=2, cols=5)
    table.cell(0, 0).merge(table.cell(0, 4)).text = "微信公号文章阅读量TOP"
    for idx, text in enumerate(["序号", "账号", "标题", "阅读量", "在看量"]):
        table.cell(1, idx).text = text
    mark_row_as_header(table.rows[0])
    mark_row_as_header(table.rows[1])
    for index, row in enumerate(rows, 1):
        cells = table.add_row().cells
        values = [index, row.get("source"), "", display_metric(row.get("spread_count"), row.get("spread_count_display")), row.get("comment_count")]
        for idx, value in enumerate(values):
            cells[idx].text = "" if value is None else str(value)
        title_paragraph = cells[2].paragraphs[0]
        add_hyperlink(title_paragraph, str(row.get("title") or ""), str(row.get("url") or ""), font_name="仿宋", size_pt=12)
    style_docx_table(table, "仿宋", 12, [35.2, 85.0, 191.3, 56.7, 54.8])
    for row in table.rows[:2]:
        for cell in row.cells:
            shade_cell(cell, "17365D")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_docx_run_font(run, "仿宋", 12, True)
                    run.font.color.rgb = RGBColor(255, 255, 255)
    return table


def display_metric(value: Any, display_value: Any = None) -> str:
    if display_value not in (None, ""):
        return str(display_value)
    return str(value if value not in (None, "") else 0)


def display_date(value: Any) -> str:
    text = str(value or "").strip()
    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:\s.*)?$", text)
    if not match:
        return text
    year, month, day = match.groups()
    return f"{year}/{int(month)}/{int(day)}"


def add_overseas_table(document: Any, rows: list[dict[str, Any]]) -> Any:
    table = document.add_table(rows=1, cols=4)
    for idx, header in enumerate(["序号", "来源", "报道日期", "超链接标题"]):
        table.rows[0].cells[idx].text = header
    mark_row_as_header(table.rows[0])
    for index, row in enumerate(rows, 1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = formal_overseas_source(row)
        cells[2].text = display_date(row.get("published_at"))
        add_hyperlink(
            cells[3].paragraphs[0],
            str(row.get("_formal_title_cn") or formal_overseas_title(row)),
            str(row.get("url") or ""),
            font_name="微软雅黑",
            size_pt=11,
        )
    style_docx_table(table, "微软雅黑", 11, [36.7, 120.5, 71.0, 203.9])
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.7
    return table


def add_paragraphs(document: Any, paragraphs: list[str]) -> None:
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)


def add_emphasized_details(paragraph: Any, details: str) -> None:
    pattern = re.compile(r"((?:微信公众号[“\"]?[^，”\"]+[”\"]?|[^，。；]{2,28}?)(?:报道称|发文称|文章称|评论称|认为|指出|建议|表示|强调|分析称|提到))")
    cursor = 0
    for match in pattern.finditer(details):
        if match.start() > cursor:
            paragraph.add_run(details[cursor:match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        cursor = match.end()
    if cursor < len(details):
        paragraph.add_run(details[cursor:])


def cluster_detail_paragraphs(cluster: dict[str, Any]) -> list[str]:
    """Layout only: keep complete attributed claims and all their conditions."""
    limit = writing_rules()['viewpoint']['paragraph_soft_max_chars']
    max_claims = writing_rules()['viewpoint']['paragraph_max_claims']
    evidence = [row for row in cluster.get('evidence') or [] if isinstance(row, dict)]
    if not evidence:
        text = clean_sentence(display_cluster_details(cluster))
        return [text + '。'] if text else []
    result, current, seen, count = [], '', set(), 0
    for row in evidence:
        sentence = evidence_sentence(row)
        signature = re.sub(r'[^0-9A-Za-z\u4e00-\u9fff]+', '', sentence).lower()
        if not signature or signature in seen:
            continue
        seen.add(signature)
        sentence = sentence.rstrip('。') + '。'
        if current and (len(current) + len(sentence) > limit or count >= max_claims):
            result.append(current); current, count = '', 0
        current += sentence; count += 1
    if current:
        result.append(current)
    return result


def add_detail_blocks(document: Any, paragraph: Any, cluster: dict[str, Any]) -> None:
    from cwh_writing_rules import formal_attribution
    entries = []
    verbs = '|'.join(map(re.escape, sorted(writing_rules()['viewpoint']['attribution_verbs'], key=len, reverse=True)))
    for row in cluster.get('evidence') or []:
        if not isinstance(row, dict):
            continue
        sentence = evidence_sentence(row).rstrip('。') + '。'
        subject = re.sub(r'\s+', '', formal_attribution(row))
        subjects = [subject] if subject else []
        if row.get('attribution_status') == 'named_person' and row.get('speaker_name'):
            name = re.sub(r'\s+', '', str(row['speaker_name']))
            roles = re.split(r'[、，,；;]', str(row.get('speaker_role') or ''))
            subjects.extend(re.sub(r'\s+', '', role) + name for role in roles if role)
        pattern = '|'.join(map(re.escape, sorted(set(subjects), key=len, reverse=True)))
        prefix = re.match(r'(?:' + pattern + r')(?:' + verbs + r')', sentence) if pattern else None
        entries.append((sentence, prefix.end() if prefix else 0))
    for index, details in enumerate(cluster_detail_paragraphs(cluster)):
        target = paragraph if index == 0 else document.add_paragraph()
        if not entries:
            # No structured speaker mapping: keep prose plain, do not guess
            # that a verb mentioned inside the claim is its attribution.
            target.add_run(details)
            continue
        remainder = details
        while remainder:
            matched = next(((sentence, prefix) for sentence, prefix in entries if remainder.startswith(sentence)), None)
            if not matched:
                target.add_run(remainder)
                break
            sentence, prefix = matched
            if prefix:
                target.add_run(sentence[:prefix]).bold = True
            target.add_run(sentence[prefix:])
            remainder = remainder[len(sentence):]


def add_cluster_paragraph(document: Any, prefix: str, cluster: dict[str, Any]) -> Any:
    paragraph = document.add_paragraph()
    summary = clean_sentence(cluster.get("summary"))
    lead = paragraph.add_run(f"{prefix}{summary}。")
    lead.bold = True
    details = clean_sentence(display_cluster_details(cluster))
    if details:
        add_detail_blocks(document, paragraph, cluster)
    else:
        evidence = evidence_text(cluster.get("evidence") or [])
        if evidence:
            paragraph.add_run(f"{evidence}。")
    if cluster_needs_attribution_review(cluster):
        run = paragraph.add_run(ATTRIBUTION_REVIEW_TEXT)
        run.bold = True
        run.font.color.rgb = RGBColor(192, 0, 0)
    return paragraph


def add_single_topic_paragraph(document: Any, index: int, item: dict[str, Any], cluster: dict[str, Any]) -> Any:
    paragraph = document.add_paragraph()
    lead = paragraph.add_run(f"{index}.{topic_heading(item)}。")
    lead.bold = True
    details = clean_sentence(display_cluster_details(cluster))
    if details:
        add_detail_blocks(document, paragraph, cluster)
    else:
        summary = clean_sentence(cluster.get("summary"))
        paragraph.add_run(f"{summary}。")
    if cluster_needs_attribution_review(cluster):
        run = paragraph.add_run(ATTRIBUTION_REVIEW_TEXT)
        run.bold = True
        run.font.color.rgb = RGBColor(192, 0, 0)
    return paragraph


def add_comment_group_paragraph(document: Any, prefix: str, name: str, rows: list[dict[str, Any]]) -> Any:
    paragraph = document.add_paragraph()
    lead = paragraph.add_run(f"{prefix}{name}。")
    lead.bold = True
    paragraph.add_run(comment_wording(rows) if rows else writing_rules()['comments']['missing_topic_notice'])
    return paragraph


def add_chart_caption(document: Any, text: str) -> Any:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run(text)
    set_docx_run_font(run, "微软雅黑", 16, True)
    return paragraph


def add_centered_picture(document: Any, image_path: str, width: Any, alt_text: str, max_height: Any = None) -> Any:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    section = document.sections[-1]
    width = min(width, section.page_width - section.left_margin - section.right_margin)
    shape = paragraph.add_run().add_picture(image_path, width=width)
    if max_height is not None and shape.height > max_height:
        shape.width = round(shape.width * max_height / shape.height)
        shape.height = max_height
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    return shape


def add_docx_title(document: Any, text: str) -> None:
    paragraph = document.add_paragraph(style='Title')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    set_docx_run_font(run, "华文中宋", 22, True)


def add_docx_heading(document: Any, text: str, level: int) -> Any:
    roles = {
        1: ("Heading 1", "黑体", 16, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 32.0),
        2: ("Heading 2", "楷体_GB2312", 16, True, WD_ALIGN_PARAGRAPH.LEFT, 32.15),
        3: ("Heading 3", "仿宋_GB2312", 16, True, WD_ALIGN_PARAGRAPH.JUSTIFY, 32.15),
    }
    style_name, font_name, size_pt, bold, alignment, first_indent = roles[level]
    paragraph = document.add_paragraph(style=style_name)
    paragraph.alignment = alignment
    paragraph.paragraph_format.first_line_indent = Pt(first_indent)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    set_docx_run_font(run, font_name, size_pt, bold)
    return paragraph


def template_path() -> Path:
    return TEMPLATE


def clear_document_body(document: Any) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def remove_nonprinting_pagination_controls(document: Any) -> None:
    """Allow body flow while retaining heading/table pagination semantics.

    Word's black paragraph squares are nonprinting UI marks, not document text.
    Removing keepNext to hide them strands real headings and table headers.
    """
    tags = ("w:keepLines", "w:pageBreakBefore")
    for root in (document._element, document.styles.element):
        for tag in tags:
            for element in list(root.iter(qn(tag))):
                parent = element.getparent()
                if parent is not None:
                    parent.remove(element)
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith('Heading'):
            paragraph.paragraph_format.keep_with_next = True
        else:
            # Body paragraphs and pictures should not create long keep chains.
            paragraph.paragraph_format.keep_with_next = False


def new_document() -> Any:
    from docx import Document

    path = template_path()
    document = Document(str(path))
    clear_document_body(document)
    return document


def section_override(data: dict[str, Any], key: str) -> list[dict[str, str]]:
    rows = (data.get("section_overrides") or {}).get(key)
    if not isinstance(rows, list):
        return []
    output = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        role = str(row.get("role") or "p")
        text = str(row.get("text") or "").strip()
        if key in {"two", "three"} and role == "p":
            text = clean_formal_comment(text)
        if text:
            output.append({"role": role, "text": text})
    return output


def add_override_block(document: Any, block: dict[str, str]) -> Any:
    role = block.get("role") or "p"
    text = str(block.get("text") or "").strip()
    if role == "h1":
        return add_docx_heading(document, text, 1)
    if role == "h2":
        return add_docx_heading(document, text, 2)
    if role == "h3":
        return add_docx_heading(document, text, 3)
    paragraph = document.add_paragraph()
    if re.match(r"^\s*\d+[.．、]\s*\S", text):
        paragraph.add_run(text).bold = True
        return paragraph
    lead = re.match(r"^((?:一|二|三|四|五|六|七|八|九|十)是.*?[。；：])", text)
    if lead:
        paragraph.add_run(lead.group(1)).bold = True
        remainder = text[lead.end():]
        if remainder:
            add_emphasized_details(paragraph, remainder)
        return paragraph
    add_emphasized_details(paragraph, text)
    return paragraph


def write_override_section(
    document: Any,
    data: dict[str, Any],
    key: str,
    blocks: list[dict[str, str]],
    docx_charts: dict[str, str],
) -> None:
    from docx.shared import Inches

    trend_inserted = False
    topic_inserted = False
    overseas_inserted = False
    wechat_inserted = False
    sentiment_visible = formal_sentiment_available(data)
    for block in blocks:
        text = str(block.get("text") or "")
        role = str(block.get("role") or "p")
        if key == "one" and role == "h2" and text.startswith("（二）") and not trend_inserted:
            chart = docx_charts.get("trend_distribution")
            if chart and Path(chart).exists():
                add_centered_picture(document, chart, width=Inches(5.8), alt_text="国务院常务会议舆情信息传播量逐日趋势图")
            trend_inserted = True
        add_override_block(document, block)
        if key == "one" and role == "h2" and text.startswith("（二）") and not topic_inserted:
            add_topic_table(document, formal_topic_rows(data, sentiment_visible), sentiment_visible, domestic_media_label(data))
            chart = docx_charts.get("topic_distribution")
            if chart and Path(chart).exists():
                add_centered_picture(document, chart, width=Inches(5.8), alt_text="各子议题信息传播总量对比图")
            topic_inserted = True
        if key == "four" and role == "h2" and text.startswith("（一）") and not overseas_inserted:
            add_overseas_table(document, formal_appendix_overseas_rows(data))
            overseas_inserted = True
        if key == "four" and role == "h2" and text.startswith("（二）") and not wechat_inserted:
            add_wechat_top_table(document, wechat_top_rows(data)[:10])
            wechat_inserted = True
    if key == "one" and not trend_inserted:
        chart = docx_charts.get("trend_distribution")
        if chart and Path(chart).exists():
            add_centered_picture(document, chart, width=Inches(5.8), alt_text="国务院常务会议舆情信息传播量逐日趋势图")
    if key == "one" and not topic_inserted:
        add_topic_table(document, formal_topic_rows(data, sentiment_visible), sentiment_visible, domestic_media_label(data))
        chart = docx_charts.get("topic_distribution")
        if chart and Path(chart).exists():
            add_centered_picture(document, chart, width=Inches(5.8), alt_text="各子议题信息传播总量对比图")
    if key == "two":
        chart = docx_charts.get("hotword_distribution")
        if chart and Path(chart).exists():
            add_centered_picture(document, chart, width=Inches(5.8), alt_text="境内舆情热词分布词云图", max_height=Inches(3.0))
    if key == "four" and not overseas_inserted:
        add_overseas_table(document, formal_appendix_overseas_rows(data))
    if key == "four" and not wechat_inserted:
        add_wechat_top_table(document, wechat_top_rows(data)[:10])


def write_docx(data: dict[str, Any], out_path: Path) -> None:
    from docx.shared import Inches

    document = new_document()
    document_rules = writing_rules()["document"]
    meeting = data["meeting"]
    month_day = month_day_text(meeting.get("date", ""))
    topics = meeting.get("topics") or []
    agenda_topics = joined_agenda_topics(topics) or meeting.get("agenda", "")
    add_docx_title(document, document_rules["title_template"].format(date=month_day))
    document.add_paragraph("")
    document.add_paragraph(opening_paragraph(meeting, month_day, agenda_topics))
    document.add_paragraph("")

    one_override = section_override(data, "one")
    two_override = section_override(data, "two")
    three_override = section_override(data, "three")
    four_override = section_override(data, "four")
    docx_charts = ensure_docx_chart_images(data, out_path.parent)
    sentiment_visible = formal_sentiment_available(data)
    if one_override:
        write_override_section(document, data, "one", one_override, docx_charts)
    else:
        add_docx_heading(document, document_rules["fixed_chapters"][0], 1)
        add_docx_heading(document, "（一）总事件传播情况", 2)
        add_paragraphs(document, total_event_paragraphs(data))
        chart = docx_charts.get("trend_distribution")
        if chart and Path(chart).exists():
            add_centered_picture(document, chart, width=Inches(5.8), alt_text="国务院常务会议舆情信息传播量逐日趋势图")
        add_docx_heading(document, "（二）子议题传播情况", 2)
        add_topic_table(document, formal_topic_rows(data, sentiment_visible), sentiment_visible, domestic_media_label(data))
        chart = docx_charts.get("topic_distribution")
        if chart and Path(chart).exists():
            add_centered_picture(document, chart, width=Inches(5.8), alt_text="各子议题信息传播总量对比图")
            document.add_paragraph("")

    if two_override:
        write_override_section(document, data, "two", two_override, docx_charts)
    else:
        add_docx_heading(document, document_rules["fixed_chapters"][1], 1)
        add_docx_heading(document, document_rules["domestic_subsections"][0], 2)
        for topic_idx, item in enumerate(report_topic_items(data), 1):
            clusters = item.get("clusters") or []
            if not clusters:
                add_docx_heading(document, f"{topic_idx}.{topic_heading(item)}", 3)
                document.add_paragraph((item.get("evidence_gap") or {}).get("notice") or "当前公开样本密度不足，尚未形成可稳定归纳的媒体自媒体观点。")
                continue
            if len(clusters) == 1:
                add_single_topic_paragraph(document, topic_idx, item, clusters[0])
                continue
            add_docx_heading(document, f"{topic_idx}.{topic_heading(item)}", 3)
            for idx, cluster in enumerate(clusters, 1):
                prefix = ordinal_prefix(idx)
                add_cluster_paragraph(document, prefix, cluster)

        add_docx_heading(document, document_rules["domestic_subsections"][1], 2)
        comments = data.get("comments", {}).get("selected") or []
        groups = report_comment_groups(data)
        if groups:
            document.add_paragraph(comment_lead(data, groups))
            for idx, (name, rows) in enumerate(groups, 1):
                prefix = ordinal_prefix(idx)
                add_comment_group_paragraph(document, prefix, name, rows)
        else:
            document.add_paragraph(empty_comment_notice(data))

        add_docx_heading(document, document_rules["domestic_subsections"][2], 2)
        document.add_paragraph(hotword_paragraph(data))
        chart = docx_charts.get("hotword_distribution")
        if chart and Path(chart).exists():
            add_centered_picture(document, chart, width=Inches(5.8), alt_text="境内舆情热词分布词云图", max_height=Inches(3.0))
            document.add_paragraph("")

    # Continue natural flow after the bounded-height cloud; forcing a fresh
    # overseas page can leave the preceding page with nothing but the cloud.
    if three_override:
        write_override_section(document, data, "three", three_override, docx_charts)
    else:
        add_docx_heading(document, document_rules["fixed_chapters"][2], 1)
        add_docx_heading(document, document_rules["overseas_subsections"][0], 2)
        add_paragraphs(document, overseas_media_body_paragraphs(data))
        add_docx_heading(document, document_rules["overseas_subsections"][1], 2)
        foreign_comment_paragraphs = overseas_comment_paragraphs(data)
        if foreign_comment_paragraphs:
            add_paragraphs(document, foreign_comment_paragraphs)
        else:
            document.add_paragraph(writing_rules()["overseas"]["no_comment_evidence_sentence"])

    # Let Word use the remaining space after a short overseas section. A forced
    # break here creates an almost-empty page whenever foreign comments are absent.
    # The appendix table can split normally and repeats its header row.
    if four_override:
        write_override_section(document, data, "four", four_override, docx_charts)
    else:
        add_docx_heading(document, document_rules["fixed_chapters"][3], 1)
        add_docx_heading(document, "（一）外媒报道列表（部分）", 2)
        add_overseas_table(document, formal_appendix_overseas_rows(data))
        add_docx_heading(document, "（二）传播量较大的公众号文章", 2)
        add_wechat_top_table(document, wechat_top_rows(data)[:10])

    end_paragraph = document.add_paragraph()
    end_paragraph.paragraph_format.first_line_indent = Pt(0)
    end_paragraph.paragraph_format.space_before = Pt(0)
    end_paragraph.paragraph_format.space_after = Pt(0)
    end_paragraph.paragraph_format.line_spacing = Pt(1)
    end_run = end_paragraph.add_run(" ")
    end_run.font.size = Pt(1)
    if not any((one_override, two_override, three_override, four_override)):
        data.setdefault('audit', {})['word_template'] = apply_template(document)
        data['audit']['word_template']['editorial_profile'] = {
            'id': editorial_profile()['profile_id'],
            'source_template_sha256': editorial_profile()['source_template_sha256'],
            'provenance': editorial_profile()['provenance'],
            'writing_rules_sha256': writing_rules_sha256(),
            'chair_origin': 'meeting_source' if meeting.get('chair_name') and meeting.get('chair_source') else 'editorial_default',
            'does_not_certify_sentiment_or_collection': True,
        }
    else:
        # Explicit user edits are a separate route, never silently labelled template-filled.
        data.setdefault('audit', {})['word_template'] = {'mode': 'explicit_section_override'}
    remove_nonprinting_pagination_controls(document)
    document.save(out_path)
def empty_comment_notice(data: dict[str, Any]) -> str:
    """Missing quote eligibility does not mean no comments were captured."""
    rows = (data.get('comments') or {}).get('selected') or []
    rules = writing_rules()['comments']
    if any(is_actual_comment_row(row) and row.get('in_monitoring_window') is not False for row in rows):
        return rules['no_ready_quotes_verified_samples']
    return rules['no_ready_quotes_unverified_candidates' if rows else 'no_comment_samples']


def audit_formal_comment_quotes(data: dict[str, Any], document_xml: bytes) -> dict[str, Any]:
    """Check actual domestic Word prose, not JSON approvals or appendix text."""
    approved = [row for row in (data.get('comments', {}).get('selected') or [])
                if row.get('ai_formal_include') is True and is_actual_comment_row(row)]
    if section_override(data, 'two'):
        return {'applicable': False, 'reason': 'explicit_domestic_section_override',
                'approved_row_count': len(approved)}
    texts = [clean_formal_comment(row.get('content') or row.get('title')) for row in approved]
    expected = list(dict.fromkeys(text for text in texts if text))
    groups = comment_groups(data.get('comments', {}).get('selected') or [])
    grouped = {clean_formal_comment(row.get('content') or row.get('title')) for _, rows in groups for row in rows}
    cap = writing_rules()['comments']['quotes_per_ready_topic'][1]
    displayed = {clean_formal_comment(row.get('content') or row.get('title'))
                 for _, rows in groups for row in rows[:cap]}
    reserved = [text for text in expected if text in grouped and text not in displayed]
    start, end = writing_rules()['document']['domestic_subsections'][1:3]
    root = ET.fromstring(document_xml)
    paragraphs = [''.join(node.text or '' for node in paragraph.iter(qn('w:t')))
                  for paragraph in root.iter(qn('w:p'))]
    content, active, found = [], False, False
    for paragraph in paragraphs:
        if paragraph.strip() == start:
            active, found = True, True
            continue
        if active and paragraph.strip() == end:
            break
        if active:
            content.append(paragraph)
    body = '\n'.join(content)
    missing_grouped = [text for text in expected if text not in grouped]
    missing_word = [text for text in expected if text not in reserved and text not in body]
    empty_approved = sum(not text for text in texts)
    return {'applicable': True,
            'passed': not missing_grouped and not missing_word and not empty_approved,
            'approved_row_count': len(approved), 'unique_approved_quote_count': len(expected),
            'grouped_unique_quote_count': len(grouped), 'domestic_section_found': found,
            'missing_from_grouping': missing_grouped, 'missing_from_word': missing_word,
            'not_displayed_due_to_template_cap': reserved,
            'empty_approved_quote_count': empty_approved,
            'scope': 'Actual Word domestic prose coverage only; not semantic or physical-pagination approval'}


def audit_formal_docx(data: dict[str, Any], docx_path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(docx_path) as archive:
        names = archive.namelist()
        document_xml = archive.read("word/document.xml")
        embedded_hashes = {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in names
            if name.startswith("word/media/")
        }
    chart_matches: dict[str, bool] = {}
    for key in ["trend_distribution", "topic_distribution", "hotword_distribution"]:
        source = Path(str(((data.get("artifacts") or {}).get("docx_charts") or {}).get(key) or ""))
        chart_matches[key] = bool(
            source.is_file()
            and hashlib.sha256(source.read_bytes()).hexdigest() in embedded_hashes
        )
    table_count = document_xml.count(b"<w:tbl>")
    hyperlink_count = document_xml.count(b"<w:hyperlink")
    image_count = sum(1 for name in names if name.startswith("word/media/"))
    comment_quote_audit = audit_formal_comment_quotes(data, document_xml)
    artifacts = data.get('artifacts') or {}
    chart_paths = artifacts.get('docx_charts') or {}
    monitoring_assets = (artifacts.get('chart_authority') in {
        'monitoring_system_embedded_assets', 'monitoring_system_assets_plus_pipeline_wordcloud'}
        and all(chart_matches.values()))
    topic_values = [(str(row.get('display') or topic_display(row.get('topic', ''))),
                     int(row.get('spread_count', 0) or row.get('total_samples', 0) or 0))
                    for row in data.get('topic_stats') or []]
    fixed_topic_fallback = (
        artifacts.get('chart_authority') == 'mixed_system_and_program_fallback'
        and all(chart_matches.values())
        and Path(chart_paths.get('trend_distribution') or '').stem == 'trend_distribution_system'
        and Path(chart_paths.get('hotword_distribution') or '').stem in {
            'hotword_distribution_system', 'hotword_distribution_pipeline'}
        and verify_topic_chart_manifest(chart_paths.get('topic_distribution') or '',
                                        (data.get('audit') or {}).get('topic_chart_style'), topic_values))
    fixed_report_visuals = (
        artifacts.get('chart_authority') == 'baseline_template_plus_reviewed_wordcloud'
        and all(chart_matches.values())
        and verify_report_charts(data, chart_paths, (data.get('audit') or {}).get('report_visual_templates') or {}))
    checks = {
        "no_unfilled_template_slots": b'CWH_' not in document_xml,
        "fixed_template_applied": (bool((data.get('audit') or {}).get('word_template', {}).get('filled_slots'))
                                   or any(section_override(data, key) for key in ('one', 'two', 'three', 'four'))),
        "uses_traceable_fixed_chart_assets": fixed_report_visuals,
        "fixed_topic_table_layout": verify_topic_table(document_xml),
        "three_required_tables": table_count == 3,
        "three_required_images": image_count >= 3,
        "appendix_hyperlinks_present": hyperlink_count >= 10,
        "approved_domestic_comment_quotes_present": (not comment_quote_audit['applicable']
                                                     or comment_quote_audit['passed']),
    }
    return {
        "passed": all(checks.values()),
        "checks": checks,
        "table_count": table_count,
        "image_count": image_count,
        "hyperlink_count": hyperlink_count,
        "system_asset_hash_matches": chart_matches,
        "chart_provenance": {"monitoring_system_assets": monitoring_assets,
                             "verified_fixed_topic_fallback": fixed_topic_fallback,
                             "verified_baseline_template": fixed_report_visuals,
                             "declared_authority": artifacts.get('chart_authority')},
        "domestic_comment_quotes": comment_quote_audit,
    }


def formalize_report(data: dict[str, Any], out_dir: Path) -> dict[str, str]:
    out_dir.mkdir(parents=True, exist_ok=True)
    apply_raw_public_top_audit(data)
    enrich_viewpoint_titles(data)
    data.setdefault("audit", {})["writing_rules_sha256"] = writing_rules_sha256()
    formal_md = out_dir / "cwh_formal_report.md"
    formal_docx = out_dir / "cwh_formal_report.docx"
    formal_md.write_text(render_formal_markdown(data, out_dir), encoding="utf-8")
    write_docx(data, formal_docx)
    docx_audit = audit_formal_docx(data, formal_docx)
    docx_audit_path = out_dir / "cwh_formal_docx_audit.json"
    docx_audit_path.write_text(json.dumps(docx_audit, ensure_ascii=False, indent=2), encoding="utf-8")
    data.setdefault("audit", {})["formal_docx_validation"] = docx_audit
    if not docx_audit["passed"]:
        acceptance = data.setdefault("audit", {}).setdefault("acceptance", {})
        acceptance["ready_for_formal_delivery"] = False
        acceptance.setdefault("blockers", []).append("Word成品未通过评论呈现、图表、表格或超链接结构校验。")
    data.setdefault("artifacts", {})["formal_report"] = str(formal_md)
    data.setdefault("artifacts", {})["formal_docx"] = str(formal_docx)
    data.setdefault("artifacts", {})["formal_docx_audit"] = str(docx_audit_path)
    return {
        "formal_report": str(formal_md),
        "formal_docx": str(formal_docx),
        "formal_docx_audit": str(docx_audit_path),
    }


def prepare_isolated_replay_artifacts(
    data: dict[str, Any], data_path: Path, out_dir: Path
) -> None:
    """Copy authoritative chart inputs for an isolated replay without touching the source run."""
    if out_dir.resolve() == data_path.parent.resolve():
        return
    artifacts = data.setdefault("artifacts", {})
    docx_charts = artifacts.get("docx_charts") or {}
    if not docx_charts:
        return
    chart_dir = out_dir / "charts"
    chart_dir.mkdir(parents=True, exist_ok=True)
    copied: dict[str, str] = {}
    for key, raw_path in docx_charts.items():
        source = Path(str(raw_path or ""))
        if key == "hotword_distribution" and (
            not source.exists() or not source.is_file() or source.stat().st_size <= 0
        ):
            for candidate_value in [
                artifacts.get("wordcloud_image"),
                data_path.parent / "charts" / "hotword_distribution_pipeline.png",
                data_path.parent / "charts" / "hotword_distribution_system.png",
            ]:
                candidate = Path(str(candidate_value or ""))
                if candidate.is_file() and candidate.stat().st_size > 0:
                    source = candidate
                    break
        if not source.exists() or not source.is_file() or source.stat().st_size <= 0:
            raise ValueError(f"隔离重放缺少非零权威图表：{key} -> {source}")
        target = chart_dir / source.name
        shutil.copy2(source, target)
        if target.stat().st_size != source.stat().st_size or target.stat().st_size <= 0:
            raise ValueError(f"隔离重放图表复制失败：{key} -> {target}")
        if hashlib.sha256(target.read_bytes()).digest() != hashlib.sha256(source.read_bytes()).digest():
            raise ValueError(f"隔离重放图表哈希不一致：{key} -> {target}")
        copied[key] = str(target.resolve())
    artifacts["docx_charts"] = copied
    if copied.get("hotword_distribution"):
        artifacts["wordcloud_image"] = copied["hotword_distribution"]


def report_data_output_path(
    data_path: Path, out_dir: Path, *, write_back: bool
) -> Path:
    if write_back or out_dir.resolve() == data_path.parent.resolve():
        return data_path
    return out_dir / "report_data.json"


def main() -> None:
    parser = argparse.ArgumentParser(description="Render formal CWH Markdown and Word report from report_data.json.")
    parser.add_argument("report_data", help="Path to report_data.json")
    parser.add_argument("--out-dir", default="", help="Output directory. Defaults beside report_data.json.")
    parser.add_argument(
        "--wordcloud-image",
        default="",
        help="Optional reviewed word-cloud PNG that overrides only the report/dashboard word cloud.",
    )
    parser.add_argument(
        "--write-back",
        action="store_true",
        help="Explicitly write updated artifact paths and audit data back to the source report_data.json.",
    )
    parser.add_argument(
        "--hotword-audit",
        default="",
        help="Optional reviewed hotword_audit.json used to synchronize report prose and dashboard terms with the word cloud.",
    )
    args = parser.parse_args()

    data_path = Path(args.report_data)
    data = json.loads(data_path.read_text(encoding="utf-8"))
    if args.hotword_audit:
        apply_hotword_audit(data, Path(args.hotword_audit))
    if args.wordcloud_image:
        data.setdefault("artifacts", {})["wordcloud_image"] = str(Path(args.wordcloud_image).resolve())
    out_dir = Path(args.out_dir) if args.out_dir else data_path.parent
    prepare_isolated_replay_artifacts(data, data_path, out_dir)
    artifacts = formalize_report(data, out_dir)
    output_data_path = report_data_output_path(data_path, out_dir, write_back=args.write_back)
    data.setdefault("artifacts", {})["report_data"] = str(output_data_path.resolve())
    output_data_path.parent.mkdir(parents=True, exist_ok=True)
    output_data_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {"status": "ok", "report_data": str(output_data_path), **artifacts},
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
