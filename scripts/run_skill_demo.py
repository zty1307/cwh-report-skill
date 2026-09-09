from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path


BAD_REPORT_MARKERS = ["需补充数据", "TODO", "placeholder"]
REQUIRED_REPORT_MARKERS = [
    "整体传播情况",
    "子议题传播情况",
    "境内媒体及自媒体观点",
    "网民评论情况",
    "候选热词",
    "境外关注情况",
    "人工审核清单",
]
REQUIRED_WORKBOOK_SHEETS = [
    "关键词",
    "总事件",
    "子事件数据汇总",
    "外媒报道列表",
    "词云",
    "公众TOP",
]


def run_orchestrator(args: argparse.Namespace) -> subprocess.CompletedProcess[str]:
    command = [
        sys.executable,
        str(Path(__file__).with_name("cwh_orchestrator.py")),
        "--out-dir",
        args.out_dir,
    ]
    if args.input_file:
        command.extend(["--input-file", args.input_file])
    if args.input:
        command.extend(["--input", args.input])
    if args.samples:
        command.extend(["--samples", args.samples])
    if args.system_workbook:
        command.extend(["--system-workbook", args.system_workbook])
    if args.system_details:
        command.extend(["--system-details", args.system_details])
    if args.analysis_bundle:
        command.extend(["--analysis-bundle", args.analysis_bundle])
    if args.data_mode:
        command.extend(["--data-mode", args.data_mode])
    if args.supplemental_collection:
        command.append("--supplemental-collection")
    if args.sentiment_results:
        command.extend(["--sentiment-results", args.sentiment_results])
    for path in args.comment_handoff:
        command.extend(["--comment-handoff", path])
    if args.sentiment_summary:
        command.extend(["--sentiment-summary", args.sentiment_summary])
    if args.trend_chart_image:
        command.extend(["--trend-chart-image", args.trend_chart_image])
    if args.topic_chart_image:
        command.extend(["--topic-chart-image", args.topic_chart_image])
    if args.no_tasks:
        command.append("--no-tasks")
    if args.platforms:
        command.extend(["--platforms", args.platforms])
    if args.collection_profile:
        command.extend(["--collection-profile", args.collection_profile])
    if args.max_posts:
        command.extend(["--max-posts", str(args.max_posts)])
    if args.max_comments:
        command.extend(["--max-comments", str(args.max_comments)])
    command.extend(["--target-samples", str(args.target_samples)])
    command.extend(["--target-comments", str(args.target_comments)])
    command.extend(["--target-comments-per-topic", str(args.target_comments_per_topic)])
    command.extend(["--target-overseas", str(args.target_overseas)])
    command.extend(["--max-low-quality-ratio", str(args.max_low_quality_ratio)])
    command.extend(["--max-collection-rounds", str(args.max_collection_rounds)])
    if args.no_agent_reach_plan:
        command.append("--no-agent-reach-plan")
    if args.agent_reach_platforms:
        command.extend(["--agent-reach-platforms", args.agent_reach_platforms])
    if args.run_agent_reach:
        command.append("--run-agent-reach")
    command.extend(["--agent-reach-limit", str(args.agent_reach_limit)])
    if args.opencli_profile:
        command.extend(["--opencli-profile", args.opencli_profile])
    if args.run_mediaspider:
        command.append("--run-mediaspider")
    if args.mediaspider_supervisor:
        command.extend(["--mediaspider-supervisor", args.mediaspider_supervisor])
    if args.media_home:
        command.extend(["--media-home", args.media_home])
    command.extend(["--mediaspider-limit", str(args.mediaspider_limit)])
    if args.mediaspider_task_timeout_sec > 0:
        command.extend(["--mediaspider-task-timeout-sec", str(args.mediaspider_task_timeout_sec)])
    if args.keep_proxy:
        command.append("--keep-proxy")
    if args.no_auto_web:
        command.append("--no-auto-web")
    if args.web_always:
        command.append("--web-always")
    command.extend(["--web-limit", str(args.web_limit)])
    if args.monitor_start:
        command.extend(["--monitor-start", args.monitor_start])
    if args.monitor_end:
        command.extend(["--monitor-end", args.monitor_end])
    return subprocess.run(
        command,
        input=sys.stdin.read() if not args.input and not args.input_file and not sys.stdin.isatty() else None,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def main() -> int:
    parser = argparse.ArgumentParser(description="Run and validate the CWH report skill demo.")
    parser.add_argument("--input", default="", help="Meeting agenda text.")
    parser.add_argument("--input-file", default="", help="UTF-8 meeting agenda file.")
    parser.add_argument("--samples", default="", help="Optional normalized/sample CSV or JSON.")
    parser.add_argument("--system-workbook", default="", help="Authoritative monitoring-system CWH workbook (.xlsx).")
    parser.add_argument("--system-details", default="", help="Optional system-exported media/comment detail CSV or JSON.")
    parser.add_argument("--analysis-bundle", default="", help="Internal AI/research evidence bundle JSON.")
    parser.add_argument("--data-mode", choices=["auto", "system", "crawler"], default="auto")
    parser.add_argument("--supplemental-collection", action="store_true", help="Allow optional crawler/search supplements while preserving system totals.")
    parser.add_argument("--sentiment-results", default="", help="Sentiment result CSV/JSON keyed by sample_id.")
    parser.add_argument("--comment-handoff", action="append", default=[], help="Strict report_comment_handoff.csv; may be repeated.")
    parser.add_argument("--sentiment-summary", default="", help="sentiment_workbook_summary.json denominator gate.")
    parser.add_argument("--trend-chart-image", default="", help="Optional reviewed trend chart PNG.")
    parser.add_argument("--topic-chart-image", default="", help="Optional reviewed subtopic chart PNG.")
    parser.add_argument("--out-dir", default="outputs/cwh_skill_demo_check", help="Output directory.")
    parser.add_argument("--platforms", default="wb,dy,ks,bili,zhihu", help="Domestic crawler task platforms; xhs, tieba and foreign platforms are excluded by default.")
    parser.add_argument("--collection-profile", choices=["demo", "formal", "deep"], default="demo", help="MediaSpider collection scale preset.")
    parser.add_argument("--max-posts", type=int, default=0, help="Override max posts per MediaSpider task.")
    parser.add_argument("--max-comments", type=int, default=0, help="Override max comments per post.")
    parser.add_argument("--target-samples", type=int, default=5000, help="Formal-delivery target for total usable samples.")
    parser.add_argument("--target-comments", type=int, default=200, help="Formal-delivery target for verifiable netizen comments.")
    parser.add_argument("--target-comments-per-topic", type=int, default=30, help="Formal-delivery target for comments per subtopic.")
    parser.add_argument("--target-overseas", type=int, default=10, help="Formal-delivery target for overseas media samples.")
    parser.add_argument("--max-low-quality-ratio", type=float, default=0.5, help="Maximum accepted low-quality sample ratio.")
    parser.add_argument("--max-collection-rounds", type=int, default=3, help="Maximum recommended collection expansion rounds.")
    parser.add_argument("--no-agent-reach-plan", action="store_true", help="Skip agent-reach lead discovery task generation.")
    parser.add_argument("--agent-reach-platforms", default="bili,v2ex", help="Comma list of domestic agent-reach lead platforms; xhs is excluded by default.")
    parser.add_argument("--run-agent-reach", action="store_true", help="Actually run agent-reach lead tasks and ingest results.")
    parser.add_argument("--agent-reach-limit", type=int, default=1, help="Maximum agent-reach tasks to run; 0 means all matched tasks.")
    parser.add_argument("--opencli-profile", default="", help="Optional OpenCLI Browser Bridge profile for agent-reach/OpenCLI commands.")
    parser.add_argument("--run-mediaspider", action="store_true", help="Actually run MediaSpider and ingest results before validating.")
    parser.add_argument("--mediaspider-supervisor", default="D:/Codex/2026-06-15/ai-1-1-https-drive-weixin/work/mediaspider-supervisor", help="mediaspider-supervisor directory.")
    parser.add_argument("--media-home", default="", help="Optional MediaSpider engine home path.")
    parser.add_argument("--mediaspider-limit", type=int, default=1, help="Maximum MediaSpider tasks to run; 0 means all matched tasks.")
    parser.add_argument("--mediaspider-task-timeout-sec", type=int, default=0, help="Per MediaSpider task timeout in seconds; 0 means no timeout.")
    parser.add_argument("--keep-proxy", action="store_true", help="Keep inherited HTTP/HTTPS proxy variables for MediaSpider.")
    parser.add_argument("--no-tasks", action="store_true", help="Skip MediaSpider task package generation.")
    parser.add_argument("--no-auto-web", action="store_true", help="Disable public web/news fallback collection.")
    parser.add_argument("--web-always", action="store_true", help="Run public web/news collection even when samples are provided.")
    parser.add_argument("--web-limit", type=int, default=5, help="Public web/news samples per topic.")
    parser.add_argument("--monitor-start", default="", help="Monitoring start date YYYY-MM-DD.")
    parser.add_argument("--monitor-end", default="", help="Monitoring end date YYYY-MM-DD.")
    args = parser.parse_args()

    proc = run_orchestrator(args)
    out_dir = Path(args.out_dir)
    report_path = out_dir / "cwh_report.md"
    data_path = out_dir / "report_data.json"
    audit_path = out_dir / "cwh_audit.json"
    formal_md_path = out_dir / "cwh_formal_report.md"
    formal_docx_path = out_dir / "cwh_formal_report.docx"
    workbook_path = out_dir / "cwh_data_workbook.xlsx"

    checks: list[dict[str, object]] = []
    checks.append({"name": "orchestrator_exit_code", "ok": proc.returncode == 0, "detail": proc.returncode})
    checks.append({"name": "report_exists", "ok": report_path.exists(), "detail": str(report_path)})
    checks.append({"name": "report_data_exists", "ok": data_path.exists(), "detail": str(data_path)})
    checks.append({"name": "audit_exists", "ok": audit_path.exists(), "detail": str(audit_path)})
    checks.append({"name": "formal_report_exists", "ok": formal_md_path.exists(), "detail": str(formal_md_path)})
    checks.append({"name": "formal_docx_exists", "ok": formal_docx_path.exists() and formal_docx_path.stat().st_size > 0, "detail": str(formal_docx_path)})
    checks.append({"name": "data_workbook_exists", "ok": workbook_path.exists() and workbook_path.stat().st_size > 0, "detail": str(workbook_path)})

    report_text = report_path.read_text(encoding="utf-8") if report_path.exists() else ""
    bad_markers = [marker for marker in BAD_REPORT_MARKERS if marker in report_text]
    checks.append({"name": "no_hollow_placeholders", "ok": not bad_markers, "detail": bad_markers})
    missing_sections = [marker for marker in REQUIRED_REPORT_MARKERS if marker not in report_text]
    checks.append({"name": "required_report_sections", "ok": not missing_sections, "detail": missing_sections})

    data = load_json(data_path) if data_path.exists() else {}
    audit = load_json(audit_path) if audit_path.exists() else {}
    topic_count = len(data.get("meeting", {}).get("topics", []))
    sample_count = data.get("statistics", {}).get("total_samples", 0)
    task_count = (data.get("collection", {}).get("mediacrawler_tasks") or {}).get("task_count", 0)
    mediaspider_count = data.get("collection", {}).get("mediaspider_sample_count", 0)
    agent_reach_count = data.get("collection", {}).get("agent_reach_sample_count", 0)
    by_topic = data.get("statistics", {}).get("by_topic", {})
    uncovered_topics = [topic for topic in data.get("meeting", {}).get("topics", []) if by_topic.get(topic, 0) == 0]
    checks.append({"name": "topics_detected", "ok": topic_count > 0, "detail": topic_count})
    checks.append({"name": "sample_count_recorded", "ok": isinstance(sample_count, int), "detail": sample_count})
    system_mode = data.get("collection", {}).get("data_mode") == "system"
    checks.append({"name": "task_package_generated", "ok": system_mode or args.no_tasks or task_count > 0, "detail": task_count})
    checks.append({"name": "system_data_loaded", "ok": not args.system_workbook or bool(data.get("system_data")), "detail": data.get("collection", {}).get("system_workbook", "")})
    checks.append({"name": "topic_sample_coverage", "ok": sample_count == 0 or len(uncovered_topics) < topic_count, "detail": uncovered_topics})
    workbook_detail: dict[str, object] = {"path": str(workbook_path), "sheets": [], "missing": []}
    workbook_ok = False
    if workbook_path.exists():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(workbook_path, read_only=True, data_only=False)
            sheets = wb.sheetnames
            required = REQUIRED_WORKBOOK_SHEETS + [f"子事件{i}" for i in range(1, topic_count + 1)]
            missing = [name for name in required if name not in sheets]
            workbook_detail = {"path": str(workbook_path), "sheets": sheets, "missing": missing}
            workbook_ok = not missing
            wb.close()
        except Exception as exc:
            workbook_detail = {"path": str(workbook_path), "error": f"{type(exc).__name__}: {exc}"}
    checks.append({"name": "data_workbook_required_sheets", "ok": workbook_ok, "detail": workbook_detail})
    quality_summary = audit.get("quality_summary", {})
    acceptance = audit.get("acceptance", {})
    expansion_plan = audit.get("expansion_plan", {})
    blockers = acceptance.get("blockers") or []
    checks.append({"name": "formal_acceptance_blockers", "ok": not blockers, "detail": blockers})
    checks.append({"name": "low_quality_ratio_under_50pct", "ok": float(quality_summary.get("low_quality_ratio") or 0) < 0.5 or sample_count == 0, "detail": quality_summary})
    checks.append({"name": "formal_outputs_registered", "ok": bool(data.get("artifacts", {}).get("formal_report")) and bool(data.get("artifacts", {}).get("formal_docx")), "detail": data.get("artifacts", {})})
    checks.append({"name": "expansion_plan_present", "ok": bool(expansion_plan), "detail": expansion_plan})

    failed = [check for check in checks if not check["ok"]]
    payload = {
        "status": "ok" if not failed else "failed",
        "out_dir": str(out_dir),
        "report": str(report_path),
        "report_data": str(data_path),
        "audit": str(audit_path),
        "formal_report": str(formal_md_path),
        "formal_docx": str(formal_docx_path),
        "data_workbook": str(workbook_path),
        "topics": topic_count,
        "samples": sample_count,
        "mediaspider_samples": mediaspider_count,
        "agent_reach_samples": agent_reach_count,
        "mediaspider_tasks": task_count,
        "acceptance": acceptance,
        "expansion_plan": expansion_plan,
        "quality_summary": quality_summary,
        "checks": checks,
        "stdout_tail": proc.stdout[-2000:],
        "stderr_tail": proc.stderr[-2000:],
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
