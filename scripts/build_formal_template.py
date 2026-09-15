from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from cwh_writing_rules import writing_rules


def set_east_asian_font(style, font_name: str, size_pt: int, bold: bool = False) -> None:
    style.font.name = font_name
    r_fonts = style._element.rPr.rFonts
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    for attr in ["eastAsiaTheme", "asciiTheme", "hAnsiTheme", "cstheme"]:
        r_fonts.attrib.pop(qn(f"w:{attr}"), None)
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


def set_paragraph_format(style, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_chars: int = 2) -> None:
    pf = style.paragraph_format
    pf.alignment = alignment
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1
    if first_line_chars:
        pf.first_line_indent = Pt(16 * first_line_chars)
    else:
        pf.first_line_indent = Pt(0)


def remove_style_borders(style) -> None:
    p_pr = style._element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def remove_nonprinting_pagination_controls(document) -> None:
    """Keep the template free of pagination marks shown as black squares in Word."""
    tags = ("w:keepNext", "w:keepLines", "w:pageBreakBefore")
    for root in (document._element, document.styles.element):
        for tag in tags:
            for element in list(root.iter(qn(tag))):
                parent = element.getparent()
                if parent is not None:
                    parent.remove(element)


def add_visible_title(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    set_run_font(run, "华文中宋", 22, True)


def set_run_font(run, font_name: str, size_pt: int, bold: bool = False) -> None:
    run.font.name = font_name
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    for attr in ["eastAsiaTheme", "asciiTheme", "hAnsiTheme", "cstheme"]:
        r_fonts.attrib.pop(qn(f"w:{attr}"), None)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_role_paragraph(document, text: str, role: str):
    roles = {
        "h1": ("Heading 1", "黑体", 16, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 32.0),
        "h2": ("Heading 2", "楷体_GB2312", 16, True, WD_ALIGN_PARAGRAPH.LEFT, 32.15),
        "h3": ("Heading 3", "仿宋_GB2312", 16, True, WD_ALIGN_PARAGRAPH.JUSTIFY, 32.15),
    }
    style_name, font_name, size_pt, bold, alignment, first_indent = roles[role]
    paragraph = document.add_paragraph(style=style_name)
    paragraph.alignment = alignment
    paragraph.paragraph_format.first_line_indent = Pt(first_indent)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run(text)
    set_run_font(run, font_name, size_pt, bold)
    return paragraph


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def set_doc_grid(section) -> None:
    sect_pr = section._sectPr
    doc_grid = sect_pr.find(qn("w:docGrid"))
    if doc_grid is None:
        doc_grid = OxmlElement("w:docGrid")
        sect_pr.append(doc_grid)
    doc_grid.set(qn("w:type"), "lines")
    doc_grid.set(qn("w:linePitch"), "312")
    doc_grid.set(qn("w:charSpace"), "0")


def style_table(table, font_name: str, size_pt: int, widths_pt: list[float]) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is not None:
        tbl_w.set(qn("w:w"), str(round(sum(widths_pt) * 20)))
        tbl_w.set(qn("w:type"), "dxa")
    for grid_col, width_pt in zip(table._tbl.tblGrid.gridCol_lst, widths_pt):
        grid_col.set(qn("w:w"), str(round(width_pt * 20)))
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        is_header = tr_pr.find(qn("w:tblHeader")) is not None
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, font_name, size_pt, is_header)


def mark_row_as_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def build_template(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)
    set_doc_grid(section)
    add_page_number(section)

    styles = doc.styles
    set_east_asian_font(styles["Normal"], "仿宋_GB2312", 16)
    set_paragraph_format(styles["Normal"])

    set_east_asian_font(styles["Title"], "华文中宋", 22)
    set_paragraph_format(styles["Title"], WD_ALIGN_PARAGRAPH.CENTER, 0)
    remove_style_borders(styles["Title"])

    set_east_asian_font(styles["Heading 1"], "黑体", 16)
    set_paragraph_format(styles["Heading 1"], WD_ALIGN_PARAGRAPH.JUSTIFY, 2)

    set_east_asian_font(styles["Heading 2"], "楷体_GB2312", 16, True)
    set_paragraph_format(styles["Heading 2"], WD_ALIGN_PARAGRAPH.JUSTIFY, 2)

    set_east_asian_font(styles["Heading 3"], "仿宋_GB2312", 16, True)
    set_paragraph_format(styles["Heading 3"], WD_ALIGN_PARAGRAPH.JUSTIFY, 2)

    doc.core_properties.title = "CWH正式舆情综述模板"
    doc.core_properties.subject = "用于cwh-report-skill渲染正式报告的可见骨架和样式模板"

    add_visible_title(doc, "【X月X日】国务院常务会议舆情综述")
    doc.add_paragraph("国务院总理李强【X月X日】主持召开国务院常务会议，【议题一】，【议题二】，【议题三】。本次国务院常务会议舆情传播情况如下：")

    add_role_paragraph(doc, "一、舆情传播情况", "h1")
    add_role_paragraph(doc, "（一）总事件传播情况", "h2")
    doc.add_paragraph("本次常务会引发境内外媒体广泛报道，境内外传播总量约【总量】条，舆论关注“李强主持召开国务院常务会议 【首个议题】等”相关话题。经发酵，舆情热度于【峰值日期】达到峰值。")
    doc.add_paragraph("境内主流媒体如人民网、新华网、央视网等均在显著位置刊文，共有相关报道【境内主流媒体量】条。")
    doc.add_paragraph("新媒体中，微信公众平台相关信息【微信量】条、微博相关信息【微博量】条，视频号相关信息【视频号量】条，新闻客户端、论坛等渠道共有相关信息【其他新媒体量】条。")
    doc.add_paragraph("境外媒体如【代表性境外媒体一】、【代表性境外媒体二】、【代表性境外媒体三】等予以关注，共有相关报道（含转载）【境外量】条。具体主流报道情况见附表。")
    doc.add_paragraph("【此处插入总事件传播走势图；横轴为监测日期，纵轴为传播量，标注峰值日。】")
    add_role_paragraph(doc, "（二）子议题传播情况", "h2")
    table = doc.add_table(rows=3, cols=9)
    for idx, text in enumerate(["序号", "标题", "境内主流媒体", "新媒体", "境外媒体", "总量"]):
        table.cell(0, idx).merge(table.cell(1, idx)).text = text
    table.cell(0, 6).merge(table.cell(0, 8)).text = "网民情感"
    for idx, text in enumerate(["正面", "中立", "负面"], 6):
        table.cell(1, idx).text = text
    for idx, text in enumerate(["1", "【子议题标题】", "【数值】", "【数值】", "【数值】", "【数值】", "【%】", "【%】", "【%】"]):
        table.cell(2, idx).text = text
    mark_row_as_header(table.rows[0])
    mark_row_as_header(table.rows[1])
    style_table(
        table,
        "微软雅黑",
        11,
        [35.5, 77.9, 47.2, 52.0, 42.5, 49.7, 42.5, 40.7, 32.8],
    )
    doc.add_paragraph("【此处插入子议题传播量对比图；按总量排序，保留全部子议题。】")

    add_role_paragraph(doc, "二、境内舆论情况", "h1")
    add_role_paragraph(doc, "（一）境内媒体自媒体情况", "h2")
    add_role_paragraph(doc, "1.【建议/认可/肯定/认为/期待类观点组标题】", "h3")
    doc.add_paragraph("【单一稳定观点簇写法】某专家称，【观点内容】。某媒体报道称，【观点内容】。微信公众号“【账号】”称，【观点内容】。")
    doc.add_paragraph("【多个稳定观点簇写法】一是【观点概括】。某专家称，【观点内容】。某媒体报道称，【观点内容】。微信公众号“【账号】”称，【观点内容】。")
    doc.add_paragraph("二是【观点概括】。某机构认为，【观点内容】。")
    doc.add_paragraph("三是【观点概括】。某媒体援引业内人士观点称，【观点内容】。")
    doc.add_paragraph("四是【观点概括】。微信公众号“【账号】”称，【观点内容】。")
    doc.add_paragraph("五是【观点概括】。【专家/媒体/机构】建议，【观点内容】。")
    add_role_paragraph(doc, "2.【下一子议题的观点组标题；按识别出的子议题动态扩展】", "h3")
    add_role_paragraph(doc, "（二）网民评论情况", "h2")
    doc.add_paragraph("网民有关本次国务院常务会议讨论的情感属性以积极正面和客观中立为主，主要有【评论关注点】等。主要评论如下：")
    doc.add_paragraph("一是【评论类别】。网民称，“【评论原话】”“【评论原话】”。")
    doc.add_paragraph("二是【评论类别】。网民称，“【评论原话】”“【评论原话】”。")
    doc.add_paragraph("三是【评论类别】。网民称，“【评论原话】”“【评论原话】”。")
    doc.add_paragraph("四是【评论类别】。网民称，“【评论原话】”“【评论原话】”。")
    doc.add_paragraph("五是【评论类别】。网民称，“【评论原话】”“【评论原话】”。")
    add_role_paragraph(doc, "（三）热词分布情况", "h2")
    doc.add_paragraph("从热词分布来看，“【热词一】”“【热词二】”“【热词三】”等词位居前列，舆论【第一关注点解释】。“【热词四】”“【热词五】”“【热词六】”等词热度较高，舆论【第二关注点解释】。“【热词七】”“【热词八】”“【热词九】”等词热传，舆论【第三关注点解释】。“【热词十】”“【热词十一】”“【热词十二】”等词受到关注，舆论【第四关注点解释】。此外，“【补充热词】”等词也受到关注，舆论【补充关注点解释】。")
    doc.add_paragraph("【此处插入词云或热词分布图；热词来自监测窗口内全部有效文本。】")

    add_role_paragraph(doc, "三、境外舆论情况", "h1")
    add_role_paragraph(doc, "（一）境外媒体情况", "h2")
    overseas_rules = writing_rules()['overseas']
    doc.add_paragraph("【仅有事实性报道时】" + overseas_rules['factual_lead_template'].format(examples='【来源】文章《【标题】》') + overseas_rules['no_interpretation_suffix'])
    doc.add_paragraph("【存在评论性解读时】数据周期内，境外媒体以事实性报道为主。如【来源】文章《【标题】》、【来源】文章《【标题】》、【来源】文章《【标题】》等，少量解读如下：")
    doc.add_paragraph("【解读主题判断】。【境外媒体】引述【专家/机构】观点称，【核心解读或风险叙事】。")
    add_role_paragraph(doc, "（二）境外网民评论", "h2")
    doc.add_paragraph(overseas_rules['no_comment_evidence_sentence'])

    add_role_paragraph(doc, "附录", "h1")
    add_role_paragraph(doc, "（一）外媒报道列表（部分）", "h2")
    table = doc.add_table(rows=2, cols=4)
    for idx, text in enumerate(["序号", "来源", "报道日期", "超链接标题"]):
        table.cell(0, idx).text = text
    for idx, text in enumerate(["1", "【来源】", "【日期】", "【标题】"]):
        table.cell(1, idx).text = text
    mark_row_as_header(table.rows[0])
    style_table(table, "微软雅黑", 11, [36.7, 120.5, 70.9, 203.9])
    add_role_paragraph(doc, "（二）传播量较大的公众号文章", "h2")
    table = doc.add_table(rows=3, cols=5)
    table.cell(0, 0).merge(table.cell(0, 4)).text = "微信公号文章阅读量TOP"
    for idx, text in enumerate(["序号", "账号", "标题", "阅读量", "在看量"]):
        table.cell(1, idx).text = text
    for idx, text in enumerate(["1", "【账号】", "【标题】", "【阅读量】", "【在看量】"]):
        table.cell(2, idx).text = text
    mark_row_as_header(table.rows[0])
    mark_row_as_header(table.rows[1])
    style_table(table, "仿宋", 12, [35.2, 85.1, 191.4, 56.7, 54.8])

    path.parent.mkdir(parents=True, exist_ok=True)
    remove_nonprinting_pagination_controls(doc)
    doc.save(path)


def main() -> None:
    target = Path(__file__).resolve().parents[1] / "templates" / "formal_report_template_complete_20260714.docx"
    build_template(target)
    print(target)


if __name__ == "__main__":
    main()
