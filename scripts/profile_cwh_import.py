from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any


def profile_csv(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        rows = []
        for index, row in enumerate(reader):
            rows.append(row)
            if index >= 50:
                break
    return {
        "type": "table",
        "rows_sampled": len(rows),
        "columns": max((len(row) for row in rows), default=0),
        "headers": rows[0] if rows else [],
    }


def profile_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        headers = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row]
            if any(values):
                headers = values
                break
        sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "first_nonempty_row": headers})
    workbook.close()
    return {"type": "workbook", "sheets": sheets}


def profile_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(value, list):
        keys = sorted({str(key) for row in value[:50] if isinstance(row, dict) for key in row})
        return {"type": "records", "records": len(value), "keys": keys}
    if isinstance(value, dict):
        return {"type": "object", "keys": sorted(map(str, value.keys()))}
    return {"type": type(value).__name__}


def profile_files(directory: Path) -> dict[str, Any]:
    files = []
    for path in sorted(directory.iterdir()):
        if not path.is_file():
            continue
        row: dict[str, Any] = {"name": path.name, "size": path.stat().st_size, "suffix": path.suffix.lower()}
        try:
            if path.suffix.lower() in {".xlsx", ".xlsm"}:
                row["structure"] = profile_xlsx(path)
            elif path.suffix.lower() in {".csv", ".tsv"}:
                row["structure"] = profile_csv(path)
            elif path.suffix.lower() == ".json":
                row["structure"] = profile_json(path)
            else:
                row["structure"] = {"type": "stored_only"}
        except Exception as exc:
            row["structure"] = {"type": "profile_error", "error": f"{type(exc).__name__}: {exc}"}
        files.append(row)
    return {"file_count": len(files), "files": files}
