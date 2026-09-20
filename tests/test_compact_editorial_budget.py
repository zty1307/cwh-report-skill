"""Soft editorial budgets must never become evidence deletion or person quotas."""
from copy import deepcopy
import json
from pathlib import Path
import sys
import pytest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import formalize_cwh_report as formal
from cwh_writing_rules import writing_rules, editorial_template_prompt, selection_audit_prompt
from cwh_article_reading import reading_prompt, reading_contract
from cwh_claim_synthesis import flat_prompt
from cwh_semantic_compiler import AUTHOR_PROMPT
from cwh_ranked_selection import ranked_selection


def test_shared_soft_targets_reach_reading_selection_and_legacy_author():
    rules = writing_rules()['viewpoint']
    assert rules['claim_char_range'] == [50, 120]
    assert rules['preferred_claim_char_range'] == [80, 100]
    assert rules['extended_claim_char_range'] == [120, 150]
    assert rules['topic_excerpt_range'] == [4, 6]
    assert rules['topic_cluster_range'] == [2, 3]
    for prompt in (reading_prompt(), reading_contract({'items': []}), AUTHOR_PROMPT,
                   editorial_template_prompt('revision')):
        assert '50—120' in prompt and '100—200' not in prompt
    for prompt in (flat_prompt(), flat_prompt(ranked=True), AUTHOR_PROMPT):
        assert rules['selection_budget_rule'] in prompt
    # Exact source and thin-evidence safety checks are not weakened by brevity.
    assert rules['minimum_claim_cjk'] == 30
    assert rules['density_gate']['minimum_independent_voices'] == 2


def test_selection_above_six_claims_is_not_mechanically_cut():
    candidates = [{'id': f'c{i}', 'speaker': f'专家{i}', 'source': '来源', 'claim': '完整观点'}
                  for i in range(7)]
    packet = {'candidates': candidates,
              'formal_selection': {'allow_reserve': True, 'max_independent_voices': 12}}
    response = {'heading': '政策判断', 'selected': [
        {'key': 'k1', 'heading': '共同判断', 'claim_ids': [r['id'] for r in candidates], 'thin_reason': ''}]}
    before = deepcopy(packet)
    result = ranked_selection(packet, response)
    assert result['selected'][0]['claim_ids'] == [r['id'] for r in candidates]
    assert packet == before


def test_synthesis_quantity_notes_do_not_require_corroborating_sources():
    for ranked in (False, True):
        prompt = flat_prompt(ranked=ranked)
        assert selection_audit_prompt() in prompt
        assert '不足4个独立主体时说明实际缺口' not in prompt
        assert 'thin_reason只陈述该组实际证据不足' not in prompt


def test_one_subject_with_distinct_judgments_survives_ranked_selection():
    packet = {'candidates': [
        {'id': f'c{i}', 'speaker': '同一专家', 'source': '来源', 'claim': f'完整判断{i}'}
        for i in range(2)],
        'formal_selection': {'allow_reserve': True, 'max_independent_voices': 12}}
    response = {'heading': '政策判断', 'selected': [
        {'key': f'k{i}', 'heading': f'独立判断{i}', 'claim_ids': [f'c{i}'],
         'thin_reason': '一位主体，提供本组独有的机制判断'} for i in range(2)],
        'shortfall_reason': '当前一位主体，保留两项不同机制，不为人数凑引文'}
    frozen = deepcopy((packet, response))
    result = ranked_selection(packet, response)
    assert [group['claim_ids'] for group in result['selected']] == [['c0']]
    assert result['reserved'][0]['id'] == 'c1'
    assert result['reserved'][0]['reason'].startswith('ranked_same_subject_topic:')
    assert (packet, response) == frozen


def test_paragraph_budget_counts_heading_without_cutting_claims():
    rows = [{'speaker_name': name, 'formal_claim': name * 130 + '。'} for name in ('甲', '乙')]
    cluster = {'evidence': rows}
    before = deepcopy(cluster)
    assert len(formal.cluster_detail_paragraphs(cluster, first_lead='短标题。')) == 1
    split = formal.cluster_detail_paragraphs(cluster, first_lead='完整判断标题' * 9)
    assert len(split) == 1
    assert all(row['formal_claim'] in ''.join(split) for row in rows)
    long = '必要条件' * 110 + '。'
    assert long in ''.join(formal.cluster_detail_paragraphs({'evidence': [
        {'speaker_name': '丙', 'formal_claim': long}]}))
    assert cluster == before


@pytest.mark.parametrize('prefix, names, visible', [
    ('', ['甲媒体'], False), ('', ['甲媒体', '乙媒体'], True),
    ('一是', ['甲媒体', '乙媒体'], True),
])
def test_single_source_natural_paragraph_does_not_repeat_summary(prefix, names, visible):
    cluster = {'summary': '认可服务拓展的作用', 'evidence': [
        {'speaker_name': name, 'attribution': name, 'attribution_status': 'media_voice',
         'formal_claim': f'{name}认为，政策可改善服务可及性，但仍需落实经费保障。'}
        for name in names]}
    frozen = deepcopy(cluster)
    doc = Document()
    formal.add_cluster_paragraph(doc, prefix, cluster)
    text = '\n\n'.join(p.text for p in doc.paragraphs)
    assert text == formal.cluster_wording(cluster, prefix=prefix)
    assert (cluster['summary'] in text) is visible
    assert all(row['formal_claim'] in text for row in cluster['evidence'])
    assert cluster == frozen


def test_summary_only_legacy_cluster_does_not_disappear():
    cluster = {'summary': '认可服务拓展的作用'}
    assert formal.cluster_display_lead(cluster) == '认可服务拓展的作用。'


@pytest.mark.parametrize('single_topic', [False, True])
def test_review_marker_follows_all_split_claims_without_breaking_mapping(single_topic):
    from normalize_cwh_analysis import assemble_cluster_details
    rows = [{'speaker_name': name, 'attribution': name,
             'attribution_status': 'media_only', 'formal_claim': name + '认为，' + '判断及必要条件' * 20 + '。'}
            for name in ('甲媒体', '乙媒体', '业内人士')]
    cluster = {'summary': '认为服务范围需要拓展', 'evidence': rows}
    cluster['details'] = assemble_cluster_details(cluster)
    frozen = deepcopy(cluster)
    assert formal.cluster_needs_attribution_review(cluster)
    doc = Document()
    if single_topic:
        formal.add_single_topic_paragraph(doc, 1, {'topic': '公共服务', 'heading': '认为服务范围需要拓展'}, cluster)
    else:
        formal.add_cluster_paragraph(doc, '一是', cluster)
    assert len(doc.paragraphs) == 1
    text = ''.join(p.text for p in doc.paragraphs)
    assert assemble_cluster_details(cluster) in text
    assert all(row['formal_claim'] in text for row in rows)
    assert text.count(formal.ATTRIBUTION_REVIEW_TEXT) == 1
    assert doc.paragraphs[-1].text.endswith(formal.ATTRIBUTION_REVIEW_TEXT)
    assert cluster == frozen


def test_size_audit_counts_claims_not_people_and_is_advisory():
    rows = [{'speaker_name': '同一专家', 'formal_claim': str(i) + '必要条件' * 50,
             'evidence_id': f'e{i}'} for i in range(7)]
    data = {'meeting': {'topics': ['议题']}, 'viewpoints': {'by_topic': [
        {'topic': '议题', 'clusters': [{'summary': '判断', 'evidence': rows}]}]}}
    doc = Document()
    doc.add_paragraph('（一）境内媒体自媒体情况')
    doc.add_paragraph('标题与署名' + '长观点' * 110)
    doc.add_paragraph('（二）网民评论情况')
    doc.add_paragraph('其他章节不纳入' * 100)
    before = deepcopy(data)
    result = formal.audit_domestic_writing_size(data, doc._element.xml.encode())
    assert result['advisory_only'] and 'passed' not in result
    assert result['structured_claim_count'] == 7  # not one distinct speaker
    assert len(result['word_paragraph_chars']) == 1
    assert {w['kind'] for w in result['warnings']} == {
        'check_additional_information_gain', 'check_long_claim_context', 'check_long_rendered_paragraph'}
    assert data == before


def test_reading_discovery_and_time_budgets_are_not_reduced():
    profiles = json.loads((ROOT / 'config/execution_policy.v1.json').read_text('utf-8'))['profiles']
    for name, articles, pages in [('bounded_40m', 16, 18), ('bounded_60m', 20, 24)]:
        research = profiles[name]['research']
        assert research['target_formal_excerpts_per_topic'] == [4, 6]
        assert 'target_independent_voices_per_topic' not in research
        assert research['max_monitoring_full_article_reviews_per_topic'] == articles
        assert research['max_full_page_fetches_per_topic'] == pages
        assert research['max_query_executions_per_topic'] == 10
        assert research['max_formal_voices_per_topic'] == 12  # safety ceiling, not target
