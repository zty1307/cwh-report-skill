from __future__ import annotations

import importlib.util
import hashlib
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
PROJECT_ROOT = Path(__file__).resolve().parents[2]
SCRIPT_PATH = PROJECT_ROOT / "cwh-report-skill" / "scripts" / "raw_system_workbook_pipeline.py"
CONFIG_PATH = PROJECT_ROOT / "cwh-report-skill" / "config" / "raw_workbook_mapping.json"
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "raw_workbook_pipeline_20260710"
METADATA_PATH = OUTPUT_DIR / "run_metadata.json"
OUTPUT_PATH = OUTPUT_DIR / "0713-7月10日国务院CWH舆情情况_自动生成.xlsx"
RAW_DIR = Path(os.environ.get("CWH_RAW_TEST_DIR", r"D:\第九周\原始数据"))
SKILL_PATH = PROJECT_ROOT / "cwh-report-skill" / "SKILL.md"
RAW_REFERENCE_PATH = PROJECT_ROOT / "cwh-report-skill" / "references" / "raw_workbook_pipeline.md"
REPORT_ORCHESTRATOR_PATH = PROJECT_ROOT / "cwh-report-skill" / "scripts" / "cwh_orchestrator.py"


def load_pipeline_module():
    spec = importlib.util.spec_from_file_location("raw_system_workbook_pipeline", SCRIPT_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法加载脚本：{SCRIPT_PATH}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@unittest.skipUnless(RAW_DIR.exists() and METADATA_PATH.exists(), "本期原始测试夹具不可用")
class CurrentFixtureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pipeline = load_pipeline_module()
        cls.config = cls.pipeline.load_json(CONFIG_PATH)
        cls.metadata = cls.pipeline.load_json(METADATA_PATH)
        cls.inputs = cls.pipeline.identify_inputs(RAW_DIR, cls.config)
        cls.normalized = cls.pipeline.build_normalized_bundle(cls.inputs, cls.config, cls.metadata)

    def test_identifies_exact_nine_file_bundle(self):
        identified = {
            self.inputs.total_heat,
            self.inputs.public_top,
            self.inputs.overseas_latest,
            *self.inputs.child_heat.values(),
        }
        self.assertEqual(6, len(self.inputs.child_heat))
        self.assertEqual(set(range(1, 7)), set(self.inputs.child_heat))
        self.assertEqual(9, len(identified))

    def test_channel_group_contract(self):
        self.assertEqual(["domestic_news"], self.config["total_groups"]["domestic_mainstream"])
        self.assertEqual(["overseas_news"], self.config["total_groups"]["overseas_media"])
        self.assertEqual(
            ["public_articles", "public_recommend", "public_comments"],
            self.config["total_groups"]["wechat"],
        )
        child_new_media = self.config["child_groups"]["new_media"]
        self.assertEqual(11, len(child_new_media))
        self.assertNotIn("domestic_news", child_new_media)
        self.assertNotIn("overseas_news", child_new_media)

    def test_current_system_totals_match_independent_recalculation(self):
        self.assertEqual(
            {
                "domestic_mainstream": 7009,
                "overseas_media": 40,
                "wechat": 29442,
                "weibo": 2106,
                "video_account": 243,
                "other_new_media": 2246,
                "total": 41086,
            },
            self.normalized["total_event"]["totals"],
        )
        self.assertEqual(
            [11852, 27877, 9699, 18891, 9732, 14694],
            [child["totals"]["total"] for child in self.normalized["children"]],
        )

    def test_semantic_lists_are_auditable(self):
        public_titles = {item["title"] for item in self.normalized["public_top"]["selected"]}
        self.assertEqual(10, len(public_titles))
        self.assertIn("利好！事关通信和算力建设！", public_titles)
        self.assertTrue(
            any(
                item["reason"] == "roundup_or_breakfast_digest" and item["decision"] == "exclude"
                for item in self.normalized["public_top"]["decisions"]
            )
        )

        overseas = self.normalized["overseas"]
        self.assertEqual(len(overseas["decisions"]), 40)
        self.assertTrue(
            any(
                item["reason"] == "preliminary_relevant_mainland_outward_report"
                for item in overseas["decisions"]
            )
        )
        self.assertEqual("ai_review_required", overseas["status"])
        self.assertFalse(overseas["counts_reconciled"])
        self.assertEqual(40, len(overseas["review_packet"]["items"]))
        self.assertFalse(self.normalized["quality_gate"]["ready_for_formal_report"])
        self.assertTrue(
            any(item["reason"] == "duplicate_recurring_series_latest" for item in overseas["decisions"])
        )

    def test_hotwords_require_ai_review_instead_of_silent_fallback(self):
        hotwords = self.normalized["hotwords"]
        self.assertEqual("ai_review_required", hotwords["status"])
        self.assertEqual("ai_review_required", hotwords["method"])
        self.assertGreater(hotwords["corpus_audit"]["deduplicated_relevant_document_count"], 0)
        self.assertGreater(hotwords["corpus_audit"]["candidate_count"], 0)
        self.assertEqual([], hotwords["selected"])
        self.assertEqual("ai_review_required", hotwords["review_packet"]["status"])


class HotwordSemanticReviewTests(unittest.TestCase):
    def setUp(self):
        self.pipeline = load_pipeline_module()
        self.metadata = {
            "topic_titles": ["建设算力基础设施", "促进数字经济发展"],
            "topic_aliases": [["算力", "算力设施"], ["数字经济"]],
            "wordcloud": {"minimum_term_count": 2, "term_count": 3},
        }
        self.rows = [
            {
                "source_row": 2,
                "account": "媒体甲",
                "url": "https://example.test/a",
                "title": "国务院常务会议部署算力设施建设",
                "content": "会议关注算力设施和数字经济，提出优化算力布局。",
            },
            {
                "source_row": 3,
                "account": "媒体乙",
                "url": "https://example.test/b",
                "title": "算力设施支撑数字经济",
                "content": "国务院常务会议研究算力设施，数字经济获得新支撑。",
            },
            {
                "source_row": 4,
                "account": "媒体丙",
                "url": "https://example.test/c",
                "title": "数字经济带动算力布局优化",
                "content": "国务院常务会议相关部署强调数字经济与算力布局。",
            },
        ]

    def review(self):
        return {
            "review_method": "ai_semantic_review",
            "second_pass_completed": True,
            "selected": [
                {
                    "term": "算力设施",
                    "topic_hits": [1],
                    "evidence_tier": "core",
                    "semantic_type": "infrastructure_or_sector",
                    "standalone_topic_label": True,
                    "selection_reason": "多篇报道将其作为会议实质议题",
                    "ai_representativeness": "高",
                },
                {
                    "term": "数字经济",
                    "topic_hits": [2],
                    "evidence_tier": "core",
                    "semantic_type": "agenda_topic",
                    "standalone_topic_label": True,
                    "selection_reason": "多来源报道围绕其政策作用展开",
                    "ai_representativeness": "高",
                },
            ],
        }

    def test_review_is_mandatory_and_packet_invites_ai_supplements(self):
        result = self.pipeline.build_hotword_payload(
            self.rows,
            self.metadata,
            ["国务院常务会议"],
        )
        self.assertEqual("ai_review_required", result["status"])
        self.assertEqual([], result["selected"])
        instructions = "".join(result["review_packet"]["instructions"])
        self.assertIn("补提", instructions)

    def test_review_packet_is_schema_driven_not_seeded_with_historical_terms(self):
        result = self.pipeline.build_hotword_payload(
            self.rows,
            self.metadata,
            ["国务院常务会议"],
        )
        packet = result["review_packet"]
        self.assertNotIn("gold_standard_style_examples", packet)
        self.assertIn("accepted_style_patterns", packet)
        self.assertEqual(2, len(packet["topics"]))
        self.assertIn("实际议题数量", "".join(packet["instructions"]))
        serialized = json.dumps(packet["accepted_style_patterns"], ensure_ascii=False)
        for historical_term in ("数字中国", "人工智能", "算力网", "自然资源保护"):
            self.assertNotIn(historical_term, serialized)

    def test_complete_ai_review_produces_only_audited_terms(self):
        result = self.pipeline.build_hotword_payload(
            self.rows,
            self.metadata,
            ["国务院常务会议"],
            review=self.review(),
        )
        self.assertEqual("ai_review_complete", result["status"])
        self.assertEqual("ai_semantic_review_with_evidence", result["method"])
        self.assertEqual({"算力设施", "数字经济"}, {item["term"] for item in result["selected"]})
        self.assertTrue(all(item["ai_reviewed"] for item in result["selected"]))
        self.assertTrue(all(item["selection_reason"] for item in result["selected"]))
        self.assertTrue(all(item["topic_index"] == item["topic_hits"][0] for item in result["selected"]))

    def test_generic_term_cannot_pass_even_when_review_file_selects_it(self):
        review = self.review()
        review["selected"][0]["term"] = "关于"
        with self.assertRaisesRegex(ValueError, "泛化词"):
            self.pipeline.build_hotword_payload(
                self.rows,
                self.metadata,
                ["国务院常务会议"],
                review=review,
            )

    def test_ai_must_fill_to_minimum_with_evidence(self):
        review = self.review()
        review["selected"] = review["selected"][:1]
        with self.assertRaisesRegex(ValueError, "继续补提"):
            self.pipeline.build_hotword_payload(
                self.rows,
                self.metadata,
                ["国务院常务会议"],
                review=review,
            )

    def test_procedural_phrase_cannot_masquerade_as_ai_topic_term(self):
        review = self.review()
        review["selected"][0]["term"] = "修改和废止"
        review["selected"][0]["semantic_type"] = "governance_mechanism"
        with self.assertRaisesRegex(ValueError, "会议程序动作"):
            self.pipeline.build_hotword_payload(
                self.rows,
                self.metadata,
                ["国务院常务会议"],
                review=review,
            )

    def test_pure_project_location_cannot_masquerade_as_topic_term(self):
        review = self.review()
        review["selected"][0]["term"] = "辽宁庄河"
        with self.assertRaisesRegex(ValueError, "纯地名或项目所在地"):
            self.pipeline.build_hotword_payload(
                self.rows,
                self.metadata,
                ["国务院常务会议"],
                review=review,
            )


@unittest.skipUnless(OUTPUT_PATH.exists(), "自动生成工作簿尚未生成")
class GeneratedWorkbookTests(unittest.TestCase):
    def test_workbook_structure_formulas_and_charts(self):
        workbook = load_workbook(OUTPUT_PATH, data_only=False, read_only=False)
        try:
            self.assertEqual(
                [
                    "关键词",
                    "总事件",
                    "子事件1",
                    "子事件2",
                    "子事件3",
                    "子事件4",
                    "子事件5",
                    "子事件6",
                    "子事件数据汇总",
                    "外媒报道列表",
                    "词云",
                    "公众TOP",
                ],
                workbook.sheetnames,
            )
            self.assertEqual('=TEXT(O4,"m/d")', workbook["总事件"]["A4"].value)
            self.assertEqual("=T4", workbook["总事件"]["B4"].value)
            self.assertEqual('=TEXT(H4,"m/d")', workbook["子事件1"]["A4"].value)
            self.assertEqual("=M4", workbook["子事件1"]["B4"].value)
            self.assertEqual(2, sum(len(sheet._charts) for sheet in workbook.worksheets))
        finally:
            workbook.close()

    def test_builder_verification_has_no_formula_or_render_errors(self):
        verification_path = OUTPUT_DIR / "run" / "builder_verification.json"
        verification = json.loads(verification_path.read_text(encoding="utf-8"))
        self.assertTrue(verification["sheet_order_ok"])
        self.assertEqual(0, verification["formula_error_count"])
        self.assertEqual(12, len(verification["rendered_sheets"]))
        self.assertEqual([], verification["render_errors"])


class StageBoundaryTests(unittest.TestCase):
    def test_skill_routes_raw_and_standard_inputs_without_forcing_report_stage(self):
        skill_text = SKILL_PATH.read_text(encoding="utf-8")
        self.assertIn("## Input Routing and Stage Boundaries", skill_text)
        self.assertIn("stop after `workbook`", skill_text)
        self.assertIn("references/raw_workbook_pipeline.md", skill_text)
        self.assertTrue(RAW_REFERENCE_PATH.exists())

    def test_raw_stage_does_not_import_report_or_dashboard_modules(self):
        source = SCRIPT_PATH.read_text(encoding="utf-8")
        for forbidden in ("cwh_orchestrator", "generate_dashboard", "formalize_cwh_report"):
            self.assertNotIn(f"import {forbidden}", source)
            self.assertNotIn(f"from {forbidden}", source)
        orchestrator_source = REPORT_ORCHESTRATOR_PATH.read_text(encoding="utf-8")
        self.assertNotIn("raw_system_workbook_pipeline", orchestrator_source)

    def test_artifact_runtime_can_be_discovered_without_project_junction(self):
        if os.name != "nt" or os.environ.get("CWH_PORTABLE_XLSX", "").strip().lower() in {
            "1",
            "true",
            "yes",
            "on",
        }:
            self.skipTest("portable workbook generation does not require @oai/artifact-tool")
        pipeline = load_pipeline_module()
        node_modules = pipeline.find_artifact_node_modules(PROJECT_ROOT / "cwh-report-skill")
        self.assertIsNotNone(node_modules)
        self.assertTrue((node_modules / "@oai" / "artifact-tool").exists())


class InputIdentificationTests(unittest.TestCase):
    def test_current_total_prefix_and_detail_aliases_are_bundled(self):
        pipeline = load_pipeline_module()
        config = pipeline.load_json(CONFIG_PATH)
        self.assertRegex("总2-示例-热度分析.xlsx", config["file_patterns"]["total_heat"])
        self.assertRegex("总2-公众文章最热样本.xlsx", config["file_patterns"]["public_top"])
        self.assertIn("账号昵称", config["detail_columns"]["source"])
        self.assertIn("账号昵称", config["detail_columns"]["account"])
        self.assertIn("在看量（推荐量）", config["detail_columns"]["recommend_count"])

    def test_filename_monitoring_window_is_audited_without_changing_values(self):
        pipeline = load_pipeline_module()
        rows = [
            {
                "file": "总事件-2026.07.31 00_00至2026.08.03 14_12-热度分析.xlsx",
                "role": "total_heat",
                "filename_monitoring_window": pipeline.filename_monitoring_window(
                    "总事件-2026.07.31 00_00至2026.08.03 14_12-热度分析.xlsx"
                ),
            },
            {
                "file": "总事件-2026.07.31 00_00至2026.08.03 13_49-公众文章.xlsx",
                "role": "public_top",
                "filename_monitoring_window": pipeline.filename_monitoring_window(
                    "总事件-2026.07.31 00_00至2026.08.03 13_49-公众文章.xlsx"
                ),
            },
        ]

        audit = pipeline.input_time_window_audit(rows, {"cutoff_consistency_tolerance_minutes": 2})

        self.assertEqual(23, audit["end_cutoff_span_minutes"])
        self.assertFalse(audit["end_cutoffs_consistent"])
        self.assertEqual("input_monitoring_windows_are_not_aligned", audit["warning"])

    def setUp(self):
        self.pipeline = load_pipeline_module()
        self.config = self.pipeline.load_json(CONFIG_PATH)
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def create_heat(self, name: str, total: int):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "日趋势"
        worksheet.append(self.pipeline.DISPLAY_DAILY_HEADERS)
        worksheet.append(["2026-01-01", total, *([0] * 12)])
        workbook.save(self.root / name)

    def create_detail(self, name: str, sheet_name: str):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        worksheet.append(
            [
                "站点名称",
                "发文链接",
                "发表时间",
                "标题",
                "正文内容",
                "昵称",
                "阅读量",
                "推荐（公号）/爱心点赞（视频号）/收藏量",
            ]
        )
        worksheet.append(["来源", "https://example.test/a", "2026-01-01", "标题", "正文", "账号", 1, 0])
        workbook.save(self.root / name)

    def test_roles_are_driven_by_sheet_schema_not_fixed_filenames(self):
        self.create_heat("附件C.xlsx", 100)
        self.create_heat("任意前缀_子1数据.xlsx", 40)
        self.create_heat("2-完全不同的名称.xlsx", 60)
        self.create_detail("附件A.xlsx", "最热公号文章")
        self.create_detail("附件B.xlsx", "境外新闻")

        inputs = self.pipeline.identify_inputs(self.root, self.config)

        self.assertEqual("附件C.xlsx", inputs.total_heat.name)
        self.assertEqual({1, 2}, set(inputs.child_heat))
        self.assertEqual("附件A.xlsx", inputs.public_top.name)
        self.assertEqual("附件B.xlsx", inputs.overseas_latest.name)
        methods = {item["role"]: item["method"] for item in inputs.classification_audit}
        self.assertEqual("sheet_and_header_schema", methods["public_top"])
        self.assertEqual("unique_largest_aggregate_fallback", methods["total_heat"])

    def test_misleading_old_style_name_cannot_override_detail_sheet_schema(self):
        self.create_heat("真实总表.xlsx", 100)
        self.create_heat("子1.xlsx", 50)
        self.create_detail("总1-看起来像热度分析.xlsx", "最热公号文章")
        self.create_detail("完全随机.xlsx", "境外新闻")

        inputs = self.pipeline.identify_inputs(self.root, self.config)

        self.assertEqual("总1-看起来像热度分析.xlsx", inputs.public_top.name)
        self.assertEqual("真实总表.xlsx", inputs.total_heat.name)

    def test_equal_heat_totals_without_total_hint_are_rejected_as_ambiguous(self):
        self.create_heat("a-子1.xlsx", 50)
        self.create_heat("b-子2.xlsx", 50)
        self.create_detail("public.xlsx", "最热公号文章")
        self.create_detail("overseas.xlsx", "境外新闻")

        with self.assertRaisesRegex(self.pipeline.PipelineError, "无法唯一识别总事件热度文件"):
            self.pipeline.identify_inputs(self.root, self.config)

    def test_child_without_auditable_index_is_rejected(self):
        self.create_heat("总事件.xlsx", 100)
        self.create_heat("无编号子表.xlsx", 50)
        self.create_detail("public.xlsx", "最热公号文章")
        self.create_detail("overseas.xlsx", "境外新闻")

        with self.assertRaisesRegex(self.pipeline.PipelineError, "未提供可验证的子事件序号"):
            self.pipeline.identify_inputs(self.root, self.config)


class OverseasSemanticReviewTests(unittest.TestCase):
    def setUp(self):
        self.pipeline = load_pipeline_module()
        self.config = {
            "meeting_anchor_patterns": ["国务院常务会议"],
            "market_promo_patterns": ["ETF"],
            "domestic_overseas_outlet_patterns": ["CGTN", "news.cgtn.com"],
        }
        self.metadata = {
            "meeting_title": "示例国务院常务会议",
            "topic_titles": ["议题一", "议题二"],
            "topic_aliases": [["议题一"], ["议题二"]],
        }
        self.rows = [
            {
                "source_row": 2,
                "source": "境外媒体甲",
                "url": "https://foreign.example/a",
                "published_at": "2026-01-01 08:00:00",
                "title": "国务院常务会议研究议题一",
                "content": "会议研究议题一。",
            },
            {
                "source_row": 3,
                "source": "CGTN",
                "url": "https://news.cgtn.com/b",
                "published_at": "2026-01-01 09:00:00",
                "title": "国务院常务会议研究两项工作",
                "content": "会议研究议题一和议题二。",
            },
            {
                "source_row": 4,
                "source": "境外财经网站",
                "url": "https://foreign.example/c",
                "published_at": "2026-01-02 10:00:00",
                "title": "ETF行情",
                "content": "偶然提及国务院常务会议。",
            },
        ]

    def review(self):
        return {
            "review_method": "ai_semantic_review",
            "items": [
                {
                    "record_id": self.pipeline.overseas_record_id(self.rows[0]),
                    "decision": "include",
                    "publisher_class": "overseas_origin_media",
                    "topic_hits": [1],
                    "review_reason": "以本次会议议题为主要内容",
                    "classification_confidence": 0.98,
                    "content_type": "factual",
                },
                {
                    "record_id": self.pipeline.overseas_record_id(self.rows[1]),
                    "decision": "include",
                    "publisher_class": "mainland_outward_media",
                    "topic_hits": [1, 2],
                    "review_reason": "境内媒体海外版对本次会议的完整报道",
                    "classification_confidence": 0.97,
                    "content_type": "factual",
                },
                {
                    "record_id": self.pipeline.overseas_record_id(self.rows[2]),
                    "decision": "exclude",
                    "publisher_class": "overseas_origin_media",
                    "topic_hits": [],
                    "review_reason": "行情推广为主，会议不是报道对象",
                    "classification_confidence": 0.96,
                    "content_type": "factual",
                },
            ],
            "supplemental_rows": [],
        }

    def test_review_separates_counted_universe_from_formal_appendix(self):
        result = self.pipeline.filter_overseas(
            self.rows,
            self.config,
            self.metadata,
            review=self.review(),
            monitoring_dates=["2026-01-01", "2026-01-02"],
        )
        self.assertEqual("ai_review_complete", result["status"])
        self.assertEqual(2, len(result["counted"]))
        self.assertEqual(1, len(result["selected"]))
        self.assertEqual("overseas_origin_media", result["selected"][0]["publisher_class"])
        self.assertEqual(1, result["summary"]["mainland_outward_total"])

    def test_review_recalculates_only_overseas_channels_and_propagates_topics(self):
        overseas = self.pipeline.filter_overseas(
            self.rows,
            self.config,
            self.metadata,
            review=self.review(),
            monitoring_dates=["2026-01-01", "2026-01-02"],
        )
        base = {
            "public_articles": 3,
            "public_recommend": 0,
            "public_comments": 0,
            "weibo": 4,
            "domestic_news": 5,
            "domestic_app": 0,
            "domestic_forum": 0,
            "other_video": 0,
            "overseas_news": 99,
            "x": 0,
            "overseas_other": 0,
            "video_account": 0,
            "douyin": 0,
        }
        total_daily = [{"date": "2026-01-01", **base}, {"date": "2026-01-02", **base}]
        child_daily = {1: [dict(row) for row in total_daily], 2: [dict(row) for row in total_daily]}
        reconciled = self.pipeline.reconcile_overseas_daily_counts(total_daily, child_daily, overseas)
        self.assertEqual([2, 0], [row["overseas_news"] for row in reconciled["total_daily"]])
        self.assertEqual([2, 0], [row["overseas_news"] for row in reconciled["child_daily"][1]])
        self.assertEqual([1, 0], [row["overseas_news"] for row in reconciled["child_daily"][2]])
        self.assertEqual([5, 5], [row["domestic_news"] for row in reconciled["total_daily"]])

    def test_incomplete_ai_review_is_rejected(self):
        review = self.review()
        review["items"] = review["items"][:-1]
        with self.assertRaisesRegex(self.pipeline.PipelineError, "未覆盖全部系统候选"):
            self.pipeline.filter_overseas(
                self.rows,
                self.config,
                self.metadata,
                review=review,
                monitoring_dates=["2026-01-01", "2026-01-02"],
            )

    def test_title_only_overseas_inclusion_is_rejected(self):
        self.rows[0]["content"] = ""
        with self.assertRaisesRegex(self.pipeline.PipelineError, "正文缺失"):
            self.pipeline.filter_overseas(self.rows, self.config, self.metadata, review=self.review(),
                                          monitoring_dates=["2026-01-01"])

    def test_interpretive_summary_cannot_certify_its_own_source(self):
        review = self.review()
        review["items"][0].update(content_type="interpretive", summary_cn_simplified="自动摘要",
                                  interpretive_verified=True, interpretive_excerpt="原文中不存在的分析")
        with self.assertRaisesRegex(self.pipeline.PipelineError, "连续原文依据"):
            self.pipeline.filter_overseas(self.rows, self.config, self.metadata, review=review,
                                          monitoring_dates=["2026-01-01"])
        review["items"][0]["interpretive_excerpt"] = self.rows[0]["content"]
        result = self.pipeline.filter_overseas(self.rows, self.config, self.metadata, review=review,
                                             monitoring_dates=["2026-01-01"])
        assert result["counted"][0]["interpretive_excerpt"] == self.rows[0]["content"]

    def test_cross_outlet_reprints_are_counted_but_same_outlet_duplicates_are_not(self):
        shared = {
            "published_at": "2026-01-01 08:00:00",
            "title": "国务院常务会议研究议题一",
            "content": "会议研究议题一。",
            "topic_hits": [1],
        }
        rows = [
            {**shared, "source_row": 1, "source": "媒体甲", "url": "https://a.example/1"},
            {**shared, "source_row": 2, "source": "媒体乙", "url": "https://b.example/1"},
            {**shared, "source_row": 3, "source": "媒体甲", "url": "https://a.example/duplicate"},
        ]
        kept, dropped = self.pipeline.deduplicate_overseas(rows)
        self.assertEqual({"媒体甲", "媒体乙"}, {row["source"] for row in kept})
        self.assertEqual(1, len(dropped))
        self.assertEqual("duplicate_same_outlet_title", dropped[0]["reason"])

    def test_same_publisher_article_id_dedupes_channel_and_mirror_urls(self):
        rows = [
            {"source": "联合新闻网", "title": "核电报道", "url": "https://udn.com/news/story/7333/9663904", "content": "短"},
            {"source": "联合新闻网", "title": "核电报道长版", "url": "https://money.udn.com/money/story/5603/9663904", "content": "更完整的正文"},
            {"source": "Money-Link财经新闻", "title": "国常会报道", "url": "https://ww2.money-link.com.tw/RealtimeNews/NewsContent.aspx?sn=2399150002", "content": "短"},
            {"source": "富联纲", "title": "国常会报道", "url": "https://moneylink.com.tw/RealtimeNews/NewsContent.aspx?SN=2399150002", "content": "更完整的正文"},
        ]
        kept, dropped = self.pipeline.deduplicate_overseas(rows)
        self.assertEqual(2, len(kept))
        self.assertEqual(2, sum(item["reason"] == "duplicate_same_publisher_article" for item in dropped))

    def test_reviewed_supplement_joins_counted_universe_but_out_of_window_row_does_not(self):
        review = self.review()
        review["supplemental_rows"] = [
            {
                "source": "境外媒体丙",
                "url": "https://foreign.example/supplement",
                "published_at": "2026-01-02 11:00:00",
                "title": "境外媒体解读国务院常务会议议题二",
                "content": "文章以本次会议议题二为主要内容。",
                "decision": "include",
                "publisher_class": "overseas_origin_media",
                "topic_hits": [2],
                "review_reason": "监测期内与本次会议直接相关",
                "classification_confidence": 0.94,
                "content_type": "interpretive",
                "summary_cn_simplified": "分析议题二的政策影响。",
                "interpretive_verified": True,
                "interpretive_excerpt": "文章以本次会议议题二为主要内容。",
            },
            {
                "source": "境外媒体丁",
                "url": "https://foreign.example/late",
                "published_at": "2026-01-03 11:00:00",
                "title": "会后跟踪报道",
                "content": "虽相关但超出监测期。",
                "decision": "include",
                "publisher_class": "overseas_origin_media",
                "topic_hits": [1],
                "review_reason": "相关但发布时间较晚",
                "classification_confidence": 0.93,
                "content_type": "factual",
            },
        ]
        result = self.pipeline.filter_overseas(
            self.rows,
            self.config,
            self.metadata,
            review=review,
            monitoring_dates=["2026-01-01", "2026-01-02"],
        )
        self.assertEqual(3, len(result["counted"]))
        self.assertEqual(2, len(result["selected"]))
        self.assertEqual(2, result["summary"]["supplement_candidate_count"])
        late = next(item for item in result["decisions"] if item.get("source") == "境外媒体丁")
        self.assertEqual("exclude", late["decision"])
        self.assertEqual("outside_monitoring_window", late["reason_code"])


class PublicTopRankingTests(unittest.TestCase):
    @staticmethod
    def public_row(source_row, account, read_count):
        return {
            "source_row": source_row,
            "account": account,
            "title": f"anchor topic article {source_row}",
            "content": "anchor topic",
            "url": f"https://example.test/{source_row}",
            "read_count": read_count,
        }

    @staticmethod
    def public_config():
        return {
            "meeting_anchor_patterns": ["anchor"],
            "public_roundup_patterns": ["roundup"],
            "public_top_n": 10,
        }

    @staticmethod
    def public_metadata():
        return {"topic_aliases": [["topic"]]}

    def test_same_account_keeps_only_highest_read_relevant_article(self):
        pipeline = load_pipeline_module()
        rows = [
            {
                "source_row": 1,
                "account": "示例公众号",
                "title": "国务院常务会议部署防汛抗洪救灾工作",
                "content": "国务院常务会议部署防汛抗洪救灾工作",
                "url": "https://example.test/a",
                "read_count": 3000,
            },
            {
                "source_row": 2,
                "account": "示例公众号 ",
                "title": "国务院常务会议研究防汛抗洪救灾工作",
                "content": "国务院常务会议研究防汛抗洪救灾工作",
                "url": "https://example.test/b",
                "read_count": 9000,
            },
            {
                "source_row": 3,
                "account": "另一公众号",
                "title": "国务院常务会议部署防汛抗洪救灾工作",
                "content": "国务院常务会议部署防汛抗洪救灾工作",
                "url": "https://example.test/c",
                "read_count": 5000,
            },
        ]
        config = {
            "meeting_anchor_patterns": ["国务院常务会议"],
            "public_roundup_patterns": ["早餐"],
            "public_top_n": 10,
        }
        metadata = {"topic_aliases": [["防汛", "抗洪", "救灾"]]}
        result = pipeline.filter_public_top(rows, config, metadata)

        self.assertEqual([2, 3], [row["source_row"] for row in result["selected"]])
        decisions = {row["source_row"]: row for row in result["decisions"]}
        self.assertEqual("duplicate_source_lower_read_count", decisions[1]["reason"])
        self.assertEqual("preliminary_selected_pending_ai_review", decisions[2]["reason"])

    def test_source_dedup_happens_before_top10_and_refills_to_ten_sources(self):
        pipeline = load_pipeline_module()
        rows = [
            self.public_row(1, "duplicate account", 20000),
            self.public_row(2, "duplicate-account", 19000),
            self.public_row(3, "duplicate account ", 18000),
        ]
        rows.extend(
            self.public_row(source_row, f"account {source_row}", 17000 - source_row)
            for source_row in range(4, 15)
        )

        result = pipeline.filter_public_top(rows, self.public_config(), self.public_metadata())

        selected = result["selected"]
        self.assertEqual(10, len(selected))
        normalized_sources = {
            pipeline.normalize_text(row["account"]).lower().replace(" ", "").replace("-", "")
            for row in selected
        }
        self.assertEqual(10, len(normalized_sources))
        self.assertEqual(1, selected[0]["source_row"])
        self.assertNotIn(2, {row["source_row"] for row in selected})
        self.assertNotIn(3, {row["source_row"] for row in selected})
        self.assertEqual(10, result["selection_summary"]["selected_count"])
        self.assertEqual(12, result["selection_summary"]["distinct_source_count"])
        self.assertEqual(2, result["selection_summary"]["duplicate_source_count"])

    def test_configured_publisher_family_uses_one_ranking_slot(self):
        pipeline = load_pipeline_module()
        rows = [
            self.public_row(1, "央视新闻", 100000),
            self.public_row(2, "央视网", 100000),
            self.public_row(3, "新闻联播", 100000),
            self.public_row(4, "中国银行保险报", 90000),
        ]
        rows[0]["recommend_count"] = 854
        rows[1]["recommend_count"] = 198
        config = {
            **self.public_config(),
            "public_publisher_families": [
                {"canonical": "央视新闻", "aliases": ["央视新闻", "央视网"]}
            ],
        }

        result = pipeline.filter_public_top(rows, config, self.public_metadata())

        self.assertEqual(
            ["央视新闻", "新闻联播", "中国银行保险报"],
            [row["account"] for row in result["selected"]],
        )
        self.assertEqual(3, result["selection_summary"]["distinct_source_count"])
        decisions = {row["source_row"]: row for row in result["decisions"]}
        self.assertEqual("duplicate_source_lower_read_count", decisions[2]["reason"])

    def test_fewer_than_ten_distinct_sources_is_the_only_short_result_case(self):
        pipeline = load_pipeline_module()
        rows = [
            self.public_row(1, "account one", 10000),
            self.public_row(2, "account one", 9000),
            self.public_row(3, "account two", 8000),
        ]

        result = pipeline.filter_public_top(rows, self.public_config(), self.public_metadata())

        self.assertEqual(2, len(result["selected"]))
        self.assertEqual(2, result["selection_summary"]["distinct_source_count"])

    def test_short_agenda_centered_headline_is_not_vetoed_for_omitting_meeting_anchor(self):
        pipeline = load_pipeline_module()
        rows = [{
            "source_row": 1,
            "account": "权威媒体",
            "title": "公积金，将有新变化",
            "content": "国务院常务会议审议住房公积金制度调整，拓宽提取和使用范围。",
            "url": "https://example.test/provident-fund",
            "read_count": 100000,
        }]
        config = {
            "meeting_anchor_patterns": ["国务院常务会议", "国常会"],
            "public_roundup_patterns": ["新闻早餐"],
            "public_top_n": 10,
        }
        metadata = {"topic_aliases": [["住房公积金", "公积金"]]}

        result = pipeline.filter_public_top(rows, config, metadata)

        self.assertEqual([1], [row["source_row"] for row in result["selected"]])
        self.assertEqual("body_meeting_anchor_and_topic", result["decisions"][0]["relevance_evidence"])

    def test_topic_only_headline_without_meeting_evidence_remains_excluded(self):
        pipeline = load_pipeline_module()
        rows = [{
            "source_row": 1,
            "account": "示例媒体",
            "title": "公积金政策观察",
            "content": "介绍某地往年公积金办理流程。",
            "url": "https://example.test/unrelated",
            "read_count": 100000,
        }]
        config = {
            "meeting_anchor_patterns": ["国务院常务会议", "国常会"],
            "public_roundup_patterns": ["新闻早餐"],
            "public_top_n": 10,
        }
        metadata = {"topic_aliases": [["住房公积金", "公积金"]]}

        result = pipeline.filter_public_top(rows, config, metadata)

        self.assertEqual([], result["selected"])
        self.assertEqual("meeting_not_primary_focus", result["decisions"][0]["reason"])


class PublicArticleEvidenceCorpusTests(unittest.TestCase):
    def setUp(self):
        self.pipeline = load_pipeline_module()
        self.config = {
            "meeting_anchor_patterns": ["国务院常务会议"],
            "public_roundup_patterns": ["新闻早餐"],
        }
        self.metadata = {
            "meeting_title": "示例国务院常务会议",
            "topic_titles": ["研究住房公积金工作"],
            "topic_aliases": [["住房公积金", "公积金"]],
        }

    def test_corpus_is_not_limited_by_read_rank_or_top_selection(self):
        rows = [
            {
                "source_row": 2,
                "account": "高阅读账号",
                "title": "国务院常务会议研究住房公积金工作",
                "content": "国务院常务会议研究住房公积金工作。",
                "url": "https://example.test/high",
                "read_count": 100000,
            },
            {
                "source_row": 1995,
                "account": "低阅读专业账号",
                "title": "公积金制度如何完善",
                "content": "国务院常务会议研究住房公积金工作。专家甲认为应扩大制度覆盖，专家乙建议完善灵活就业人员缴存机制。",
                "url": "https://example.test/low",
                "read_count": 17,
            },
        ]

        result = self.pipeline.build_public_article_evidence_corpus(rows, self.config, self.metadata)

        self.assertEqual(2, result["candidate_count"])
        low = next(item for item in result["candidates"] if item["source_row"] == 1995)
        self.assertIn("专家甲", low["content"])
        self.assertIn("专家乙", low["content"])

    def test_exact_duplicates_and_nonmeeting_rows_keep_auditable_decisions(self):
        base = {
            "account": "账号甲",
            "title": "国务院常务会议研究住房公积金工作",
            "content": "国务院常务会议研究住房公积金工作。",
            "url": "https://example.test/article?utm_source=test",
            "read_count": 100,
        }
        rows = [
            {"source_row": 2, **base},
            {"source_row": 3, **base, "url": "https://example.test/article"},
            {
                "source_row": 4,
                "account": "账号乙",
                "title": "往年公积金办事指南",
                "content": "介绍往年办理流程。",
                "url": "https://example.test/unrelated",
                "read_count": 50,
            },
        ]

        result = self.pipeline.build_public_article_evidence_corpus(rows, self.config, self.metadata)

        self.assertEqual(1, result["candidate_count"])
        self.assertEqual(2, len(result["exclusions"]))
        reasons = {item["reason"] for item in result["exclusions"]}
        self.assertEqual({"duplicate_url", "meeting_not_primary_focus"}, reasons)
        duplicate = next(item for item in result["exclusions"] if item["reason"] == "duplicate_url")
        self.assertTrue(duplicate["duplicate_of"].startswith("public-evidence:2:"))


class HotwordMaskTests(unittest.TestCase):
    def test_default_cloud_assets_match_authoritative_illustrator_shape(self):
        assets = PROJECT_ROOT / "cwh-report-skill" / "assets"
        expected = {
            "wordcloud_cloud_shape.ai": "7a3797761141ba525edcc16e63c033690a852d87c45b86c5d7236393c78928e3",
            "wordcloud_cloud_mask.png": "16de66ad19dec615e84d1734a6e82eda35cc11d442e8160cbce7f3345afe9aa1",
        }
        for name, expected_hash in expected.items():
            path = assets / name
            self.assertTrue(path.exists())
            self.assertEqual(expected_hash, hashlib.sha256(path.read_bytes()).hexdigest())

    def test_default_cloud_mask_uses_alpha_and_preserves_shape_aspect(self):
        load_pipeline_module()
        from cwh_hotword_pipeline import cloud_mask

        mask = cloud_mask(1600, 1000)
        rows, columns = mask.nonzero()
        content_width = int(columns.max() - columns.min() + 1)
        content_height = int(rows.max() - rows.min() + 1)
        self.assertGreater(content_width, 1500)
        self.assertEqual(1000, content_height)
        self.assertGreater(content_width / content_height, 1.55)
        self.assertLess(content_width / content_height, 1.65)
        self.assertFalse(bool(mask[0, 0]))
        self.assertTrue(bool(mask[500, 800]))

    def test_default_font_scale_has_clear_headword_hierarchy(self):
        load_pipeline_module()
        from cwh_hotword_pipeline import (
            DEFAULT_REPEAT_REFERENCE_EXPONENT,
            DEFAULT_REPEAT_REFERENCE_MAX,
            DEFAULT_REPEAT_REFERENCE_MIN,
            DEFAULT_REPEAT_SCALES,
            DEFAULT_MAX_PLACED_INSTANCES,
            primary_font_size,
        )

        largest = primary_font_size(1.0)
        middle = primary_font_size(0.5)
        smallest = primary_font_size(0.0)
        self.assertEqual(150, largest)
        self.assertEqual(22, smallest)
        self.assertGreaterEqual(largest / middle, 3.5)
        self.assertGreater(middle, smallest)

        largest_repeat = round(
            primary_font_size(
                1.0,
                DEFAULT_REPEAT_REFERENCE_MIN,
                DEFAULT_REPEAT_REFERENCE_MAX,
                DEFAULT_REPEAT_REFERENCE_EXPONENT,
            )
            * DEFAULT_REPEAT_SCALES[0]
        )
        self.assertEqual(36, largest_repeat)
        self.assertEqual(12, len(DEFAULT_REPEAT_SCALES))
        self.assertEqual(289, DEFAULT_MAX_PLACED_INSTANCES)


if __name__ == "__main__":
    unittest.main()
