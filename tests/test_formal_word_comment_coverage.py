import sys
from pathlib import Path
from xml.etree import ElementTree as ET
from docx import Document
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import formalize_cwh_report as report


def source():
    return {'comments': {'selected': [{'content': '希望更多劳动者都能使用这一制度。',
        'content_type': 'actual_comment', 'ai_formal_include': True, 'ai_comment_heading': '希望扩大制度覆盖',
        'system_topic': '政策甲', 'ai_semantic_quality': 'substantive', 'quote_verified': True,
        'url': 'https://example.test/comments', 'comment_id': 'c1', 'evidence_mode': 'verbatim_public_comment'}]}}


def xml(quote=None, appendix_quote=None):
    doc = Document()
    doc.add_paragraph(report.writing_rules()['document']['domestic_subsections'][1])
    if quote:
        para = doc.add_paragraph()
        para.add_run(quote[:5])
        para.add_run(quote[5:])
    doc.add_paragraph(report.writing_rules()['document']['domestic_subsections'][2])
    doc.add_paragraph('热词内容。')
    if appendix_quote:
        doc.add_paragraph(appendix_quote)
    return ET.tostring(doc.element)


def test_actual_quote_coverage_across_word_runs():
    data = source()
    quote = data['comments']['selected'][0]['content']
    result = report.audit_formal_comment_quotes(data, xml(quote))
    assert result['passed'] and result['unique_approved_quote_count'] == 1


def test_appendix_or_json_quote_does_not_mask_missing_domestic_prose():
    data = source()
    quote = data['comments']['selected'][0]['content']
    result = report.audit_formal_comment_quotes(data, xml(appendix_quote=quote))
    assert not result['passed'] and result['missing_from_word'] == [report.clean_formal_comment(quote)]


def test_invalid_shared_heading_falls_back_to_native_individual_heading():
    data = source()
    data['comments']['selected'][0]['topic_comment_heading'] = '政策甲部署情况'
    result = report.audit_formal_comment_quotes(data, xml(data['comments']['selected'][0]['content']))
    assert result['passed'] and not result['missing_from_grouping']


def test_native_approval_cannot_hide_a_grouping_loss():
    data = source()
    data['comments']['selected'][0]['ai_comment_heading'] = '政策甲部署情况'
    quote = data['comments']['selected'][0]['content']
    result = report.audit_formal_comment_quotes(data, xml(quote))
    assert not result['passed'] and result['missing_from_grouping'] == [report.clean_formal_comment(quote)]


def test_duplicate_approved_text_is_audited_once_without_changing_decisions():
    data = source()
    data['comments']['selected'].append(dict(data['comments']['selected'][0]))
    result = report.audit_formal_comment_quotes(data, xml(data['comments']['selected'][0]['content']))
    assert result['passed'] and result['approved_row_count'] == 2 and result['unique_approved_quote_count'] == 1
    assert len(data['comments']['selected']) == 2


def test_no_approved_quotes_requires_no_invented_prose():
    result = report.audit_formal_comment_quotes({}, xml())
    assert result['passed'] and result['approved_row_count'] == 0


def test_fixed_quote_cap_is_explicitly_accounted_for_not_called_missing():
    data = source()
    base = data['comments']['selected'][0]
    data['comments']['selected'] = [{**base, 'comment_id': str(i),
        'content': '希望更多劳动者都能使用这一制度，具体诉求' + str(i)} for i in range(4)]
    body = '；'.join(row['content'] for row in data['comments']['selected'][:3])
    result = report.audit_formal_comment_quotes(data, xml(body))
    assert result['passed'] and result['not_displayed_due_to_template_cap'] == [data['comments']['selected'][3]['content']]


def test_explicit_user_override_is_not_falsely_certified(monkeypatch):
    monkeypatch.setattr(report, 'section_override', lambda *args: '用户指定正文')
    result = report.audit_formal_comment_quotes(source(), xml())
    assert not result['applicable'] and 'passed' not in result


def test_different_reviewed_headings_in_one_topic_preserve_approved_quotes():
    import copy
    data = source()
    base = data['comments']['selected'][0]
    data['comments']['selected'].append({**base, 'comment_id': 'c2',
        'content': '具体管理办法是什么？什么时候公示',
        'ai_comment_heading': '希望明确管理办法及公示时间'})
    original = copy.deepcopy(data)
    groups = report.report_comment_groups(data)
    assert len(groups) == 1 and len(groups[0][1]) == 2
    assert groups[0][0] == '希望扩大制度覆盖；希望明确管理办法及公示时间'
    doc = Document()
    doc.add_paragraph(report.writing_rules()['document']['domestic_subsections'][1])
    report.add_comment_group_paragraph(doc, '一是', *groups[0])
    doc.add_paragraph(report.writing_rules()['document']['domestic_subsections'][2])
    audit = report.audit_formal_comment_quotes(data, ET.tostring(doc.element))
    assert audit['passed'] and audit['grouped_unique_quote_count'] == 2
    wording = report.comment_wording(groups[0][1])
    assert all(report.clean_formal_comment(row['content']) in wording for row in data['comments']['selected'])
    assert data == original
