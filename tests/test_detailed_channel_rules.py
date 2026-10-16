"""The new channel contract must change category totals without changing volume."""
import copy
import json
from pathlib import Path
import sys

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import raw_system_workbook_pipeline as pipeline

CONFIG = json.loads((Path(__file__).resolve().parents[1] / "config/raw_workbook_mapping.json").read_text("utf-8"))


def make_input(path, *, missing=None, extra=None, include_detail=True):
    wb = Workbook()
    legacy = wb.active
    legacy.title = "日趋势"
    legacy.append([names[0] for names in CONFIG["daily_columns"].values()])
    legacy.append(["2030-01-02"] + [9999] * (len(CONFIG["daily_columns"]) - 1))
    values = dict.fromkeys(CONFIG["detailed_daily_columns"], 0)
    values.update(date="2030-01-02", public_articles=2, public_recommend=3, public_comments=5,
        weibo=7, domestic_news=11, domestic_app=13, baijiahao=17, toutiao=19,
        forum_douban=2, forum_other=3, forum_xiaohongshu=5, forum_zhihu=7,
        video_huoshan=2, video_long=1, video_other_short=4, video_xigua=8,
        overseas_news=9, reddit=2, x=23, facebook=2, tiktok=3, youtube=4, video_account=29, douyin=31)
    if include_detail:
        ws = wb.create_sheet("细分渠道-日趋势")
        keys = list(reversed([key for key in values if key != missing]))
        ws.append([CONFIG["detailed_daily_columns"][key][0] for key in keys])
        ws.append([values[key] for key in keys])
        if extra:
            ws.cell(1, len(keys) + 1, extra)
            ws.cell(2, len(keys) + 1, 99999)
    wb.save(path)
    wb.close()
    return values


def test_detailed_sheet_and_labels_own_counts_not_legacy_sheet_or_column_positions(tmp_path):
    path = tmp_path / "total.xlsx"
    source = make_input(path, extra="境内新闻")
    audit = []
    rows = pipeline.read_daily(path, CONFIG, audit=audit)
    actual = pipeline.summarize_daily(rows, CONFIG["total_groups"])[0]
    assert actual == {"date": "2030-01-02", "domestic_mainstream": 11, "overseas_media": 9,
        "wechat": 10, "weibo": 7, "video_account": 29, "other_new_media": 146, "total": 212}
    assert rows[0]["domestic_app"] == 13
    assert rows[0]["baijiahao"] == 17 and rows[0]["toutiao"] == 19
    assert actual["total"] == sum(v for k, v in source.items() if k != "date")
    assert audit[0]["source_sheet"] == "细分渠道-日趋势"
    assert audit[0]["ignored_summary_columns"] == ["境内新闻"]
    assert audit[0]["rows"][0]["values"] == source
    assert pipeline.profile_raw_workbook(path, CONFIG)["heat_total"] == 212
    assert pipeline.summarize_daily(rows, CONFIG["child_groups"])[0]["new_media"] == 192


def test_missing_detail_sheet_never_falls_back_to_legacy(tmp_path):
    path = tmp_path / "total.xlsx"
    make_input(path, include_detail=False)
    with pytest.raises(pipeline.PipelineError, match="细分渠道"):
        pipeline.read_daily(path, CONFIG)
    with pytest.raises(pipeline.PipelineError, match="不能回退"):
        pipeline.profile_raw_workbook(path, CONFIG)


@pytest.mark.parametrize("missing", ["toutiao", "reddit", "domestic_news"])
def test_missing_channel_is_not_silently_zero(tmp_path, missing):
    path = tmp_path / "total.xlsx"
    make_input(path, missing=missing)
    with pytest.raises(pipeline.PipelineError, match="缺少字段"):
        pipeline.read_daily(path, CONFIG)


def test_new_unmapped_channel_requires_explicit_mapping(tmp_path):
    path = tmp_path / "total.xlsx"
    make_input(path, extra="其他视频-新增平台")
    with pytest.raises(pipeline.PipelineError, match="未映射渠道"):
        pipeline.read_daily(path, CONFIG)


def test_duplicate_rollup_assignment_is_rejected():
    config = copy.deepcopy(CONFIG)
    config["detailed_daily_rollup"]["domestic_news"].append("baijiahao")
    with pytest.raises(pipeline.PipelineError, match="遗漏或重复"):
        pipeline.daily_source_contract(config)


def test_same_channel_cannot_enter_two_report_categories():
    with pytest.raises(pipeline.PipelineError, match="重复"):
        pipeline.summarize_daily([], {"news": ["domestic_news"], "other": ["domestic_news"]})


def test_known_source_channel_cannot_be_omitted_from_report_groups():
    with pytest.raises(pipeline.PipelineError, match="覆盖所有"):
        pipeline.summarize_daily([{"date": "2030-01-02", "domestic_news": 2, "toutiao": 3}],
                                 {"news": ["domestic_news"]})


def test_total_row_cannot_be_counted_as_an_extra_day(tmp_path):
    path = tmp_path / "total.xlsx"
    make_input(path)
    wb = load_workbook(path)
    ws = wb["细分渠道-日趋势"]
    date_column = next(c.column for c in ws[1] if c.value == "日期")
    ws.cell(2, date_column, "合计")
    wb.save(path)
    wb.close()
    with pytest.raises(pipeline.PipelineError, match="有效日期"):
        pipeline.read_daily(path, CONFIG)


def test_overseas_reconciliation_preserves_all_domestic_detailed_counts(tmp_path):
    path = tmp_path / "total.xlsx"
    make_input(path)
    rows = pipeline.read_daily(path, CONFIG)
    original = copy.deepcopy(rows)
    unreviewed = pipeline.reconcile_overseas_daily_counts(rows, {1: rows}, {"counts_reconciled": False})
    assert not unreviewed["applied"] and unreviewed["total_daily"] == rows
    result = pipeline.reconcile_overseas_daily_counts(rows, {1: rows}, {
        "counts_reconciled": True, "counted": [{"published_at": "2030-01-02", "topic_hits": [1]}]})
    assert result["total_daily"][0] == {**original[0], "overseas_news": 1}
    assert result["child_daily"][1][0] == result["total_daily"][0]
    assert rows == original


def test_portable_builder_writes_every_source_channel_to_total_and_child(tmp_path):
    from PIL import Image
    from raw_system_workbook_portable import write_total_event, write_child
    path = tmp_path / "total.xlsx"
    make_input(path)
    rows = pipeline.read_daily(path, CONFIG)
    total = pipeline.summarize_daily(rows, CONFIG["total_groups"])[0]
    child = pipeline.summarize_daily(rows, CONFIG["child_groups"])[0]
    keys = pipeline.daily_output_columns(CONFIG)
    headers = [str(key) for key in keys]
    chart = tmp_path / "chart.png"
    Image.new("RGB", (20, 20), "white").save(chart)
    wb = Workbook()
    write_total_event(wb, {"total_event": {"daily": rows, "summary": [total], "totals": total},
        "canonical_daily_columns": keys, "daily_display_headers": headers}, chart)
    write_child(wb, {"index": 1, "title": "test", "daily": rows, "summary": [child], "totals": child}, headers, keys)
    for sheet, start in [("总事件", 15), ("子事件1", 8)]:
        for offset, key in enumerate(keys):
            assert wb[sheet].cell(3, start + offset).value == key
            if key != "date":
                assert wb[sheet].cell(4, start + offset).value == rows[0][key]
    assert wb["总事件"]["G4"].value == 146
    assert wb["子事件1"]["D4"].value == 192
    wb.close()
