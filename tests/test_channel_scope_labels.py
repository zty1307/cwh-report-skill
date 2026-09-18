"""Channel display scope must not silently promote news aggregates to mainstream."""
import copy
from datetime import datetime
import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import ingest_monitoring_workbook as ingest
import cwh_orchestrator as orchestrator
import formalize_cwh_report as formal
import generate_dashboard as dashboard
from cwh_writing_rules import domestic_media_label


@pytest.mark.parametrize('label', ['境内新闻', '境内主流媒体', '境内媒体'])
def test_ingestion_and_authority_keep_explicit_channel_scope_and_same_counts(label):
    wb = Workbook()
    ws = wb.active
    for column, value in enumerate(['日期', label, '境外媒体', '微信公众号', '新浪微博',
                                    '视频号', '新闻客户端、论坛等', '信息传播量'], 1):
        ws.cell(3, column, value)
    for column, value in enumerate([datetime(2026, 1, 2), 10, 1, 20, 3, 2, 4, 40], 1):
        ws.cell(4, column, value)
    parsed = ingest.parse_trend_sheet(ws)
    wb.close()
    system = {'overall': parsed, 'channel_labels': parsed['channel_labels'], 'subevents': []}
    original = copy.deepcopy(system)
    stats, _ = orchestrator.apply_system_authority(system, [], {})
    assert domestic_media_label({'statistics': stats}) == label
    assert stats['by_source_bucket']['domestic_media'] == 10
    assert stats['total_spread'] == 40
    assert system == original


@pytest.mark.parametrize('sentiment', [False, True])
def test_word_table_uses_same_explicit_label_for_both_shapes(sentiment):
    doc = Document()
    table = formal.add_topic_table(doc, [], sentiment, '境内新闻')
    assert table.cell(0, 2).text == '境内新闻'
    assert len(table.columns) == 9  # Missing ratios never change the fixed layout.


def test_markdown_body_table_and_dashboard_keep_scope_without_reclassifying_counts(tmp_path):
    data = {'meeting': {'date': '2026-01-02', 'topics': ['研究公共服务']},
            'statistics': {'total_spread': 40, 'by_source_bucket': {'domestic_media': 10},
                           'source_bucket_labels': {'domestic_media': '境内新闻'}},
            'topic_stats': [], 'comments': {'selected': []}, 'viewpoints': {'by_topic': []}}
    text = formal.render_formal_markdown(data, tmp_path)
    assert '境内新闻如人民网、新华网、央视网等均在显著位置刊文，共有相关报道10条。' in text
    assert '| 境内新闻 |' in text
    assert '境内主流媒体' not in text
    prepared = dashboard.prepare_dashboard_data(copy.deepcopy(data), tmp_path)
    assert prepared['domestic_media_label'] == '境内新闻'
    assert data['statistics']['by_source_bucket']['domestic_media'] == 10


def test_unknown_nonempty_scope_does_not_promote_or_echo_arbitrary_markup():
    assert domestic_media_label({'statistics': {'source_bucket_labels': {'domestic_media': '<未知>'}}}) == '境内媒体'
    assert domestic_media_label({'statistics': {}}) == '境内主流媒体'  # legacy compatibility, not certification
