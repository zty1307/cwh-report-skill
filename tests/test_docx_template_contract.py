from copy import deepcopy
from pathlib import Path
import sys
from unittest.mock import patch
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import formalize_cwh_report as formal
from cwh_docx_template import TEMPLATE, SLOTS, validate_template, import_reviewed_template
from cwh_chart_style import topic_chart_spec, valid_image, write_topic_chart
from cwh_writing_rules import opening_paragraph, writing_rules, editorial_profile


def test_shipped_template_has_unique_slots_and_only_explicit_editorial_defaults():
    document = Document(TEMPLATE)
    validate_template(document)
    text = '\n'.join(p.text for p in document.paragraphs)
    for stale in ('7月31', '7月10', '5月15', '2026', '这里有几个', '依然是有几个', '{{判断、依据及适用范围}}'):
        assert stale not in text
    assert editorial_profile()['document']['approved_chair_default'] in text
    assert '本次常务会引发境内外媒体广泛报道' in text
    assert '人民网、新华网、央视网等均在显著位置刊文' in text
    assert '{{具体观点}}' in text
    assert document.styles['Title'].font.name == '华文中宋'
    assert document.styles['Normal'].font.name == '仿宋_GB2312'


def test_missing_or_duplicate_slot_cannot_silently_ship():
    document = Document(TEMPLATE)
    paragraph = next(p for p in document.paragraphs if p._p.xpath('.//w:bookmarkStart'))
    document._element.body.insert(0, deepcopy(paragraph._p))
    with pytest.raises(ValueError, match='slots'):
        validate_template(document)


def test_unbound_editing_instruction_cannot_leak_into_report():
    document = Document(TEMPLATE)
    document.add_paragraph('这里是人工修改说明，不能进入正式报告')
    with pytest.raises(ValueError, match='Unbound'):
        validate_template(document)


def test_import_never_overwrites_or_accepts_unread_human_edit(tmp_path):
    source = tmp_path / 'human.docx'
    Document(TEMPLATE).save(source)
    before = source.read_bytes()
    with pytest.raises(ValueError, match='never be overwritten'):
        import_reviewed_template(source, source)
    with pytest.raises(ValueError, match='source changed'):
        import_reviewed_template(source, tmp_path / 'new.docx')
    assert source.read_bytes() == before


def test_approved_default_is_configurable_and_sourced_chair_overrides_it():
    assert opening_paragraph({}, '2月3日', '研究公共服务工作').startswith('国务院总理李强2月3日主持召开')
    assert opening_paragraph({'chair_name': '某姓名', 'chair_title': '国务院总理', 'chair_source': '本期通稿'},
                             '2月3日', '研究公共服务工作').startswith('国务院总理某姓名2月3日主持召开')


def test_all_current_topics_keep_their_slots_without_fabricated_comments():
    data = {'meeting': {'topics': ['议题甲', '议题乙', '议题丙']},
            'viewpoints': {'by_topic': [{'topic': '议题乙', 'clusters': []}]}, 'comments': {'selected': []}}
    before = deepcopy(data)
    assert [row['topic'] for row in formal.report_topic_items(data)] == ['议题甲', '议题乙', '议题丙']
    assert formal.report_comment_groups(data) == [('议题甲', []), ('议题乙', []), ('议题丙', [])]
    assert data == before


def test_editorial_wording_does_not_relabel_sentiment_or_invent_counts():
    data = {'statistics': {'total_spread': 41000}, 'comments': {'sentiment': {'negative': 20}}}
    before = deepcopy(data)
    assert formal.total_event_paragraphs(data)[0].startswith('本次常务会引发境内外媒体广泛报道，境内外传播总量约4.1万条')
    assert '以积极正面和客观中立为主' in formal.comment_lead(data, [('建议改善服务', [{'content': '已审原话'}])])
    assert data == before
    assert writing_rules()['viewpoint']['claim_char_range'] == [100, 200]
    assert writing_rules()['viewpoint']['minimum_claim_cjk'] == 30


def test_template_filling_preserves_all_claims_tables_and_relationships(tmp_path):
    rows = [{'speaker_name': f'机构{index}', 'formal_claim': f'第{index}项判断：' + '应保留实际实施条件和作用范围，' * 6,
             'evidence_id': f'e{index}'} for index in range(15)]
    cluster = {'summary': '建议完善实际实施条件', 'evidence': rows}
    data = {'meeting': {'date': '2030-02-03', 'topics': ['研究公共服务工作']},
            'statistics': {}, 'topic_stats': [],
            'viewpoints': {'by_topic': [{'topic': '公共服务', 'heading': '建议完善公共服务', 'clusters': [cluster]}]},
            'comments': {'selected': []}}
    original = deepcopy(data['viewpoints'])
    picture = tmp_path / 'image.png'; Image.new('RGB', (60, 40), 'navy').save(picture)
    output = tmp_path / 'report.docx'
    with patch.object(formal, 'ensure_docx_chart_images', return_value={'topic_distribution': str(picture)}):
        formal.write_docx(data, output)
    document = Document(output)
    text = '\n'.join(p.text for p in document.paragraphs)
    assert '{{' not in text and 'CWH_' not in text
    assert data['audit']['word_template']['filled_slots'] == list(SLOTS)
    assert len(document.tables) == 3 and len(document.inline_shapes) == 1
    assert max(len(p.text) for p in document.paragraphs) < 550
    assert data['viewpoints'] == original
    for row in rows:
        assert text.count(row['formal_claim'].rstrip('，。')) == 1
    with zipfile.ZipFile(output) as archive:
        assert b'r:embed' in archive.read('word/document.xml')
    assert document.paragraphs[0].style.name == 'Title'


def test_long_single_claim_retains_every_condition_without_hard_cut():
    claim = '如果' + '必要条件、' * 110 + '得到满足才可实施。'
    assert claim in ''.join(formal.cluster_detail_paragraphs({'evidence': [{'speaker_name': '甲', 'formal_claim': claim}]}))


def test_bold_attribution_uses_source_fields_not_verbs_inside_the_claim():
    document = Document()
    paragraph = document.add_paragraph()
    row = {'speaker_name': '测试账号', 'attribution_status': 'self_media',
           'url': 'https://mp.weixin.qq.com/s/example',
           'formal_claim': '原文提到实施条件尚需落实，不能将建议直接写为已经实现。'}
    formal.add_detail_blocks(document, paragraph, {'evidence': [row]})
    assert [run.text for run in paragraph.runs if run.bold] == ['微信公众号“测试账号”称']
    assert paragraph.text == formal.evidence_sentence(row).rstrip('。') + '。'


def test_fixed_chart_order_units_style_and_invalid_images(tmp_path):
    style, rows = topic_chart_spec({'小议题': 9000, '大议题': 42000, '同量议题': 42000})
    assert [row[0] for row in rows] == ['大议题', '同量议题', '小议题']
    assert [row[2] for row in rows] == ['4.2万', '4.2万', '0.9万']
    assert style['internal_title'] is False and style['bar_color'] == '#17365D'
    broken = tmp_path / 'broken.png'; broken.write_bytes(b'not a png')
    assert not valid_image(broken)
    path = tmp_path / 'topic.png'
    audit = write_topic_chart(path, {'完整政策标题不应截断' * 4: 42000, '乙议题': 9000})
    assert valid_image(path) and audit['origin'] == 'program_fallback'
    with Image.open(path) as image:
        assert image.width == 1100
        assert (23, 54, 93) in {rgb for _, rgb in image.getcolors(image.width * image.height)}


def test_raw_workbook_chart_uses_same_style_without_reordering_or_merging_input(tmp_path):
    from raw_system_workbook_portable import render_topic_chart
    children = [{'title': '重复标签', 'totals': {'total': 9000}},
                {'title': '完整政策名称' * 8, 'totals': {'total': 42000}},
                {'title': '重复标签', 'totals': {'total': 12000}}]
    before = deepcopy(children)
    path = tmp_path / 'raw-topic.png'
    audit = render_topic_chart(children, path)
    assert children == before
    assert audit['style'] == 'chart_style.v1/topic_distribution'
    assert audit['origin'] == 'program_raw_workbook_chart'
    assert [row['value'] for row in audit['ordered_values']] == [42000, 12000, 9000]
    assert audit['ordered_values'][0]['label'] == children[1]['title']
    assert [row['display_value'] for row in audit['ordered_values']] == ['4.2万', '1.2万', '0.9万']
    with Image.open(path) as image:
        assert image.width == 1100
        colors = {rgb for _, rgb in image.getcolors(image.width * image.height)}
        assert (23, 54, 93) in colors
        assert (47, 111, 115) not in colors


def test_native_workbook_chart_paths_read_the_shared_style_contract():
    root = Path(__file__).resolve().parents[1]
    builder = (root / 'scripts/raw_system_workbook_builder.mjs').read_text('utf-8')
    finalizer = (root / 'scripts/finalize_cwh_workbook_charts.ps1').read_text('utf-8-sig')
    assert 'config/chart_style.v1.json' in builder and 'series.fill = topicChartStyle.bar_color' in builder
    assert 'config/chart_style.v1.json' in finalizer and 'Fill.ForeColor.RGB = $topicOleColor' in finalizer
    assert "Substring(0, 4)" not in finalizer
    assert "$point.DataLabel.NumberFormat = '0.0' + $tenThousandFormatSuffix" in finalizer


def test_chart_config_is_utf8_in_production_windows_powershell():
    # Production invokes Windows PowerShell 5.1, whose Get-Content defaults to
    # the system ANSI code page, unlike pwsh. Execute the real assignment.
    import base64
    import shutil
    import subprocess

    root = Path(__file__).resolve().parents[1]
    source = (root / 'scripts/finalize_cwh_workbook_charts.ps1').read_text('utf-8-sig')
    assignment = next(line for line in source.splitlines() if line.startswith('$topicChartStyle ='))
    assert '-Encoding UTF8' in assignment
    shell = shutil.which('powershell.exe')
    if not shell:
        pytest.skip('Windows PowerShell production host is unavailable')
    scripts = str(root / 'scripts').replace("'", "''")
    code = ("$ErrorActionPreference='Stop'; $PSScriptRoot='" + scripts + "'; " + assignment
            + "; if ($topicChartStyle.label_unit -ne [string][char]19975) { throw 'Wrong unit encoding' };"
            + " if ($topicChartStyle.bar_color -ne '#17365D') { throw 'Wrong chart color' }")
    encoded = base64.b64encode(code.encode('utf-16le')).decode('ascii')
    result = subprocess.run([shell, '-NoProfile', '-EncodedCommand', encoded],
                            capture_output=True, timeout=30)
    assert result.returncode == 0, result.stderr.decode('utf-8', errors='replace')


def test_fixed_opening_accepts_sourced_role_and_learning_agenda():
    text = opening_paragraph({'chair_title': '国务院总理', 'chair_name': '测试姓名', 'chair_source': '本期通稿'},
                             '2月3日', '学习贯彻有关讲话精神，研究公共服务工作')
    assert text.startswith('国务院总理测试姓名2月3日主持召开国务院常务会议，学习贯彻')
    assert '涉及学习贯彻' not in text


def test_permitted_topic_fallback_requires_source_values_config_and_embedded_image(tmp_path):
    data = {'topic_stats': [{'topic': '公共服务', 'spread_count': 12000},
                            {'topic': '公共服务', 'spread_count': 8000}]}
    values = [('公共服务', 12000), ('公共服务', 8000)]
    topic = tmp_path / 'topic_distribution.png'
    manifest = write_topic_chart(topic, values)
    trend = tmp_path / 'trend_distribution_system.png'
    cloud = tmp_path / 'hotword_distribution_pipeline.png'
    Image.new('RGB', (20, 20), 'red').save(trend)
    Image.new('RGB', (20, 20), 'green').save(cloud)
    data['artifacts'] = {'chart_authority': 'mixed_system_and_program_fallback',
                        'docx_charts': dict(zip(('trend_distribution', 'topic_distribution', 'hotword_distribution'),
                                               map(str, (trend, topic, cloud))))}
    data['audit'] = {'topic_chart_style': manifest}
    doc = Document()
    for path in (trend, topic, cloud):
        doc.add_picture(str(path))
    path = tmp_path / 'chart_provenance.docx'
    doc.save(path)
    audit = formal.audit_formal_docx(data, path)
    assert audit['checks']['uses_traceable_fixed_chart_assets']
    assert audit['chart_provenance']['verified_fixed_topic_fallback']
    assert not audit['chart_provenance']['monitoring_system_assets']
    for field in ('image_sha256', 'style_sha256', 'ordered_values'):
        changed = deepcopy(data)
        changed['audit']['topic_chart_style'].pop(field)
        assert not formal.audit_formal_docx(changed, path)['checks']['uses_traceable_fixed_chart_assets']
    changed = deepcopy(data)
    changed['topic_stats'][1]['spread_count'] += 1
    assert not formal.audit_formal_docx(changed, path)['checks']['uses_traceable_fixed_chart_assets']
    Image.new('RGB', (20, 20), 'blue').save(topic)
    assert not formal.audit_formal_docx(data, path)['checks']['uses_traceable_fixed_chart_assets']


def test_propagation_does_not_invent_a_headline():
    data = {'meeting': {'focus_title': '未经核实的标题'}, 'statistics': {'total_spread': 41000}}
    assert '未经核实' not in formal.total_event_paragraphs(data)[0]
    data['meeting']['focus_source'] = '本期已核实来源'
    assert '“未经核实的标题”相关话题' in formal.total_event_paragraphs(data)[0]


def test_fixed_focus_can_use_current_topics_but_never_benchmark_facts():
    data = {'meeting': {'topics': ['本期议题甲', '本期议题乙']}, 'statistics': {'total_spread': 41000}}
    assert '舆论关注“本期议题甲”“本期议题乙”相关话题' in formal.total_event_paragraphs(data)[0]
    assert opening_paragraph({'agenda_text': '听取完整工作汇报，研究配套工作。', 'agenda_source': '本期通稿'},
                             '2月3日', '简略标签').endswith('听取完整工作汇报，研究配套工作。本次国务院常务会议舆情传播情况如下：')
