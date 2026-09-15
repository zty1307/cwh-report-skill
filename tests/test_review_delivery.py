import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from build_cwh_review_delivery import accepted_path, build_review_delivery
from cwh_pipeline_runtime import sha256_file
from run_cwh_inline_review import compact_hotword_packet, hotword_shortfall_packet, merge_hotword_supplement, response_object, normalize_hotword_transport, normalize_topic_hit_transport, validate_transport_result
from docx import Document


def test_review_snapshot_never_changes_gate_and_missing_counts_stay_unknown(tmp_path):
    state = {"status": "waiting_ai", "input_contract": {"agenda": "本次会议"},
             "stages": [{"stage_id": "workbook", "status": "waiting_ai", "message": "需要审核"}]}
    state_path = tmp_path / "pipeline_state.json"
    state_path.write_text(json.dumps(state), encoding="utf-8")
    original = state_path.read_bytes()
    result = build_review_delivery(tmp_path)
    assert state_path.read_bytes() == original
    assert result["status"] == "review_draft"
    payload = json.loads((tmp_path / "review_delivery/report_data.json").read_text(encoding="utf-8"))
    assert payload["audit"]["acceptance"]["ready_for_formal_delivery"] is False
    assert "暂不展示传播量" in "\n".join(p.text for p in Document(result["word"]).paragraphs)
    html = Path(result["dashboard"]).read_text(encoding="utf-8")
    assert "待审核稿 未通过正式交付" in html
    assert '"total_spread":null' in html


def test_modified_or_unaccepted_workbook_is_not_trusted(tmp_path):
    target = tmp_path / "data.xlsx"
    target.write_bytes(b"accepted")
    state = {"stages": [{"stage_id": "workbook", "status": "succeeded", "artifacts": {"workbook": str(target)},
                         "artifact_hashes": {"workbook": sha256_file(target)}}]}
    assert accepted_path(tmp_path, state, "workbook", "workbook") == target
    target.write_bytes(b"changed")
    assert accepted_path(tmp_path, state, "workbook", "workbook") is None


def test_fallback_lists_available_drafts_without_promoting_them_to_conclusions(tmp_path):
    artifacts = tmp_path / 'artifacts'
    artifacts.mkdir()
    draft = artifacts / 'analysis_bundle.json'
    draft.write_text('{"unverified_claim":"不能进入结论"}', encoding='utf-8')
    (tmp_path / 'pipeline_state.json').write_text(json.dumps({'status': 'failed', 'stages': []}), encoding='utf-8')
    result = build_review_delivery(tmp_path)
    manifest = json.loads((tmp_path / 'review_delivery/manifest.json').read_text(encoding='utf-8'))
    source = manifest['available_source_artifacts'][0]
    assert source['path'] == str(draft.resolve())
    assert source['status'] == 'unaccepted_source_not_report_conclusion'
    assert source['sha256'] == sha256_file(draft)
    text = '\n'.join(p.text for p in Document(result['word']).paragraphs)
    assert 'analysis_bundle.json' in text and '未验收过程材料' in text
    assert '不能进入结论' not in text


def test_fallback_keeps_existing_nested_builder_images_available(tmp_path):
    run = tmp_path / 'artifacts/run'
    run.mkdir(parents=True)
    picture = run / 'cwh_wordcloud.png'
    picture.write_bytes(b'unit index fixture, not rendered picture')
    (tmp_path / 'pipeline_state.json').write_text(json.dumps({'status': 'failed', 'stages': []}), encoding='utf-8')
    build_review_delivery(tmp_path)
    manifest = json.loads((tmp_path / 'review_delivery/manifest.json').read_text(encoding='utf-8'))
    assert any(item['path'] == str(picture.resolve()) for item in manifest['available_source_artifacts'])


def test_hotword_compaction_keeps_candidates_and_only_exact_source_windows():
    text = "前文" * 100 + "政策工具" + "后文" * 200
    packet = {"candidates": [{"term": "政策工具"}, {"term": "无证据"}],
              "document_samples": [{"source_row": 3, "content_excerpt": text}]}
    result = compact_hotword_packet(packet)
    assert result["candidates"] == packet["candidates"]
    hit = result["candidate_source_windows"][0]["source_windows"][0]
    assert hit["excerpt"] == text[hit["excerpt_start_in_packet"]:hit["excerpt_end_in_packet"]]
    assert result["candidate_source_windows"][1]["source_windows"] == []
    assert "document_samples" in packet


def test_inline_transport_accepts_cli_result_not_unrelated_tool_text():
    payload = {"review_method": "ai_semantic_review", "items": [{"record_id": "real-id"}]}
    result = {"type": "result", "result": json.dumps(payload)}
    assert response_object("notice\n" + json.dumps(result)) == payload
    assert response_object(json.dumps({"result": json.dumps({"output_shape": payload})})) == payload


def test_fixed_hotword_rejection_is_audited_not_silently_padded():
    original = {"review_method": "ai_semantic_review", "selected": [
        {"term": "城市更新"}, {"term": "中央城市工作会议"},
        {"term": "核准四个核电项目"}, {"term": "辽宁庄河核电"},
    ]}
    result = normalize_hotword_transport(original)
    assert [x["term"] for x in result["selected"]] == ["城市更新"]
    assert result["transport_exclusions"][0]["term"] == "中央城市工作会议"
    assert [x["reason"] for x in result["transport_exclusions"][1:]] == [
        "fixed_procedural_marker_rejected", "fixed_geography_rule_rejected"]
    assert len(original["selected"]) == 4
    import pytest
    with pytest.raises(ValueError, match="below"):
        validate_transport_result("hotword", {"minimum_term_count": 36}, result)


def test_topic_hit_transport_accepts_indices_exact_titles_and_unique_aliases():

    validate_transport_result("hotword", {
        "minimum_term_count": 36, "delivery_policy": "deliver_available_with_gaps"
    }, {"review_method": "ai_semantic_review", "selected": [{"term": "城市更新"}]})
    packet = {
        "topic_titles": ["新型电网建设", "核准四个核电项目"],
        "topic_aliases": [["新型电网"], ["核电项目"]],
    }
    result = {"items": [
        {"record_id": "a", "topic_hits": ["核电项目", "2", 2]},
        {"record_id": "b", "topic_hits": ["新型电网建设"]},
    ]}
    normalized = normalize_topic_hit_transport(packet, result, "overseas")
    assert normalized["items"][0]["topic_hits"] == [2]
    assert normalized["items"][1]["topic_hits"] == [1]
    assert result["items"][0]["topic_hits"] == ["核电项目", "2", 2]


def test_topic_hit_transport_rejects_ambiguous_or_unknown_labels():
    import pytest
    packet = {"topic_titles": ["议题一", "议题二"], "topic_aliases": [["共同词"], ["共同词"]]}
    with pytest.raises(ValueError, match="unknown or ambiguous"):
        normalize_topic_hit_transport(packet, {"items": [{"topic_hits": ["共同词"]}]}, "overseas")


def test_hotword_shortfall_packet_only_sends_unselected_valid_evidence_candidates():
    packet = {
        "topic_titles": ["议题一"], "minimum_term_count": 2, "target_term_count": 3,
        "candidates": [{"term": "城市更新"}, {"term": "政策工具"}, {"term": "召开"}, {"term": "无窗口"}],
        "candidate_source_windows": [
            {"term": "城市更新", "source_windows": [{"excerpt": "城市更新"}]},
            {"term": "政策工具", "source_windows": [{"excerpt": "政策工具"}]},
            {"term": "召开", "source_windows": [{"excerpt": "召开"}]},
        ],
    }
    result = hotword_shortfall_packet(packet, {"selected": [{"term": "城市更新"}]})
    assert [row["term"] for row in result["remaining_candidates"]] == ["政策工具"]
    assert result["already_selected_terms"] == ["城市更新"]


def test_hotword_supplement_is_capped_and_overflow_is_audited():
    original = {"selected": [{"term": "原词"}]}
    supplement = {"selected": [{"term": "补词一"}, {"term": "补词二"}, {"term": "越界词"}]}
    result = merge_hotword_supplement(original, supplement, {"补词一", "补词二", "越界词"}, 3)
    assert [row["term"] for row in result["selected"]] == ["原词", "补词一", "补词二"]
    assert result["transport_exclusions"][0]["term"] == "越界词"
    assert result["transport_exclusions"][0]["reason"] == "fixed_target_cap"
    assert len(original["selected"]) == 1


def test_hotword_shortfall_preserves_semantic_contract_and_topic_mapping():
    packet = {"topics": [{"index": 1, "title": "议题一", "aliases": ["别名"]}],
              "instructions": ["不得凑数"], "allowed_semantic_types": ["policy_tool"],
              "accepted_style_patterns": ["完整短语"], "rejected_style_patterns": ["泛词"]}
    result = hotword_shortfall_packet(packet, {"selected": []})
    assert result["topic_titles"] == ["议题一"]
    assert result["topic_aliases"] == [["别名"]]
    for key in ("instructions", "allowed_semantic_types", "accepted_style_patterns", "rejected_style_patterns"):
        assert result[key] == packet[key]


def test_overseas_span_transport_preserves_original_and_rejects_cross_record_ids():
    import pytest
    from run_cwh_inline_review import overseas_span_packet, resolve_overseas_spans
    packet = {"items": [{"record_id": "a", "content": "政策��點。繁體原文，不能改。"},
                        {"record_id": "b", "content": "另一條原文。"}]}
    sent = overseas_span_packet(packet)
    assert "content" not in sent["items"][0]
    assert "".join(x["text"] for x in sent["items"][0]["interpretive_segments"]) == packet["items"][0]["content"]
    review = {"items": [{"record_id": "a", "interpretive_verified": True, "interpretive_range": ["o1/1", "o1/2"]}]}
    result = resolve_overseas_spans(packet, review)
    assert result["items"][0]["interpretive_excerpt"] == packet["items"][0]["content"]
    assert "interpretive_excerpt" not in review["items"][0]
    with pytest.raises(ValueError, match="belong"):
        resolve_overseas_spans(packet, {"items": [{"record_id": "a", "interpretive_range": ["o2/1", "o2/1"]}]})


def test_missing_or_title_only_overseas_body_is_structurally_excluded_and_audited():
    from run_cwh_inline_review import resolve_overseas_spans
    packet = {"items": [{"record_id": "a", "title": "國務院常務會議", "content": "国 务 院 常 务 会 议"}]}
    review = {"items": [{"record_id": "a", "decision": "include", "review_reason": "模型认为有关"}]}
    result = resolve_overseas_spans(packet, review)
    assert result["items"][0]["decision"] == "exclude"
    assert result["items"][0]["transport_exclusions"][0]["original_review"] == review["items"][0]
    assert review["items"][0]["decision"] == "include"
