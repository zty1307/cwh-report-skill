"""Deterministic font/layout checks; not a rendered Word pagination certificate."""
from pathlib import Path
import hashlib
import shutil
import sys

from docx import Document
from docx.shared import Inches
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import cwh_hotword_pipeline as hotwords
from formalize_cwh_report import add_centered_picture


def test_process_font_override_precedes_package_fallback(monkeypatch, tmp_path):
    approved = tmp_path / 'approved-font.otf'
    approved.touch()
    monkeypatch.setenv('CWH_CJK_FONT', str(approved))
    assert hotwords.resolve_font(None) == approved
    assert '第九周' not in str(hotwords.DEFAULT_FONT_PATH)


def test_packaged_tencent_font_renders_after_relocation(monkeypatch, tmp_path):
    original = Path(hotwords.__file__).resolve().parents[1] / 'assets/fonts/TencentFont.otf'
    assert original.is_file(), 'The distributed Skill must include its default Tencent font'
    relocated = tmp_path / '另一台机器的 skill' / 'assets/fonts/TencentFont.otf'
    relocated.parent.mkdir(parents=True)
    shutil.copy2(original, relocated)
    monkeypatch.delenv('CWH_CJK_FONT', raising=False)
    monkeypatch.setattr(hotwords, 'BUNDLED_CJK_FONT', relocated)
    monkeypatch.setattr(hotwords, 'BUNDLED_CJK_FALLBACK_FONT', tmp_path / 'no-fallback.otf')
    assert hotwords.resolve_font(None) == relocated
    payload = {
        'settings': {'width': 800, 'height': 600, 'primary_font_min': 24,
                     'primary_font_max': 48, 'max_placed_instances': 2},
        'selected': [{'term': '公共服务', 'weight': 100}, {'term': '经济发展', 'weight': 80}],
    }
    output = tmp_path / 'wordcloud.png'
    hotwords.render_wordcloud_png(payload, output)
    audit = payload['render_audit']
    assert audit['font_family'] == ['TTTGB', 'Medium']
    assert audit['font_sha256'] == hashlib.sha256(original.read_bytes()).hexdigest()
    assert audit['font_sha256'] == '6dff84548a2bb4ad7ba30e46e86f2f02bcf896e367b38080b3ae46836ed81be0'
    with Image.open(output) as rendered:
        assert rendered.getbbox() is not None


def test_inline_cloud_preserves_aspect_ratio_and_caps_height(tmp_path):
    fixture = tmp_path / 'image.png'
    Image.new('RGB', (1600, 1000), 'white').save(fixture)
    doc = Document()
    shape = add_centered_picture(doc, str(fixture), Inches(5.8), '词云', max_height=Inches(3.0))
    assert shape.height == Inches(3.0)
    assert abs(shape.width / shape.height - 1.6) < 1e-5
    assert doc.paragraphs[-1].paragraph_format.line_spacing == 1
    assert not doc.paragraphs[-1]._p.xpath('.//w:br')


def test_cloud_limits_do_not_enlarge_small_graphics(tmp_path):
    fixture = tmp_path / 'small.png'
    Image.new('RGB', (1600, 1000), 'white').save(fixture)
    doc = Document()
    shape = add_centered_picture(doc, str(fixture), Inches(2), '词云', max_height=Inches(3.0))
    assert shape.width == Inches(2)
