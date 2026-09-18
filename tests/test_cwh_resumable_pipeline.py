from __future__ import annotations

import importlib.util
import hashlib
import json
import socket
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook


SCRIPT_DIR = Path(__file__).resolve().parents[1] / "scripts"
for value in (str(SCRIPT_DIR),):
    if value not in sys.path:
        sys.path.insert(0, value)
SCRIPT = SCRIPT_DIR / "run_cwh_resumable_pipeline.py"
SPEC = importlib.util.spec_from_file_location("run_cwh_resumable_pipeline", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


class CwhResumablePipelineTests(unittest.TestCase):
    def test_sentiment_validation_accepts_header_only_handoff_after_audited_zero_results(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            summary = root / "summary.json"
            handoff = root / "handoff.csv"
            topics = ["议题甲"]
            registry = json.loads(MODULE.SOURCE_REGISTRY_PATH.read_text(encoding="utf-8"))
            required = [row["id"] for row in registry["sources"]
                        if row.get("must_check") and row.get("tier") == "comment_platform"]
            topic_checks = []
            global_checks = []
            for source_id in required:
                status = "no_relevant_result" if source_id == "toutiao_public_comments" else "access_failed"
                check = {"source_id": source_id, "status": status,
                         "execution_mode": "fixed_test", "queries_or_seed_urls": ["真实查询"],
                         "result_count": 0, "eligible_comment_ids": []}
                if status == "access_failed":
                    check["blocker"] = "no authorized login"
                topic_checks.append(check)
                global_checks.append({"source_id": source_id, "status": status})
            summary.write_text(json.dumps({
                "collection_audit": {"registry_version": registry["version"],
                                     "terminal_status": "bounded_checks_completed",
                                     "waiting_login_terminal": False,
                                     "checks": global_checks,
                                     "coverage_by_topic": [{"topic": "议题甲", "checks": topic_checks}]},
                "topics": [{"title": "议题甲", "status": "no_public_evidence", "denominator": 0}],
            }, ensure_ascii=False), encoding="utf-8")
            handoff.write_text(
                "sample_id,topic,content,url,ai_formal_include,topic_comment_heading\n",
                encoding="utf-8-sig",
            )
            self.assertEqual([], MODULE.validate_sentiment(summary, handoff, topics))

    def test_sentiment_validation_allows_audited_zero_or_insufficient_topics(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            summary = root / "summary.json"
            handoff = root / "handoff.csv"
            topics = ["议题甲", "议题乙"]
            summary.write_text(
                json.dumps(
                    {
                        "collection_audit": {
                            "registry_version": json.loads(
                                MODULE.SOURCE_REGISTRY_PATH.read_text(encoding="utf-8")
                            )["version"],
                            "checks": [
                                {"source_id": "toutiao_public_comments", "status": "hit"},
                                {"source_id": "weibo_comments", "status": "access_failed"},
                                {"source_id": "douyin_comments", "status": "access_failed"},
                                {"source_id": "bilibili_comments", "status": "access_failed"},
                            ],
                            "coverage_by_topic": [
                                {
                                    "topic": topic,
                                    "checks": [
                                        {
                                            "source_id": "toutiao_public_comments",
                                            "status": "hit" if topic == "议题乙" else "no_relevant_result",
                                            "execution_mode": "reviewed_seed_then_public_comment_api",
                                            "queries_or_seed_urls": [f"https://example.test/{topic}"],
                                            "result_count": 1 if topic == "议题乙" else 0,
                                            "eligible_comment_ids": ["c1"] if topic == "议题乙" else [],
                                        },
                                        {"source_id": "weibo_comments", "status": "access_failed", "blocker": "no approved login"},
                                        {"source_id": "douyin_comments", "status": "access_failed", "blocker": "no approved login"},
                                        {"source_id": "bilibili_comments", "status": "access_failed", "blocker": "no approved login"},
                                    ],
                                }
                                for topic in topics
                            ],
                        },
                        "topics": [
                            {"title": "议题甲", "status": "pending", "denominator": 0},
                            {"title": "议题乙", "status": "insufficient_sample", "denominator": 3},
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            handoff.write_text(
                "sample_id,topic,content,url,ai_formal_include,topic_comment_heading\n"
                "c1,议题乙,示例评论,https://example.test/comment,true,认可政策方向\n",
                encoding="utf-8-sig",
            )
            self.assertEqual([], MODULE.validate_sentiment(summary, handoff, topics))

            payload = json.loads(summary.read_text(encoding="utf-8"))
            payload["collection_audit"]["waiting_login_terminal"] = True
            payload["collection_audit"]["terminal_status"] = "waiting_login"
            summary.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            problems = MODULE.validate_sentiment(summary, handoff, topics)
            self.assertTrue(any("waiting_login当成终点" in item for item in problems))

    def test_foreign_validation_accepts_verified_completed_zero_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audit = root / "foreign_audit.json"
            supplements = root / "supplements.json"
            audit.write_text(
                json.dumps(
                    {
                        "dry_run": False,
                        "command_exit_code": 4,
                        "collection_completed": True,
                        "collection_status": "completed_no_rows",
                        "zero_result": True,
                        "collector_exit_code": 0,
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            supplements.write_text(
                json.dumps({"media": [], "comments": []}, ensure_ascii=False),
                encoding="utf-8",
            )
            self.assertEqual([], MODULE.validate_foreign(audit, supplements))

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.workbook = self.root / "system.xlsx"
        self.make_workbook(self.workbook)
        self.job = self.root / "job"
        self.contract = {
            "agenda": "2026年7月10日国务院常务会议，听取数字中国建设情况汇报",
            "system_workbook": str(self.workbook),
            "raw_input_dir": "",
            "metadata": "",
            "analysis_bundle": "",
            "comment_handoff": "",
            "sentiment_results": "",
            "sentiment_summary": "",
            "overseas_supplements": "",
            "foreign_collection_audit": "",
            "hotword_audit": "",
            "ai_worker_command": [],
            "source_contracts": {"system_workbook": MODULE.file_contract(str(self.workbook))},
        }

    def tearDown(self) -> None:
        self.temp.cleanup()

    @staticmethod
    def domestic_research_audit(topics: list[str]) -> dict:
        registry = MODULE.read_json(MODULE.SOURCE_REGISTRY_PATH)
        source_text = "测试媒体认为，数字基础设施建设应与安全治理同步推进，并保留需求研判和风险防控机制。"
        required = [
            row
            for row in registry["sources"]
            if row.get("must_check") and row.get("region") == "domestic" and row.get("tier") != "comment_platform"
        ]
        platform_tiers = set(registry["execution"].get("platform_specific_search_required_for_tiers") or [])
        return {
            "domestic_media_research": {
                "registry_version": registry["version"],
                "open_search_completed": True,
                "pool_summary": {
                    "raw_monitoring_candidate_count": 1,
                    "fixed_registry_web_candidate_count": 0,
                    "open_web_candidate_count": 1,
                    "public_platform_web_candidate_count": 0,
                },
                "coverage_by_topic": [
                    {
                        "topic": topic,
                        "checks": [
                            {
                                "source_id": source["id"],
                                "status": "no_relevant_result",
                                "execution_mode": (
                                    "platform_specific"
                                    if source.get("tier") in platform_tiers
                                    else "site_restricted_search"
                                ),
                                "queries": [f"site:{(source.get('domains') or ['example.com'])[0]} {topic}"],
                            }
                            for source in required
                        ],
                    }
                    for topic in topics
                ],
                "candidate_pool_by_topic": [
                    {
                        "topic": topic,
                        "candidates": [
                            {
                                "candidate_id": "candidate-1",
                                "discovery_origin": "open_web",
                                "discovery_query_id": "query-1",
                                "first_seen_round": 1,
                                "viewpoint_cluster_key": "development_and_security",
                                "source": "测试媒体",
                                "title": "数字基础设施建设应统筹发展与安全",
                                "url": "https://example.com/a",
                                "published_at": "2026-07-10",
                                "published_at_source_text": "2026-07-10 09:30",
                                "published_at_verified_from_source": True,
                                "discovery_route": "open_search:test_query",
                                "content_summary": "数字基础设施建设应与安全治理同步推进，并保留需求研判、风险防控和长效评估机制。",
                                "source_type": "mainstream_media",
                                "source_tier": "mainstream",
                                "decision": "eligible",
                                "decision_reason": "原文包含可核验的独立媒体观点。",
                                "source_snapshot": {
                                    "snapshot_id": "snapshot-candidate-1",
                                    "url": "https://example.com/a",
                                    "title": "数字基础设施建设应统筹发展与安全",
                                    "captured_at": "2026-07-10T12:00:00+08:00",
                                    "capture_method": "test_fixture",
                                    "source_text": source_text,
                                    "source_text_sha256": hashlib.sha256(source_text.encode("utf-8")).hexdigest(),
                                },
                            }
                        ],
                        "saturation": {
                            "completed": True,
                            "stop_reason": MODULE.DOMESTIC_SATURATION_STOP,
                            "route_coverage": [
                                {"route": "open_web", "status": "completed"},
                                {"route": "public_platform", "status": "completed"},
                            ],
                            "rounds": [
                                {
                                    "round": 1,
                                    "queries_or_sources": [f"{topic} 专家 解读"],
                                    "executions": [{
                                        "query_id": "query-1",
                                        "query": f"{topic} 专家 解读",
                                        "backend": "test_search",
                                        "route": "open_web",
                                        "status": "completed",
                                        "executed_at": "2026-07-10T12:00:00+08:00",
                                        "result_count": 1,
                                        "result_urls": ["https://example.com/a"],
                                        "retained_candidate_ids": ["candidate-1"],
                                    }],
                                    "new_candidates": 1,
                                    "new_eligible_candidates": 1,
                                    "new_independent_viewpoints": 1,
                                },
                                {
                                    "round": 2,
                                    "queries_or_sources": [f"{topic} 媒体 评论 建议"],
                                    "executions": [{
                                        "query_id": "query-2",
                                        "query": f"{topic} 媒体 评论 建议",
                                        "backend": "test_search",
                                        "route": "open_web",
                                        "status": "completed",
                                        "executed_at": "2026-07-10T12:05:00+08:00",
                                        "result_count": 0,
                                        "result_urls": [],
                                        "retained_candidate_ids": [],
                                    }],
                                    "new_candidates": 0,
                                    "new_eligible_candidates": 0,
                                    "new_independent_viewpoints": 0,
                                },
                                {
                                    "round": 3,
                                    "queries_or_sources": [f"site:toutiao.com {topic} 观点"],
                                    "executions": [{
                                        "query_id": "query-3",
                                        "query": f"site:toutiao.com {topic} 观点",
                                        "backend": "test_search",
                                        "route": "public_platform",
                                        "status": "completed",
                                        "executed_at": "2026-07-10T12:10:00+08:00",
                                        "result_count": 0,
                                        "result_urls": [],
                                        "retained_candidate_ids": [],
                                    }],
                                    "new_candidates": 0,
                                    "new_eligible_candidates": 0,
                                    "new_independent_viewpoints": 0,
                                },
                            ],
                        },
                    }
                    for topic in topics
                ],
            }
        }

    @staticmethod
    def make_workbook(path: Path) -> None:
        workbook = Workbook()
        keywords = workbook.active
        keywords.title = "关键词"
        keywords.cell(2, 2, "2026年7月10日国务院常务会议")
        keywords.cell(4, 3, "子事件")
        keywords.cell(5, 1, 1)
        keywords.cell(5, 2, "听取数字中国建设情况汇报")
        total = workbook.create_sheet("总事件")
        total.cell(2, 1, "2026年7月10日国务院常务会议")
        headers = ["日期", "境内主流媒体", "境外媒体", "微信公众号", "微博", "视频号", "新媒体", "总量"]
        for column, value in enumerate(headers, 1):
            total.cell(4, column, value)
        values = ["2026-07-10", 1, 1, 2, 3, 4, 5, 16]
        for column, value in enumerate(values, 1):
            total.cell(5, column, value)
        summary = workbook.create_sheet("子事件数据汇总")
        summary.append(["序号", "标题", "境内主流媒体", "新媒体", "境外媒体", "总量", "正面", "中立", "负面"])
        summary.append([1, "听取数字中国建设情况汇报", 1, 2, 1, 4, None, None, None])
        workbook.save(path)
        workbook.close()

    def test_stops_at_ai_stage_and_resumes_without_rebuilding_workbook(self) -> None:
        pipeline = MODULE.CwhPipeline(self.job, self.contract)
        first = pipeline.runner.run()
        self.assertEqual("waiting_ai", first["status"])
        self.assertEqual("domestic_viewpoints", first["current_stage"])
        workbook_state = pipeline.runner.stage_state("workbook")
        workbook_hash = workbook_state["output_fingerprint"]
        viewpoint_task = self.job / "tasks/domestic_viewpoints.json"
        self.assertTrue(viewpoint_task.exists())
        task_rules = " ".join(MODULE.read_json(viewpoint_task)["rules"])
        self.assertIn("候选池", task_rules)
        self.assertIn("查询/抓取上限", task_rules)
        self.assertIn("每议题常规4至6处引述", task_rules)
        self.assertIn("最多12个独立主体仅是异常兜底上限", task_rules)
        self.assertIn("formal_use=reserve", task_rules)

        analysis = {
            "metadata": {
                "method": "ai_semantic_source_review",
                "monitoring_start": "2026-07-10",
                "monitoring_end": "2026-07-10",
                "evidence_mapping_version": "1.0",
                "authoring_run_id": "author-run-test",
            },
            "research_audit": self.domestic_research_audit(["听取数字中国建设情况汇报"]),
            "viewpoints": {
                "by_topic": [
                    {
                        "topic": "听取数字中国建设情况汇报",
                        "heading": "认可数字中国建设统筹发展与安全",
                        "single_cluster_exception": {
                            "reason": "完整候选池仅形成一个独立观点家族。",
                            "search_evidence": "query-1至query-3",
                            "reviewed_by": "test",
                        },
                        "clusters": [
                            {
                                "summary": "认为数字中国建设应统筹发展与安全",
                                "details": "测试媒体认为，数字基础设施建设应与安全治理同步推进，并保留需求研判和风险防控机制。",
                                "thin_cluster_exception": {
                                    "reason": "监测期内仅检得一个可追溯独立声音。",
                                    "search_evidence": "query-1至query-3",
                                    "reviewed_by": "test",
                                },
                                "evidence": [
                                    {
                                        "source": "测试媒体",
                                        "url": "https://example.com/a",
                                        "candidate_id": "candidate-1",
                                        "selection_reason": "与议题直接相关且包含独立、完整的媒体判断。",
                                        "attribution": "测试媒体",
                                        "attribution_status": "media_voice",
                                        "source_excerpt": "测试媒体认为，数字基础设施建设应与安全治理同步推进，并保留需求研判和风险防控机制。",
                                        "formal_claim": "数字基础设施建设应与安全治理同步推进，并保留需求研判和风险防控机制。",
                                        "wording_fidelity": "faithful_paraphrase",
                                        "evidence_id": "evidence-candidate-1-test-media",
                                        "source_snapshot_id": "snapshot-candidate-1",
                                        "article_title": "数字基础设施建设应统筹发展与安全",
                                        "speaker_name": "测试媒体",
                                        "speaker_role": "",
                                        "source_excerpt_start": 0,
                                        "source_excerpt_end": 41,
                                        "semantic_review": {
                                            "verdict": "fully_supported",
                                            "reviewed_by": "test",
                                            "reviewed_at": "2026-07-10T12:30:00+08:00",
                                            "rationale": "观点逐字来自连续原文。",
                                            "propositions": [
                                                {
                                                    "text": "数字基础设施建设应与安全治理同步推进，并保留需求研判和风险防控机制。",
                                                    "verdict": "fully_supported",
                                                    "source_quote": "测试媒体认为，数字基础设施建设应与安全治理同步推进，并保留需求研判和风险防控机制。",
                                                    "source_quote_start": 0,
                                                    "source_quote_end": 41,
                                                    "rationale": "连续原文直接支持。",
                                                }
                                            ],
                                        },
                                    }
                                ],
                            }
                        ],
                    }
                ]
            },
        }
        MODULE.atomic_write_json(self.job / "artifacts/analysis_bundle.json", analysis)
        second = pipeline.runner.run()
        self.assertEqual("waiting_ai", second["status"])
        self.assertEqual("domestic_evidence_verification", second["current_stage"])
        verification_task = MODULE.read_json(self.job / "tasks/domestic_evidence_verification.json")
        self.assertIn("独立第二遍", " ".join(verification_task["rules"]))
        self.assertEqual(workbook_hash, pipeline.runner.stage_state("workbook")["output_fingerprint"])
        self.assertEqual(1, pipeline.runner.stage_state("workbook")["attempts"])

        # Exercise the automatic worker return path with the same evidence,
        # without author self-certification. It must reach the review stage.
        for item in analysis["viewpoints"]["by_topic"]:
            for cluster in item["clusters"]:
                for evidence in cluster["evidence"]:
                    evidence.pop("semantic_review", None)
        automatic = MODULE.CwhPipeline(self.job / "automatic", self.contract)

        def write_draft(runner, spec, task, output):
            MODULE.atomic_write_json(output, analysis)
            return None

        with mock.patch.object(MODULE, "topic_titles", return_value=["听取数字中国建设情况汇报"]), mock.patch.object(MODULE, "maybe_run_ai_worker", side_effect=write_draft):
            result = automatic.domestic_viewpoints(automatic.runner, automatic.runner.spec_by_id["domestic_viewpoints"])
        self.assertEqual("succeeded", result.status, result.message)
        unreviewed = automatic.artifacts / "analysis_bundle.json"
        self.assertTrue(MODULE.validate_analysis_bundle(unreviewed, ["听取数字中国建设情况汇报"], require_semantic_review=True))

    def test_data_workbook_gate_rejects_comment_inflated_spread_total(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            path = root / "cwh_data_workbook.xlsx"
            workbook = Workbook()
            summary = workbook.active
            summary.title = "子事件数据汇总"
            summary.append(["序号", "标题", "境内主流媒体", "新媒体", "境外媒体", "总量", "正面", "中立", "负面", "合计"])
            summary.append([])
            summary.append([1, "议题甲", 10, 22, 1, 33, "", "", "", ""])
            total = workbook.create_sheet("总事件")
            child = workbook.create_sheet("子事件1")
            workbook.save(path)
            workbook.close()
            report_data = {
                "topic_stats": [{
                    "domestic_media": 10,
                    "self_media": 20,
                    "overseas_media": 1,
                    "spread_count": 31,
                    "comments": 2,
                    "sentiment_formal_ready": False,
                }],
                "system_data": {
                    "overall": {},
                    "subevents": [{"index": 1, "daily": [], "totals": {}}],
                },
            }
            problems = MODULE.validate_data_workbook(path, report_data)
            self.assertTrue(any("汇总C:F" in item for item in problems))
            self.assertTrue(any("缺少权威逐日数据" in item for item in problems))

    def test_waiting_state_does_not_consume_retry_budget(self) -> None:
        pipeline = MODULE.CwhPipeline(self.job, self.contract)
        for _ in range(5):
            state = pipeline.runner.run()
            self.assertEqual("waiting_ai", state["status"])
        stage = pipeline.runner.stage_state("domestic_viewpoints")
        self.assertEqual(0, stage["attempts"])
        self.assertEqual(5, stage["wait_count"])

    def test_comment_worker_may_return_to_login_wait_without_false_failure(self) -> None:
        runner = mock.Mock()
        runner.root = self.job
        runner.input_contract = {"ai_worker_command": ["worker", "{task}"]}
        runner.remaining_budget_seconds.return_value = None
        runner.run_command.return_value = (0, self.job / "worker.log")
        spec = MODULE.StageSpec("domestic_comments_sentiment", "comments")
        task = self.job / "tasks/comments.json"
        output = self.job / "artifacts/sentiment_workbook_summary.json"

        outcome = MODULE.maybe_run_ai_worker(
            runner,
            spec,
            task,
            output,
            allow_missing_output=True,
        )

        self.assertIsNone(outcome)

    def test_analysis_validator_rejects_missing_topic_and_links(self) -> None:
        path = self.root / "analysis.json"
        path.write_text(
            json.dumps(
                {
                    "viewpoints": {
                        "by_topic": [
                            {
                                "topic": "听取数字中国建设情况汇报",
                                "clusters": [{"summary": "认为应完善建设", "details": "证据", "evidence": [{}]}],
                            }
                        ]
                    }
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        problems = MODULE.validate_analysis_bundle(
            path,
            ["听取数字中国建设情况汇报", "研究自然资源保护利用有关工作"],
        )
        self.assertTrue(any("没有原始链接" in item for item in problems))
        self.assertTrue(any("缺少子议题观点" in item for item in problems))

    def test_skill_declares_one_durable_pipeline_entry(self) -> None:
        skill = (SCRIPT_DIR.parent / "SKILL.md").read_text(encoding="utf-8")
        reference = (SCRIPT_DIR.parent / "references" / "resumable_pipeline.md").read_text(encoding="utf-8")
        self.assertIn("run_cwh_resumable_pipeline.py", skill)
        self.assertIn("references/resumable_pipeline.md", skill)
        self.assertIn("waiting_ai", reference)
        self.assertIn("能用脚本完成", reference)

    def test_research_plan_uses_versioned_sources_then_open_search(self) -> None:
        sys.path.insert(0, str(SCRIPT_DIR))
        from build_research_plan import build_plan

        plan = build_plan(str(self.workbook), self.contract["agenda"])
        self.assertEqual("registry_first_then_open_search", plan["source_registry"]["mode"])
        source_ids = {row["source_id"] for row in plan["topics"][0]["stable_source_tasks"]}
        self.assertIn("toutiao_articles", source_ids)
        self.assertIn("baijiahao", source_ids)
        self.assertIn("lane_authoritative", source_ids)
        self.assertIn("lane_mainstream", source_ids)
        self.assertIn("lane_industry_expert", source_ids)
        self.assertTrue(plan["research_audit_contract"]["open_search_required"])
        self.assertTrue(plan["research_audit_contract"]["candidate_pool_required"])
        self.assertTrue(plan["research_audit_contract"]["saturation_required"])
        contract = plan["topics"][0]["candidate_pool_contract"]
        self.assertEqual("priority_seed_not_allowlist", contract["registry_role"])
        self.assertEqual(1, contract["saturation_rule"]["required_zero_new_rounds"])
        self.assertEqual(
            ["open_web", "public_platform"],
            contract["saturation_rule"]["required_route_coverage"],
        )
        toutiao = next(row for row in plan["topics"][0]["stable_source_tasks"] if row["source_id"] == "toutiao_articles")
        self.assertEqual("platform_specific", toutiao["execution_mode"])
        self.assertEqual(1, len(toutiao["query_families"]))
        baijiahao = next(row for row in plan["topics"][0]["stable_source_tasks"] if row["source_id"] == "baijiahao")
        self.assertTrue(baijiahao["must_check"])
        self.assertEqual("platform_specific", baijiahao["execution_mode"])
        self.assertIn("executions", contract["saturation_rule"]["round_fields"])
        self.assertIn("result_urls", contract["saturation_rule"]["execution_fields"])
        self.assertIsNone(plan["topics"][0]["minimum_evidence"]["fixed_result_target"])
        self.assertEqual("bounded_60m", plan["execution_profile"])
        self.assertEqual(12, contract["max_formal_voices_per_topic"])
        self.assertEqual(
            "bounded_selected_eligible_with_audited_reserve",
            plan["topics"][0]["minimum_evidence"]["formal_sources_per_mature_cluster"],
        )
        self.assertEqual(
            "normally_two_complementary_voices_with_audited_exceptions",
            plan["topics"][0]["domestic_viewpoint_contract"]["topic_density"]["independent_voices_per_cluster"],
        )
        self.assertEqual("all_workbook_topics", plan["global_tasks"]["public_comments"]["target_topics"])
        self.assertEqual([2, 3], plan["global_tasks"]["public_comments"]["target_quotes"])
        self.assertEqual(6, plan["global_tasks"]["public_comments"]["max_query_executions_per_topic"])
        self.assertNotIn("reddit", plan["global_tasks"]["overseas"]["mediaspider_foreign"]["platforms"])
        self.assertIn("reddit", plan["global_tasks"]["overseas"]["mediaspider_foreign"]["paused_platforms"])

        exhaustive = build_plan(str(self.workbook), self.contract["agenda"], "exhaustive")
        exhaustive_ids = {row["source_id"] for row in exhaustive["topics"][0]["stable_source_tasks"]}
        self.assertIn("xinhua", exhaustive_ids)
        self.assertEqual(2, exhaustive["topics"][0]["candidate_pool_contract"]["saturation_rule"]["required_zero_new_rounds"])
        self.assertEqual(
            "all_eligible_independent_samples_no_upper_cap",
            exhaustive["topics"][0]["minimum_evidence"]["formal_sources_per_mature_cluster"],
        )

    def test_analysis_validator_rejects_generic_search_miss_as_platform_no_result(self) -> None:
        topic = "听取数字中国建设情况汇报"
        audit = self.domestic_research_audit([topic])
        checks = audit["domestic_media_research"]["coverage_by_topic"][0]["checks"]
        toutiao = next(row for row in checks if row["source_id"] == "toutiao_articles")
        toutiao["execution_mode"] = "generic_web_search"
        toutiao["queries"] = []
        path = self.root / "bad_platform_miss.json"
        path.write_text(
            json.dumps(
                {
                    "research_audit": audit,
                    "viewpoints": {
                        "by_topic": [
                            {
                                "topic": topic,
                                "heading": "认为数字基础设施建设应统筹发展与安全",
                                "clusters": [],
                            }
                        ]
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        problems = MODULE.validate_analysis_bundle(path, [topic])

        self.assertTrue(any("没有平台定向检索及查询留痕" in item for item in problems))

    def test_analysis_validator_rejects_selected_only_bundle_without_candidate_pool(self) -> None:
        topic = "听取数字中国建设情况汇报"
        path = self.root / "selected_only.json"
        path.write_text(
            json.dumps(
                {
                    "research_audit": {
                        "domestic_media_research": {
                            **self.domestic_research_audit([topic])["domestic_media_research"],
                            "candidate_pool_by_topic": [],
                        }
                    },
                    "viewpoints": {
                        "by_topic": [
                            {
                                "topic": topic,
                                "heading": "认为数字基础设施建设应统筹发展与安全",
                                "clusters": [
                                    {
                                        "summary": "认为数字基础设施建设应统筹发展与安全",
                                        "details": "测试媒体认为，数字基础设施建设应统筹规划、网络安全、数据治理和长效评估，形成覆盖建设、运行、监督和风险处置的完整机制。",
                                        "evidence": [{"url": "https://example.com/a"}],
                                    }
                                ],
                            }
                        ]
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        problems = MODULE.validate_analysis_bundle(path, [topic])
        self.assertTrue(any("缺少全网候选池" in item for item in problems))

    def test_analysis_validator_rejects_raw_candidate_used_as_registry_web_hit(self) -> None:
        topic = "听取数字中国建设情况汇报"
        audit = self.domestic_research_audit([topic])
        pool = audit["domestic_media_research"]["candidate_pool_by_topic"][0]
        candidate = pool["candidates"][0]
        candidate["discovery_origin"] = "raw_monitoring"
        checks = audit["domestic_media_research"]["coverage_by_topic"][0]["checks"]
        xinhua = next(row for row in checks if row["source_id"] == "xinhua")
        xinhua.update({"status": "hit", "candidate_ids": [candidate["candidate_id"]]})
        path = self.root / "raw_as_web_hit.json"
        path.write_text(
            json.dumps({"research_audit": audit, "viewpoints": {"by_topic": []}}, ensure_ascii=False),
            encoding="utf-8",
        )

        problems = MODULE.validate_analysis_bundle(path, [topic])

        self.assertTrue(any("错误地用原始监测池候选证明全网命中" in item for item in problems))

    def test_analysis_validator_rejects_search_result_url_without_candidate_decision(self) -> None:
        topic = "听取数字中国建设情况汇报"
        audit = self.domestic_research_audit([topic])
        pool = audit["domestic_media_research"]["candidate_pool_by_topic"][0]
        execution = pool["saturation"]["rounds"][0]["executions"][0]
        execution["result_count"] = 2
        execution["result_urls"].append("https://example.com/silently-dropped")
        path = self.root / "silent_search_drop.json"
        path.write_text(
            json.dumps({"research_audit": audit, "viewpoints": {"by_topic": []}}, ensure_ascii=False),
            encoding="utf-8",
        )

        problems = MODULE.validate_analysis_bundle(path, [topic])

        self.assertTrue(any("结果URL未进入候选判定" in item for item in problems))

    def test_analysis_validator_rejects_unsaturated_candidate_pool(self) -> None:
        topic = "听取数字中国建设情况汇报"
        audit = self.domestic_research_audit([topic])
        pool = audit["domestic_media_research"]["candidate_pool_by_topic"][0]
        pool["saturation"]["rounds"][-1]["new_independent_viewpoints"] = 1
        path = self.root / "unsaturated.json"
        path.write_text(
            json.dumps(
                {
                    "research_audit": audit,
                    "viewpoints": {
                        "by_topic": [
                            {
                                "topic": topic,
                                "heading": "认为数字基础设施建设应统筹发展与安全",
                                "clusters": [
                                    {
                                        "summary": "认为数字基础设施建设应统筹发展与安全",
                                        "details": "测试媒体认为，数字基础设施建设应统筹规划、网络安全、数据治理和长效评估，形成覆盖建设、运行、监督和风险处置的完整机制。",
                                        "evidence": [
                                            {
                                                "candidate_id": "candidate-1",
                                                "url": "https://example.com/a",
                                                "attribution": "测试媒体",
                                                "attribution_status": "media_voice",
                                                "source_excerpt": "数字基础设施建设应统筹发展与安全。",
                                                "formal_claim": "数字基础设施建设应统筹规划、网络安全、数据治理和长效评估。",
                                                "wording_fidelity": "faithful_paraphrase",
                                                "selection_reason": "与议题直接相关。",
                                            }
                                        ],
                                    }
                                ],
                            }
                        ]
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        problems = MODULE.validate_analysis_bundle(path, [topic])
        self.assertTrue(any("缺少2轮" in item for item in problems))

    def test_analysis_validator_rejects_unverified_or_out_of_window_candidate_date(self) -> None:
        topic = "听取数字中国建设情况汇报"
        audit = self.domestic_research_audit([topic])
        candidate = audit["domestic_media_research"]["candidate_pool_by_topic"][0]["candidates"][0]
        candidate["published_at"] = "2026-07-11"
        candidate["published_at_verified_from_source"] = False
        path = self.root / "bad_candidate_date.json"
        path.write_text(
            json.dumps(
                {
                    "metadata": {"monitoring_start": "2026-07-10", "monitoring_end": "2026-07-10"},
                    "research_audit": audit,
                    "viewpoints": {"by_topic": []},
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        problems = MODULE.validate_analysis_bundle(path, [topic])

        self.assertTrue(any("发布时间未从原始页面" in item for item in problems))
        self.assertTrue(any("超出监测期" in item for item in problems))

    def test_analysis_validator_rejects_eligible_candidate_missing_from_formal_clusters(self) -> None:
        topic = "听取数字中国建设情况汇报"
        audit = self.domestic_research_audit([topic])
        pool = audit["domestic_media_research"]["candidate_pool_by_topic"][0]
        second = dict(pool["candidates"][0])
        second.update(
            {
                "candidate_id": "candidate-2",
                "url": "https://example.com/b",
                "source": "测试媒体二",
                "title": "数字基础设施需要完善长效治理",
                "viewpoint_cluster_key": "development_and_security",
                "discovery_query_id": "query-1",
            }
        )
        pool["candidates"].append(second)
        execution = pool["saturation"]["rounds"][0]["executions"][0]
        execution["result_count"] = 2
        execution["result_urls"].append(second["url"])
        execution["retained_candidate_ids"].append(second["candidate_id"])
        path = self.root / "missing_eligible_formal.json"
        path.write_text(
            json.dumps(
                {
                    "metadata": {"monitoring_start": "2026-07-10", "monitoring_end": "2026-07-10"},
                    "research_audit": audit,
                    "viewpoints": {
                        "by_topic": [
                            {
                                "topic": topic,
                                "single_cluster_exception": {
                                    "reason": "一个观点家族",
                                    "search_evidence": "query-1至query-3",
                                    "reviewed_by": "test",
                                },
                                "clusters": [
                                    {
                                        "summary": "认为数字基础设施建设应统筹发展与安全",
                                        "details": "测试媒体认为，数字基础设施建设应统筹规划、网络安全、数据治理和长效评估，形成覆盖建设、运行、监督和风险处置的完整机制。",
                                        "thin_cluster_exception": {
                                            "reason": "测试",
                                            "search_evidence": "query-1至query-3",
                                            "reviewed_by": "test",
                                        },
                                        "evidence": [
                                            {
                                                "candidate_id": "candidate-1",
                                                "source": "测试媒体",
                                                "url": "https://example.com/a",
                                                "attribution": "测试媒体",
                                                "attribution_status": "media_voice",
                                                "source_excerpt": "数字基础设施建设应统筹发展与安全。",
                                                "formal_claim": "数字基础设施建设应统筹规划、网络安全、数据治理和长效评估。",
                                                "wording_fidelity": "faithful_paraphrase",
                                                "selection_reason": "与议题直接相关。",
                                            }
                                        ],
                                    }
                                ],
                            }
                        ]
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        problems = MODULE.validate_analysis_bundle(path, [topic])

        self.assertTrue(any("合格独立样本未归入正式观点和正文" in item for item in problems))

    def test_domestic_supplemental_publishers_have_explicit_evidence_rules(self) -> None:
        registry = MODULE.read_json(MODULE.SOURCE_REGISTRY_PATH)
        sources = {row["id"]: row for row in registry["sources"]}
        expected = {
            "baijiahao",
            "sogou_weixin_index",
            "netease_news",
            "netease_media",
            "sina_news",
        }
        self.assertTrue(expected.issubset(sources))
        self.assertEqual("search_resolve_original", sources["sogou_weixin_index"]["reader"])
        self.assertIn("/dy/article/", sources["netease_media"]["article_url_patterns"])
        for source_id in expected:
            self.assertEqual("domestic", sources[source_id]["region"])
            self.assertTrue(sources[source_id].get("evidence_policy"))

    def test_workbench_launch_waits_for_http_success(self) -> None:
        dashboard = self.job / "report" / "cwh_dashboard.html"
        dashboard.parent.mkdir(parents=True)
        dashboard.write_text("<html></html>", encoding="utf-8")
        process = mock.Mock(pid=12345)
        process.poll.return_value = None
        response = mock.MagicMock()
        response.__enter__.return_value.status = 200
        with mock.patch.object(MODULE.subprocess, "Popen", return_value=process), mock.patch.object(
            MODULE.urllib.request, "urlopen", return_value=response
        ):
            result = MODULE.launch_local_workbench(self.job)
        self.assertTrue(result["launched"])
        self.assertEqual(200, result["http_status"])
        self.assertEqual({"dashboard": 200, "archive": 200}, result["required_endpoints"])

    def test_workbench_launch_rejects_partial_first_screen_health(self) -> None:
        dashboard = self.job / "report" / "cwh_dashboard.html"
        dashboard.parent.mkdir(parents=True)
        dashboard.write_text("<html></html>", encoding="utf-8")
        process = mock.Mock(pid=12345)
        process.poll.return_value = None
        response = mock.MagicMock()
        response.__enter__.return_value.status = 200

        def open_endpoint(url: str, timeout: float = 0.5):
            if str(url).endswith("/api/archive"):
                raise MODULE.urllib.error.HTTPError(str(url), 500, "archive failed", {}, None)
            return response

        with mock.patch.object(MODULE.subprocess, "Popen", return_value=process), mock.patch.object(
            MODULE.urllib.request, "urlopen", side_effect=open_endpoint
        ), mock.patch.object(MODULE.time, "sleep"):
            result = MODULE.launch_local_workbench(self.job)
        self.assertFalse(result["launched"])
        self.assertTrue(result["partially_available"])
        self.assertEqual(200, result["required_endpoints"]["dashboard"])
        self.assertEqual(500, result["required_endpoints"]["archive"])

    def test_workbench_launch_uses_one_long_archive_probe_after_dashboard_ready(self) -> None:
        dashboard = self.job / "report" / "cwh_dashboard.html"
        dashboard.parent.mkdir(parents=True)
        dashboard.write_text("<html></html>", encoding="utf-8")
        process = mock.Mock(pid=12345)
        process.poll.return_value = None
        response = mock.MagicMock()
        response.__enter__.return_value.status = 200
        calls = []

        def open_endpoint(url: str, timeout: float = 0.5):
            calls.append((str(url), float(timeout)))
            return response

        with mock.patch.object(MODULE.subprocess, "Popen", return_value=process), mock.patch.object(
            MODULE.urllib.request, "urlopen", side_effect=open_endpoint
        ):
            result = MODULE.launch_local_workbench(self.job)

        archive_calls = [item for item in calls if item[0].endswith("/api/archive")]
        self.assertTrue(result["launched"])
        self.assertEqual(1, len(archive_calls))
        self.assertGreaterEqual(archive_calls[0][1], 60.0)
        self.assertEqual({"dashboard": 200, "archive": 200}, result["required_endpoints"])

    def test_workbench_launch_uses_next_free_port(self) -> None:
        dashboard = self.job / "report" / "cwh_dashboard.html"
        dashboard.parent.mkdir(parents=True)
        dashboard.write_text("<html></html>", encoding="utf-8")
        process = mock.Mock(pid=12345)
        process.poll.return_value = None
        response = mock.MagicMock()
        response.__enter__.return_value.status = 200
        occupied = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        reserved_port = None
        for candidate in range(8788, 8817):
            try:
                occupied.bind(("127.0.0.1", candidate))
                reserved_port = candidate
                break
            except OSError:
                continue
        self.assertIsNotNone(reserved_port)
        with occupied:
            with mock.patch.object(MODULE.subprocess, "Popen", return_value=process), mock.patch.object(
                MODULE.urllib.request, "urlopen", return_value=response
            ):
                result = MODULE.launch_local_workbench(self.job)
        selected_port = int(result["url"].split(":")[2].split("/")[0])
        self.assertGreater(selected_port, reserved_port)

    def test_workbench_launch_reports_early_server_exit(self) -> None:
        dashboard = self.job / "report" / "cwh_dashboard.html"
        dashboard.parent.mkdir(parents=True)
        dashboard.write_text("<html></html>", encoding="utf-8")
        process = mock.Mock(pid=12345)
        process.poll.return_value = 3
        with mock.patch.object(MODULE.subprocess, "Popen", return_value=process):
            result = MODULE.launch_local_workbench(self.job)
        self.assertFalse(result["launched"])
        self.assertEqual("server_exited_before_http_ready:3", result["reason"])

    def test_with_runtime_never_auto_opens_local_workbench(self) -> None:
        state = {"status": "succeeded"}
        self.assertFalse(
            MODULE.should_launch_local_workbench(
                state, no_open=False, environment={"WITH_PROJECT_ID": "q0eb32dw"}
            )
        )
        self.assertFalse(
            MODULE.should_launch_local_workbench(
                state, no_open=False, environment={"CWH_RUNTIME": "with"}
            )
        )

    def test_no_open_disables_local_workbench(self) -> None:
        self.assertFalse(
            MODULE.should_launch_local_workbench(
                {"status": "succeeded"}, no_open=True, environment={}
            )
        )

    def test_foreign_gate_requires_complete_reviewed_simplified_fields(self) -> None:
        audit = self.root / "foreign_audit.json"
        supplements = self.root / "foreign_supplements.json"
        audit.write_text(json.dumps({"attempted": True, "exit_code": 0}), encoding="utf-8")
        supplements.write_text(
            json.dumps(
                {
                    "media": [
                        {
                            "formal_include": True,
                            "title_cn_simplified": "国务院常务会议部署相关工作",
                        }
                    ],
                    "comments": [],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        problems = MODULE.validate_foreign(audit, supplements)
        self.assertTrue(any("source_cn_simplified" in item for item in problems))
        self.assertTrue(any("简体中文AI复核" in item for item in problems))

        supplements.write_text(
            json.dumps(
                {
                    "media": [
                        {
                            "formal_include": True,
                            "title_cn_simplified": "国务院常务会议部署相关工作",
                            "source_cn_simplified": "联合早报",
                            "summary_cn_simplified": "报道转述会议议程和相关工作部署。",
                            "simplified_chinese_reviewed": True,
                        }
                    ],
                    "comments": [],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        self.assertEqual([], MODULE.validate_foreign(audit, supplements))


if __name__ == "__main__":
    unittest.main()
