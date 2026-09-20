from copy import deepcopy
from pathlib import Path
import sys

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import formalize_cwh_report as formal
from normalize_cwh_analysis import assemble_cluster_details, terminated_sentence


@pytest.mark.parametrize('verb', ['说', '称', '表示'])
def test_literal_postposed_named_attribution_is_not_duplicated(verb):
    row = {'speaker_name': '张甲', 'speaker_role': '某研究院研究员', 'attribution': '某研究院研究员张甲',
           'attribution_status': 'named_person',
           'formal_claim': f'“扩大服务覆盖仍需配套人员经费。”某研究院研究员张甲{verb}。'}
    frozen = deepcopy(row)
    text = assemble_cluster_details({'evidence': [row]})
    assert text == row['formal_claim']
    assert text.count('张甲') == 1
    doc = Document()
    paragraph = doc.add_paragraph()
    formal.add_detail_blocks(doc, paragraph, {'evidence': [row]})
    assert paragraph.text == text
    assert [r.text for r in paragraph.runs if r.bold] == [f'某研究院研究员张甲{verb}']
    assert row == frozen


@pytest.mark.parametrize('text,expected', [
    ('完整判断', '完整判断。'), ('完整判断。', '完整判断。'),
    ('“具体判断。”', '“具体判断。”'), ('“条件满足了吗？”', '“条件满足了吗？”'),
    ('“明确建议！”', '“明确建议！”'), ('应改善“办理效率”', '应改善“办理效率”。'),
    ('', ''),
])
def test_only_missing_display_stop_is_added(text, expected):
    assert terminated_sentence(text) == expected


def test_quoted_claims_stay_identical_in_cluster_and_word():
    rows = [{'speaker_name': '甲专家', 'speaker_role': '研究员',
             'attribution_status': 'named_person', 'formal_claim': '“需先满足公开征求意见条件。”'},
            {'speaker_name': '乙专家', 'speaker_role': '研究员',
             'attribution_status': 'named_person', 'formal_claim': '应保留具体适用范围和“执行条件”'}]
    cluster = {'summary': '需明确条件', 'evidence': rows}
    original = deepcopy(cluster)
    expected = assemble_cluster_details(cluster)
    assert '。”。' not in expected
    assert '“执行条件”。' in expected
    assert ''.join(formal.cluster_detail_paragraphs(cluster)) == expected
    document = Document()
    paragraph = document.add_paragraph()
    formal.add_detail_blocks(document, paragraph, cluster)
    assert ''.join(p.text for p in document.paragraphs) == expected
    assert cluster == original
    for row in rows:
        assert row['formal_claim'] in expected
    assert sum(bool(run.bold) for p in document.paragraphs for run in p.runs) == 2
