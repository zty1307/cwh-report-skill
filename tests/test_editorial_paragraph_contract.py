import copy
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from cwh_editorial_selection import enforce_author_layout
from cwh_writing_rules import attribution_verb, judgment_heading, formal_claim_char_count
from formalize_cwh_report import add_cluster_paragraph, cluster_wording
from report_rules import approved_source_layout_issues


def decision(groups):
    return {'heading': '建议完善公共服务', 'clusters': [
        {'key': f'k{i}', 'heading': f'建议落实条件{i}'} for i in range(len(groups))],
        'items': [{'id': f'r{i}-{j}', 'decision': 'eligible', 'reason': '原文有独立观点',
                   'claims': [{'speaker': name, 'speaker_type': 'named_person', 'role': '研究员',
                               'claim': '只有落实配套资金，服务扩围才可能持续。', 'quote': '完整原文保留。',
                               'cluster': f'k{i}'}]}
                  for i, names in enumerate(groups) for j, name in enumerate(names)]}


def test_three_sources_stay_in_one_numbered_paragraph_without_truncation():
    cluster = {'summary': '建议落实配套条件', 'evidence': [
        {'speaker_name': name, 'formal_claim': '条件与依据' * 25 + '。'} for name in ['甲', '乙', '丙']]}
    original = copy.deepcopy(cluster)
    doc = Document()
    add_cluster_paragraph(doc, '一是', cluster)
    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0].text.startswith('一是')
    assert doc.paragraphs[0].text == cluster_wording(cluster, prefix='一是')
    assert all(row['formal_claim'] in doc.paragraphs[0].text for row in cluster['evidence'])
    assert cluster == original


def test_neutral_attribution_does_not_change_heading_or_specific_stance():
    for verb in ['认为', '表示', '说', '称']:
        row = {'attribution_verb': verb}
        assert attribution_verb(row) == '称'
        assert row == {'attribution_verb': verb}
    for verb in ['建议', '预计', '担忧', '呼吁', '指出']:
        assert attribution_verb({'attribution_verb': verb}) == verb
    assert judgment_heading('资金安排影响服务扩围') == '认为资金安排影响服务扩围'


def test_thin_group_is_reported_without_silently_deleting_valid_claims():
    before = decision([['甲', '乙'], ['丙'], ['丁', '戊']])
    frozen = copy.deepcopy(before)
    after, audit = enforce_author_layout(before)
    assert before == frozen
    assert after['clusters'] == before['clusters']
    assert after['items'] == before['items']
    assert audit['status'] == 'pending_editorial_completion'
    assert audit['reserves'] == []
    assert audit['layout_shortfalls'] == [{'key': 'k1', 'sources': 1, 'required': 2}]
    assert enforce_author_layout(after)[0] == after


def test_all_singletons_keep_all_claims_without_false_editorial_success():
    before = decision([['甲'], ['乙'], ['丙']])
    after, audit = enforce_author_layout(before)
    assert after['clusters'] == before['clusters']
    assert after['items'] == before['items']
    assert audit['reserves'] == []
    assert len(audit['layout_shortfalls']) == 3
    assert '未达到写作标准' in after['shortfall_reason']


def test_150_character_boundary_and_existing_reserves_survive():
    before = decision([['甲', '乙']])
    before['items'][0]['claims'][0]['claim'] = '甲' * 150
    before['items'][1]['claims'][0]['claim'] = '乙' * 151
    after, audit = enforce_author_layout(before)
    assert after['items'][0]['claims'][0].get('formal_use') != 'reserve'
    assert after['items'][1]['claims'][0]['formal_use'] == 'reserve'
    assert audit['reserves'][0]['original_claim']['claim'] == '乙' * 151
    assert before['items'][1]['claims'][0].get('formal_use') is None


def test_length_excludes_exact_attribution_but_not_unknown_role_or_other_speaker():
    row = {'speaker_name': '张某', 'speaker_role': '研究员', 'formal_claim': '研究员张某认为，' + '甲' * 150}
    before = copy.deepcopy(row)
    assert formal_claim_char_count(row) == 150
    assert row == before
    row['formal_claim'] = '张某所在城市' + '甲' * 150
    assert formal_claim_char_count(row) > 150
    row['formal_claim'] = '未知职务张某认为，' + '甲' * 150
    assert formal_claim_char_count(row) > 150
    row['formal_claim'] = '研究员李某认为，' + '甲' * 150
    assert formal_claim_char_count(row) > 150


def test_real_final_gate_still_reports_post_review_thin_group_and_long_sample():
    data = {'viewpoints': {'by_topic': [{'topic': '公共服务', 'clusters': [
        {'summary': '判断一', 'evidence': [{'speaker_name': '甲', 'formal_claim': '甲' * 151}]},
        {'summary': '判断二', 'evidence': [{'speaker_name': n, 'formal_claim': '完整观点'} for n in ['乙', '丙']]}]}]}}
    codes = {row['code'] for row in approved_source_layout_issues(data)}
    assert {'formal_claim_over_limit', 'numbered_cluster_source_shortfall'} <= codes
