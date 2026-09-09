from __future__ import annotations

import csv
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

import cwh_orchestrator as ORCHESTRATOR  # noqa: E402
import formalize_cwh_report as FORMALIZE  # noqa: E402
import generate_dashboard as DASHBOARD  # noqa: E402
import ingest_monitoring_workbook as INGEST  # noqa: E402


class CwhCommentReportIntegrationTests(unittest.TestCase):
    def test_data_workbook_keeps_comments_out_of_spread_and_hides_tiny_sentiment(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "cwh_data_workbook.xlsx"
            daily = [{
                "date": "2026-05-15",
                "domestic_mainstream": 10,
                "overseas_media": 1,
                "new_media": 20,
                "total_spread": 31,
                "details": {},
            }]
            data = {
                "meeting": {"date": "2026-05-15", "topics": ["议题甲"]},
                "samples": [],
                "system_data": {
                    "master_event": {},
                    "topics": [{"index": 1, "title": "议题甲"}],
                    "overall": {
                        "daily": [{**daily[0], "wechat_public": 10, "weibo": 5, "video_account": 2}],
                        "totals": {"domestic_mainstream": 10, "overseas_media": 1, "wechat_public": 10, "weibo": 5, "video_account": 2, "new_media": 3, "total_spread": 31},
                    },
                    "subevents": [{"index": 1, "daily": daily, "totals": {"domestic_mainstream": 10, "overseas_media": 1, "new_media": 20, "total_spread": 31}}],
                },
                "topic_stats": [{
                    "topic": "议题甲", "display": "议题甲", "domestic_media": 10,
                    "self_media": 20, "overseas_media": 1, "spread_count": 31,
                    "comments": 2, "sentiment": {"positive": 2, "neutral": 0, "negative": 0},
                    "sentiment_formal_ready": False,
                }],
                "hotwords": [], "appendices": {"wechat_top": []},
            }
            ORCHESTRATOR.write_cwh_data_workbook(data, path)
            workbook = load_workbook(path, data_only=False)
            try:
                summary = workbook["子事件数据汇总"]
                self.assertEqual([10, 20, 1, 31], [summary.cell(3, col).value for col in range(3, 7)])
                self.assertEqual([None, None, None, None], [summary.cell(3, col).value for col in range(7, 11)])
                self.assertEqual(31, workbook["子事件1"]["E4"].value)
            finally:
                workbook.close()

    def test_formal_delivery_gate_is_not_only_an_audit_warning(self) -> None:
        blocked = {"audit": {"acceptance": {"ready_for_formal_delivery": False}}}
        ready = {"audit": {"acceptance": {"ready_for_formal_delivery": True}}}
        self.assertFalse(ORCHESTRATOR.formal_delivery_ready(blocked))
        self.assertTrue(ORCHESTRATOR.formal_delivery_ready(ready))

    def test_reviewed_hotword_schema_supplies_report_example(self) -> None:
        reviewed = {
            "word": "三代核电技术",
            "sample_titles": ["我国新核准八台核电机组"],
        }
        self.assertEqual(ORCHESTRATOR.hotword_example(reviewed), "我国新核准八台核电机组")
        self.assertEqual(ORCHESTRATOR.hotword_example({"word": "健康优先"}), "")

    def test_short_display_date_uses_workbook_title_year(self) -> None:
        self.assertEqual(INGEST.date_text("7/10", 2026), "2026-07-10")

    def test_subevent_year_is_inferred_from_right_hand_raw_date_column(self) -> None:
        from datetime import datetime

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "子事件1"
        sheet["A2"] = "子事件1-议题甲"
        for col, value in enumerate(["日期", "境内主流媒体", "境外媒体", "新媒体", "信息传播量"], 1):
            sheet.cell(3, col, value)
        for col, value in enumerate(["日期", "公众文章", "公号-在看量", "公号-精选评论量", "新浪微博", "境内新闻", "境内APP", "境内论坛", "其他视频", "境外新闻", "推特", "境外其他", "视频号发文", "抖音"], 8):
            sheet.cell(3, col, value)
        for col, value in enumerate(["5/15", 10, 1, 20, 31], 1):
            sheet.cell(4, col, value)
        sheet.cell(4, 8, datetime(2026, 5, 15))
        parsed = INGEST.parse_trend_sheet(sheet)
        workbook.close()
        self.assertEqual("2026-05-15", parsed["daily"][0]["date"])
        self.assertEqual(31, parsed["daily"][0]["total_spread"])

    def write_handoff(self, path: Path) -> None:
        row = {
            "id": "comment-1",
            "sample_id": "comment-1",
            "topic": "议题一",
            "platform": "今日头条",
            "source_type": "netizen_comment",
            "region": "domestic",
            "source": "用户甲",
            "title": "支持",
            "content": "支持",
            "url": "https://example.test/article/1",
            "published_at": "2026-07-11T10:00:00+08:00",
            "is_comment": "true",
            "quote_verified": "true",
            "evidence_mode": "verbatim_public_comment",
            "comment_id": "c1",
            "reply_id": "",
            "parent_comment_id": "",
            "raw_file": "comment_detail.json",
            "sentiment": "positive",
            "sentiment_source": "ai_reviewed",
            "sentiment_status": "classified",
            "in_sentiment_denominator": "true",
        }
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(row))
            writer.writeheader()
            writer.writerow(row)

    def test_handoff_survives_report_and_dashboard_traceability_gates(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "report_comment_handoff.csv"
            self.write_handoff(path)
            samples = ORCHESTRATOR.load_comment_handoffs([str(path)], ["议题一"])

            self.assertEqual(len(samples), 1)
            self.assertTrue(ORCHESTRATOR.is_actual_platform_comment(samples[0]))
            selected = ORCHESTRATOR.select_comments(samples)
            self.assertEqual(len(selected), 1, "reviewed short stance comments must remain visible")
            self.assertEqual(selected[0]["comment_id"], "c1")
            self.assertEqual(selected[0]["raw_file"], "comment_detail.json")

            dashboard_rows = DASHBOARD.dashboard_comments_by_topic(
                {"comments": {"selected": selected}},
                [{"topic": "议题一", "display": "议题一"}],
            )
            self.assertEqual(len(dashboard_rows[0]["items"]), 1)
            self.assertEqual(dashboard_rows[0]["items"][0]["content"], "支持")

    def test_sentiment_summary_separates_observation_from_formal_percentage(self) -> None:
        topic_stats = [{
            "topic": "议题一",
            "sentiment": {"positive": 1, "neutral": 0, "negative": 0},
            "sentiment_authority": "large_scale_sentiment_analysis",
        }]
        summary = {"topics": [{"title": "议题一", "status": "insufficient_sample", "denominator": 1, "counts": {"positive": 1}}]}
        rows = ORCHESTRATOR.apply_sentiment_summary(topic_stats, summary)

        self.assertFalse(rows[0]["sentiment_formal_ready"])
        self.assertEqual(rows[0]["sentiment_denominator"], 1)
        lead = FORMALIZE.netizen_sentiment_lead({
            "comments": {"sentiment": {"positive": 1, "neutral": 0, "negative": 0}},
            "topic_stats": rows,
        })
        self.assertIn("本批共取得并复核1条", lead)
        self.assertIn("样本尚不足", lead)

    def test_reviewed_exclusion_absent_from_handoff_is_not_reported_as_data_loss(self) -> None:
        samples = [{
            "id": "comment-keep",
            "source_type": "netizen_comment",
            "is_comment": True,
            "platform": "bili",
            "content": "需要说明修改内容",
        }]
        with tempfile.TemporaryDirectory() as temp:
            results = Path(temp) / "review.csv"
            with results.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["sample_id", "label", "label_source", "in_sentiment_denominator", "exclusion_reason", "needs_review"],
                )
                writer.writeheader()
                writer.writerow({"sample_id": "comment-keep", "label": "negative", "label_source": "ai_reviewed", "in_sentiment_denominator": "true", "exclusion_reason": "", "needs_review": "false"})
                writer.writerow({"sample_id": "comment-excluded", "label": "", "label_source": "ai_reviewed", "in_sentiment_denominator": "false", "exclusion_reason": "与会议无关", "needs_review": "false"})
            matched, gaps = ORCHESTRATOR.apply_sentiment_results(samples, str(results))
        self.assertEqual(matched, 1)
        self.assertEqual(gaps, [])

    def test_comments_on_same_article_are_not_deduped_by_url_or_text(self) -> None:
        base = {
            "topic": "议题一",
            "platform": "今日头条",
            "source_type": "netizen_comment",
            "region": "domestic",
            "source": "用户",
            "title": "支持",
            "content": "支持",
            "url": "https://example.test/article/1",
            "published_at": "2026-07-11T10:00:00+08:00",
            "is_comment": True,
            "quote_verified": True,
            "evidence_mode": "verbatim_public_comment",
            "sentiment": "positive",
            "sentiment_source": "ai_reviewed",
            "sentiment_status": "classified",
            "in_sentiment_denominator": True,
            "quality_flags": [],
        }
        rows = [
            {**base, "id": "comment-1", "comment_id": "c1", "raw_file": "comments.json"},
            {**base, "id": "comment-2", "comment_id": "c2", "raw_file": "comments.json"},
        ]
        deduped = ORCHESTRATOR.dedupe_samples(rows)
        self.assertEqual(len(deduped), 2)
        self.assertEqual(len(ORCHESTRATOR.select_comments(deduped)), 2)

    def test_topic_dash_variants_map_to_authoritative_workbook_topic(self) -> None:
        topic = "审议通过《全民健身计划（2026－2030年）》"
        row = {
            "id": "comment-fitness",
            "sample_id": "comment-fitness",
            "topic": "审议通过《全民健身计划（2026—2030年）》",
            "platform": "今日头条",
            "source_type": "netizen_comment",
            "region": "domestic",
            "source": "用户甲",
            "title": "支持全民健身",
            "content": "支持全民健身",
            "url": "https://example.test/article/fitness",
            "published_at": "2026-07-12T10:00:00+08:00",
            "is_comment": True,
            "quote_verified": True,
            "evidence_mode": "verbatim_public_comment",
            "comment_id": "fitness-1",
        }
        sample = ORCHESTRATOR.normalize_sample(row, [topic], 1)
        self.assertIsNotNone(sample)
        self.assertEqual(sample["topic"], topic)

        dashboard_rows = DASHBOARD.dashboard_comments_by_topic(
            {"comments": {"selected": [sample]}},
            [{"topic": topic, "display": "全民健身计划（2026－2030年）"}],
        )
        self.assertEqual(len(dashboard_rows[0]["items"]), 1)

        topic_stats = [{"topic": topic, "sentiment": {}}]
        summary = {
            "topics": [{
                "title": "审议通过《全民健身计划（2026—2030年）》",
                "status": "insufficient_sample",
                "denominator": 14,
                "counts": {"positive": 12, "neutral": 2},
            }]
        }
        result = ORCHESTRATOR.apply_sentiment_summary(topic_stats, summary)
        self.assertEqual(result[0]["sentiment_denominator"], 14)

    def test_reaction_emoji_stays_in_dashboard_but_not_formal_report(self) -> None:
        row = {
            "topic": "议题一",
            "platform": "今日头条",
            "source_type": "netizen_comment",
            "source": "用户甲",
            "content": "政策方向很好[赞][心][福][点亮平安灯]，期待落地[呲牙]🌹",
            "url": "https://example.test/article/1",
            "published_at": "2026-07-11",
            "quote_verified": True,
            "evidence_mode": "verbatim_public_comment",
            "comment_id": "c1",
            "report_order": 1,
            "ai_formal_include": True,
            "ai_semantic_quality": "substantive",
            "ai_formal_reason": "表达明确支持和期待",
            "comment_heading": "支持政策方向并期待尽快落地",
        }
        dashboard_rows = DASHBOARD.dashboard_comments_by_topic(
            {"comments": {"selected": [row]}},
            [{"topic": "议题一", "display": "议题一"}],
        )
        self.assertEqual(dashboard_rows[0]["items"][0]["content"], row["content"])

        groups = FORMALIZE.comment_groups([row])
        self.assertEqual(groups[0][1][0]["content"], "政策方向很好，期待落地")
        wording = FORMALIZE.comment_wording(groups[0][1])
        self.assertNotIn("[赞]", wording)
        self.assertNotIn("[心]", wording)
        self.assertNotIn("[福]", wording)
        self.assertNotIn("[点亮平安灯]", wording)
        self.assertNotIn("[呲牙]", wording)
        self.assertNotIn("🌹", wording)

        only_reactions = {**row, "content": "[心][福][点亮平安灯][呲牙]🌹", "comment_id": "c2"}
        self.assertEqual(FORMALIZE.comment_groups([only_reactions]), [])

    def test_article_like_comment_remains_data_but_is_not_used_as_formal_quote(self) -> None:
        base = {
            "topic": "住房政策议题",
            "platform": "bili",
            "source_type": "netizen_comment",
            "url": "https://example.test/video/1",
            "quote_verified": True,
            "evidence_mode": "verbatim_public_comment",
            "ai_formal_include": True,
            "ai_semantic_quality": "substantive",
            "ai_formal_reason": "已完成正式引用审核",
            "comment_heading": "期待明确说明政策修改内容",
        }
        article_like = {
            **base,
            "comment_id": "long-1",
            "content": "省流：核心调整分为四大块：1.扩大覆盖范围；2.拓宽提取用途；3.提升跨区域服务；4.完善风险管理。" * 3,
        }
        opinion = {
            **base,
            "comment_id": "short-1",
            "content": "能不能直接说明到底修改了哪些条款",
        }
        self.assertTrue(FORMALIZE.comment_is_report_quote_suitable(opinion["content"]))
        self.assertFalse(FORMALIZE.comment_is_report_quote_suitable(article_like["content"]))
        groups = FORMALIZE.comment_groups([article_like, opinion])
        self.assertEqual(len(groups), 1)
        self.assertEqual([row["comment_id"] for row in groups[0][1]], ["short-1"])

    def test_dashboard_uses_report_aligned_module_tabs(self) -> None:
        template = (Path(__file__).resolve().parents[1] / "assets" / "cwh_dashboard_template.html").read_text(encoding="utf-8")
        for label in (
            "（一）总事件传播情况",
            "（二）子议题传播情况",
            "（一）境内媒体自媒体情况",
            "（二）网民评论情况",
            "（三）热词分布情况",
            "（一）境外媒体情况",
            "（二）境外网民评论",
        ):
            self.assertIn(label, template)
        self.assertIn("data-module-group=\"one\"", template)
        self.assertIn("data-module-group=\"three\"", template)
        self.assertNotIn("识别文件结构", template)
        self.assertNotIn("个成文观点", template)

    def test_hotwords_require_traceable_domestic_evidence_and_count_comments(self) -> None:
        topic = "研究数字中国建设有关工作"
        samples = [
            {
                "id": "media-1",
                "topic": topic,
                "topic_hits": [topic],
                "source_type": "mainstream_media",
                "source": "媒体甲",
                "title": "算力网络支撑数字中国建设",
                "content": "算力网络成为数字中国建设的重要基础。",
                "url": "https://example.test/media",
            },
            {
                "id": "comment-1",
                "topic": topic,
                "topic_hits": [topic],
                "source_type": "netizen_comment",
                "platform": "微博",
                "source": "用户甲",
                "title": "",
                "content": "希望算力网络尽快覆盖更多地区。",
                "url": "https://example.test/comment",
                "is_comment": True,
                "quote_verified": True,
                "evidence_mode": "platform_comment",
                "comment_id": "c1",
            },
        ]
        hotwords = ORCHESTRATOR.build_evidence_hotwords(topic and [topic], samples, ["算力网络", "凭空热词", "务院常"])
        by_word = {row["word"]: row for row in hotwords}

        self.assertIn("算力网络", by_word)
        self.assertEqual(1, by_word["算力网络"]["media_sample_count"])
        self.assertEqual(1, by_word["算力网络"]["comment_sample_count"])
        self.assertEqual(
            {
                "independent_sample_coverage",
                "title_or_lead_salience",
                "source_diversity",
                "ai_semantic_representativeness",
                "verified_comment_mentions",
                "subtopic_balance",
            },
            set(by_word["算力网络"]["score_components"]),
        )
        self.assertNotIn("凭空热词", by_word)
        self.assertNotIn("务院常", by_word)

    def test_hotword_single_row_requires_agenda_anchor_exception(self) -> None:
        topic = "研究数字中国建设有关工作"
        samples = [
            {
                "id": "media-1",
                "topic": topic,
                "topic_hits": [topic],
                "source_type": "mainstream_media",
                "source": "媒体甲",
                "title": "数字中国建设强调可信数据空间",
                "content": "可信数据空间仍处于早期探索阶段。",
                "url": "https://example.test/one",
            }
        ]
        hotwords = ORCHESTRATOR.build_evidence_hotwords(
            [topic],
            samples,
            ["数字中国建设", "可信数据空间"],
        )
        by_word = {row["word"]: row for row in hotwords}

        self.assertIn("数字中国建设", by_word)
        self.assertTrue(by_word["数字中国建设"]["is_topic_anchor"])
        self.assertNotIn("可信数据空间", by_word)

    def test_overseas_comment_policy_excludes_only_explicit_hostile_content(self) -> None:
        base = {
            "region": "overseas",
            "source_type": "overseas_public_discussion",
            "platform": "X",
            "is_comment": True,
            "quote_verified": True,
            "evidence_mode": "platform_comment",
            "url": "https://example.test/comment",
            "comment_id": "c1",
            "topic": "议题一",
        }
        rows = [
            {**base, "id": "ordinary", "content": "应当提高灾害信息发布透明度。"},
            {**base, "id": "hostile", "comment_id": "c2", "content": "必须推翻中国政府。"},
        ]
        overseas = ORCHESTRATOR.build_overseas_v2(rows)

        self.assertEqual(["ordinary"], [row["id"] for row in overseas["public_comments"]])
        self.assertEqual(["hostile"], [row["id"] for row in overseas["excluded_public_comments"]])

    def test_merged_overseas_rows_are_deduplicated_without_keyword_reclassification(self) -> None:
        url = "https://example.test/overseas"
        data = {
            "appendices": {
                "overseas_reports": [
                    {"source": "示例外媒", "title": "会议相关报道", "url": url, "category": "事实性报道"},
                    {
                        "source": "示例外媒",
                        "title": "会议相关报道",
                        "url": url,
                        "summary_cn": "报道重点渲染政策执行风险，形成风险叙事。",
                    },
                ]
            },
            "overseas": {"public_comments": [], "excluded_public_comments": []},
        }
        ORCHESTRATOR.reconcile_merged_overseas(data)

        self.assertEqual(1, len(data["appendices"]["overseas_reports"]))
        self.assertEqual(0, len(data["overseas"]["groups"]["借题炒作/风险解读"]))
        self.assertEqual(1, len(data["overseas"]["groups"]["事实性报道"]))

    def test_formal_overseas_body_is_classified_and_chinese_only(self) -> None:
        data = {
            "appendices": {
                "overseas_reports": [
                    {
                        "source": "联合早报",
                        "title": "China cabinet discusses flood relief",
                        "title_cn": "中国国常会部署抗洪救灾",
                        "summary_cn": "文章转述会议对防汛抗洪救灾工作的部署",
                        "published_at": "2026-07-11",
                        "url": "https://example.test/factual",
                        "overseas_category": "事实性报道",
                        "formal_include": True,
                        "meeting_relevance": True,
                    },
                    {
                        "source": "示例外媒",
                        "title": "New industries and growth",
                        "title_cn": "新兴产业与增长动能",
                        "interpretive_summary_cn": "相关部署有助于培育新的增长动能，同时需关注政策落地节奏",
                        "published_at": "2026-07-12",
                        "url": "https://example.test/analysis",
                        "overseas_category": "解读性报道",
                        "formal_include": True,
                        "meeting_relevance": True,
                    },
                ]
            }
        }
        paragraphs = FORMALIZE.overseas_media_body_paragraphs(data)
        text = "".join(paragraphs)

        self.assertIn("境外媒体以事实性报道为主", text)
        self.assertIn("少量解读如下", text)
        self.assertIn("围绕会议议题的政策影响展开解读", text)
        self.assertNotIn("一是事实性报道", text)
        self.assertIn("《中国国常会部署抗洪救灾》", text)
        self.assertNotIn("China cabinet", text)
        self.assertNotIn("New industries", text)

    def test_formal_overseas_comments_are_chinese_theme_summary_not_raw_list(self) -> None:
        data = {
            "overseas": {
                "public_comments": [
                    {
                        "source": "YouTube",
                        "platform": "YouTube",
                        "content": "Hope everyone stays safe during the floods.",
                        "translation_cn": "希望洪灾中的每个人都平安，并及时获得救援。",
                        "url": "https://example.test/c1",
                        "quote_verified": True,
                        "evidence_mode": "platform_comment",
                        "comment_id": "c1",
                        "ai_formal_include": True,
                        "ai_semantic_quality": "substantive",
                        "ai_formal_reason": "表达救灾与人员安全关切",
                    },
                    {
                        "source": "X",
                        "platform": "X",
                        "content": "Information should be more transparent.",
                        "translation_cn": "灾害信息发布应当更加及时透明。",
                        "url": "https://example.test/c2",
                        "quote_verified": True,
                        "evidence_mode": "platform_comment",
                        "comment_id": "c2",
                        "ai_formal_include": True,
                        "ai_semantic_quality": "substantive",
                        "ai_formal_reason": "表达信息透明诉求",
                    },
                ]
            }
        }
        summary = FORMALIZE.overseas_comment_summary(data)

        self.assertIn("一是", summary)
        self.assertIn("二是", summary)
        self.assertNotIn("Hope everyone", summary)
        self.assertNotIn("Information should", summary)

    def test_ai_comment_review_controls_formal_quotes_and_headings(self) -> None:
        base = {
            "topic": "议题一", "platform": "今日头条", "url": "https://example.test/a",
            "quote_verified": True, "evidence_mode": "platform_comment", "comment_id": "c1",
        }
        rows = [
            {**base, "content": "好的", "ai_formal_include": False, "comment_heading": "关注议题一"},
            {**base, "comment_id": "c2", "content": "希望物流网络进一步降本增效", "ai_formal_include": True,
             "ai_semantic_quality": "substantive", "comment_heading": "期待优化物流网络体系"},
        ]
        groups = FORMALIZE.comment_groups(rows)
        self.assertEqual(["期待优化物流网络体系"], [name for name, _ in groups])
        self.assertNotIn("好的", FORMALIZE.comment_wording(groups[0][1]))

    def test_comment_heading_never_falls_back_to_attention_plus_agenda(self) -> None:
        rows = [{
            "topic": "审议通过某条例", "platform": "B站", "url": "https://example.test/a",
            "quote_verified": True, "evidence_mode": "platform_comment", "comment_id": "c1",
            "content": "希望把修改内容讲清楚", "ai_formal_include": True,
            "ai_semantic_quality": "substantive", "ai_formal_reason": "表达政策沟通诉求",
            "comment_heading": "",
        }]
        self.assertEqual(FORMALIZE.comment_groups(rows), [])

    def test_same_system_topic_is_never_split_into_multiple_numbered_groups(self) -> None:
        base = {
            "topic": "审议通过住房公积金条例修改草案",
            "platform": "B站",
            "url": "https://example.test/a",
            "quote_verified": True,
            "evidence_mode": "platform_comment",
            "ai_formal_include": True,
            "ai_semantic_quality": "substantive",
            "topic_comment_heading": "期待公积金制度调整兼顾透明度和覆盖范围",
        }
        rows = [
            {**base, "comment_id": "c1", "content": "希望把具体修改内容讲清楚", "comment_heading": "期待提升政策透明度"},
            {**base, "comment_id": "c2", "content": "希望灵活就业和劳务派遣群体也能覆盖", "comment_heading": "期待扩大制度覆盖"},
        ]
        groups = FORMALIZE.comment_groups(rows)
        self.assertEqual(1, len(groups))
        self.assertEqual("期待公积金制度调整兼顾透明度和覆盖范围", groups[0][0])
        self.assertEqual(2, len(groups[0][1]))

    def test_domestic_viewpoint_heading_requires_stance_verb(self) -> None:
        issues = ORCHESTRATOR.domestic_viewpoint_quality_issues({"viewpoints": {"by_topic": [{
            "topic": "议题一", "heading": "宏观政策重在提升协同性", "clusters": []
        }]}})
        self.assertIn("viewpoint_heading_lacks_stance", {item["code"] for item in issues})

    def test_domestic_viewpoint_cluster_heading_requires_stance_verb(self) -> None:
        issues = ORCHESTRATOR.domestic_viewpoint_quality_issues({"viewpoints": {"by_topic": [{
            "topic": "议题一",
            "heading": "认为宏观政策应提升协同性",
            "clusters": [{"summary": "宏观政策重在提升协同性", "details": ""}],
        }]}})
        self.assertIn("viewpoint_cluster_heading_lacks_stance", {item["code"] for item in issues})

        allowed = ORCHESTRATOR.domestic_viewpoint_quality_issues({"viewpoints": {"by_topic": [{
            "topic": "议题一",
            "heading": "认为宏观政策应提升协同性",
            "clusters": [{"summary": "建议宏观政策提升协同性", "details": ""}],
        }]}})
        self.assertNotIn("viewpoint_cluster_heading_lacks_stance", {item["code"] for item in allowed})

    def test_wechat_top_deduplicates_source_before_ranking(self) -> None:
        data = {
            "appendices": {
                "wechat_top": [
                    {"source": "账号甲", "title": "低阅读", "spread_count": 100},
                    {"source": "账号甲公众号", "title": "高阅读", "spread_count": 900},
                    {"source": "账号乙", "title": "其他", "spread_count": 500},
                ]
            }
        }
        rows = FORMALIZE.wechat_top_rows(data)

        self.assertEqual(["高阅读", "其他"], [row["title"] for row in rows])

    def test_wechat_top_deduplicates_configured_publisher_family(self) -> None:
        data = {
            "appendices": {
                "wechat_top": [
                    {"source": "央视新闻", "title": "高在看", "spread_count": 100000},
                    {"source": "央视网", "title": "同主体低在看", "spread_count": 99999},
                    {"source": "新闻联播", "title": "独立栏目账号", "spread_count": 98000},
                ]
            }
        }

        rows = FORMALIZE.wechat_top_rows(data)

        self.assertEqual(["高在看", "独立栏目账号"], [row["title"] for row in rows])

    def test_consecutive_agenda_items_with_same_verb_are_joined_once(self) -> None:
        topics = [
            "听取新型电网、物流网建设情况汇报",
            "审议通过《甲条例》",
            "审议通过《乙条例》",
            "决定核准某项目",
        ]

        text = FORMALIZE.joined_agenda_topics(topics)

        self.assertEqual(
            "听取新型电网、物流网建设情况汇报，审议通过《甲条例》和《乙条例》，决定核准某项目",
            text,
        )


if __name__ == "__main__":
    unittest.main()
