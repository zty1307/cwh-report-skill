"""Generic variable-data templates, not regression assertions tied to one meeting."""
from copy import deepcopy
from pathlib import Path
import sys
import json
import pytest
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import cwh_report_visuals as visuals
import formalize_cwh_report as formal


def example(n=6):
    return {'meeting':{'date':'2031-02-03'},
        'statistics':{'total_samples':7,'total_spread':42000,'by_date':{'2031-02-03':12000,'2031-02-04':30000}},
        'topic_stats':[{'topic':'实施公共服务政策和具体配套措施'+str(i),'spread_count':(i+1)*1000} for i in range(n)]}


@pytest.mark.parametrize('n',[4,6,9])
def test_charts_and_table_accept_variable_topics_without_period_constants(tmp_path,n):
    data=example(n); original=deepcopy(data)
    paths,manifests=visuals.write_report_charts(data,tmp_path)
    assert visuals.verify_report_charts(data,paths,manifests)
    assert manifests['trend_distribution']['title']=='2月3日国务院常务会议舆情走势'
    assert manifests['topic_distribution']['title']=='2月3日国务院常务会议子议题传播量'
    assert manifests['trend_distribution']['input']['total']==42000  # not seven evidence samples
    assert manifests['topic_distribution']['input']['topics'][0]['value']==n*1000
    with Image.open(paths['topic_distribution']) as im:
        assert (0,51,102) in {c for _,c in im.getcolors(im.width*im.height)}
    doc=Document();table=formal.add_topic_table(doc,formal.formal_topic_rows(data,False),False,'境内新闻')
    assert len(table.columns)==9 and len(table.rows)==n+2
    assert [int(c.get(qn('w:w'))) for c in table._tbl.tblGrid]==visuals.style()['table_grid_twips']
    assert table.cell(0,6).text=='网民情感'
    assert table.cell(1,8).text=='负面'
    assert all(c.text=='' for row in table.rows[2:] for c in row.cells[6:])
    assert all(c.paragraphs[0].paragraph_format.first_line_indent.pt==0 for row in table.rows for c in row.cells)
    assert all(c.paragraphs[0].paragraph_format.line_spacing.pt==31.2 for row in table.rows for c in row.cells)
    assert data==original


def test_manifest_rejects_changed_counts_dates_and_images(tmp_path):
    data=example(); paths,manifests=visuals.write_report_charts(data,tmp_path)
    for change in ('date','total','topic','trend','manifest'):
        d,m=deepcopy(data),deepcopy(manifests)
        if change=='date':d['meeting']['date']='2031-04-05'
        if change=='total':d['statistics']['total_spread']+=1
        if change=='topic':d['topic_stats'][0]['spread_count']+=1
        if change=='trend':d['statistics']['by_date']['2031-02-03']+=1
        if change=='manifest':m['topic_distribution']['config_sha256']='altered'
        assert not visuals.verify_report_charts(d,paths,m)
    Image.new('RGB',(20,20),'red').save(paths['trend_distribution'])
    assert not visuals.verify_report_charts(data,paths,manifests)


def test_empty_zero_duplicate_and_long_labels_preserved(tmp_path):
    data=example(0)
    data['statistics']['by_date']={}
    data['topic_stats']=[{'topic':'不能截断的原始标题'*16,'spread_count':0},
                         {'topic':'相同标题','spread_count':10000},
                         {'topic':'相同标题','spread_count':10000}]
    paths,m=visuals.write_report_charts(data,tmp_path)
    assert len(m['topic_distribution']['input']['topics'])==3
    assert m['topic_distribution']['input']['topics'][-1]['label']==data['topic_stats'][0]['topic']
    assert visuals.ten_thousands(10000)=='1'
    assert visuals.ten_thousands(0)=='0'
    assert visuals.verify_report_charts(data,paths,m)


def test_system_chart_cannot_override_new_template_and_audit_checks_embedding(tmp_path):
    data=example()
    system=tmp_path/'arbitrary.png'; Image.new('RGB',(20,20),'red').save(system)
    cloud=tmp_path/'cloud.png'; Image.new('RGB',(20,20),'green').save(cloud)
    data['artifacts']={'trend_chart_image':str(system),'topic_chart_image':str(system),'wordcloud_image':str(cloud)}
    paths=formal.ensure_docx_chart_images(data,tmp_path/'report')
    assert all(Path(paths[k]).name.endswith('_template.png') for k in ('trend_distribution','topic_distribution'))
    assert system.read_bytes()!=Path(paths['trend_distribution']).read_bytes()
    doc=Document()
    for key in ('trend_distribution','topic_distribution','hotword_distribution'):doc.add_picture(paths[key])
    path=tmp_path/'report.docx';doc.save(path)
    assert formal.audit_formal_docx(data,path)['checks']['uses_traceable_fixed_chart_assets']


def test_table_retains_verified_ratios_and_never_promotes_missing_review():
    doc=Document();table=formal.add_topic_table(doc,[[1,'议题',2,3,4,9,'40%','50%','10%'],
        [2,'另一议题',1,1,1,3,'待分析','待分析','待分析']],True)
    assert [c.text for c in table.rows[2].cells[6:]]==['40%','50%','10%']
    assert [c.text for c in table.rows[3].cells[6:]]==['','','']


@pytest.mark.parametrize('fault',['grid','merge','indent','font','line_spacing'])
def test_table_audit_rejects_format_drift(fault):
    doc=Document();table=formal.add_topic_table(doc,[[1,'任意本期政策',2,3,4,9]],False)
    assert visuals.verify_topic_table(doc._element.xml.encode())
    if fault=='grid':table._tbl.tblGrid[0].set(qn('w:w'),'800')
    if fault=='merge':table._tbl.tr_lst[0].tc_lst[6].tcPr.gridSpan.val=2
    if fault=='indent':table.cell(2,1).paragraphs[0].paragraph_format.first_line_indent=100000
    if fault=='font':table.cell(2,1).paragraphs[0].runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    if fault=='line_spacing':table.cell(2,1).paragraphs[0].paragraph_format.line_spacing=1
    assert not visuals.verify_topic_table(doc._element.xml.encode())
