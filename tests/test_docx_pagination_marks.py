from __future__ import annotations

import importlib.util
import sys
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "formalize_cwh_report.py"
sys.path.insert(0, str(SCRIPT.parent))
SPEC = importlib.util.spec_from_file_location("formalize_cwh_report_for_pagination_test", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
SPEC.loader.exec_module(MODULE)


class DocxPaginationMarkTests(unittest.TestCase):
    def test_body_flows_but_headings_and_table_headers_stay_with_content(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("测试正文")
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.page_break_before = True
        document.styles["Heading 1"].paragraph_format.keep_with_next = True
        heading = document.add_heading('一级标题', 1)
        table = MODULE.add_table(document, ['表头'], [['数据']], widths_pt=[100])

        MODULE.remove_nonprinting_pagination_controls(document)

        for root in (document._element, document.styles.element):
            self.assertEqual([], list(root.iter(qn("w:keepLines"))))
            self.assertEqual([], list(root.iter(qn("w:pageBreakBefore"))))
        self.assertFalse(paragraph.paragraph_format.keep_with_next)
        self.assertTrue(heading.paragraph_format.keep_with_next)
        self.assertTrue(table.cell(0, 0).paragraphs[0].paragraph_format.keep_with_next)
        self.assertFalse(table.cell(1, 0).paragraphs[0].paragraph_format.keep_with_next)


if __name__ == "__main__":
    unittest.main()
